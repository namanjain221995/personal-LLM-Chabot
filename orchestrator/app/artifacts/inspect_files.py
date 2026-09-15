"""What the PRODUCED files say — never what the renderer was asked to do.

WHY A SECOND READER. The self-check (artifacts/selfcheck.py) must be able to
say "the headings are dark blue" only when the bytes the person downloads
make them dark blue. Reading the spec, the ResolvedStyle or the CSS the
renderer generated would check the renderer's INPUT: a renderer bug, a
style the model reviewer dropped, a font the container substituted, all
pass. So everything here opens the rendered files themselves:

  DOCX  the package XML (document.xml, styles.xml with the basedOn chain
        and docDefaults, theme1.xml for theme fonts, header/footer parts,
        every .rels part) — run properties are RESOLVED the way Word
        resolves them: docDefaults < paragraph style chain < character
        style chain < direct run formatting.
  XLSX  openpyxl with data_only=False (so formulas are formulas), cell
        fills/fonts, conditional-format rules and priorities, freeze
        panes, number formats, native charts and the cells their series
        reference.
  PPTX  python-pptx: slide titles, runs, table header fills, native charts
        with their cached values.
  PDF   pypdfium2's C API: per text object the fill colour, the font's
        base name, weight and flags, the effective size (font size x the
        object matrix), and the filled paths behind the text (table header
        fills), plus MediaBox per page.
  CSV   the raw text: header, row count, neutralised formula leads.
  PNG   dimensions and dominant non-neutral colours (Pillow).
  SVG   safety scan plus the fill/stroke colours it declares.

THE VOCABULARY is requirements.py's: an Observation carries (target,
property, value) with target in title|subtitle|heading1..3|paragraph|
table_header|table_total|column:<name>|row:<n>|cell:<A1>|slide_title|chart|
page|file|document|sheet|security, so the evaluator compares like with like.
A property the format cannot carry, or this reader cannot see, is simply
not observed; the evaluator turns "no observation" into `unverifiable`,
never into a pass.

Every function is synchronous and CPU-bound: callers run inspect() through
asyncio.to_thread (the Fast-mode event-loop lesson, 2026-09-05).
"""
from __future__ import annotations

import csv as _csv
import io
import logging
import math
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

log = logging.getLogger(__name__)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_W = "{%s}" % W_NS
_A = "{%s}" % A_NS

#: Zip members read per package, and bytes per member — a hostile or broken
#: file must not make the checker the slow part of a job.
_MAX_MEMBER_BYTES = 64 * 1024 * 1024
#: Observations kept per (format, target kind); a 400-block document has
#: thousands of runs and the evaluator needs proportions, not every run.
_MAX_PER_TARGET = 600

SAFE_LINK_SCHEMES = ("http://", "https://", "mailto:")
_FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")
_PLAIN_NUMBER_RE = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")
_RENDERER_FORMULA_RE = re.compile(r"^=(SUM|SUBTOTAL|AVERAGE|COUNT|COUNTA|MIN|MAX)\((\d{1,3},\s*)?\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+\)$", re.I)


@dataclass
class Observation:
    target: str
    property: str
    value: Any
    format: str
    locator: Dict[str, Any] = field(default_factory=dict)
    verifiable: bool = True

    def to_dict(self) -> dict:
        value = self.value
        if isinstance(value, str) and len(value) > 200:
            value = value[:200] + "…"
        if isinstance(value, (list, tuple)) and len(value) > 20:
            value = list(value[:20]) + ["…"]
        return {"target": self.target, "property": self.property, "value": value, "format": self.format,
                "locator": dict(self.locator), "verifiable": self.verifiable}


# ------------------------------------------------------------ helpers --


def norm_hex(value: Any) -> Optional[str]:
    """'1f3864' / '#1F3864' / 'FF1F3864' (ARGB) / (31, 56, 100) -> '#1F3864'."""
    if value is None:
        return None
    if isinstance(value, (tuple, list)) and len(value) >= 3:
        try:
            return "#%02X%02X%02X" % tuple(max(0, min(255, int(v))) for v in value[:3])
        except (TypeError, ValueError):
            return None
    text = str(value).strip().lstrip("#")
    if len(text) == 8 and re.fullmatch(r"[0-9A-Fa-f]{8}", text):
        text = text[2:]
    if len(text) == 3 and re.fullmatch(r"[0-9A-Fa-f]{3}", text):
        text = "".join(c * 2 for c in text)
    if re.fullmatch(r"[0-9A-Fa-f]{6}", text):
        return "#" + text.upper()
    return None


def norm_text(text: Any) -> str:
    """Case- and whitespace-folded text with markdown/typographic noise off,
    for coverage comparisons (a heading 'Risk & Controls' in markdown is
    'risk & controls' in the file)."""
    s = str(text or "")
    s = s.replace("\u00a0", " ").replace("\u2019", "'").replace("\u2018", "'").replace("\u201c", '"').replace("\u201d", '"')
    s = s.replace("\u2013", "-").replace("\u2014", "-").replace("\u200b", "").replace("\u200c", "").replace("\u200d", "")
    s = re.sub(r"[*_`#>|]+", " ", s)
    return " ".join(s.split()).casefold()


def _read_member(zf: zipfile.ZipFile, name: str) -> Optional[bytes]:
    try:
        info = zf.getinfo(name)
    except KeyError:
        return None
    if info.file_size > _MAX_MEMBER_BYTES:
        return None
    return zf.read(name)


def _xml(data: Optional[bytes]):
    if not data:
        return None
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False, load_dtd=False)
    try:
        return etree.fromstring(data, parser)
    except etree.XMLSyntaxError:
        return None


def contrast_ratio(a: str, b: str) -> float:
    def lum(h: str) -> float:
        h = h.lstrip("#")
        rgb = [int(h[i:i + 2], 16) / 255.0 for i in (0, 2, 4)]
        lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in rgb]
        return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]

    la, lb = lum(a), lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def script_of(text: str) -> Dict[str, int]:
    counts = {"devanagari": 0, "gujarati": 0, "latin": 0}
    for ch in text or "":
        o = ord(ch)
        if 0x0900 <= o <= 0x097F:
            counts["devanagari"] += 1
        elif 0x0A80 <= o <= 0x0AFF:
            counts["gujarati"] += 1
        elif ch.isascii() and ch.isalpha():
            counts["latin"] += 1
    return counts


#: Fonts that cover an Indic script (by family-name fragment). Used for the
#: PDF tofu check: text in Devanagari set in a font whose name carries none
#: of these is reported as not covered.
_SCRIPT_FONTS = {
    "devanagari": ("devanagari", "lohit", "mangal", "nirmala", "gargi", "kalimati", "sarai", "hind", "mukta", "annapurna", "noto sans dev", "notosansdevanagari", "freesans", "arial unicode"),
    "gujarati": ("gujarati", "lohit", "shruti", "nirmala", "rekha", "padmaa", "notosansgujarati", "freesans", "arial unicode"),
}


# --------------------------------------------------------------- DOCX --


def _rpr_props(rpr, theme_fonts: Dict[str, str]) -> Dict[str, Any]:
    """The properties one w:rPr sets (None where it sets nothing)."""
    out: Dict[str, Any] = {}
    if rpr is None:
        return out
    for tag, key in (("b", "bold"), ("i", "italic")):
        el = rpr.find(_W + tag)
        if el is not None:
            val = el.get(_W + "val")
            out[key] = val not in ("0", "false", "off")
    u = rpr.find(_W + "u")
    if u is not None:
        out["underline"] = (u.get(_W + "val") or "single") not in ("none", "0", "false")
    color = rpr.find(_W + "color")
    if color is not None and color.get(_W + "val") and color.get(_W + "val") != "auto":
        out["color"] = norm_hex(color.get(_W + "val"))
    sz = rpr.find(_W + "sz")
    if sz is not None and (sz.get(_W + "val") or "").isdigit():
        out["size_pt"] = int(sz.get(_W + "val")) / 2.0
    fonts = rpr.find(_W + "rFonts")
    if fonts is not None:
        name = fonts.get(_W + "ascii") or fonts.get(_W + "hAnsi")
        theme = fonts.get(_W + "asciiTheme") or fonts.get(_W + "hAnsiTheme")
        if theme:
            name = theme_fonts.get("major" if theme.startswith("major") else "minor") or name
        if name:
            out["font_family"] = name
    spacing = rpr.find(_W + "spacing")
    if spacing is not None and spacing.get(_W + "val") not in (None, "0"):
        try:
            out["letter_spacing"] = int(spacing.get(_W + "val")) / 20.0
        except ValueError:
            pass
    shd = rpr.find(_W + "shd")
    if shd is not None and shd.get(_W + "fill") not in (None, "auto"):
        out["background"] = norm_hex(shd.get(_W + "fill"))
    return out


class _DocxStyles:
    def __init__(self, styles_root, theme_fonts: Dict[str, str]) -> None:
        self.theme_fonts = theme_fonts
        self.by_id: Dict[str, Any] = {}
        self.default_para: Optional[str] = None
        self.defaults: Dict[str, Any] = {}
        if styles_root is None:
            return
        dd = styles_root.find(f"{_W}docDefaults/{_W}rPrDefault/{_W}rPr")
        self.defaults = _rpr_props(dd, theme_fonts)
        for st in styles_root.findall(_W + "style"):
            sid = st.get(_W + "styleId")
            if not sid:
                continue
            self.by_id[sid] = st
            if st.get(_W + "type") == "paragraph" and st.get(_W + "default") in ("1", "true"):
                self.default_para = sid

    def name(self, sid: Optional[str]) -> str:
        st = self.by_id.get(sid or "")
        if st is None:
            return ""
        el = st.find(_W + "name")
        return str(el.get(_W + "val") if el is not None else sid)

    def chain_props(self, sid: Optional[str]) -> Dict[str, Any]:
        chain: List[Any] = []
        seen = set()
        while sid and sid not in seen and sid in self.by_id:
            seen.add(sid)
            st = self.by_id[sid]
            chain.append(st)
            based = st.find(_W + "basedOn")
            sid = based.get(_W + "val") if based is not None else None
        props: Dict[str, Any] = {}
        for st in reversed(chain):
            props.update(_rpr_props(st.find(_W + "rPr"), self.theme_fonts))
        return props

    def para_border(self, sid: Optional[str]) -> bool:
        seen = set()
        while sid and sid not in seen and sid in self.by_id:
            seen.add(sid)
            st = self.by_id[sid]
            ppr = st.find(_W + "pPr")
            if ppr is not None and ppr.find(_W + "pBdr") is not None:
                return True
            based = st.find(_W + "basedOn")
            sid = based.get(_W + "val") if based is not None else None
        return False


def _docx_theme_fonts(zf: zipfile.ZipFile) -> Dict[str, str]:
    root = _xml(_read_member(zf, "word/theme/theme1.xml"))
    out: Dict[str, str] = {}
    if root is None:
        return out
    for kind in ("major", "minor"):
        el = root.find(f".//{_A}{kind}Font/{_A}latin")
        if el is not None and el.get("typeface"):
            out[kind] = el.get("typeface")
    return out


def _para_text(p) -> str:
    parts: List[str] = []
    for el in p.iter(_W + "t", _W + "tab", _W + "br"):
        if el.tag == _W + "t":
            parts.append(el.text or "")
        else:
            parts.append(" ")
    return "".join(parts)


def _classify_style(name: str) -> str:
    low = name.strip().lower()
    if low == "title":
        return "title"
    if low == "subtitle":
        return "subtitle"
    m = re.fullmatch(r"heading ([1-9])", low)
    if m:
        return f"heading{min(int(m.group(1)), 3)}"
    if low in ("normal", "body", "body text", "bullet", "number", "list bullet", "list number", "list paragraph", "callout", "lede", "quote"):
        return "paragraph"
    return ""


def _is_list_style(name: str) -> bool:
    return name.strip().lower() in ("bullet", "number", "list bullet", "list number", "list paragraph")


def _weighted(runs: Sequence[Tuple[int, Dict[str, Any]]], key: str) -> Any:
    """The value of `key` carried by the most characters of a paragraph."""
    tally: Dict[Any, int] = {}
    for n, props in runs:
        if key in props:
            tally[props[key]] = tally.get(props[key], 0) + max(1, n)
    if not tally:
        return None
    return max(tally.items(), key=lambda kv: kv[1])[0]


_STYLE_KEYS = ("color", "font_family", "size_pt", "bold", "italic", "underline", "background", "letter_spacing")


def _docx_runs(p, styles: _DocxStyles, para_props: Dict[str, Any]) -> List[Tuple[int, Dict[str, Any]]]:
    runs: List[Tuple[int, Dict[str, Any]]] = []
    for r in p.iter(_W + "r"):
        text = "".join(t.text or "" for t in r.findall(_W + "t"))
        if not text.strip():
            continue
        rpr = r.find(_W + "rPr")
        props = dict(para_props)
        if rpr is not None:
            rstyle = rpr.find(_W + "rStyle")
            if rstyle is not None:
                props.update(styles.chain_props(rstyle.get(_W + "val")))
            props.update(_rpr_props(rpr, styles.theme_fonts))
        runs.append((len(text), props))
    return runs


def _page_class(w: float, h: float) -> Tuple[str, str]:
    """(orientation, size name) from a page's width and height in mm."""
    orientation = "landscape" if w > h + 1 else "portrait"
    short, long_ = sorted((w, h))
    size = "other"
    for name, (a, b) in (("A4", (210.0, 297.0)), ("Letter", (215.9, 279.4)), ("Legal", (215.9, 355.6)), ("A3", (297.0, 420.0))):
        if abs(short - a) <= 2.5 and abs(long_ - b) <= 2.5:
            size = name
            break
    return orientation, size


def inspect_docx(path: Path) -> List[Observation]:
    fmt = "docx"
    obs: List[Observation] = []
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        theme_fonts = _docx_theme_fonts(zf)
        styles = _DocxStyles(_xml(_read_member(zf, "word/styles.xml")), theme_fonts)
        doc = _xml(_read_member(zf, "word/document.xml"))
        if doc is None:
            return [Observation("file", "readable", False, fmt)]
        body = doc.find(_W + "body")
        base = dict(styles.defaults)

        def para_props(p) -> Tuple[str, str, Dict[str, Any]]:
            ppr = p.find(_W + "pPr")
            sid = None
            if ppr is not None and ppr.find(_W + "pStyle") is not None:
                sid = ppr.find(_W + "pStyle").get(_W + "val")
            sid = sid or styles.default_para
            props = dict(base)
            props.update(styles.chain_props(sid))
            return sid or "", styles.name(sid), props

        headings: List[Tuple[int, str]] = []
        sections: Dict[str, List[str]] = {}
        current = "(before the first heading)"
        text_lines: List[str] = []
        counts: Dict[str, int] = {}
        cells: List[str] = []
        table_index = 0
        # The heading path each element sits under (level 1..3) and the
        # ordinal of its level-1 section, so a style scoped to one section
        # ("the Risks section paragraphs in italic") is judged on that
        # section only.
        path: List[str] = ["", "", ""]
        h1_ordinal = 0

        def where() -> Dict[str, Any]:
            return {"sections": [s for s in path if s], "section_index": h1_ordinal}

        def emit(target: str, props: Dict[str, Any], locator: Dict[str, Any]) -> None:
            counts[target] = counts.get(target, 0) + 1
            if counts[target] > _MAX_PER_TARGET:
                return
            locator = {**where(), **locator}
            for key in _STYLE_KEYS:
                if key != "letter_spacing" and key in props and props[key] is not None:
                    obs.append(Observation(target, key, props[key], fmt, dict(locator)))

        for child in (body if body is not None else []):
            if child.tag == _W + "p":
                sid, sname, pprops = para_props(child)
                text = _para_text(child).strip()
                if not text:
                    continue
                text_lines.append(text)
                kind = _classify_style(sname)
                runs = _docx_runs(child, styles, pprops)
                agg = {k: _weighted(runs, k) for k in _STYLE_KEYS}
                agg = {k: v for k, v in agg.items() if v is not None}
                for k in ("bold", "italic", "underline"):
                    agg.setdefault(k, False)
                if kind.startswith("heading"):
                    level = int(kind[-1])
                    headings.append((level, text))
                    if level == 1:
                        h1_ordinal += 1
                    path[level - 1] = text
                    for deeper in range(level, 3):
                        path[deeper] = ""
                    emit(kind, agg, {"text": text[:120], "index": len(headings) - 1, "level": level})
                    if level == 1:
                        current = text
                elif kind in ("title", "subtitle"):
                    emit(kind, agg, {"text": text[:120]})
                    if kind == "title":
                        obs.append(Observation("title", "letter_spacing", float(agg.get("letter_spacing") or 0.0), fmt))
                        obs.append(Observation("title", "border", styles.para_border(sid), fmt))
                elif kind == "paragraph":
                    emit("paragraph", agg, {"chars": len(text), "list": _is_list_style(sname)})
                sections.setdefault(current, []).append(text)
            elif child.tag == _W + "tbl":
                rows = child.findall(_W + "tr")
                header_names: List[str] = []
                for r_i, tr in enumerate(rows):
                    row_texts: List[str] = []
                    for c_i, tc in enumerate(tr.findall(_W + "tc")):
                        ctext = " ".join(_para_text(p).strip() for p in tc.findall(_W + "p")).strip()
                        row_texts.append(ctext)
                        cells.append(ctext)
                        tcpr = tc.find(_W + "tcPr")
                        fill = None
                        if tcpr is not None and tcpr.find(_W + "shd") is not None:
                            f = tcpr.find(_W + "shd").get(_W + "fill")
                            fill = norm_hex(f) if f and f != "auto" else None
                        run_props: List[Tuple[int, Dict[str, Any]]] = []
                        for p in tc.findall(_W + "p"):
                            _, _, pp = para_props(p)
                            run_props.extend(_docx_runs(p, styles, pp))
                        agg = {k: _weighted(run_props, k) for k in _STYLE_KEYS}
                        agg = {k: v for k, v in agg.items() if v is not None}
                        for k in ("bold", "italic", "underline"):
                            agg.setdefault(k, False)
                        agg["background"] = fill or agg.get("background") or "#FFFFFF"
                        loc = {"table": table_index, "row": r_i, "col": c_i}
                        if r_i == 0:
                            header_names.append(ctext)
                            emit("table_header", agg, {**loc, "text": ctext[:60]})
                        else:
                            if c_i < len(header_names) and header_names[c_i]:
                                emit("column:" + norm_text(header_names[c_i]), agg, loc)
                            if row_texts and norm_text(row_texts[0]).startswith("total"):
                                emit("table_total", agg, loc)
                    text_lines.append(" | ".join(row_texts))
                    sections.setdefault(current, []).append(" | ".join(row_texts))
                obs.append(Observation("table", "columns", header_names, fmt, {"table": table_index}))
                table_index += 1

        # Section geometry: every sectPr (paragraph-level ones end a section).
        for sect in doc.iter(_W + "sectPr"):
            pgsz = sect.find(_W + "pgSz")
            pgmar = sect.find(_W + "pgMar")
            if pgsz is not None:
                try:
                    w_mm = int(pgsz.get(_W + "w")) / 1440 * 25.4
                    h_mm = int(pgsz.get(_W + "h")) / 1440 * 25.4
                except (TypeError, ValueError):
                    continue
                orientation, size = _page_class(w_mm, h_mm)
                if pgsz.get(_W + "orient") == "landscape" and w_mm < h_mm:
                    orientation = "inconsistent"  # orient says landscape, geometry says portrait
                obs.append(Observation("page", "orientation", orientation, fmt))
                obs.append(Observation("page", "page_size", size, fmt, {"w_mm": round(w_mm, 1), "h_mm": round(h_mm, 1)}))
            if pgmar is not None:
                try:
                    margins = {k: round(int(pgmar.get(_W + k)) / 1440 * 25.4, 1) for k in ("top", "bottom", "left", "right")}
                    obs.append(Observation("page", "margins_mm", margins, fmt))
                except (TypeError, ValueError):
                    pass

        # Page numbers live in header/footer parts as PAGE fields.
        has_page_field = False
        for name in names:
            if re.fullmatch(r"word/(header|footer)\d*\.xml", name):
                data = _read_member(zf, name) or b""
                root = _xml(data)
                if root is None:
                    continue
                instr = " ".join((el.text or "") for el in root.iter(_W + "instrText"))
                instr += " " + " ".join(el.get(_W + "instr") or "" for el in root.iter(_W + "fldSimple"))
                if re.search(r"\bPAGE\b", instr):
                    has_page_field = True
        obs.append(Observation("page", "page_numbers", has_page_field, fmt))

        # External relationships (security) and images.
        external: List[str] = []
        for name in names:
            if not name.endswith(".rels"):
                continue
            root = _xml(_read_member(zf, name))
            if root is None:
                continue
            for rel in root:
                if rel.get("TargetMode") == "External":
                    external.append(str(rel.get("Target") or ""))
        obs.append(Observation("security", "external_targets", external, fmt))
        obs.append(Observation("security", "unsafe_external_targets", [t for t in external if not t.lower().startswith(SAFE_LINK_SCHEMES)], fmt))
        media = [n for n in names if n.startswith("word/media/")]
        obs.append(Observation("document", "image_count", len(media), fmt))
        # Charts reach a DOCX as PNGs: their dominant colours are the only
        # file-level evidence of a series colour there.
        for n in media[:12]:
            if not n.lower().endswith(".png"):
                continue
            data = _read_member(zf, n)
            if not data:
                continue
            try:
                for o in _png_observations(io.BytesIO(data), n):
                    obs.append(Observation(o.target, o.property, o.value, fmt, o.locator))
            except Exception:  # noqa: BLE001
                continue
        obs.append(Observation("document", "headings", [t for _, t in headings], fmt, {"levels": [lvl for lvl, _ in headings]}))
        obs.append(Observation("document", "table_cells", cells, fmt))
        obs.append(Observation("document", "text", "\n".join(text_lines), fmt))
        obs.append(Observation("document", "sections", {k: "\n".join(v) for k, v in sections.items()}, fmt))
        obs.append(Observation("document", "fonts_cover_script", None, fmt, verifiable=False))
    return obs


# --------------------------------------------------------------- XLSX --


def _xlsx_colour(color) -> Optional[str]:
    """An openpyxl Color as hex, or None when it is a theme/indexed colour
    this reader does not resolve (reported unverifiable, never guessed)."""
    if color is None:
        return None
    try:
        if getattr(color, "type", "rgb") == "rgb" and isinstance(color.rgb, str):
            return norm_hex(color.rgb)
    except Exception:  # noqa: BLE001
        return None
    return None


def _cf_formula_true(formula: str, cell, ws) -> Optional[bool]:
    """Evaluate the small, closed set of conditional-format expressions the
    renderers write. None = not evaluable here (the fill is unverifiable)."""
    f = (formula or "").strip().lstrip("=").replace(" ", "").upper()
    if f in ("TRUE", "1", "TRUE()"):
        return True
    if f in ("FALSE", "0", "FALSE()"):
        return False
    m = re.fullmatch(r"MOD\(ROW\(\),2\)=([01])", f)
    if m:
        return cell.row % 2 == int(m.group(1))
    m = re.fullmatch(r"ROW\(\)=(\d+)", f)
    if m:
        return cell.row == int(m.group(1))
    m = re.fullmatch(r'\$?([A-Z]{1,3})\$?(\d+)="((?:[^"]|"")*)"', (formula or "").strip().lstrip("=").replace(" ", "").upper())
    if m:
        col, _, lit = m.groups()
        other = ws[f"{col}{cell.row}"].value
        return str(other or "").upper() == lit.replace('""', '"')
    return None


def _cf_fill_for(cell, ws, rules: List[Tuple[int, Any, Any]]) -> Tuple[Optional[str], bool]:
    """(fill hex from conditional formatting, verifiable). `rules` is
    (priority, CellRange set, rule) sorted by priority (1 = highest)."""
    for _, ranges, rule in rules:
        if not any(cell.coordinate in rng for rng in ranges):
            continue
        dxf = getattr(rule, "dxf", None)
        fill = getattr(dxf, "fill", None) if dxf is not None else None
        colour = None
        if fill is not None:
            colour = _xlsx_colour(getattr(fill, "bgColor", None)) or _xlsx_colour(getattr(fill, "fgColor", None))
        if colour is None:
            continue
        rtype = getattr(rule, "type", "")
        if rtype == "expression":
            truth = _cf_formula_true((rule.formula or [""])[0], cell, ws)
        elif rtype == "cellIs" and getattr(rule, "operator", "") == "equal" and rule.formula:
            lit = str(rule.formula[0]).strip()
            truth = str(cell.value) == (lit[1:-1].replace('""', '"') if lit.startswith('"') else lit)
        elif rtype == "containsText":
            truth = str(getattr(rule, "text", "") or "").lower() in str(cell.value or "").lower()
        else:
            truth = None
        if truth is None:
            return None, False
        if truth:
            return colour, True
    return None, True


def _srgb(choice) -> Optional[str]:
    """An openpyxl ColorChoice's srgbClr (a str on read, an RGB object on
    some write paths) as hex; None for scheme/system colours."""
    if choice is None:
        return None
    value = getattr(choice, "srgbClr", None)
    if value is None:
        return None
    return norm_hex(value if isinstance(value, str) else getattr(value, "val", None))


def _chart_obs(chart, fmt: str, locator: Dict[str, Any]) -> List[Observation]:
    out: List[Observation] = []
    tag = getattr(chart, "tagname", "")
    kind = {"barChart": "bar", "bar3DChart": "bar", "lineChart": "line", "line3DChart": "line", "pieChart": "pie",
            "pie3DChart": "pie", "doughnutChart": "donut", "scatterChart": "scatter", "areaChart": "area",
            "radarChart": "radar", "bubbleChart": "bubble"}.get(tag, tag or "unknown")
    grouping = str(getattr(chart, "grouping", "") or "")
    if kind == "bar" and grouping in ("stacked", "percentStacked"):
        kind = "stacked_bar"
    out.append(Observation("chart", "type", kind, fmt, {**locator, "bar_dir": str(getattr(chart, "barDir", "") or "")}))
    title = ""
    try:
        rich = chart.title.tx.rich
        title = "".join((r.t or "") for p in rich.p for r in (p.r or []))
    except Exception:  # noqa: BLE001
        title = ""
    out.append(Observation("chart", "title", title, fmt, dict(locator)))
    colours: List[Optional[str]] = []
    labels = bool(getattr(chart, "dataLabels", None) is not None)
    trend = False
    for ser in list(getattr(chart, "series", None) or getattr(chart, "ser", None) or []):
        gp = getattr(ser, "graphicalProperties", None) or getattr(ser, "spPr", None)
        colour = None
        try:
            colour = _srgb(gp.solidFill) if gp is not None else None
            if colour is None and gp is not None and gp.line is not None:
                colour = _srgb(gp.line.solidFill)
        except Exception:  # noqa: BLE001
            colour = None
        colours.append(colour)
        if getattr(ser, "dLbls", None) is not None:
            labels = True
        if getattr(ser, "trendline", None) is not None:
            trend = True
    out.append(Observation("chart", "series_colors", colours, fmt, dict(locator), verifiable=any(colours)))
    out.append(Observation("chart", "data_labels", labels, fmt, dict(locator)))
    out.append(Observation("chart", "trendline", trend, fmt, dict(locator)))
    legend = getattr(chart, "legend", None)
    pos = {"b": "bottom", "t": "top", "r": "right", "l": "left", "tr": "top_right"}.get(str(getattr(legend, "position", "") or ""), "none" if legend is None else "right")
    out.append(Observation("chart", "legend_position", pos, fmt, dict(locator)))
    return out


def _ref_values(wb, formula: str) -> Optional[List[Any]]:
    m = re.fullmatch(r"'?(.+?)'?!\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?", (formula or "").strip())
    if not m:
        return None
    sheet, c1, r1, c2, r2 = m.groups()
    if sheet not in wb.sheetnames:
        return None
    ws = wb[sheet]
    rng = f"{c1}{r1}:{c2 or c1}{r2 or r1}"
    values: List[Any] = []
    for row in ws[rng] if ":" in rng else [[ws[rng]]]:
        for cell in (row if isinstance(row, tuple) else row):
            values.append(cell.value)
    return values


_AUXILIARY_SHEET_RE = re.compile(r"(?:notes(?:-\d{1,2})?|chart data)", re.I)


def _spec_sheet_titles(spec) -> Optional[set]:
    try:
        body = spec.body
        if getattr(body, "sheets", None) is None:
            return None
        from .render.xlsx import sheet_titles

        return {t.casefold() for t in sheet_titles(body)}
    except Exception:  # noqa: BLE001 — no spec, or not a workbook: judge by name only
        return None


def inspect_xlsx(path: Path, spec=None) -> List[Observation]:
    import openpyxl
    from openpyxl.worksheet.cell_range import CellRange

    fmt = "xlsx"
    obs: List[Observation] = []
    wb = openpyxl.load_workbook(str(path), data_only=False)
    all_headers: List[str] = []
    formula_flags: List[str] = []
    unexpected_formulas: List[str] = []
    data_titles = _spec_sheet_titles(spec)
    for s_index, ws in enumerate(wb.worksheets):
        loc_sheet = {"sheet": ws.title, "sheet_index": s_index}
        # The renderer's own "Notes" and "Chart data" sheets are not tables:
        # their A1 is a bold label ("Assumptions", a chart title), never a
        # header row. Read as tables they failed "a filled header row" on
        # eight of eight live workbooks with notes or native charts
        # (2026-09-15). A spec sheet that happens to be called "Notes" is
        # still a table.
        auxiliary = bool(_AUXILIARY_SHEET_RE.fullmatch(ws.title.strip())) and (data_titles is None or ws.title.casefold() not in data_titles)
        headers: List[str] = []
        for cell in (ws[1] if ws.max_row >= 1 and not auxiliary else []):
            if cell.value is None or str(cell.value).strip() == "":
                break
            headers.append(str(cell.value))
        all_headers.extend(headers)
        obs.append(Observation("sheet", "columns", headers, fmt, dict(loc_sheet)))
        obs.append(Observation("sheet", "freeze_panes", ws.freeze_panes, fmt, dict(loc_sheet)))
        rules: List[Tuple[int, Any, Any]] = []
        try:
            for cf in ws.conditional_formatting:
                ranges = [CellRange(str(r)) for r in cf.sqref.ranges]
                for rule in cf.rules:
                    rules.append((int(rule.priority or 9999), ranges, rule))
        except Exception:  # noqa: BLE001
            rules = []
        rules.sort(key=lambda t: t[0])
        obs.append(Observation("sheet", "conditional_rules", [
            {"priority": p, "type": getattr(r, "type", ""), "ranges": [str(x) for x in rg], "formula": list(getattr(r, "formula", None) or [])[:2],
             "stop": bool(getattr(r, "stopIfTrue", False)),
             "fill": (_xlsx_colour(getattr(getattr(getattr(r, "dxf", None), "fill", None), "bgColor", None)) if getattr(r, "dxf", None) is not None else None)}
            for p, rg, r in rules[:60]], fmt, dict(loc_sheet)))

        def cell_style(cell) -> Dict[str, Any]:
            props: Dict[str, Any] = {}
            font = cell.font
            if font is not None:
                props["bold"] = bool(font.b)
                props["italic"] = bool(font.i)
                props["underline"] = bool(font.u and font.u != "none")
                if font.name:
                    props["font_family"] = font.name
                if font.sz:
                    props["size_pt"] = float(font.sz)
                colour = _xlsx_colour(font.color)
                if colour:
                    props["color"] = colour
            static = None
            if cell.fill is not None and cell.fill.fill_type == "solid":
                static = _xlsx_colour(cell.fill.fgColor)
            cf_fill, verifiable = _cf_fill_for(cell, ws, rules) if rules else (None, True)
            if not verifiable:
                props["background"] = None
            else:
                props["background"] = cf_fill or static or "#FFFFFF"
            return props

        def emit(target: str, cell, extra: Optional[Dict[str, Any]] = None) -> None:
            props = cell_style(cell)
            loc = {**loc_sheet, "cell": cell.coordinate, **(extra or {})}
            for key, value in props.items():
                obs.append(Observation(target, key, value, fmt, dict(loc), verifiable=value is not None))

        for c_i, name in enumerate(headers):
            emit("table_header", ws.cell(row=1, column=c_i + 1), {"text": name})
        data_rows = 0
        total_rows: List[int] = []
        emitted = 0
        width = max(1, len(headers))
        for r in range(2, ws.max_row + 1 if not auxiliary else 2):
            values = [ws.cell(row=r, column=c + 1).value for c in range(width)]
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            first = str(values[0] or "").strip().lower()
            has_formula = any(isinstance(v, str) and v.startswith("=") and ws.cell(row=r, column=i + 1).data_type == "f" for i, v in enumerate(values))
            if first.startswith("total") and has_formula:
                total_rows.append(r)
                for c in range(width):
                    if values[c] not in (None, ""):
                        emit("table_total", ws.cell(row=r, column=c + 1))
                continue
            data_rows += 1
            if emitted < _MAX_PER_TARGET:
                for c in range(width):
                    emit("column:" + norm_text(headers[c]) if c < len(headers) else "cell", ws.cell(row=r, column=c + 1), {"row": r})
                    emitted += 1
        if not auxiliary:
            obs.append(Observation("sheet", "row_count", data_rows, fmt, dict(loc_sheet)))
            obs.append(Observation("sheet", "total_rows", total_rows, fmt, dict(loc_sheet)))
        # Row and cell-range fills are read on demand by the evaluator from
        # these per-row observations (bounded to the used range).
        max_r = min(ws.max_row, 500)
        max_c = min(ws.max_column, 60)
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                cell = ws.cell(row=r, column=c)
                if not auxiliary:  # no table rows or ranges there; its formula text is still scanned below
                    props = cell_style(cell)
                    obs.append(Observation("cell:" + cell.coordinate, "background", props.get("background"), fmt, {**loc_sheet, "row": r}, verifiable=props.get("background") is not None))
                if c <= width and not auxiliary:
                    # "row 5 yellow" is the table's row, not the helper
                    # columns a native chart reads beside it.
                    obs.append(Observation("row:" + str(r), "background", props.get("background"), fmt, {**loc_sheet, "cell": cell.coordinate}, verifiable=props.get("background") is not None))
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    if not _RENDERER_FORMULA_RE.match(str(cell.value)):
                        unexpected_formulas.append(f"{ws.title}!{cell.coordinate}")
                elif isinstance(cell.value, str) and cell.value.startswith(_FORMULA_LEADS) and not _PLAIN_NUMBER_RE.match(cell.value) and not cell.quotePrefix:
                    formula_flags.append(f"{ws.title}!{cell.coordinate}")
        # Number formats per header column.
        for c_i, name in enumerate(headers):
            fmts = {ws.cell(row=r, column=c_i + 1).number_format for r in range(2, min(ws.max_row, 50) + 1)}
            obs.append(Observation("column:" + norm_text(name), "number_formats", sorted(fmts), fmt, dict(loc_sheet)))
        for ch_i, chart in enumerate(getattr(ws, "_charts", []) or []):
            loc = {**loc_sheet, "chart": ch_i}
            obs.extend(_chart_obs(chart, fmt, loc))
            series_values: List[List[Any]] = []
            categories: Optional[List[Any]] = None
            for ser in list(getattr(chart, "series", []) or []):
                try:
                    f = ser.val.numRef.f
                except Exception:  # noqa: BLE001
                    f = None
                series_values.append(_ref_values(wb, f) if f else [])
                if categories is None:
                    try:
                        cf = (ser.cat.numRef or ser.cat.strRef).f
                        categories = _ref_values(wb, cf)
                    except Exception:  # noqa: BLE001
                        categories = None
            obs.append(Observation("chart", "values", {"categories": categories or [], "series": series_values}, fmt, loc))
    obs.append(Observation("security", "unprefixed_formula_text", formula_flags, fmt))
    obs.append(Observation("security", "unexpected_formulas", unexpected_formulas, fmt))
    obs.append(Observation("document", "headers", all_headers, fmt))
    return obs


# ---------------------------------------------------------------- CSV --


def inspect_csv(path: Path) -> List[Observation]:
    fmt = "csv"
    raw = Path(path).read_bytes()
    text = raw.decode("utf-8-sig", errors="replace")
    rows = list(_csv.reader(io.StringIO(text)))
    header = rows[0] if rows else []
    data = [r for r in rows[1:] if any(str(v).strip() for v in r)]
    flags = []
    for r_i, row in enumerate(data, start=2):
        for c_i, value in enumerate(row):
            if value.startswith(_FORMULA_LEADS) and not _PLAIN_NUMBER_RE.match(value):
                flags.append(f"R{r_i}C{c_i + 1}")
    return [
        Observation("sheet", "columns", header, fmt, {"file": Path(path).name}),
        Observation("sheet", "row_count", len(data), fmt, {"file": Path(path).name}),
        Observation("security", "unprefixed_formula_text", flags, fmt, {"file": Path(path).name}),
        Observation("document", "headers", header, fmt),
        Observation("document", "table_cells", [v for r in data[:5000] for v in r], fmt),
    ]


# --------------------------------------------------------------- PPTX --


def inspect_pptx(path: Path, spec=None) -> List[Observation]:
    from pptx import Presentation
    from pptx.util import Emu

    fmt = "pptx"
    spec_titles = {norm_text(t) for t in _spec_texts(spec)["slide_titles"]}
    obs: List[Observation] = []
    prs = Presentation(str(path))
    w_mm = Emu(prs.slide_width).mm if prs.slide_width else 0
    h_mm = Emu(prs.slide_height).mm if prs.slide_height else 0
    obs.append(Observation("page", "orientation", "landscape" if w_mm > h_mm else "portrait", fmt))
    titles: List[str] = []
    text_lines: List[str] = []
    cells: List[str] = []

    def run_props(run, para) -> Dict[str, Any]:
        props: Dict[str, Any] = {}
        font = run.font
        if font.name:
            props["font_family"] = font.name
        if font.size is not None:
            props["size_pt"] = float(font.size.pt)
        if font.bold is not None:
            props["bold"] = bool(font.bold)
        if font.italic is not None:
            props["italic"] = bool(font.italic)
        if font.underline is not None:
            props["underline"] = bool(font.underline)
        try:
            if font.color is not None and font.color.type is not None and font.color.rgb is not None:
                props["color"] = norm_hex(str(font.color.rgb))
        except (AttributeError, TypeError):
            pass
        return props

    for s_i, slide in enumerate(prs.slides):
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:  # noqa: BLE001
            title_shape = None
        for shape in slide.shapes:
            is_title = shape is title_shape or (title_shape is None and "title" in (shape.name or "").lower())
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if not text:
                        continue
                    text_lines.append(text)
                    # The renderer draws titles as text boxes on a custom
                    # layout (no title placeholder): a paragraph whose text
                    # is the slide's title in the spec IS the slide title.
                    sizes_here = [r.font.size.pt for r in para.runs if r.font.size is not None]
                    title_here = is_title or (norm_text(text) in spec_titles and (not sizes_here or max(sizes_here) >= 16))
                    target = "slide_title" if title_here else "paragraph"
                    if title_here:
                        titles.append(text)
                    for run in para.runs:
                        if not run.text.strip():
                            continue
                        for key, value in run_props(run, para).items():
                            obs.append(Observation(target, key, value, fmt, {"slide": s_i, "text": text[:80]}))
            if getattr(shape, "has_table", False) and shape.has_table:
                table = shape.table
                for r_i, row in enumerate(table.rows):
                    for c_i, cell in enumerate(row.cells):
                        cells.append(cell.text)
                        if r_i == 0:
                            try:
                                fill = norm_hex(str(cell.fill.fore_color.rgb)) if cell.fill.type == 1 else None
                            except (AttributeError, TypeError):
                                fill = None
                            obs.append(Observation("table_header", "background", fill, fmt, {"slide": s_i, "col": c_i}, verifiable=fill is not None))
                            for para in cell.text_frame.paragraphs:
                                for run in para.runs:
                                    for key, value in run_props(run, para).items():
                                        obs.append(Observation("table_header", key, value, fmt, {"slide": s_i, "col": c_i}))
            if getattr(shape, "has_chart", False) and shape.has_chart:
                chart = shape.chart
                loc = {"slide": s_i}
                ct = str(chart.chart_type or "").lower()
                kind = "unknown"
                for key, name in (("stacked", "stacked_bar"), ("bar", "bar"), ("column", "bar"), ("line", "line"), ("pie", "pie"),
                                  ("doughnut", "donut"), ("xy_scatter", "scatter"), ("area", "area"), ("radar", "radar"), ("bubble", "bubble")):
                    if key in ct:
                        kind = name
                        break
                obs.append(Observation("chart", "type", kind, fmt, loc))
                title = ""
                try:
                    if chart.has_title:
                        title = chart.chart_title.text_frame.text
                except Exception:  # noqa: BLE001
                    title = ""
                obs.append(Observation("chart", "title", title, fmt, loc))
                colours: List[Optional[str]] = []
                series: List[List[Any]] = []
                labels = False
                for plot in chart.plots:
                    try:
                        labels = labels or bool(plot.has_data_labels)
                    except Exception:  # noqa: BLE001
                        pass
                    for ser in plot.series:
                        series.append(list(ser.values))
                        try:
                            colours.append(norm_hex(str(ser.format.fill.fore_color.rgb)) if ser.format.fill.type == 1 else None)
                        except (AttributeError, TypeError):
                            colours.append(None)
                cats: List[Any] = []
                try:
                    cats = list(chart.plots[0].categories)
                except Exception:  # noqa: BLE001
                    cats = []
                obs.append(Observation("chart", "series_colors", colours, fmt, loc, verifiable=any(colours)))
                obs.append(Observation("chart", "data_labels", labels, fmt, loc))
                legend = "none"
                try:
                    if chart.has_legend:
                        legend = {1: "bottom", 2: "corner", 3: "top", 4: "right", 5: "left", -4107: "bottom", -4160: "top", -4152: "right", -4131: "left"}.get(int(chart.legend.position), "right")
                except Exception:  # noqa: BLE001
                    legend = "right"
                obs.append(Observation("chart", "legend_position", legend, fmt, loc))
                obs.append(Observation("chart", "values", {"categories": cats, "series": series}, fmt, loc))
    obs.append(Observation("document", "headings", titles, fmt))
    obs.append(Observation("document", "table_cells", cells, fmt))
    obs.append(Observation("document", "text", "\n".join(text_lines + cells), fmt))
    obs.append(Observation("page", "page_numbers", None, fmt, verifiable=False))
    return obs


# ---------------------------------------------------------------- PDF --


@dataclass
class PdfTextRun:
    page: int
    text: str
    font: str
    size_pt: float
    color: Optional[str]
    bold: bool
    italic: bool
    bbox: Tuple[float, float, float, float]
    background: Optional[str] = None
    underline: bool = False


def _pdf_font_family(base: str) -> str:
    name = base.split("+", 1)[1] if "+" in base[:8] else base
    name = re.sub(r"[-,](Bold|Italic|Oblique|Regular|BoldItalic|BoldOblique|Medium|Semibold|Light)+$", "", name, flags=re.I)
    name = re.sub(r"(MT|PS)$", "", name)
    name = name.replace("-", " ")
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name).strip()


def pdf_runs(path: Path, *, max_pages: int = 60) -> Tuple[List[PdfTextRun], List[Tuple[float, float]], List[dict]]:
    """Every text object of the first `max_pages` pages, with the filled
    paths they sit on. Returns (runs, page sizes in pt, paths)."""
    import ctypes

    import pypdfium2 as pdfium
    import pypdfium2.raw as R

    runs: List[PdfTextRun] = []
    sizes: List[Tuple[float, float]] = []
    all_paths: List[dict] = []
    pdf = pdfium.PdfDocument(str(path))
    try:
        for p_i in range(min(len(pdf), max_pages)):
            page = pdf[p_i]
            left, bottom, right, top = page.get_mediabox()
            sizes.append((right - left, top - bottom))
            textpage = page.get_textpage()
            paths: List[dict] = []
            strokes: List[dict] = []
            texts: List[PdfTextRun] = []
            count = R.FPDFPage_CountObjects(page.raw)
            for i in range(count):
                obj = R.FPDFPage_GetObject(page.raw, i)
                kind = R.FPDFPageObj_GetType(obj)
                r, g, b, a = (ctypes.c_uint() for _ in range(4))
                R.FPDFPageObj_GetFillColor(obj, r, g, b, a)
                L, B, Rr, T = (ctypes.c_float() for _ in range(4))
                R.FPDFPageObj_GetBounds(obj, L, B, Rr, T)
                bbox = (L.value, B.value, Rr.value, T.value)
                if kind == R.FPDF_PAGEOBJ_PATH:
                    fillmode = ctypes.c_int()
                    stroke = ctypes.c_int()
                    R.FPDFPath_GetDrawMode(obj, fillmode, stroke)
                    # Rectangles only: a frame drawn as an even-odd ring
                    # (10 segments) contains its text but is not behind it.
                    if fillmode.value and a.value > 0 and R.FPDFPath_CountSegments(obj) <= 6:
                        paths.append({"bbox": bbox, "color": norm_hex((r.value, g.value, b.value)), "area": max(0.0, (bbox[2] - bbox[0]) * (bbox[3] - bbox[1]))})
                    elif not fillmode.value and stroke.value and R.FPDFPath_CountSegments(obj) <= 2 and (bbox[3] - bbox[1]) <= 6.5:
                        # A stroked horizontal rule: WeasyPrint draws
                        # text-decoration underline this way, and its
                        # bounds (stroke width included) sit INSIDE the
                        # glyph box, above the descenders.
                        strokes.append({"bbox": bbox})
                    continue
                if kind != R.FPDF_PAGEOBJ_TEXT:
                    continue
                n = R.FPDFTextObj_GetText(obj, textpage.raw, None, 0)
                if n <= 0:
                    continue
                buf = ctypes.create_string_buffer(n * 2)
                R.FPDFTextObj_GetText(obj, textpage.raw, ctypes.cast(buf, ctypes.POINTER(ctypes.c_ushort)), n)
                text = buf.raw.decode("utf-16-le", "ignore").rstrip("\x00")
                if not text.strip():
                    continue
                font = R.FPDFTextObj_GetFont(obj)
                fl = R.FPDFFont_GetBaseFontName(font, None, 0) if font else 0
                fbuf = ctypes.create_string_buffer(max(1, fl))
                if font and fl:
                    R.FPDFFont_GetBaseFontName(font, fbuf, fl)
                base = fbuf.value.decode("latin-1", "ignore")
                size = ctypes.c_float()
                R.FPDFTextObj_GetFontSize(obj, size)
                matrix = R.FS_MATRIX()
                scale = 1.0
                if R.FPDFPageObj_GetMatrix(obj, ctypes.byref(matrix)):
                    scale = math.sqrt(abs(matrix.a * matrix.d - matrix.b * matrix.c)) or 1.0
                weight = R.FPDFFont_GetWeight(font) if font else -1
                flags = R.FPDFFont_GetFlags(font) if font else 0
                angle = ctypes.c_int()
                has_angle = bool(font and R.FPDFFont_GetItalicAngle(font, ctypes.byref(angle)))
                bold = bool(re.search(r"bold|black|heavy|semibold", base, re.I)) or weight >= 600
                italic = bool(re.search(r"italic|oblique", base, re.I)) or bool(flags & 64) or (has_angle and angle.value != 0)
                texts.append(PdfTextRun(p_i, text, base, round(size.value * scale, 2), norm_hex((r.value, g.value, b.value)), bold, italic, bbox))
            for run in texts:
                cx = (run.bbox[0] + run.bbox[2]) / 2
                # Behind = a filled rectangle that spans the run's whole
                # height, not a rule line under a heading.
                behind = [pth for pth in paths if pth["bbox"][0] - 0.5 <= cx <= pth["bbox"][2] + 0.5
                          and pth["bbox"][1] <= run.bbox[1] + 1.0 and pth["bbox"][3] >= run.bbox[3] - 1.0 and pth["area"] > 4]
                if behind:
                    run.background = min(behind, key=lambda pth: pth["area"])["color"]
                width = run.bbox[2] - run.bbox[0]
                for pth in paths:
                    pb = pth["bbox"]
                    if (pb[3] - pb[1]) <= 2.5 and (pb[2] - pb[0]) >= 0.8 * width and pb[0] <= run.bbox[0] + 2 and pb[3] <= run.bbox[1] + 3 and pb[3] >= run.bbox[1] - 4:
                        run.underline = True
                        break
                if not run.underline:
                    # Below the baseline band (never a strike-through) and
                    # no wider than the text (never a table rule or border).
                    lower = run.bbox[1] + 0.3 * (run.bbox[3] - run.bbox[1])
                    for st in strokes:
                        sb = st["bbox"]
                        mid = (sb[1] + sb[3]) / 2
                        if 0.8 * width <= (sb[2] - sb[0]) <= width + 12 and sb[0] <= run.bbox[0] + 4 and sb[2] >= run.bbox[2] - 4 and run.bbox[1] - 2 <= mid <= lower:
                            run.underline = True
                            break
            runs.extend(texts)
            all_paths.extend({**pth, "page": p_i} for pth in paths)
    finally:
        pdf.close()
    return runs, sizes, all_paths


def _spec_texts(spec) -> Dict[str, Any]:
    """The texts a PDF's runs are classified against: the title, subtitle,
    headings with levels, table header cells, slide titles."""
    out: Dict[str, Any] = {"title": "", "subtitle": "", "headings": [], "table_headers": set(), "slide_titles": [], "totals": True}
    if spec is None:
        return out
    try:
        body = spec.body
    except Exception:  # noqa: BLE001
        return out
    out["title"] = str(getattr(body, "title", "") or "")
    out["subtitle"] = str(getattr(body, "subtitle", "") or getattr(body, "purpose", "") or "")
    out["list_items"] = []
    out["prose"] = []
    for block in list(getattr(body, "blocks", None) or []):
        if getattr(block, "type", "") in ("bullets", "numbered"):
            out["list_items"].extend(norm_text(t) for t in block.items)
        elif getattr(block, "type", "") == "paragraph":
            out["prose"].append(norm_text(block.text))
        if getattr(block, "type", "") == "heading":
            out["headings"].append((int(block.level), str(block.text)))
        elif getattr(block, "type", "") == "table":
            out["table_headers"].update(norm_text(c) for c in block.table.columns)
    for slide in list(getattr(body, "slides", None) or []):
        if slide.title:
            out["slide_titles"].append(str(slide.title))
        if slide.table is not None:
            out["table_headers"].update(norm_text(c) for c in slide.table.columns)
    for sheet in list(getattr(body, "sheets", None) or []):
        out["table_headers"].update(norm_text(c.name) for c in sheet.columns)
    return out


def _pdf_list_run(text: str, texts: Dict[str, Any]) -> bool:
    """A body run that is (a wrapped line of) a bullet or numbered item and
    of no paragraph: its text is inside a list item of the spec only."""
    t = norm_text(re.sub(r"^\s*(?:[•◦▪‣\-–]|\d{1,3}[.)])\s*", "", text))
    if len(t) < 3 or not texts.get("list_items"):
        return False
    return any(t in item for item in texts["list_items"]) and not any(t in prose for prose in texts.get("prose") or [])


def _matches(run_text: str, target: str) -> bool:
    a, b = norm_text(run_text), norm_text(target)
    return bool(a) and bool(b) and (a == b or (len(a) >= 4 and b.startswith(a)) or (len(b) >= 4 and a.startswith(b)))


def inspect_pdf(path: Path, spec=None) -> List[Observation]:
    fmt = "pdf"
    obs: List[Observation] = []
    runs, sizes, _paths = pdf_runs(path)
    obs.append(Observation("document", "page_count", len(sizes), fmt))
    for p_i, (w, h) in enumerate(sizes):
        orientation, size = _page_class(w / 72 * 25.4, h / 72 * 25.4)
        obs.append(Observation("page", "orientation", orientation, fmt, {"page": p_i}))
        obs.append(Observation("page", "page_size", size, fmt, {"page": p_i}))
    texts = _spec_texts(spec)
    heading_left = list(texts["headings"])
    pdf_headings: List[str] = []
    title_seen = subtitle_seen = False
    counts: Dict[str, int] = {}

    # The heading path of the run being read (see inspect_docx): only when
    # the spec's headings are known, so a PDF read without its spec never
    # claims a section.
    path: List[str] = ["", "", ""]
    h1_ordinal = 0

    def emit(target: str, run: PdfTextRun, locator: Dict[str, Any]) -> None:
        counts[target] = counts.get(target, 0) + 1
        if counts[target] > _MAX_PER_TARGET:
            return
        loc = {"page": run.page, "text": run.text[:80], **locator}
        if texts["headings"]:
            loc.update(sections=[s for s in path if s], section_index=h1_ordinal)
        obs.append(Observation(target, "color", run.color, fmt, dict(loc)))
        obs.append(Observation(target, "font_family", _pdf_font_family(run.font), fmt, dict(loc)))
        obs.append(Observation(target, "size_pt", run.size_pt, fmt, dict(loc)))
        obs.append(Observation(target, "bold", run.bold, fmt, dict(loc)))
        obs.append(Observation(target, "italic", run.italic, fmt, dict(loc)))
        obs.append(Observation(target, "underline", run.underline, fmt, dict(loc)))
        obs.append(Observation(target, "background", run.background or "#FFFFFF", fmt, dict(loc)))

    page_numbers_by_page: Dict[int, bool] = {}
    slide_titles = [norm_text(t) for t in texts["slide_titles"]]
    body_sizes: List[float] = []
    for run in runs:
        w, h = sizes[run.page]
        # The header/footer band is measured on the LONG side: a landscape
        # A4 page is 595 pt tall, and 7.5 % of that (45 pt) put the footer
        # "Page 1 of 2" at 46–55 pt outside the band — every landscape PDF
        # read as having no page numbers (live 2026-09-15, "make it
        # landscape" edits v4–v7).
        band = 0.075 * max(w, h)
        in_margin = run.bbox[3] < band or run.bbox[1] > h - band
        stripped = run.text.strip()
        number = re.search(r"\bpage\s+(\d+)", stripped, re.I) or re.fullmatch(r"(\d{1,4})(\s*(/|of)\s*\d{1,4})?", stripped)
        # The number must be THIS page's (a cover may go unnumbered): the
        # wider band reaches the last body line on a landscape page, where a
        # table cell "12" is not a page number (verifier 2026-09-15).
        if in_margin and number and int(number.group(1)) in (run.page + 1, run.page):
            page_numbers_by_page[run.page] = True
            continue
        if in_margin:
            continue
        if not title_seen and texts["title"] and _matches(stripped, texts["title"]) and run.size_pt >= 14:
            title_seen = True
            emit("title", run, {})
            continue
        if not subtitle_seen and texts["subtitle"] and _matches(stripped, texts["subtitle"]) and not _matches(stripped, texts["title"]):
            subtitle_seen = True
            emit("subtitle", run, {})
            continue
        hit = next(((lvl, t) for lvl, t in heading_left if _matches(stripped, t)), None)
        if hit is not None and run.size_pt >= 10:
            heading_left.remove(hit)
            level = min(max(int(hit[0]), 1), 3)
            if level == 1:
                h1_ordinal += 1
            path[level - 1] = hit[1]
            for deeper in range(level, 3):
                path[deeper] = ""
            emit(f"heading{min(hit[0], 3)}", run, {"level": hit[0]})
            # The heading as the page shows it (a spec heading that never
            # reached the page is never listed).
            pdf_headings.append(stripped if len(stripped) >= len(hit[1]) else hit[1])
            continue
        if slide_titles and norm_text(stripped) in slide_titles:
            emit("slide_title", run, {})
            continue
        if norm_text(stripped) in texts["table_headers"] and run.background and run.background != "#FFFFFF":
            emit("table_header", run, {})
            continue
        if run.page < 2 and run.size_pt >= 22 and not texts["title"]:
            continue
        emit("paragraph", run, {"chars": len(stripped), "list": _pdf_list_run(stripped, texts)})
        body_sizes.append(run.size_pt)
    obs.append(Observation("page", "page_numbers", bool(sizes) and len(page_numbers_by_page) >= max(1, len(sizes) - 1), fmt,
                           {"pages_with_numbers": len(page_numbers_by_page), "pages": len(sizes)}))
    full = "\n".join(r.text for r in runs)
    obs.append(Observation("document", "text", full, fmt))
    obs.append(Observation("document", "headings", pdf_headings, fmt))
    # Tofu: Indic text set in a font that covers none of the scripts.
    uncovered = []
    for script, frags in _SCRIPT_FONTS.items():
        for run in runs:
            if script_of(run.text)[script] and not any(f in run.font.lower().replace("-", " ").replace("_", " ") or f.replace(" ", "") in run.font.lower() for f in frags):
                uncovered.append(f"{script}:{_pdf_font_family(run.font)}")
    has_indic = any(script_of(r.text)["devanagari"] or script_of(r.text)["gujarati"] for r in runs)
    obs.append(Observation("document", "fonts_cover_script", not uncovered, fmt, {"uncovered": sorted(set(uncovered))[:5], "indic": has_indic}))
    obs.append(Observation("document", "fonts", sorted({_pdf_font_family(r.font) for r in runs})[:20], fmt))
    return obs


# ---------------------------------------------------------- PNG / SVG --


def _png_observations(fileobj, name: str) -> List[Observation]:
    from PIL import Image

    fmt = "png"
    with Image.open(fileobj) as im:
        w, h = im.size
        # NEAREST keeps real pixel colours; a smoothing resize would invent
        # blends of a bar and its background.
        small = im.convert("RGB").resize((max(1, w // 4), max(1, h // 4)), Image.NEAREST)
        colours = small.getcolors(maxcolors=1 << 20) or []
    dominant: List[str] = []
    for count, (r, g, b) in sorted(colours, reverse=True):
        if max(r, g, b) - min(r, g, b) < 18:  # white, black, greys: background, text, gridlines
            continue
        hx = norm_hex((r, g, b))
        if hx and hx not in dominant:
            dominant.append(hx)
        if len(dominant) >= 8:
            break
    return [
        Observation("image", "size", [w, h], fmt, {"file": name}),
        Observation("chart", "dominant_colors", dominant, fmt, {"file": name}),
    ]


def inspect_png(path: Path) -> List[Observation]:
    with open(path, "rb") as fh:
        return _png_observations(io.BytesIO(fh.read()), Path(path).name)


_SVG_REFUSE = (
    (re.compile(rb"<\s*script", re.I), "script element"),
    (re.compile(rb"<\s*foreignObject", re.I), "foreignObject"),
    (re.compile(rb"\son[a-z]+\s*=", re.I), "event handler attribute"),
    (re.compile(rb"<!DOCTYPE|<!ENTITY", re.I), "DOCTYPE/ENTITY"),
    (re.compile(rb"(?:xlink:)?href\s*=\s*[\"'](?!#)", re.I), "non-fragment href"),
    (re.compile(rb"url\(\s*['\"]?(?!#)", re.I), "url() to a non-fragment"),
    (re.compile(rb"@import", re.I), "@import"),
)


def svg_problems(data: bytes) -> List[str]:
    try:
        from .render.validate import validate_svg_bytes  # charts track, when merged

        return list(validate_svg_bytes(data))
    except ImportError:
        return [label for rx, label in _SVG_REFUSE if rx.search(data)]


def inspect_svg(path: Path) -> List[Observation]:
    fmt = "svg"
    data = Path(path).read_bytes()
    colours = []
    for m in re.finditer(rb"(?:fill|stroke)\s*[:=]\s*[\"']?\s*(#[0-9a-fA-F]{3,6})", data):
        hx = norm_hex(m.group(1).decode())
        if hx and hx not in colours:
            colours.append(hx)
    return [
        Observation("security", "svg_problems", svg_problems(data), fmt, {"file": Path(path).name}),
        Observation("chart", "dominant_colors", colours[:12], fmt, {"file": Path(path).name}),
    ]


# ------------------------------------------------------------ dispatch --


def format_of(path: Path) -> str:
    return Path(path).suffix.lower().lstrip(".")


def inspect(paths: Dict[str, Path], spec=None) -> List[Observation]:
    """Observations over every file in `paths` ({name: path}). A file that
    cannot be read yields one `file/readable=False` observation for its
    format — the evaluator fails a format item on it, and the self-check
    never raises out of here."""
    out: List[Observation] = []
    for name, p in paths.items():
        p = Path(p)
        fmt = format_of(p)
        try:
            if fmt == "docx":
                out.extend(inspect_docx(p))
            elif fmt == "xlsx":
                out.extend(inspect_xlsx(p, spec))
            elif fmt == "pptx":
                out.extend(inspect_pptx(p, spec))
            elif fmt == "pdf":
                out.extend(inspect_pdf(p, spec))
            elif fmt == "csv":
                out.extend(inspect_csv(p))
            elif fmt == "png":
                out.extend(inspect_png(p))
            elif fmt == "svg":
                out.extend(inspect_svg(p))
            else:
                continue
            out.append(Observation("file", "format", fmt, fmt, {"name": p.name}))
        except Exception as exc:  # noqa: BLE001 — a file the checker cannot read is an observation, not a crash
            log.warning("inspect_files: %s could not be read: %s", p.name, type(exc).__name__)
            out.append(Observation("file", "readable", False, fmt, {"name": p.name, "error": type(exc).__name__}))
    return out


def extract_text(path: Path) -> str:
    """The plain text a reader sees in one file (for faithfulness and
    preservation): paragraphs and table rows, in order."""
    for o in inspect({Path(path).name: Path(path)}):
        if o.target == "document" and o.property == "text":
            return str(o.value or "")
    return ""


__all__ = [
    "Observation", "inspect", "extract_text", "inspect_docx", "inspect_xlsx", "inspect_pptx", "inspect_pdf",
    "inspect_csv", "inspect_png", "inspect_svg", "pdf_runs", "norm_hex", "norm_text", "contrast_ratio", "script_of",
    "svg_problems", "format_of",
]
