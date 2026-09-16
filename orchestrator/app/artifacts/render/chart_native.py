"""Native charts in XLSX (openpyxl) and PPTX (python-pptx), from computed values.

A person can click these charts and edit them in Excel or PowerPoint. The
numbers are the ones chart_data computed from the bound table:

XLSX. Every chart's values are written by code to a visible "Chart data"
sheet (kept last, tab colour #5F6B7A) and the chart references those cells —
an aggregated chart never points at 150 raw rows, and a pie never gets 150
slices. Text cells go through the formula-neutralising writer (a category
named "=HYPERLINK(...)" is text with quotePrefix, never a formula).

PPTX. python-pptx embeds the values in the chart part (c:numCache) and in an
embedded workbook; labels have formula leads stripped first (the embedded
workbook is written by XlsxWriter, which turns '=' strings into formulas).
Combo charts and trend lines are added to the chart XML directly, because
python-pptx has no API for a second plot or a trendline.

IMAGE FALLBACK. Types a container has no native form for (chart_spec.
CONTAINER_SUPPORT: box, heatmap, waterfall, funnel, gantt) are placed as a
PNG drawn by render/charts.py from the same values; the result dict says
mode='image' and carries a note, so the caller can count and report it.

Styling: series colours, per-point colours, axis number format '#,##0'
(sourceLinked=0, or the chart's number_format), fonts (txPr), legend
position, data labels, axis min/max/log, gridlines, marker + dash styles for
3+ line series. Libraries are imported lazily.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional, Sequence, Tuple

from .. import chart_spec as CS
from . import charts as R

CHART_DATA_SHEET = "Chart data"
CHART_DATA_TAB = "5F6B7A"
_FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")

EXCEL_NUMBER_FORMATS: Dict[str, str] = {
    "integer": "#,##0", "decimal1": "#,##0.0", "decimal2": "#,##0.00", "percent": "0.0%",
    "currency_INR": '"₹"#,##,##0', "currency_USD": '"$"#,##0', "currency_EUR": '"€"#,##0', "currency_GBP": '"£"#,##0',
    "compact": '#,##0,"K"',
}
_XL_MARKERS = ("circle", "square", "triangle", "diamond", "x", "star", "plus", "dash")
_XL_DASHES = ("solid", "dash", "dashDot", "sysDot", "lgDash", "lgDashDot", "sysDash", "lgDashDotDot")
_LEGEND_XL = {"bottom": "b", "top": "t", "left": "l", "right": "r"}


def _hex(colour: str) -> str:
    return colour.lstrip("#").upper()


def _number_format(chart: CS.Chart) -> str:
    st = chart.style
    if st and st.number_format:
        if st.number_format == "percent":
            vals = [abs(v) for s in chart.series for v in s.values]
            return "0.0%" if vals and max(vals) <= 1.0 else '0.0"%"'
        return EXCEL_NUMBER_FORMATS[st.number_format]
    vals = [v for s in chart.series for v in s.values]
    if vals and any(abs(v - round(v)) > 1e-9 for v in vals):
        return "#,##0.0"
    return "#,##0"


def support(fmt: str, chart_type: str) -> str:
    return CS.support_for(fmt, chart_type)


# ------------------------------------------------------------------- XLSX --


def _write_text(cell, text: Any) -> None:
    """A text cell Excel reads as text whatever it starts with."""
    value = "" if text is None else str(text)
    cell.value = value
    cell.data_type = "s"
    if value.startswith(_FORMULA_LEADS):
        cell.quotePrefix = True


def chart_data_sheet(wb) -> Any:
    """The workbook's "Chart data" sheet, created visible and kept last."""
    if CHART_DATA_SHEET in wb.sheetnames:
        ws = wb[CHART_DATA_SHEET]
    else:
        ws = wb.create_sheet(CHART_DATA_SHEET)
        ws.sheet_properties.tabColor = CHART_DATA_TAB
        ws.column_dimensions["A"].width = 28
    keep_chart_data_last(wb)
    return ws


def keep_chart_data_last(wb) -> None:
    if CHART_DATA_SHEET in wb.sheetnames:
        ws = wb[CHART_DATA_SHEET]
        idx = wb.sheetnames.index(CHART_DATA_SHEET)
        if idx != len(wb.sheetnames) - 1:
            wb.move_sheet(ws, offset=len(wb.sheetnames) - 1 - idx)


def _next_free_row(ws) -> int:
    if ws.max_row == 1 and ws["A1"].value is None:
        return 1
    return ws.max_row + 2


def _write_block(ws, chart: CS.Chart) -> Dict[str, Any]:
    """Title row, header row, then one row per category (or per point for
    XY charts). Returns the block's geometry."""
    from openpyxl.styles import Font, PatternFill

    top = _next_free_row(ws)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="5F6B7A")
    title = ws.cell(row=top, column=1)
    _write_text(title, chart.title or "Chart data")
    title.font = Font(bold=True, size=12, color="1F2937")
    hdr = top + 1
    first = hdr + 1
    geometry: Dict[str, Any] = {"sheet": ws.title, "top": top, "header_row": hdr, "first_row": first, "series_cols": [], "x_cols": [], "size_cols": []}
    if chart.type in CS.XY_TYPES:
        col = 1
        n = 0
        for s in chart.series:
            names = [f"{s.name} x", s.name] + ([f"{s.name} size"] if s.sizes is not None else [])
            for j, name in enumerate(names):
                c = ws.cell(row=hdr, column=col + j)
                _write_text(c, name)
                c.font, c.fill = head_font, head_fill
            for i, y in enumerate(s.values):
                ws.cell(row=first + i, column=col, value=float((s.x or [])[i]))
                ws.cell(row=first + i, column=col + 1, value=float(y))
                if s.sizes is not None:
                    ws.cell(row=first + i, column=col + 2, value=float(s.sizes[i]))
            geometry["x_cols"].append(col)
            geometry["series_cols"].append(col + 1)
            geometry["size_cols"].append(col + 2 if s.sizes is not None else None)
            geometry.setdefault("lengths", []).append(len(s.values))
            n = max(n, len(s.values))
            col += len(names)
        geometry["rows"] = n
        return geometry
    c = ws.cell(row=hdr, column=1)
    _write_text(c, chart.x_label or "Category")
    c.font, c.fill = head_font, head_fill
    for k, s in enumerate(chart.series):
        c = ws.cell(row=hdr, column=2 + k)
        _write_text(c, s.name)
        c.font, c.fill = head_font, head_fill
        geometry["series_cols"].append(2 + k)
    nf = _number_format(chart)
    for i, cat in enumerate(chart.categories):
        _write_text(ws.cell(row=first + i, column=1), cat)
        for k, s in enumerate(chart.series):
            cell = ws.cell(row=first + i, column=2 + k, value=float(s.values[i]))
            cell.number_format = nf
    geometry["rows"] = len(chart.categories)
    return geometry


def _rich(size_pt: float, colour: str, bold: bool = False, font: Optional[str] = None, italic: bool = False):
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties

    cp = CharacterProperties(sz=int(round(size_pt * 100)), b=bold, i=italic, solidFill=_hex(colour))
    if font:
        cp.latin = DFont(typeface=font)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _xl_title(text: str, size_pt: float, colour: str, bold: bool, font: Optional[str]):
    from openpyxl.chart.text import RichText, Text
    from openpyxl.chart.title import Title
    from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun

    cp = CharacterProperties(sz=int(round(size_pt * 100)), b=bold, solidFill=_hex(colour))
    if font:
        cp.latin = DFont(typeface=font)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def _style_axes_xl(c, chart: CS.Chart, d: R.ChartStyleDefaults, *, value_axis=None, category_axis=None) -> None:
    from openpyxl.chart.axis import ChartLines
    from openpyxl.chart.data_source import NumFmt
    from openpyxl.chart.shapes import GraphicalProperties

    st = chart.style or CS.ChartStyle()
    font = st.font_family or d.font_family
    axis_size = (st.axis.size_pt if st.axis and st.axis.size_pt else d.axis_size_pt)
    axis_colour = (st.axis.color if st.axis and st.axis.color else d.axis_text_color)
    va = value_axis if value_axis is not None else c.y_axis
    ca = category_axis if category_axis is not None else c.x_axis
    for ax in (va, ca):
        ax.delete = False
        ax.txPr = _rich(axis_size, axis_colour, font=font)
    va.numFmt = NumFmt(formatCode="0%" if chart.type == "percent_stacked_bar" else _number_format(chart), sourceLinked=False)
    if st.gridlines:
        gp = GraphicalProperties()
        gp.line.solidFill = _hex(d.grid_color)
        va.majorGridlines = ChartLines(spPr=gp)
    else:
        va.majorGridlines = None
    if st.y_min is not None:
        va.scaling.min = st.y_min
    elif chart.type not in ("line",) and chart.type in CS.BAR_FAMILY + ("area", "stacked_area", "histogram", "combo"):
        va.scaling.min = min(0.0, min((v for s in chart.series for v in s.values), default=0.0))
    if st.y_max is not None:
        va.scaling.max = st.y_max
    if st.log_y:
        va.scaling.logBase = 10
    if chart.y_label:
        va.title = _xl_title(chart.y_label, axis_size, axis_colour, False, font)
    # NO_AXIS_TYPES, not PART_OF_WHOLE_TYPES: the set of types drawn with no
    # x/y axis pair is the one that must not print a category-axis title.
    # They differ by "sunburst" only, which CONTAINER_SUPPORT never sends
    # here today, so this is a contract fix rather than a behaviour change.
    if chart.x_label and chart.type not in CS.NO_AXIS_TYPES:
        ca.title = _xl_title(chart.x_label, axis_size, axis_colour, False, font)


def add_xlsx_chart(ws, chart: Any, anchor: str, resolved: Any = None, helper_ws: Any = None) -> Dict[str, Any]:
    """Add `chart` to worksheet `ws` at `anchor` (e.g. "H2"). The values go to
    `helper_ws` (default: the workbook's "Chart data" sheet). Returns
    {mode: native|image, type, helper: geometry, note}."""
    from openpyxl.chart import AreaChart, BarChart, BubbleChart, DoughnutChart, LineChart, PieChart, RadarChart, Reference, ScatterChart
    from openpyxl.chart import Series as XLSeries
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.marker import DataPoint, Marker
    from openpyxl.chart.trendline import Trendline

    c_spec = R._as_chart(chart) if not isinstance(chart, CS.Chart) else chart
    if not c_spec.series:
        raise ValueError("the chart has no computed values; resolve it first")
    d = R.defaults_from(resolved)
    st = c_spec.style or CS.ChartStyle()
    helper = helper_ws if helper_ws is not None else chart_data_sheet(ws.parent)
    geo = _write_block(helper, c_spec)
    t = c_spec.type
    width_cm = (st.width_in * 2.54) if st.width_in else 16.0
    height_cm = (st.height_in * 2.54) if st.height_in else 8.5
    result: Dict[str, Any] = {"mode": "native", "type": t, "helper": geo, "note": ""}

    if support("xlsx", t) != "native":
        from openpyxl.drawing.image import Image as XLImage

        png = R.render_png(c_spec, resolved, int(width_cm / 2.54 * 150), int(height_cm / 2.54 * 150))
        img = XLImage(io.BytesIO(png))
        img.width, img.height = int(width_cm / 2.54 * 96), int(height_cm / 2.54 * 96)
        img.anchor = anchor
        ws.add_image(img)
        result.update(mode="image", note=f"{c_spec.title or 'A chart'} ({t.replace('_', ' ')}) is a picture in the Excel file: Excel has no native {t.replace('_', ' ')} chart here; its numbers are on the Chart data sheet.")
        keep_chart_data_last(ws.parent)
        return result

    n = geo["rows"]
    first, hdr = geo["first_row"], geo["header_row"]
    cats = Reference(helper, min_col=1, min_row=first, max_row=first + n - 1) if t not in CS.XY_TYPES else None

    def data_ref(col: int) -> Any:
        return Reference(helper, min_col=col, min_row=hdr, max_row=first + n - 1)

    secondary = None
    if t in CS.BAR_FAMILY or t == "histogram" or t == "combo":
        c = BarChart()
        c.type = "bar" if t in ("horizontal_bar", "stacked_horizontal_bar") else "col"
        c.grouping = {"stacked_bar": "stacked", "stacked_horizontal_bar": "stacked", "percent_stacked_bar": "percentStacked"}.get(t, "clustered")
        if c.grouping != "clustered":
            c.overlap = 100
        if t == "histogram":
            c.gapWidth = 0
        bar_cols = [col for k, col in enumerate(geo["series_cols"]) if t != "combo" or (c_spec.series[k].kind or "bar") == "bar" and c_spec.series[k].axis == "primary"]
        for col in bar_cols:
            c.add_data(data_ref(col), titles_from_data=True)
        c.set_categories(cats)
        if t == "combo":
            line_idx = [k for k, s in enumerate(c_spec.series) if geo["series_cols"][k] not in bar_cols]
            if line_idx:
                ln = LineChart()
                for k in line_idx:
                    ln.add_data(data_ref(geo["series_cols"][k]), titles_from_data=True)
                ln.set_categories(cats)
                if any(c_spec.series[k].axis == "secondary" for k in line_idx):
                    ln.y_axis.axId = 200
                    ln.y_axis.crosses = "max"
                    ln.y_axis.delete = False
                    ln.y_axis.numFmt = None
                    if c_spec.y2_label:
                        ln.y_axis.title = _xl_title(c_spec.y2_label, d.axis_size_pt, d.axis_text_color, False, st.font_family or d.font_family)
                    ln.y_axis.txPr = _rich(d.axis_size_pt, d.axis_text_color, font=st.font_family or d.font_family)
                    ln.y_axis.majorGridlines = None
                for j, s in enumerate(ln.series):
                    k = line_idx[j]
                    colour = _hex(_series_colour(c_spec, k, d))
                    s.graphicalProperties.line.solidFill = colour
                    s.graphicalProperties.line.width = 28575
                    s.marker = Marker(symbol=_XL_MARKERS[(j + 1) % len(_XL_MARKERS)], size=6)
                    s.marker.graphicalProperties = _gp_fill(colour)
                    s.smooth = False
                secondary = ln
    elif t in ("line",):
        c = LineChart()
        for col in geo["series_cols"]:
            c.add_data(data_ref(col), titles_from_data=True)
        c.set_categories(cats)
    elif t in ("area", "stacked_area"):
        c = AreaChart()
        c.grouping = "stacked" if t == "stacked_area" else "standard"
        for col in geo["series_cols"]:
            c.add_data(data_ref(col), titles_from_data=True)
        c.set_categories(cats)
    elif t in ("pie", "donut"):
        c = DoughnutChart(holeSize=50) if t == "donut" else PieChart()
        c.add_data(data_ref(geo["series_cols"][0]), titles_from_data=True)
        c.set_categories(cats)
    elif t == "radar":
        c = RadarChart()
        c.type = "marker"
        for col in geo["series_cols"]:
            c.add_data(data_ref(col), titles_from_data=True)
        c.set_categories(cats)
    elif t == "scatter":
        c = ScatterChart()
        c.style = 13
        for k, s in enumerate(c_spec.series):
            ln_n = geo["lengths"][k]
            xref = Reference(helper, min_col=geo["x_cols"][k], min_row=first, max_row=first + ln_n - 1)
            yref = Reference(helper, min_col=geo["series_cols"][k], min_row=hdr, max_row=first + ln_n - 1)
            xs = XLSeries(yref, xref, title_from_data=True)
            xs.marker = Marker(symbol=_XL_MARKERS[k % len(_XL_MARKERS)] if len(c_spec.series) >= 3 else "circle", size=5)
            colour = _hex(_series_colour(c_spec, k, d))
            xs.marker.graphicalProperties = _gp_fill(colour)
            xs.graphicalProperties.line.noFill = True
            if c_spec.extra and any(tr.series == s.name for tr in c_spec.extra.trendlines):
                xs.trendline = Trendline(trendlineType="linear", dispRSqr=True, dispEq=False)
            c.series.append(xs)
    elif t == "bubble":
        c = BubbleChart()
        for k, s in enumerate(c_spec.series):
            ln_n = geo["lengths"][k]
            xref = Reference(helper, min_col=geo["x_cols"][k], min_row=first, max_row=first + ln_n - 1)
            yref = Reference(helper, min_col=geo["series_cols"][k], min_row=first, max_row=first + ln_n - 1)
            zref = Reference(helper, min_col=geo["size_cols"][k], min_row=first, max_row=first + ln_n - 1)
            bs = XLSeries(values=yref, xvalues=xref, zvalues=zref, title=s.name)
            bs.graphicalProperties.solidFill = _hex(_series_colour(c_spec, k, d))
            c.series.append(bs)
    else:  # pragma: no cover - CONTAINER_SUPPORT keeps this unreachable
        raise ValueError(f"no native xlsx form for {t}")

    c.width, c.height = width_cm, height_cm
    title_ts = st.title
    if c_spec.title:
        c.title = _xl_title(c_spec.title, title_ts.size_pt if title_ts and title_ts.size_pt else d.title_size_pt,
                            title_ts.color if title_ts and title_ts.color else d.ink, True if not title_ts or title_ts.bold is None else title_ts.bold,
                            (title_ts.font_family if title_ts and title_ts.font_family else None) or st.font_family or d.font_family)

    # Series and point colours, line styles.
    if t in ("pie", "donut"):
        s0 = c.series[0]
        for i, cat in enumerate(c_spec.categories):
            pt = DataPoint(idx=i)
            pt.graphicalProperties = _gp_fill(_hex(st.category_colors.get(cat) or (st.palette or list(d.palette))[i % len(st.palette or d.palette)]))
            s0.dPt.append(pt)
    elif t not in CS.XY_TYPES:
        series_list = list(c.series)
        idx_map = list(range(len(series_list))) if t != "combo" else [k for k, s in enumerate(c_spec.series) if (s.kind or "bar") == "bar" and s.axis == "primary"]
        many = len(series_list) >= 3
        for j, xs in enumerate(series_list):
            k = idx_map[j] if j < len(idx_map) else j
            colour = _hex(_series_colour(c_spec, k, d))
            if t in ("line", "radar"):
                xs.graphicalProperties.line.solidFill = colour
                xs.graphicalProperties.line.width = 28575
                xs.marker = Marker(symbol=_XL_MARKERS[j % len(_XL_MARKERS)] if many else "circle", size=6)
                xs.marker.graphicalProperties = _gp_fill(colour)
                if many:
                    xs.graphicalProperties.line.dashStyle = _XL_DASHES[j % len(_XL_DASHES)]
                xs.smooth = False
            else:
                xs.graphicalProperties.solidFill = colour
                xs.graphicalProperties.line.solidFill = colour
                if len(series_list) == 1 and st.category_colors:
                    for i, cat in enumerate(c_spec.categories):
                        if cat in st.category_colors:
                            pt = DataPoint(idx=i)
                            pt.graphicalProperties = _gp_fill(_hex(st.category_colors[cat]))
                            xs.dPt.append(pt)

    # Data labels: outside in ink; stacked segments inside only where
    # white/ink reaches 4.5:1 on the series colour.
    labels_on = _labels_on(c_spec)
    if labels_on and t in CS.STACKED_TYPES:
        for k, xs in enumerate(c.series):
            fill = _series_colour(c_spec, k, d)
            colour = d.label_color_for(fill)
            if R.contrast_ratio(colour, fill) < 4.5:
                continue
            dl = DataLabelList(showVal=True)
            dl.position = "ctr"
            dl.showSerName = dl.showCatName = dl.showLegendKey = dl.showPercent = False
            dl.txPr = _rich(max(7.0, d.axis_size_pt - 2), colour, font=st.font_family or d.font_family)
            xs.dLbls = dl
    elif labels_on:
        dl = DataLabelList()
        if t in ("pie", "donut"):
            dl.showPercent = True
            dl.showVal = False
            dl.position = "outEnd" if t == "pie" else None
        else:
            dl.showVal = True
            if t in ("bar", "horizontal_bar", "histogram", "combo"):
                dl.position = "outEnd"
        dl.showSerName = dl.showCatName = dl.showLegendKey = False
        dl.txPr = _rich(max(7.0, d.axis_size_pt - 2), d.ink, font=st.font_family or d.font_family)
        c.dataLabels = dl
    if t == "bubble":
        c.bubbleScale = 35

    # Axes, legend.
    if t not in ("pie", "donut"):
        _style_axes_xl(c, c_spec, d)
    legend_pos = st.legend_position
    n_series = len(c_spec.series)
    if legend_pos == "none" or (n_series < 2 and t not in ("pie", "donut")):
        c.legend = None
    else:
        c.legend.position = _LEGEND_XL.get(legend_pos, "b")
        c.legend.txPr = _rich(max(7.0, d.axis_size_pt - 1), d.axis_text_color, font=st.font_family or d.font_family)
    if secondary is not None:
        c += secondary
    ws.add_chart(c, anchor)
    keep_chart_data_last(ws.parent)
    return result


def _labels_on(chart: CS.Chart) -> bool:
    st = chart.style or CS.ChartStyle()
    if st.data_labels == "on":
        return chart.type not in CS.XY_TYPES
    if st.data_labels == "off":
        return False
    return len(chart.categories) <= R.LABEL_MAX_CATEGORIES and chart.type not in CS.XY_TYPES + ("line", "radar", "area", "stacked_area")


def _gp_fill(colour_hex: str):
    from openpyxl.chart.shapes import GraphicalProperties

    gp = GraphicalProperties(solidFill=colour_hex)
    gp.line.solidFill = colour_hex
    return gp


def _series_colour(chart: CS.Chart, k: int, d: R.ChartStyleDefaults) -> str:
    st = chart.style or CS.ChartStyle()
    s = chart.series[k]
    if s.color:
        return s.color
    if s.name in st.series_colors:
        return st.series_colors[s.name]
    if k == 0 and st.color and (len(chart.series) == 1 or chart.type == "combo"):
        return st.color
    palette = st.palette or list(d.palette)
    return palette[k % len(palette)]


# ------------------------------------------------------------------- PPTX --


def _label(text: Any) -> str:
    """A label as TEXT for the embedded workbook: formula leads stripped,
    tabs and CRs removed, bounded."""
    out = str(text or "").replace("\t", " ").replace("\r", " ").strip()
    while out and out[0] in _FORMULA_LEADS:
        out = out[1:].lstrip()
    return (out or "-")[:120]


def _pptx_type(t: str):
    from pptx.enum.chart import XL_CHART_TYPE as X

    return {
        "bar": X.COLUMN_CLUSTERED, "horizontal_bar": X.BAR_CLUSTERED, "stacked_bar": X.COLUMN_STACKED,
        "stacked_horizontal_bar": X.BAR_STACKED, "percent_stacked_bar": X.COLUMN_STACKED_100, "line": X.LINE_MARKERS,
        "area": X.AREA, "stacked_area": X.AREA_STACKED, "pie": X.PIE, "donut": X.DOUGHNUT, "scatter": X.XY_SCATTER,
        "histogram": X.COLUMN_CLUSTERED, "combo": X.COLUMN_CLUSTERED, "radar": X.RADAR_MARKERS, "bubble": X.BUBBLE,
    }[t]


def add_pptx_chart(slide, chart: Any, box: Tuple[float, float, float, float], resolved: Any = None) -> Dict[str, Any]:
    """Add `chart` to `slide` in `box` = (left, top, width, height) inches.
    Returns {mode: native|image, type, note}."""
    from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.util import Inches, Pt

    c_spec = R._as_chart(chart) if not isinstance(chart, CS.Chart) else chart
    if not c_spec.series:
        raise ValueError("the chart has no computed values; resolve it first")
    d = R.defaults_from(resolved)
    st = c_spec.style or CS.ChartStyle()
    t = c_spec.type
    x, y, w, h = box
    result: Dict[str, Any] = {"mode": "native", "type": t, "note": ""}
    if support("pptx", t) != "native":
        png = R.render_png(c_spec, resolved, int(w * 200), int(h * 200), orientation="landscape")
        slide.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), Inches(w), Inches(h))
        result.update(mode="image", note=f"{c_spec.title or 'A chart'} ({t.replace('_', ' ')}) is a picture in the PowerPoint file: PowerPoint has no native {t.replace('_', ' ')} chart here.")
        return result

    nf = _number_format(c_spec)
    if t == "scatter":
        data = XyChartData()
        for s in c_spec.series:
            ser = data.add_series(_label(s.name), number_format=nf)
            for xv, yv in zip(s.x or [], s.values):
                ser.add_data_point(float(xv), float(yv))
    elif t == "bubble":
        data = BubbleChartData()
        for s in c_spec.series:
            ser = data.add_series(_label(s.name), number_format=nf)
            for xv, yv, zv in zip(s.x or [], s.values, s.sizes or []):
                ser.add_data_point(float(xv), float(yv), float(zv))
    else:
        data = CategoryChartData(number_format=nf)
        data.categories = [_label(cat) for cat in c_spec.categories]
        for s in c_spec.series:
            data.add_series(_label(s.name), [float(v) for v in s.values])
    frame = slide.shapes.add_chart(_pptx_type(t), Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = frame.chart
    font = st.font_family or d.font_family
    ch.font.name = font
    ch.font.size = Pt(d.axis_size_pt + 2)
    ch.font.color.rgb = RGBColor.from_string(_hex(d.axis_text_color))
    ch.has_title = bool(c_spec.title)
    if c_spec.title:
        ch.chart_title.text_frame.text = c_spec.title
        run = ch.chart_title.text_frame.paragraphs[0].runs[0]
        ts = st.title
        run.font.size = Pt(ts.size_pt if ts and ts.size_pt else d.title_size_pt + 6)
        run.font.bold = True if not ts or ts.bold is None else ts.bold
        run.font.color.rgb = RGBColor.from_string(_hex(ts.color if ts and ts.color else d.ink))
        run.font.name = (ts.font_family if ts and ts.font_family else None) or font
    n_series = len(c_spec.series)
    ch.has_legend = st.legend_position != "none" and (n_series >= 2 or t in ("pie", "donut"))
    if ch.has_legend:
        ch.legend.position = {"bottom": XL_LEGEND_POSITION.BOTTOM, "top": XL_LEGEND_POSITION.TOP, "left": XL_LEGEND_POSITION.LEFT,
                              "right": XL_LEGEND_POSITION.RIGHT}.get(st.legend_position, XL_LEGEND_POSITION.BOTTOM)
        ch.legend.include_in_layout = False
    plot = ch.plots[0]
    if t == "histogram":
        plot.gap_width = 0
    if t in ("stacked_bar", "stacked_horizontal_bar", "percent_stacked_bar"):
        plot.overlap = 100
    many = n_series >= 3
    markers = (XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.SQUARE, XL_MARKER_STYLE.TRIANGLE, XL_MARKER_STYLE.DIAMOND, XL_MARKER_STYLE.X,
               XL_MARKER_STYLE.STAR, XL_MARKER_STYLE.PLUS, XL_MARKER_STYLE.DASH)
    dashes = (MSO_LINE_DASH_STYLE.SOLID, MSO_LINE_DASH_STYLE.DASH, MSO_LINE_DASH_STYLE.DASH_DOT, MSO_LINE_DASH_STYLE.ROUND_DOT,
              MSO_LINE_DASH_STYLE.LONG_DASH, MSO_LINE_DASH_STYLE.LONG_DASH_DOT, MSO_LINE_DASH_STYLE.SQUARE_DOT, MSO_LINE_DASH_STYLE.DASH_DOT_DOT)
    if t in ("pie", "donut"):
        pts = plot.series[0].points
        palette = st.palette or list(d.palette)
        for i, cat in enumerate(c_spec.categories):
            pts[i].format.fill.solid()
            pts[i].format.fill.fore_color.rgb = RGBColor.from_string(_hex(st.category_colors.get(cat) or palette[i % len(palette)]))
    else:
        for k, ser in enumerate(plot.series):
            colour = RGBColor.from_string(_hex(_series_colour(c_spec, k, d)))
            if t in ("line", "radar", "scatter"):
                if t != "scatter":
                    ser.format.line.color.rgb = colour
                    ser.format.line.width = Pt(2.25)
                    if many:
                        ser.format.line.dash_style = dashes[k % len(dashes)]
                ser.marker.style = markers[k % len(markers)] if many or t == "scatter" and n_series >= 3 else XL_MARKER_STYLE.CIRCLE
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = colour
                ser.marker.format.line.color.rgb = colour
                if t != "scatter":
                    ser.smooth = False
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = colour
                if n_series == 1 and st.category_colors:
                    for i, cat in enumerate(c_spec.categories):
                        if cat in st.category_colors:
                            ser.points[i].format.fill.solid()
                            ser.points[i].format.fill.fore_color.rgb = RGBColor.from_string(_hex(st.category_colors[cat]))
    labels_on = _labels_on(c_spec)
    if labels_on and t in CS.STACKED_TYPES:
        for k, ser in enumerate(plot.series):
            fill = _series_colour(c_spec, k, d)
            colour = d.label_color_for(fill)
            if R.contrast_ratio(colour, fill) < 4.5:
                continue
            sdl = ser.data_labels
            sdl.show_value = True
            sdl.number_format = "0" if t == "percent_stacked_bar" else nf
            sdl.number_format_is_linked = False
            sdl.position = XL_LABEL_POSITION.CENTER
            sdl.font.size = Pt(max(9.0, d.axis_size_pt))
            sdl.font.color.rgb = RGBColor.from_string(_hex(colour))
    elif labels_on:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(max(9.0, d.axis_size_pt))
        dl.font.color.rgb = RGBColor.from_string(_hex(d.ink))
        if t in ("pie", "donut"):
            dl.number_format = "0%"
            dl.number_format_is_linked = False
            dl.show_percentage = True
            dl.show_value = False
            if t == "pie":
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
        else:
            dl.number_format = nf
            dl.number_format_is_linked = False
            if t in ("bar", "horizontal_bar", "histogram", "combo"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
    if t == "bubble":
        plot.bubble_scale = 35
    if t not in ("pie", "donut"):
        va = ch.value_axis
        va.has_major_gridlines = bool(st.gridlines)
        if st.gridlines:
            va.major_gridlines.format.line.color.rgb = RGBColor.from_string(_hex(d.grid_color))
        va.tick_labels.number_format = "0%" if t == "percent_stacked_bar" else nf
        va.tick_labels.number_format_is_linked = False
        if st.y_min is not None:
            va.minimum_scale = st.y_min
        elif t in CS.BAR_FAMILY + ("histogram", "combo", "area", "stacked_area"):
            va.minimum_scale = min(0.0, min((v for s in c_spec.series for v in s.values), default=0.0))
        if st.y_max is not None:
            va.maximum_scale = st.y_max
        if c_spec.y_label:
            va.has_title = True
            va.axis_title.text_frame.text = c_spec.y_label
        try:
            ca = ch.category_axis
            if c_spec.x_label:
                ca.has_title = True
                ca.axis_title.text_frame.text = c_spec.x_label
        except Exception:
            pass
        if st.log_y:
            _pptx_log_axis(ch)
    if t == "combo":
        _pptx_combo(ch, c_spec, nf)
        line_idx = [k for k, s in enumerate(c_spec.series) if (s.kind or "bar") == "line" or s.axis == "secondary"]
        if len(ch.plots) > 1:
            for j, ser in enumerate(ch.plots[1].series):
                k = line_idx[j]
                colour = RGBColor.from_string(_hex(_series_colour(c_spec, k, d)))
                ser.format.line.color.rgb = colour
                ser.format.line.width = Pt(2.25)
                ser.format.line.dash_style = dashes[(j + 1) % len(dashes)]
                ser.marker.style = markers[(j + 1) % len(markers)]
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = colour
    if t == "scatter" and c_spec.extra and c_spec.extra.trendlines:
        _pptx_trendlines(ch, c_spec)
    return result


def _c(tag: str, **attrs: str):
    from lxml import etree
    from pptx.oxml.ns import qn

    el = etree.Element(qn(f"c:{tag}"))
    for k, v in attrs.items():
        el.set(k, v)
    return el


def _pptx_log_axis(ch) -> None:
    from pptx.oxml.ns import qn

    for va in ch._chartSpace.chart.plotArea.findall(qn("c:valAx")):
        scaling = va.find(qn("c:scaling"))
        if scaling is not None and scaling.find(qn("c:logBase")) is None:
            scaling.insert(0, _c("logBase", val="10"))


def _pptx_trendlines(ch, chart: CS.Chart) -> None:
    from pptx.oxml.ns import qn

    names = {tr.series for tr in chart.extra.trendlines} if chart.extra else set()
    plot = ch._chartSpace.chart.plotArea.find(qn("c:scatterChart"))
    if plot is None:
        return
    for k, ser in enumerate(plot.findall(qn("c:ser"))):
        if k >= len(chart.series) or chart.series[k].name not in names:
            continue
        tl = _c("trendline")
        tl.append(_c("trendlineType", val="linear"))
        tl.append(_c("dispRSqr", val="1"))
        tl.append(_c("dispEq", val="0"))
        anchor = ser.find(qn("c:xVal"))
        if anchor is None:
            ser.append(tl)
        else:
            anchor.addprevious(tl)


def _pptx_combo(ch, chart: CS.Chart, nf: str) -> None:
    """Move the line series out of the bar plot into a c:lineChart; put them
    on secondary axes when any is marked secondary."""
    from pptx.oxml.ns import qn

    plot_area = ch._chartSpace.chart.plotArea
    bar = plot_area.find(qn("c:barChart"))
    if bar is None:
        return
    sers = bar.findall(qn("c:ser"))
    line_idx = [k for k, s in enumerate(chart.series) if (s.kind or "bar") == "line" or s.axis == "secondary"]
    if not line_idx:
        return
    bar_ax = [e.get("val") for e in bar.findall(qn("c:axId"))]
    secondary = any(chart.series[k].axis == "secondary" for k in line_idx)
    line = _c("lineChart")
    line.append(_c("grouping", val="standard"))
    line.append(_c("varyColors", val="0"))
    for k in line_idx:
        ser = sers[k]
        bar.remove(ser)
        inv = ser.find(qn("c:invertIfNegative"))
        if inv is not None:
            ser.remove(inv)
        marker = _c("marker")
        marker.append(_c("symbol", val="circle"))
        marker.append(_c("size", val="6"))
        after = ser.find(qn("c:spPr"))
        if after is None:
            after = ser.find(qn("c:tx"))
        after.addnext(marker)
        ser.append(_c("smooth", val="0"))
        line.append(ser)
    line.append(_c("marker", val="1"))
    ids = ("50001", "50002") if secondary else tuple(bar_ax[:2])
    for i in ids:
        line.append(_c("axId", val=i))
    bar.addnext(line)
    if not secondary:
        return
    last_ax = plot_area.findall(qn("c:valAx"))[-1]
    cat = _c("catAx")
    for el in (_c("axId", val=ids[0]), _scaling(), _c("delete", val="1"), _c("axPos", val="b"), _c("majorTickMark", val="none"),
               _c("minorTickMark", val="none"), _c("tickLblPos", val="nextTo"), _c("crossAx", val=ids[1]), _c("crosses", val="autoZero"),
               _c("auto", val="1"), _c("lblAlgn", val="ctr"), _c("lblOffset", val="100"), _c("noMultiLvlLbl", val="0")):
        cat.append(el)
    val = _c("valAx")
    for el in (_c("axId", val=ids[1]), _scaling(), _c("delete", val="0"), _c("axPos", val="r"), _c("numFmt", formatCode=nf, sourceLinked="0"),
               _c("majorTickMark", val="out"), _c("minorTickMark", val="none"), _c("tickLblPos", val="nextTo"), _c("crossAx", val=ids[0]),
               _c("crosses", val="max"), _c("crossBetween", val="between")):
        val.append(el)
    last_ax.addnext(cat)
    cat.addnext(val)
    if chart.y2_label:
        from pptx.oxml.ns import qn as _qn  # noqa: F401 - the title is optional chrome; values are what matter


def _scaling():
    s = _c("scaling")
    s.append(_c("orientation", val="minMax"))
    return s


def expected_native_counts(charts: Sequence[CS.Chart], fmt: str) -> Tuple[int, int]:
    """(native, image) counts a container should hold for these charts."""
    native = sum(1 for c in charts if support(fmt, c.type) == "native")
    return native, len(charts) - native


__all__ = [
    "CHART_DATA_SHEET", "EXCEL_NUMBER_FORMATS", "add_xlsx_chart", "add_pptx_chart", "chart_data_sheet", "keep_chart_data_last",
    "expected_native_counts", "support",
]
