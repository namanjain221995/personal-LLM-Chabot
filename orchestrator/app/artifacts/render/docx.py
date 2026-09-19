"""DocumentSpec → .docx with python-docx, following the same plan as the PDF.

WHY python-docx AND NOT PANDOC. The legacy report path converts Markdown with
pandoc and ships pandoc's reference.docx styles (CURRENT_STATE.md: "no
--reference-doc"). Here the DOCX is built from the plan html.py made — the
same hoisted KPI row, the same heading numbers, the same citation numbers,
the same sources and assumptions sections — so a person who opens the .docx
next to the .pdf sees the same document, block for block. Every paragraph
gets a NAMED style (Title, Subtitle, Heading 1-3, Body, Caption, Callout,
Bullet, Number) so the file behaves like one a person made: the navigation
pane works, "update table of contents" works, and restyling is one click.

THE CLASSY DEFAULT (style guide §3, artifacts/style.py). The named styles
are written from the ResolvedStyle — TechSara Classic unless the spec's
style says otherwise: a 28 pt #1F3864 title with no letter spacing and a
2 pt rule under the title block (the template's Title border and character
spacing are removed), H1 18 pt with a hairline rule, H2 in the accent, H3
in ink, every heading kept with the next block; docDefaults and the THEME
fonts (theme1.xml major/minor) set to the resolved families, so text Word
creates later matches too, with 'Nirmala UI' in the complex-script slot
when the document has Devanagari or Gujarati; a 'TS Table' style with a
header fill, hairline rows and zebra bands; KPI tiles with a top bar and
callouts with a left bar drawn as cell borders; margins, paper size and
orientation from the page style. A rule that names ONE element (the
'Recommendations' heading, the first paragraph, a column) becomes run and
paragraph properties on that element only, so the named style stays clean.

WHAT WORD DOES THAT WE CANNOT. Page numbers are FIELDS (PAGE / NUMPAGES)
that Word evaluates on open — python-docx cannot know the page count, and a
number typed into the footer would be a lie. The contents page is a TOC
field for the same reason. The field instructions are fixed strings in this
file; header/footer text a person asked for goes into plain runs, never
into a field.

NO EXTERNAL RELATIONSHIPS. Source URLs are text; nothing is linked or
fetched, and no remote image is ever embedded (validate.py refuses an
External target).

python-docx is imported lazily: tests/test_imports.py asserts the app
imports without heavy or optional libraries.

THE TABULAR DOCUMENT (`render_workbook_docx`, CONTRACT-2 §1). A workbook
delivered as Word: the title and a methodology note, then one SECTION per
sheet — python-docx sections carry their own orientation, so a nine-column
audit sheet is landscape (page size swapped, not just the flag, or Word
draws portrait) while a three-column sheet stays portrait — with the full
table: the header row marked `w:tblHeader` so Word repeats it on every
page, grid borders, the header shaded as the Excel file shades it, the
highlighted column in the same red pair, cells formatted by
theme.format_cell, the totals row computed, rows that do not split across
pages. The plan (orientation, column shares, colours) comes from html.py so
the PDF twin agrees.
"""
from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from .. import spec as S
from .. import style as ST
from . import theme
from .html import DocumentPlan, cell_text, chart_filename, document_orientation, plan_document

_OXML = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_INDIC = ("Devanagari", "Gujarati", "Bengali", "Gurmukhi", "Odia", "Tamil", "Telugu", "Kannada", "Malayalam")


def _qn(tag: str) -> str:
    from docx.oxml.ns import qn

    return qn(tag)


def _rgb(colour: str):
    from docx.shared import RGBColor

    return RGBColor(*theme.hex_to_rgb(colour))


def _hex6(colour: str) -> str:
    return colour.lstrip("#").upper()


def _el(tag: str, **attrs: str):
    from docx.oxml import OxmlElement

    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(_qn(k), v)
    return el


#: Child order the OOXML schema requires (ECMA-376 §17): Word refuses a
#: part whose property children are out of sequence, so every element this
#: module adds goes through `_put`, never a bare append.
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
              "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
              "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
              "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign",
               "hideMark", "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"]
_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing",
                "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription", "tblPrChange"]
_TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc",
               "hidden", "ins", "del", "trPrChange"]


def _put(parent, child, order: Sequence[str]) -> None:
    """Replace `parent`'s child of the same tag with `child`, at the position
    the schema sequence `order` requires."""
    local = child.tag.split("}", 1)[-1]
    for old in parent.findall(child.tag):
        parent.remove(old)
    rank = order.index(local) if local in order else len(order)
    for i, existing in enumerate(list(parent)):
        name = existing.tag.split("}", 1)[-1]
        if name in order and order.index(name) > rank:
            existing.addprevious(child)
            return
    parent.append(child)


def _shade(cell, fill: str) -> None:
    """Cell background: python-docx has no API for w:shd, so write the element."""
    _put(cell._tc.get_or_add_tcPr(), _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": _hex6(fill)}), _TCPR_ORDER)


def _cell_border(cell, **edges: tuple) -> None:
    """`edges`: top/bottom/left/right → (size in eighths of a point, colour)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(_qn("w:tcBorders"))
    if borders is None:
        borders = _el("w:tcBorders")
        _put(tc_pr, borders, _TCPR_ORDER)
    for edge, (size, colour) in edges.items():
        for old in borders.findall(_qn(f"w:{edge}")):
            borders.remove(old)
        if size <= 0:
            borders.append(_el(f"w:{edge}", **{"w:val": "nil"}))
        else:
            borders.append(_el(f"w:{edge}", **{"w:val": "single", "w:sz": str(int(size)), "w:space": "0", "w:color": _hex6(colour)}))
    _sort_children(borders, _BORDER_ORDER)


_BORDER_ORDER = ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV", "tl2br", "tr2bl", "between", "bar"]


def _sort_children(parent, order: Sequence[str]) -> None:
    children = list(parent)
    children.sort(key=lambda c: order.index(c.tag.split("}", 1)[-1]) if c.tag.split("}", 1)[-1] in order else len(order))
    for c in children:
        parent.remove(c)
        parent.append(c)


def _para_border(paragraph_or_style_element, edge: str, size_eighths: int, colour: str, space: int = 4) -> None:
    ppr = paragraph_or_style_element.get_or_add_pPr()
    pbdr = ppr.find(_qn("w:pBdr"))
    if pbdr is None:
        pbdr = _el("w:pBdr")
        _put(ppr, pbdr, _PPR_ORDER)
    for old in pbdr.findall(_qn(f"w:{edge}")):
        pbdr.remove(old)
    pbdr.append(_el(f"w:{edge}", **{"w:val": "single", "w:sz": str(size_eighths), "w:space": str(space), "w:color": _hex6(colour)}))
    _sort_children(pbdr, _BORDER_ORDER)


def _para_shade(paragraph, fill: str) -> None:
    _put(paragraph._p.get_or_add_pPr(), _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": _hex6(fill)}), _PPR_ORDER)


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    """A complex field (PAGE, NUMPAGES, TOC ...): begin / instrText / separate
    / placeholder result / end. Word recomputes it on open. `instruction` is
    always a constant in this module."""

    def run_with(child_tag: str, text: Optional[str] = None, **attrs: str):
        run = paragraph.add_run()
        el = _el(child_tag, **attrs)
        if text is not None:
            el.text = text
        run._r.append(el)
        return run

    run_with("w:fldChar", **{"w:fldCharType": "begin"})
    run_with("w:instrText", instruction, **{"xml:space": "preserve"})
    run_with("w:fldChar", **{"w:fldCharType": "separate"})
    paragraph.add_run(placeholder)
    run_with("w:fldChar", **{"w:fldCharType": "end"})


# ----------------------------------------------------------------- styles --


def _set_rfonts(rpr, font: str, cs_font: Optional[str] = None) -> None:
    rfonts = rpr.find(_qn("w:rFonts"))
    if rfonts is None:
        rfonts = _el("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(_qn(attr), font)
    rfonts.set(_qn("w:cs"), cs_font or font)
    # The default template's Title/Heading styles name THEME fonts, which
    # win over an explicit family (measured: the title came out in the
    # theme's Calibri-Light substitute, not the family asked for). Drop them.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(_qn(attr), None)


def _apply_font(font, rpr, ts: ST.TextStyle, R: ST.ResolvedStyle, cs_font: Optional[str]) -> None:
    from docx.shared import Pt

    face = R.face(ts.font_family)
    font.name = face.office_name
    _set_rfonts(rpr, face.office_name, cs_font)
    if ts.size_pt:
        font.size = Pt(ts.size_pt)
    font.bold = bool(ts.bold)
    font.italic = bool(ts.italic)
    font.underline = bool(ts.underline)
    if ts.color:
        font.color.rgb = _rgb(ts.color)


def _clean_template_style(st) -> None:
    """The default template's Title has a bottom border and wide character
    spacing, Subtitle has spacing too; the classy title has neither."""
    el = st.element
    ppr = el.pPr
    if ppr is not None:
        for tag in ("w:pBdr",):
            for old in ppr.findall(_qn(tag)):
                ppr.remove(old)
    rpr = el.rPr
    if rpr is not None:
        for tag in ("w:spacing", "w:kern", "w:color"):
            for old in rpr.findall(_qn(tag)):
                rpr.remove(old)


def _ensure_styles(document, t: theme.TypeScale, R: Optional[ST.ResolvedStyle] = None, *, cs_font: Optional[str] = None) -> None:
    """Named styles, created or restyled from the ResolvedStyle. python-docx's
    default template already defines Title/Subtitle/Heading n/Caption/List
    Bullet/List Number; Body, Callout, Lede, KPI and Table styles are ours."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt

    R = R or ST.resolve(None, type_scale=t)
    styles = document.styles
    tokens = R.tokens

    def para_style(name: str, base: Optional[str], ts: ST.TextStyle, *, space_before: float = 0, space_after: float = 6,
                   line: Optional[float] = None):
        try:
            st = styles[name]
        except KeyError:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            if base:
                st.base_style = styles[base]
        _clean_template_style(st)
        _apply_font(st.font, st.element.get_or_add_rPr(), ts, R, cs_font)
        pf = st.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = line if line is not None else t.line_height
        ppr = st.element.get_or_add_pPr()
        for old in ppr.findall(_qn("w:shd")):
            ppr.remove(old)
        if ts.background:
            _put(ppr, _el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": _hex6(ts.background)}), _PPR_ORDER)
        if ts.align:
            pf.alignment = _ALIGN[ts.align]
        return st

    def el(kind: str, **kw: Any) -> ST.TextStyle:
        return R.element(kind, generic_only=True, **kw)

    body = el("paragraph")
    normal = para_style("Normal", None, body)
    para_style("Body", "Normal", body)
    title = para_style("Title", "Normal", el("title"), space_after=4, line=1.1)
    para_style("Subtitle", "Normal", el("subtitle"), space_after=6)
    for level, (before, after) in ((1, (18, 6)), (2, (14, 4)), (3, (10, 3))):
        st = para_style(f"Heading {level}", "Normal", el("heading", level=level), space_before=before, space_after=after, line=1.15)
        st.paragraph_format.keep_with_next = True
        if level == 1:
            _para_border(st.element, "bottom", 6, tokens.hairline, space=2)
    para_style("Caption", "Normal", el("caption"), space_after=10)
    para_style("Callout", "Normal", el("callout"), space_after=4)
    para_style("Bullet", "List Bullet", el("bullet"), space_after=3)
    para_style("Number", "List Number", el("bullet"), space_after=3)
    lede = el("paragraph")
    para_style("Lede", "Normal", ST.TextStyle.model_construct(**{**lede.model_dump(), "size_pt": (lede.size_pt or t.body) + 1.5, "color": tokens.accent_text}), space_after=10)
    para_style("KPI Value", "Normal", el("kpi_value"), space_after=0, line=1.0)
    para_style("KPI Label", "Normal", el("kpi_label"), space_after=0, line=1.0)
    hf = el("header_footer")
    para_style("Header", "Normal", hf, space_after=0)
    para_style("Footer", "Normal", hf, space_after=0)
    para_style("Table Text", "Normal", el("table_body"), space_after=0, line=1.1)
    normal.paragraph_format.widow_control = True
    title.paragraph_format.keep_with_next = True
    _table_style(document, R)
    _doc_defaults(document, R, cs_font)


def _table_style(document, R: ST.ResolvedStyle) -> None:
    """'TS Table': hairline rows, no verticals, 4/6 pt cell padding."""
    from docx.enum.style import WD_STYLE_TYPE

    styles = document.styles
    try:
        st = styles["TS Table"]
    except KeyError:
        st = styles.add_style("TS Table", WD_STYLE_TYPE.TABLE)
    el = st.element
    tbl_pr = el.find(_qn("w:tblPr"))
    if tbl_pr is None:
        tbl_pr = _el("w:tblPr")
        el.append(tbl_pr)
    borders = _el("w:tblBorders")
    hair = _hex6(R.tokens.hairline)
    for edge in ("top", "bottom", "insideH"):
        borders.append(_el(f"w:{edge}", **{"w:val": "single", "w:sz": "4", "w:space": "0", "w:color": hair}))
    for edge in ("left", "right", "insideV"):
        borders.append(_el(f"w:{edge}", **{"w:val": "nil"}))
    _sort_children(borders, _BORDER_ORDER)
    _put(tbl_pr, borders, _TBLPR_ORDER)
    mar = _el("w:tblCellMar")
    for edge, twips in (("top", "80"), ("left", "120"), ("bottom", "80"), ("right", "120")):
        mar.append(_el(f"w:{edge}", **{"w:w": twips, "w:type": "dxa"}))
    _put(tbl_pr, mar, _TBLPR_ORDER)


def _doc_defaults(document, R: ST.ResolvedStyle, cs_font: Optional[str]) -> None:
    """docDefaults rFonts and theme1.xml major/minor fonts → the resolved
    families, so what Word creates later (a new paragraph, a comment) is
    the same family as the document."""
    styles_el = document.styles.element
    defaults = styles_el.find(_qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(_qn("w:rPrDefault"))
        rpr = rpr_default.find(_qn("w:rPr")) if rpr_default is not None else None
        if rpr is not None:
            _set_rfonts(rpr, R.body_face.office_name, cs_font)
    try:
        from lxml import etree
    except ImportError:  # pragma: no cover - lxml ships with python-docx
        return
    for part in document.part.package.iter_parts():
        if str(part.partname).endswith("/theme/theme1.xml"):
            try:
                root = etree.fromstring(part.blob)
            except Exception:  # noqa: BLE001 — a theme we cannot read is left alone
                return
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            heading_face = R.face(R.element("heading", generic_only=True, level=1).font_family)
            body_face = R.face(R.element("paragraph", generic_only=True).font_family)
            for which, face in (("majorFont", heading_face), ("minorFont", body_face)):
                latin = root.find(f".//a:fontScheme/a:{which}/a:latin", ns)
                if latin is not None:
                    latin.set("typeface", face.office_name)
                    for attr in ("panose", "pitchFamily", "charset"):
                        latin.attrib.pop(attr, None)
                if cs_font:
                    cs = root.find(f".//a:fontScheme/a:{which}/a:cs", ns)
                    if cs is not None:
                        cs.set("typeface", cs_font)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            return


def _apply_overrides(paragraph, specific: ST.TextStyle, generic: ST.TextStyle, R: ST.ResolvedStyle, cs_font: Optional[str]) -> None:
    """Run and paragraph properties for a rule that names THIS element: only
    the properties that differ from its named style."""
    from docx.shared import Pt

    diff = {k: v for k, v in specific.model_dump().items() if v != getattr(generic, k)}
    if not diff:
        return
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        if "font_family" in diff and specific.font_family:
            face = R.face(specific.font_family)
            run.font.name = face.office_name
            _set_rfonts(rpr, face.office_name, cs_font)
        if "size_pt" in diff and specific.size_pt:
            run.font.size = Pt(specific.size_pt)
        if "bold" in diff:
            run.font.bold = bool(specific.bold)
        if "italic" in diff:
            run.font.italic = bool(specific.italic)
        if "underline" in diff:
            run.font.underline = bool(specific.underline)
        if "color" in diff and specific.color:
            run.font.color.rgb = _rgb(specific.color)
    if "background" in diff and specific.background:
        _para_shade(paragraph, specific.background)
    if "align" in diff and specific.align:
        paragraph.alignment = _ALIGN[specific.align]


def _align_map():
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    return {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


class _LazyAlign(dict):
    def __missing__(self, key):
        self.update(_align_map())
        return dict.__getitem__(self, key)


_ALIGN: Dict[str, Any] = _LazyAlign()


def _style_run(run, ts: ST.TextStyle, R: ST.ResolvedStyle, cs_font: Optional[str]) -> None:
    """Every property of `ts` on one run (table cells, tiles)."""
    _apply_font(run.font, run._r.get_or_add_rPr(), ts, R, cs_font)


# ------------------------------------------------------------------- page --


def _set_page(section, R: ST.ResolvedStyle, orientation: str, grid: theme.PageGrid) -> None:
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm

    w, h = ST.PAGE_SIZES_MM.get(R.page.size, (210, 297))
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        w, h = max(w, h), min(w, h)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        w, h = min(w, h), max(w, h)
    section.page_width, section.page_height = Mm(w), Mm(h)
    top, bottom, left, right = R.page.margins_mm or (grid.top_mm, grid.bottom_mm, grid.left_mm, grid.right_mm)
    section.top_margin = Mm(top)
    section.bottom_margin = Mm(bottom)
    section.left_margin = Mm(left)
    section.right_margin = Mm(right)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)


def _set_margins(document, grid: theme.PageGrid, orientation: str, R: Optional[ST.ResolvedStyle] = None) -> None:
    _set_page(document.sections[0], R or ST.resolve(None), orientation, grid)


def _right_tab(paragraph, section) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT

    from docx.shared import Twips

    usable = section.page_width - section.left_margin - section.right_margin
    stops = paragraph.paragraph_format.tab_stops
    # The template's Header/Footer styles carry centre (3.25 in) and right
    # (6.5 in) stops; a '\t' would land on the centre one. Clear both.
    for twips in (4680, 9360):
        stops.add_tab_stop(Twips(twips), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)


def _header_footer(document, spec: S.DocumentSpec, plan: DocumentPlan, R: Optional[ST.ResolvedStyle] = None, cs_font: Optional[str] = None) -> None:
    """Header: the title (or the requested header text) left, the date or
    CONFIDENTIAL on a right tab stop, a hairline under it. Footer: the
    author (or the requested footer text) left, 'Page X of Y' right. Text a
    person supplied goes into plain runs; the fields are fixed code."""
    R = R or ST.resolve(None)
    section = document.sections[0]
    hf = R.element("header_footer")
    if plan.cover:
        section.different_first_page_header_footer = True  # the cover carries none
    if plan.grid.header:
        p = section.header.paragraphs[0]
        p.style = document.styles["Header"]
        _right_tab(p, section)
        run = p.add_run(R.page.header_text or spec.title)
        run.bold = True if not R.page.header_text else bool(hf.bold)
        right = "CONFIDENTIAL" if spec.confidential else spec.date
        if right:
            r2 = p.add_run("\t" + right)
            if spec.confidential:
                r2.font.color.rgb = _rgb(ST.STATUS_PAIRS["danger"][1])
                r2.bold = True
        _para_border(p._p, "bottom", 4, R.tokens.hairline, space=3)
    if plan.grid.footer:
        p = section.footer.paragraphs[0]
        p.style = document.styles["Footer"]
        _right_tab(p, section)
        left = R.page.footer_text or " · ".join(b for b in (spec.author,) if b)
        if left:
            p.add_run(left)
        if R.page.page_numbers:
            p.add_run("\tPage ")
            _field(p, "PAGE", "1")
            p.add_run(" of ")
            _field(p, "NUMPAGES", "1")


# ------------------------------------------------------------------ blocks --


def _cover(document, spec: S.DocumentSpec, plan: DocumentPlan, R: Optional[ST.ResolvedStyle] = None, cs_font: Optional[str] = None) -> None:
    """The cover: a primary band carrying the kicker, the title in white and
    the subtitle, an accent bar under it, then the meta lines."""
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    R = R or ST.resolve(None)
    t = R.tokens
    kicker = document.add_paragraph(plan.template_id.replace("_", " ").upper(), style="KPI Label")
    kicker.paragraph_format.space_before = Pt(96)
    kicker.runs[0].font.color.rgb = _rgb(t.accent_text)
    title = document.add_paragraph(spec.title, style="Title")
    title_ts = R.element("title")
    band = t.cover_band
    if not R.matching_rules("title"):
        _para_shade(title, band)
        for run in title.runs:
            run.font.color.rgb = _rgb(ST.readable_on(band))
            run.font.size = Pt(32)
    else:
        if title_ts.background:
            _para_shade(title, title_ts.background)
    last = title
    if spec.subtitle:
        sub = document.add_paragraph(spec.subtitle, style="Subtitle")
        if not R.matching_rules("subtitle"):
            _para_shade(sub, band)
            for run in sub.runs:
                run.font.color.rgb = _rgb("#DCE6F2" if ST.contrast_ratio("#DCE6F2", band) >= 4.5 else ST.readable_on(band))
                run.font.size = Pt(16)
        last = sub
    _para_border(last._p, "bottom", 32, t.accent, space=6)
    meta = document.add_paragraph(style="Body")
    meta.paragraph_format.space_before = Pt(48)
    for label, value in (("Prepared for", spec.audience), ("Prepared by", spec.author), ("Date", spec.date), ("Purpose", spec.purpose)):
        if value:
            r = meta.add_run(f"{label}:  ")
            r.font.color.rgb = _rgb(t.muted)
            meta.add_run(value + "\n")
    if spec.confidential:
        c = document.add_paragraph(style="KPI Label")
        c.add_run("CONFIDENTIAL").font.color.rgb = _rgb(ST.STATUS_PAIRS["danger"][1])
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _kpis(document, row: S.KPIRow, R: Optional[ST.ResolvedStyle] = None, cs_font: Optional[str] = None) -> None:
    """KPI tiles: band fill, a 3 pt accent bar on top, the value large, the
    label small and uppercase."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    R = R or ST.resolve(None)
    table = document.add_table(rows=1, cols=len(row.items))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    value_ts, label_ts = R.element("kpi_value"), R.element("kpi_label")
    for cell, kpi in zip(table.rows[0].cells, row.items):
        _shade(cell, value_ts.background or R.tokens.band)
        _cell_border(cell, top=(24, R.tokens.kpi_bar), left=(0, ""), right=(0, ""), bottom=(0, ""))
        p = cell.paragraphs[0]
        p.style = document.styles["KPI Value"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_run(p.add_run(kpi.value), value_ts, R, cs_font)
        lab = cell.add_paragraph(style="KPI Label")
        lab.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_run(lab.add_run(kpi.label.upper()), label_ts, R, cs_font)
        if kpi.note:
            note = cell.add_paragraph(style="KPI Label")
            note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_run(note.add_run(kpi.note), R.element("caption"), R, cs_font)
    document.add_paragraph(style="Body")


def _table(document, table: S.Table, cite: str, R: Optional[ST.ResolvedStyle] = None, *, index: int = 1, cs_font: Optional[str] = None) -> None:
    """A data table: header fill with contrasting bold text and a 1 pt rule
    under it, repeated on every page; zebra bands; numbers right-aligned;
    rows that never split."""
    R = R or ST.resolve(None)
    numeric = set(table.numeric_columns)
    t = document.add_table(rows=1, cols=len(table.columns))
    t.style = document.styles["TS Table"]
    head_row = t.rows[0]
    for i, (cell, name) in enumerate(zip(head_row.cells, table.columns)):
        ts = R.element("table_header", table=index, column=name, column_index=i)
        if ts.background:
            _shade(cell, ts.background)
        _cell_border(cell, bottom=(8, R.tokens.primary))
        p = cell.paragraphs[0]
        p.style = document.styles["Table Text"]
        _style_run(p.add_run(name), ts, R, cs_font)
        p.alignment = _ALIGN[ts.align] if ts.align else (_ALIGN["right"] if i in numeric else None)
    for r_index, row in enumerate(table.rows):
        new_row = t.add_row()
        _keep_row_whole(new_row)
        for i, (cell, value) in enumerate(zip(new_row.cells, row)):
            ts = R.element("table_body", table=index, column=table.columns[i], column_index=i, row=r_index + 1)
            p = cell.paragraphs[0]
            p.style = document.styles["Table Text"]
            _style_run(p.add_run(cell_text(value, i in numeric)), ts, R, cs_font)
            p.alignment = _ALIGN[ts.align] if ts.align else (_ALIGN["right"] if i in numeric else None)
            if ts.background:
                _shade(cell, ts.background)
            elif R.banded and r_index % 2 == 1:
                _shade(cell, R.tokens.band)
    _repeat_header(head_row)
    if table.caption or cite:
        document.add_paragraph((table.caption + cite).strip(), style="Caption")
    else:
        document.add_paragraph(style="Body")


def _chart(document, chart: S.Chart, image: Optional[str], cite: str, missing: List[str], width_in: float = 6.3) -> None:
    from docx.shared import Inches

    if image and Path(image).is_file():
        document.add_picture(str(image), width=Inches(width_in))
    else:
        missing.append(chart.title or "chart")
        document.add_paragraph(f"[Chart not available: {chart.title or 'untitled'}]", style="Caption")
    if chart.caption or cite:
        document.add_paragraph((chart.caption + cite).strip(), style="Caption")


_CALLOUT_STATUS = {"note": "info", "tip": "success", "warning": "warning", "quote": "neutral"}


def _callout(document, c: S.Callout, R: Optional[ST.ResolvedStyle] = None, cs_font: Optional[str] = None) -> None:
    """A callout: the status fill of its kind with a 3 pt left bar in the
    status text colour."""
    from docx.shared import Pt

    R = R or ST.resolve(None)
    fill, text = ST.STATUS_PAIRS[_CALLOUT_STATUS.get(c.kind, "info")]
    ts = R.element("callout")
    table = document.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    _shade(cell, ts.background or fill)
    _cell_border(cell, left=(24, text), top=(0, ""), right=(0, ""), bottom=(0, ""))
    p = cell.paragraphs[0]
    p.style = document.styles["Callout"]
    if c.kind != "quote" or c.title:
        title = p.add_run((c.title or c.kind.capitalize()).upper() + "\n")
        title.bold = True
        title.font.size = Pt(max(8.0, (ts.size_pt or 11) - 2))
        title.font.color.rgb = _rgb(text)
    body = p.add_run(c.text)
    _style_run(body, ts, R, cs_font)
    if c.kind == "quote" and not R.matching_rules("callout"):
        body.italic = True
        serif = ST.font_face("Cambria")
        body.font.name = serif.office_name if serif else theme.OFFICE_SERIF
    document.add_paragraph(style="Body")


def _list_number_abstract_id(document) -> Optional[int]:
    """The abstractNumId behind the built-in 'List Number' style, or None
    when the template has no numbering for it."""
    style = document.styles["List Number"]
    ppr = style.element.pPr
    num_pr = ppr.numPr if ppr is not None else None
    if num_pr is None or num_pr.numId is None:
        return None
    numbering = document.part.numbering_part.numbering_definitions._numbering
    num = numbering.num_having_numId(num_pr.numId.val)
    return int(num.abstractNumId.val)


def _restart_numbering(document, paragraphs: Sequence, abstract_id: Optional[int]) -> None:
    """Give one Numbered block its OWN <w:num> so it starts at 1. Every
    paragraph in the 'Number' style otherwise shares the single numbering
    instance of 'List Number', and the second list in an SOP continued
    "3. 4." where the PDF (a fresh <ol>) showed "1. 2." (review 2026-09-11)."""
    if abstract_id is None or not paragraphs:
        return
    numbering = document.part.numbering_part.numbering_definitions._numbering
    num = numbering.add_num(abstract_id)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    for p in paragraphs:
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num.numId


def _cite_text(ids: Sequence[str], index: Dict[str, int]) -> str:
    nums = sorted({index[i] for i in ids if i in index})
    return f" [{', '.join(str(n) for n in nums)}]" if nums else ""


def _cs_font_for(text: str) -> Optional[str]:
    return "Nirmala UI" if any(s in _INDIC for s in theme.unsupported_scripts(text)) else None


def render_docx(spec: S.DocumentSpec, out_path: str | Path, chart_dir: str | Path, *,
                plan: Optional[DocumentPlan] = None, warnings: Optional[List[str]] = None,
                resolved: Optional[ST.ResolvedStyle] = None) -> Path:
    """Write the document to `out_path`. Chart PNGs are read from
    `chart_dir` as `chart-<n>.png` (rendered first by charts.py)."""
    from docx import Document
    from docx.enum.text import WD_BREAK

    plan = plan or plan_document(spec)
    R = resolved or ST.resolve(spec, type_scale=plan.type)
    warnings = warnings if warnings is not None else []
    index = plan.citation_index
    chart_dir = Path(chart_dir)
    missing: List[str] = []
    cs_font = _cs_font_for(S.text_of(S.ArtifactSpec(kind="document", document=spec)))

    document = Document()
    _ensure_styles(document, plan.type, R, cs_font=cs_font)
    orientation = document_orientation(spec, R)
    _set_page(document.sections[0], R, orientation, plan.grid)
    _header_footer(document, spec, plan, R, cs_font)

    props = document.core_properties
    props.title = spec.title
    props.subject = theme.core_property(spec.subtitle or spec.purpose)
    props.author = spec.author or "TechSara Local AI"
    props.created = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None)
    props.comments = "Generated by TechSara Local AI Artifact Studio"

    generic_cache: Dict[tuple, ST.TextStyle] = {}

    def generic(kind: str, level: Optional[int] = None) -> ST.TextStyle:
        key = (kind, level)
        if key not in generic_cache:
            generic_cache[key] = R.element(kind, generic_only=True, level=level)
        return generic_cache[key]

    if plan.cover:
        _cover(document, spec, plan, R, cs_font)
    else:
        tp = document.add_paragraph(spec.title, style="Title")
        _apply_overrides(tp, R.element("title"), generic("title"), R, cs_font)
        last = tp
        if spec.subtitle:
            sp = document.add_paragraph(spec.subtitle, style="Subtitle")
            _apply_overrides(sp, R.element("subtitle"), generic("subtitle"), R, cs_font)
            last = sp
        byline = " · ".join(b for b in (spec.author, spec.date, spec.audience and f"For {spec.audience}") if b)
        if byline:
            last = document.add_paragraph(byline, style="Caption")
        # The 2 pt rule under the title block (style guide §2).
        _para_border(last._p, "bottom", 16, R.tokens.primary, space=6)

    if plan.toc:
        document.add_paragraph("Contents", style="Heading 1")
        p = document.add_paragraph(style="Body")
        _field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    if plan.control_strip:
        rows = [("Document", spec.title)]
        for k, v in (("Owner", spec.author), ("Date", spec.date), ("Audience", spec.audience), ("Purpose", spec.purpose)):
            if v:
                rows.append((k, v))
        t = document.add_table(rows=0, cols=2)
        t.style = document.styles["TS Table"]
        for k, v in rows:
            cells = t.add_row().cells
            _shade(cells[0], R.tokens.band)
            cells[0].paragraphs[0].style = document.styles["Table Text"]
            cells[0].paragraphs[0].add_run(k.upper()).font.color.rgb = _rgb(R.tokens.muted)
            cells[1].paragraphs[0].style = document.styles["Table Text"]
            cells[1].paragraphs[0].add_run(v)
        document.add_paragraph(style="Body")

    if plan.hoisted_kpis is not None:
        _kpis(document, plan.hoisted_kpis, R, cs_font)

    heading_by_index = {h.index: h for h in plan.headings}
    list_number_abstract = _list_number_abstract_id(document)
    chart_ordinal = 0
    table_ordinal = 0
    heading_ordinal = 0
    level_ordinals = [0, 0, 0]
    sections: List[str] = ["", "", ""]  # the current heading text per level
    first_paragraph_done = False
    chart_width = 9.7 if orientation == "landscape" else 6.3
    for i, b in enumerate(plan.blocks):
        if isinstance(b, S.Heading):
            h = heading_by_index[i]
            heading_ordinal += 1
            level_ordinals[b.level - 1] += 1
            sections[b.level - 1] = b.text
            for deeper in range(b.level, 3):
                sections[deeper] = ""
            text = f"{h.number}  {b.text}" if h.number else b.text
            p = document.add_paragraph(text, style=f"Heading {b.level}")
            if R.has_specific_rules("heading"):
                _apply_overrides(p, R.element("heading", level=b.level, text=b.text, index=heading_ordinal, level_index=level_ordinals[b.level - 1]),
                                 generic("heading", b.level), R, cs_font)
        elif isinstance(b, S.Paragraph):
            style = "Lede" if i == plan.lede_index else "Body"
            p = document.add_paragraph(b.text + _cite_text(b.sources, index), style=style)
            if R.has_specific_rules("paragraph"):
                specific = R.element("paragraph", first=not first_paragraph_done, sections=[s for s in sections if s])
                _apply_overrides(p, specific, generic("paragraph"), R, cs_font)
            first_paragraph_done = True
        elif isinstance(b, S.Numbered):
            _restart_numbering(document, [document.add_paragraph(item, style="Number") for item in b.items], list_number_abstract)
            if b.sources:
                document.add_paragraph(_cite_text(b.sources, index).strip(), style="Caption")
        elif isinstance(b, S.Bullets):
            for item in b.items:
                document.add_paragraph(item, style="Bullet")
            if b.sources:
                document.add_paragraph(_cite_text(b.sources, index).strip(), style="Caption")
        elif isinstance(b, S.TableBlock):
            table_ordinal += 1
            _table(document, b.table, _cite_text(b.table.sources, index), R, index=table_ordinal, cs_font=cs_font)
        elif isinstance(b, S.ChartBlock):
            chart_ordinal += 1
            _chart(document, b.chart, str(chart_dir / chart_filename(chart_ordinal)), _cite_text(b.chart.sources, index), missing, chart_width)
        elif isinstance(b, S.Callout):
            _callout(document, b, R, cs_font)
        elif isinstance(b, S.KPIRow):
            _kpis(document, b, R, cs_font)
        elif isinstance(b, S.PageBreak):
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    if spec.sources:
        document.add_paragraph("References" if plan.template_id == "research_report" else "Sources", style="Heading 2")
        for n, c in enumerate(spec.sources, start=1):
            bits = [c.title]
            if c.url:
                bits.append(c.url)
            if c.retrieved_at:
                bits.append(f"retrieved {c.retrieved_at}")
            if c.note:
                bits.append(c.note)
            document.add_paragraph(f"[{n}] " + " — ".join(bits), style="Body")
    if spec.assumptions:
        document.add_paragraph("Assumptions", style="Heading 2")
        for a in spec.assumptions:
            document.add_paragraph(a, style="Bullet")

    if missing:
        warnings.append(f"{len(missing)} chart image{'s' if len(missing) != 1 else ''} could not be placed in the DOCX.")
    out = Path(out_path)
    document.save(str(out))
    return out


# ------------------------------------------------------ tabular document --


def _section_orientation(section, orientation: str, margin_mm: float = 14, R: Optional[ST.ResolvedStyle] = None) -> None:
    """Set a section's orientation AND swap its page size: python-docx's
    `orientation` is a flag Word reads only when the size agrees."""
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm

    size = R.page.size if R is not None else "A4"
    w, h = ST.PAGE_SIZES_MM.get(size, (210, 297))
    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(max(w, h)), Mm(min(w, h))
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Mm(min(w, h)), Mm(max(w, h))
    margins = R.page.margins_mm if R is not None else None
    if margins:
        section.top_margin, section.bottom_margin, section.left_margin, section.right_margin = (Mm(m) for m in margins)
    else:
        section.top_margin = section.bottom_margin = Mm(16)
        section.left_margin = section.right_margin = Mm(margin_mm)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)


def _repeat_header(row) -> None:
    _put(row._tr.get_or_add_trPr(), _el("w:tblHeader", **{"w:val": "true"}), _TRPR_ORDER)


def _keep_row_whole(row) -> None:
    """`w:cantSplit`: a row never breaks across pages (a wrapped comment
    stays with its id)."""
    _put(row._tr.get_or_add_trPr(), _el("w:cantSplit", **{"w:val": "true"}), _TRPR_ORDER)


def _fixed_layout(table) -> None:
    """`w:tblLayout fixed` so the column widths set below are honoured
    instead of Word's autofit, which lets a long comment column squeeze a
    date column to one character per line."""
    _put(table._tbl.tblPr, _el("w:tblLayout", **{"w:type": "fixed"}), _TBLPR_ORDER)


def _no_borders(table) -> None:
    """Every edge off, for `style.borders == "none"` (Table Grid draws all)."""
    borders = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(f"w:{edge}", **{"w:val": "nil"}))
    _put(table._tbl.tblPr, borders, _TBLPR_ORDER)


def _grid_borders(table, colour: str) -> None:
    """Thin grid lines in the grid colour (never black)."""
    borders = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(f"w:{edge}", **{"w:val": "single", "w:sz": "4", "w:space": "0", "w:color": _hex6(colour)}))
    _put(table._tbl.tblPr, borders, _TBLPR_ORDER)


def _tabular_styles(document, R: Optional[ST.ResolvedStyle] = None, cs_font: Optional[str] = None) -> None:
    """The two styles the tabular document adds to the theme's: 9.5 pt
    cells and a bold header at the same size."""
    from docx.shared import Pt

    from .html import TABULAR_CELL_PT

    R = R or ST.resolve(None, type_scale=theme.DOCUMENT_TYPE)
    _ensure_styles(document, theme.DOCUMENT_TYPE, R, cs_font=cs_font)
    for name, bold in (("Table Cell", False), ("Table Head", True)):
        try:
            st = document.styles[name]
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE

            st = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = document.styles["Table Text"]
        st.font.size = Pt(TABULAR_CELL_PT)
        st.font.bold = bold
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.line_spacing = 1.15


def _sheet_table(document, sheet: S.Sheet, section, R: Optional[ST.ResolvedStyle] = None, cs_font: Optional[str] = None) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Emu

    from .html import TABULAR_CELL_PT, TABULAR_HEADER_FILLS, TABULAR_HIGHLIGHT, column_shares, highlight_columns, sheet_totals

    R = R or ST.resolve(None)
    style = sheet.style if sheet.style is not None else S.SheetStyle()
    numeric = {i for i, c in enumerate(sheet.columns) if c.type != "text"}
    highlight = highlight_columns(sheet)
    shares = column_shares(sheet)
    usable = int(section.page_width - section.left_margin - section.right_margin)
    widths = [Emu(int(usable * sh)) for sh in shares]
    legacy_fill = sheet.style is not None and style.header_fill != "dark"
    fill, ink = TABULAR_HEADER_FILLS.get(style.header_fill, TABULAR_HEADER_FILLS["dark"])

    t = document.add_table(rows=1, cols=len(sheet.columns))
    t.style = document.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _fixed_layout(t)
    if style.borders != "thin":
        _no_borders(t)
    else:
        _grid_borders(t, R.tokens.hairline)
    head = t.rows[0]
    _repeat_header(head)

    def sized(ts: ST.TextStyle, base_kind: str) -> ST.TextStyle:
        if ts.size_pt == R.base(base_kind).size_pt:
            return ST.TextStyle.model_construct(**{**ts.model_dump(), "size_pt": TABULAR_CELL_PT})
        return ts

    for j, (cell, col) in enumerate(zip(head.cells, sheet.columns)):
        cell.width = widths[j]
        ts = sized(R.element("table_header", sheet=sheet.name, column=col.name, column_index=j), "table_header")
        if legacy_fill and not any(r.style.background for r in R.matching_rules("table_header", sheet=sheet.name)):
            ts = ST.TextStyle.model_construct(**{**ts.model_dump(), "background": fill or None, "color": ink})
        if sheet.style is not None and not style.header_bold:
            ts = ST.TextStyle.model_construct(**{**ts.model_dump(), "bold": False})
        colour = highlight.get(j)
        p = cell.paragraphs[0]
        p.style = document.styles["Table Head" if (ts.bold or ts.bold is None) else "Table Cell"]
        run = p.add_run(col.name)
        _style_run(run, ts, R, cs_font)
        if colour:
            h_fill, h_ink, _, _ = TABULAR_HIGHLIGHT[colour]
            _shade(cell, h_fill)
            run.font.color.rgb = _rgb(h_ink)
        elif ts.background:
            _shade(cell, ts.background)
        if style.borders == "thin":
            _cell_border(cell, bottom=(8, R.tokens.primary))
        p.alignment = _ALIGN[ts.align] if ts.align else (_ALIGN["right"] if j in numeric else None)
    col_styles = [sized(R.element("table_body", sheet=sheet.name, column=c.name, column_index=j), "table_body") for j, c in enumerate(sheet.columns)]
    row_rules = {r.target.index: r.style for r in R.rules if r.target.kind == "row" and r.target.index is not None}
    for r_index, row in enumerate(sheet.rows):
        new_row = t.add_row()
        _keep_row_whole(new_row)
        excel_row = r_index + 2
        for j, (cell, value) in enumerate(zip(new_row.cells, row)):
            cell.width = widths[j]
            ts = col_styles[j]
            if excel_row in row_rules:
                ts = row_rules[excel_row].over(ts)
            p = cell.paragraphs[0]
            p.style = document.styles["Table Cell"]
            run = p.add_run(theme.format_cell(value, sheet.columns[j]) if j in numeric else cell_text(value, False))
            _style_run(run, ts, R, cs_font)
            colour = highlight.get(j)
            if colour:
                _, _, c_fill, c_ink = TABULAR_HIGHLIGHT[colour]
                _shade(cell, c_fill)
                run.font.color.rgb = _rgb(c_ink)
            elif ts.background:
                _shade(cell, ts.background)
            elif R.banded and r_index % 2 == 1:
                _shade(cell, R.tokens.band)
            p.alignment = _ALIGN[ts.align] if ts.align else (_ALIGN["right"] if j in numeric else None)
    totals = sheet_totals(sheet)
    if totals:
        total_row = t.add_row()
        _keep_row_whole(total_row)
        tts = sized(R.element("table_total", sheet=sheet.name), "table_total")
        for j, cell in enumerate(total_row.cells):
            cell.width = widths[j]
            p = cell.paragraphs[0]
            p.style = document.styles["Table Cell"]
            _style_run(p.add_run(totals.get(j, "")), tts, R, cs_font)
            if tts.background:
                _shade(cell, tts.background)
            _cell_border(cell, top=(8, tts.color or R.tokens.primary))
            p.alignment = _ALIGN[tts.align] if tts.align else (_ALIGN["right"] if j in numeric else None)
    # Column widths live on the grid too, or Word ignores the cell widths
    # under a fixed layout.
    for j, width in enumerate(widths):
        t.columns[j].width = width


def render_workbook_docx(spec: S.WorkbookSpec, out_path: str | Path, *, warnings: Optional[List[str]] = None,
                         transform: Optional[dict] = None, resolved: Optional[ST.ResolvedStyle] = None) -> Path:
    """Write the workbook as a tabular Word document at `out_path`: one
    section per sheet with its own orientation, the whole table each."""
    from docx import Document
    from docx.enum.section import WD_SECTION

    from .html import methodology_note, sheet_orientation

    warnings = warnings if warnings is not None else []
    R = resolved or ST.resolve(spec, type_scale=theme.DOCUMENT_TYPE)
    cs_font = _cs_font_for(S.text_of(S.ArtifactSpec(kind="workbook", workbook=spec)))
    document = Document()
    _tabular_styles(document, R, cs_font)
    first = document.sections[0]

    def orient(sheet: S.Sheet) -> str:
        return R.page.orientation if R.page.orientation_explicit else sheet_orientation(sheet)

    _section_orientation(first, orient(spec.sheets[0]) if spec.sheets else "portrait", R=R)

    # Running header (the title) and footer (Page X of Y) on every section:
    # sections inherit the first one's header and footer unless unlinked.
    hp = first.header.paragraphs[0]
    hp.style = document.styles["Header"]
    run = hp.add_run(R.page.header_text or spec.title)
    run.bold = True
    _para_border(hp._p, "bottom", 4, R.tokens.hairline, space=3)
    fp = first.footer.paragraphs[0]
    fp.style = document.styles["Footer"]
    _right_tab(fp, first)
    if R.page.footer_text:
        fp.add_run(R.page.footer_text)
    if R.page.page_numbers:
        fp.add_run("\tPage ")
        _field(fp, "PAGE", "1")
        fp.add_run(" of ")
        _field(fp, "NUMPAGES", "1")

    props = document.core_properties
    props.title = spec.title
    props.subject = theme.core_property(spec.purpose)
    props.author = "TechSara Local AI"
    props.created = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None)
    props.comments = "Generated by TechSara Local AI Artifact Studio"

    document.add_paragraph(spec.title, style="Title")
    if spec.purpose:
        document.add_paragraph(spec.purpose, style="Subtitle")
    total_rows = sum(len(sh.rows) for sh in spec.sheets)
    document.add_paragraph(f"{len(spec.sheets)} sheet{'s' if len(spec.sheets) != 1 else ''} · {total_rows:,} rows", style="Caption")
    document.add_paragraph(methodology_note(spec, transform), style="Caption")

    section = first
    for k, sheet in enumerate(spec.sheets):
        if k > 0:
            section = document.add_section(WD_SECTION.NEW_PAGE)
            _section_orientation(section, orient(sheet), R=R)
        document.add_paragraph(sheet.name, style="Heading 2")
        if sheet.notes:
            document.add_paragraph(sheet.notes, style="Caption")
        _sheet_table(document, sheet, section, R, cs_font)
        document.add_paragraph(f"{len(sheet.rows):,} row{'s' if len(sheet.rows) != 1 else ''} · {len(sheet.columns)} columns", style="Caption")

    if spec.sources or spec.assumptions:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        _section_orientation(section, "portrait", R=R)
    if spec.sources:
        document.add_paragraph("Sources", style="Heading 2")
        for n, c in enumerate(spec.sources, start=1):
            bits = [c.title]
            if c.url:
                bits.append(c.url)
            if c.retrieved_at:
                bits.append(f"retrieved {c.retrieved_at}")
            if c.note:
                bits.append(c.note)
            document.add_paragraph(f"[{n}] " + " — ".join(bits), style="Body")
    if spec.assumptions:
        document.add_paragraph("Assumptions", style="Heading 2")
        for a in spec.assumptions:
            document.add_paragraph(a, style="Bullet")

    out = Path(out_path)
    document.save(str(out))
    return out


__all__ = ["render_docx", "render_workbook_docx"]
