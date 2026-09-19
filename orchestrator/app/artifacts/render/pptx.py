"""PresentationSpec → .pptx with python-pptx, one slide per planned slide.

THE DECK IS DRAWN, NOT TEMPLATED. python-pptx's default template has
placeholder layouts whose fonts, positions and colours we do not control and
whose theme is Office's. Every slide here uses the BLANK layout and our own
shapes, placed on the same geometry (theme.SLIDE_*_IN) the PDF preview uses,
so the .pptx a person opens is the deck they saw in the viewer. Charts are
NATIVE (pptx.chart.data.CategoryChartData): a person can click the chart in
PowerPoint and edit the numbers. Tables are native tables. Timelines are
shapes and connectors. Speaker notes go into the notes slide.

WHAT FITS. Body text never drops below 14 pt; instead of overflowing, the
plan (html.plan_deck) trims bullets over the template's cap and the slide's
line budget, cuts tables to the rows a slide holds, and records one warning
per slide. This writer trusts the plan and draws what it says.

STYLE. Fonts, colours and sizes come from style.resolve(spec), the same
ResolvedStyle the PDF preview reads (html.deck_html): the title band is the
template's colour unless a palette was chosen, slide titles and body text
take the resolved families and colours, a rule on one slide's title (by
slide number) or on a table column/row applies to that element only, table
headers take the resolved fill, charts take the chart title/axis/legend/
label styles, and a requested slide background fills every content slide.

python-pptx is imported lazily (tests/test_imports.py). Its native charts
embed an .xlsx workbook written with XlsxWriter — that is why XlsxWriter is
a transitive requirement; the workbook holds only the chart's numbers.
"""
from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from .. import spec as S
from .. import style as ST
from .. import chart_colours as CC
from . import theme
from .html import BULLET_GAP_EM, DeckPlan, PlannedSlide, cell_text, deck_band, plan_deck

#: The deck's ResolvedStyle while render_pptx runs (one render per process
#: call; set and cleared by render_pptx, read by the drawing helpers).
_R: Optional[ST.ResolvedStyle] = None


def _style() -> ST.ResolvedStyle:
    return _R if _R is not None else ST.resolve(None)


def _run_style(run, ts: ST.TextStyle, *, size: Optional[float] = None) -> None:
    from pptx.util import Pt

    R = _style()
    run.font.name = R.face(ts.font_family).office_name
    run.font.size = Pt(max(size if size is not None else (ts.size_pt or 14), 1))
    run.font.bold = bool(ts.bold)
    run.font.italic = bool(ts.italic)
    run.font.underline = bool(ts.underline)
    if ts.color:
        run.font.color.rgb = _rgb(ts.color)


def _size_or(ts: ST.TextStyle, kind: str, template_size: float) -> float:
    """The template's size unless a rule changed the element's size."""
    base = _style().base(kind)
    return float(ts.size_pt) if ts.size_pt is not None and ts.size_pt != base.size_pt else float(template_size)

# python-pptx's blank layout index in the default template.
_BLANK_LAYOUT = 6


def _rgb(colour: str):
    from pptx.dml.color import RGBColor

    return RGBColor(*theme.hex_to_rgb(colour))


def _text_box(slide, x, y, w, h, text: str, *, size: float, bold: bool = False, colour: str = theme.INK,
              align=None, anchor=None, font: str = theme.OFFICE_SANS, wrap: bool = True):
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    if anchor is not None:
        tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(max(size, 1))
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = _rgb(colour)
    return box


def _styled_box(slide, x, y, w, h, text: str, ts: ST.TextStyle, *, size: float, align=None, anchor=None, wrap: bool = True):
    """A text box whose one run carries every property of `ts`."""
    from pptx.enum.text import PP_ALIGN

    box = _text_box(slide, x, y, w, h, text, size=size, align=align, anchor=anchor, wrap=wrap)
    p = box.text_frame.paragraphs[0]
    _run_style(p.runs[0], ts, size=size)
    if ts.align:
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}[ts.align]
    if ts.background:
        box.fill.solid()
        box.fill.fore_color.rgb = _rgb(ts.background)
    return box


def _rect(slide, x, y, w, h, fill: str, *, line: Optional[str] = None):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = _rgb(line)
    shape.shadow.inherit = False
    # No text in a plain rectangle: keep the frame empty so nothing can be
    # read from it by an accessibility tool.
    return shape


def _bullets(slide, x, y, w, h, items: Sequence[str], size: float, *, numbered: bool = False, colour: Optional[str] = None):
    from pptx.util import Inches, Pt

    R = _style()
    bullet_ts = R.element("slide_body")
    for rule in R.matching_rules("bullet"):
        bullet_ts = rule.style.over(bullet_ts)
    size = max(_size_or(bullet_ts, "slide_body", size), theme.SLIDE_MIN_BODY_PT)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * BULLET_GAP_EM)  # the gap the plan's line budget assumes
        marker = f"{i + 1}.  " if numbered else "•  "
        run = p.add_run()
        run.text = marker
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.name = R.face(bullet_ts.font_family).office_name
        run.font.color.rgb = _rgb(R.tokens.accent_text if numbered else R.tokens.accent)
        run = p.add_run()
        run.text = item
        _run_style(run, bullet_ts if colour is None else bullet_ts.model_copy(update={"color": colour}), size=size)
    return box


def _slide_title(slide, s: PlannedSlide, plan: DeckPlan) -> None:
    from pptx.enum.text import MSO_ANCHOR

    t = plan.type
    m = theme.SLIDE_MARGIN_IN
    R = _style()
    ts = R.element("slide_title", index=s.number)
    _styled_box(slide, m, theme.SLIDE_TITLE_TOP_IN, theme.SLIDE_W_IN - 2 * m, theme.SLIDE_TITLE_H_IN, s.title, ts,
                size=_size_or(ts, "slide_title", t.title), anchor=MSO_ANCHOR.BOTTOM)
    _rect(slide, m, theme.SLIDE_TITLE_TOP_IN + theme.SLIDE_TITLE_H_IN + 0.02, theme.SLIDE_W_IN - 2 * m, 0.03, R.tokens.accent)


def _footer(slide, s: PlannedSlide, plan: DeckPlan, total: int) -> None:

    t = plan.type
    m = theme.SLIDE_MARGIN_IN
    y = theme.SLIDE_FOOTER_TOP_IN
    w = theme.SLIDE_W_IN - 2 * m
    R = _style()
    ts = R.element("header_footer")
    size = _size_or(ts, "header_footer", t.small)
    _rect(slide, m, y, w, 0.01, R.tokens.hairline)
    label = (R.page.footer_text or plan.spec.title) + ("   CONFIDENTIAL" if plan.spec.confidential else "")
    _styled_box(slide, m, y + 0.03, w * 0.75, theme.SLIDE_FOOTER_H_IN, label, ts, size=size)
    if R.page.page_numbers:
        _styled_box(slide, m + w * 0.75, y + 0.03, w * 0.25, theme.SLIDE_FOOTER_H_IN, f"{s.number} / {total}", ts.model_copy(update={"align": "right"}), size=size)


def _band_slide(prs, s: PlannedSlide, plan: DeckPlan) -> None:
    """title / section / closing: a full-bleed band in the template colour."""
    spec = plan.spec
    t = plan.type
    R = _style()
    band = deck_band(plan, R)
    on_band = ST.readable_on(band)
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(band)
    title = s.title or (spec.title if s.layout == "title" else ("Thank you" if s.layout == "closing" else ""))
    subtitle = s.subtitle or (spec.subtitle if s.layout == "title" else "")
    _rect(slide, 0.9, 2.0, 1.2, 0.08, R.tokens.accent)

    def on(kind: str, template_size: float) -> Tuple[ST.TextStyle, float]:
        ts = R.element(kind)
        user_colour = any(r.style.color for r in R.matching_rules(kind))
        ts = ts.model_copy(update={"color": ts.color if user_colour else on_band, "background": None})
        return ts, _size_or(ts, kind, template_size)

    if s.layout == "section":
        ts, size = on("title", t.section)
        _styled_box(slide, 0.9, 2.9, theme.SLIDE_W_IN - 1.8, 1.4, title, ts, size=size)
        if subtitle:
            ts, size = on("subtitle", t.body + 4)
            _styled_box(slide, 0.9, 4.4, theme.SLIDE_W_IN - 1.8, 1.0, subtitle, ts, size=size)
    else:
        ts, size = on("title", t.cover_title)
        _styled_box(slide, 0.9, 2.3, theme.SLIDE_W_IN - 1.8, 1.5, title, ts, size=size)
        if subtitle:
            ts, size = on("subtitle", t.body + 4)
            _styled_box(slide, 0.9, 3.9, theme.SLIDE_W_IN - 1.8, 1.0, subtitle, ts, size=size)
        meta_bits: List[str] = []
        if s.layout == "title":
            meta_bits = [b for b in (spec.author, spec.date, spec.audience and f"Prepared for {spec.audience}") if b]
        elif s.bullets:
            meta_bits = list(s.bullets)
        if meta_bits:
            _text_box(slide, 0.9, 6.4, theme.SLIDE_W_IN - 1.8, 0.6, " · ".join(meta_bits), size=t.body, colour=on_band,
                      font=R.body_face.office_name)
    _notes(slide, s)


def _notes(slide, s: PlannedSlide) -> None:
    if s.notes:
        slide.notes_slide.notes_text_frame.text = s.notes


def _chart_type(chart: S.Chart):
    from pptx.enum.chart import XL_CHART_TYPE

    return {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "horizontal_bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }[chart.type]


_FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")


def _label(text: str) -> str:
    """A chart label as TEXT for the embedded workbook: leading formula
    characters are stripped (a label "=Total" reads "Total"), control
    characters removed, and the length bounded."""
    out = str(text or "").replace("\t", " ").replace("\r", " ").strip()
    while out and out[0] in _FORMULA_LEADS:
        out = out[1:].lstrip()
    return (out or "-")[:120]


def _legacy_colour(scheme, index: int, name: str, *, category: bool = False) -> str:
    """The scheme's colour for a legacy (unstyled) deck chart, falling back
    to the legacy painter's own palette when no rule claims the mark."""
    got = scheme.category_colour(index, name) if category else scheme.series_colour(index, name)
    return got or theme.series_colour(index)


def _native_chart(slide, x, y, w, h, chart: S.Chart, body_pt: float) -> None:
    # --- AS3 integration: Chart v2 types and requested chart styling go to
    # the charts track's writer (a picture where PowerPoint has no such chart).
    if (str(chart.type) not in ("bar", "horizontal_bar", "line", "pie") or getattr(chart, "style", None) is not None) and chart.series:
        from . import chart_native as CN

        CN.add_pptx_chart(slide, chart, (x, y, w, h), _style())
        return
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    # Legacy deck charts used theme.PALETTE (teal first) while documents used
    # the artifact palette (blue first), so the same chart was two different
    # colours in two files. They read the same scheme now.
    scheme = CC.scheme_for(chart, plan=getattr(_style(), "chart_plan", None))
    data = CategoryChartData()
    # python-pptx writes the chart's data into an EMBEDDED workbook with
    # XlsxWriter, whose write() turns any string starting with '=' into a
    # formula cell — so a category or series named "=HYPERLINK(...)" would
    # be live the moment a reader chose Edit Data (review, 2026-09-11). A
    # label is text: the formula lead is dropped, never written.
    data.categories = [_label(c) for c in chart.categories]
    for series in chart.series:
        data.add_series(_label(series.name), [float(v) for v in series.values])
    frame = slide.shapes.add_chart(_chart_type(chart), Inches(x), Inches(y), Inches(w), Inches(h), data)
    c = frame.chart
    R = _style()
    axis_ts = R.element("chart_axis")
    c.font.size = Pt(_size_or(axis_ts, "chart_axis", max(body_pt - 4, theme.SLIDE_MIN_BODY_PT - 2)))
    c.font.name = R.face(axis_ts.font_family).office_name
    c.font.bold = bool(axis_ts.bold)
    c.font.italic = bool(axis_ts.italic)
    c.font.color.rgb = _rgb(axis_ts.color or R.tokens.muted)
    c.has_title = bool(chart.title)
    if chart.title:
        title_ts = R.element("chart_title")
        c.chart_title.text_frame.text = chart.title
        para = c.chart_title.text_frame.paragraphs[0]
        _run_style(para.runs[0], title_ts, size=_size_or(title_ts, "chart_title", body_pt))
    c.has_legend = len(chart.series) > 1 or chart.type == "pie"
    if c.has_legend:
        c.legend.position = XL_LEGEND_POSITION.BOTTOM
        c.legend.include_in_layout = False
        legend_ts = R.element("chart_legend")
        c.legend.font.name = R.face(legend_ts.font_family).office_name
        c.legend.font.size = Pt(_size_or(legend_ts, "chart_legend", max(body_pt - 4, theme.SLIDE_MIN_BODY_PT - 2)))
        c.legend.font.bold = bool(legend_ts.bold)
        c.legend.font.italic = bool(legend_ts.italic)
        if legend_ts.color:
            c.legend.font.color.rgb = _rgb(legend_ts.color)
    plot = c.plots[0]
    if chart.type == "pie":
        plot.has_data_labels = True
        plot.data_labels.number_format = "0.0%"
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False
        points = plot.series[0].points
        for i in range(len(chart.categories)):
            points[i].format.fill.solid()
            points[i].format.fill.fore_color.rgb = _rgb(_legacy_colour(scheme, i, str(chart.categories[i]), category=True))
    else:
        if len(chart.categories) <= theme.CHART_LABEL_MAX_CATEGORIES and chart.type != "line":
            plot.has_data_labels = True
            labels_ts = R.element("chart_labels")
            plot.data_labels.font.size = Pt(_size_or(labels_ts, "chart_labels", max(body_pt - 6, 10)))
            plot.data_labels.font.name = R.face(labels_ts.font_family).office_name
            plot.data_labels.font.bold = bool(labels_ts.bold)
            plot.data_labels.font.italic = bool(labels_ts.italic)
            plot.data_labels.font.color.rgb = _rgb(labels_ts.color or R.tokens.ink)
        for i, series in enumerate(plot.series):
            name = chart.series[i].name if i < len(chart.series) else ""
            colour = _rgb(_legacy_colour(scheme, i, name))
            fill = series.format.line if chart.type == "line" else series.format.fill
            if chart.type == "line":
                fill.color.rgb = colour
                fill.width = Pt(2.25)
                series.smooth = False
            else:
                fill.solid()
                fill.fore_color.rgb = colour
                if len(chart.series) == 1 and scheme.by_category:
                    for j, cat in enumerate(chart.categories):
                        series.points[j].format.fill.solid()
                        series.points[j].format.fill.fore_color.rgb = _rgb(_legacy_colour(scheme, j, str(cat), category=True))
        if chart.type != "pie":
            va = c.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = _rgb(R.tokens.grid)
            va.format.line.fill.background()
            if chart.y_label:
                va.has_title = True
                va.axis_title.text_frame.text = chart.y_label
                va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(max(body_pt - 6, 10))
            c.category_axis.format.line.color.rgb = _rgb(R.tokens.hairline)


def _native_table(slide, x, y, w, h, table: S.Table, body_pt: float, ordinal: int = 0) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    R = _style()

    numeric = set(table.numeric_columns)
    rows = len(table.rows) + 1
    cols = len(table.columns)
    # Rows past what fits at the minimum body size are not drawn smaller —
    # the plan caps slide tables to what a slide can hold (html.fit_table).
    row_h = min(0.45, h / rows)
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(row_h * rows))
    tbl = shape.table
    size = Pt(max(body_pt - 2, theme.SLIDE_MIN_BODY_PT))

    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}

    def put(cell, text: str, ts: ST.TextStyle, kind: str, *, right: bool = False):
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = text
        _run_style(run, ts, size=_size_or(ts, kind, size.pt))
        if ts.align:
            p.alignment = aligns[ts.align]
        elif right:
            p.alignment = PP_ALIGN.RIGHT
        cell.margin_left = cell.margin_right = Inches(0.08)

    for j, name in enumerate(table.columns):
        cell = tbl.cell(0, j)
        ts = R.element("table_header", table=ordinal, column=name, column_index=j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(ts.background or R.tokens.header_fill)
        put(cell, name, ts, "table_header", right=j in numeric)
    for i, row in enumerate(table.rows, start=1):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            ts = R.element("table_body", table=ordinal, column=table.columns[j], column_index=j, row=i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(ts.background or (R.tokens.band if (R.banded and i % 2 == 0) else theme.WHITE))
            put(cell, cell_text(value, j in numeric), ts, "table_body", right=j in numeric)


def _kpi_boxes(slide, s: PlannedSlide, plan: DeckPlan) -> None:
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    t = plan.type
    m = theme.SLIDE_MARGIN_IN
    n = max(1, len(s.kpis))
    gap = 0.3
    w = (theme.SLIDE_W_IN - 2 * m - gap * (n - 1)) / n
    y = theme.SLIDE_BODY_TOP_IN + 0.7
    h = 2.4
    R = _style()
    value_ts, label_ts = R.element("kpi_value"), R.element("kpi_label")
    for i, k in enumerate(s.kpis):
        x = m + i * (w + gap)
        _rect(slide, x, y, w, h, value_ts.background or R.tokens.band)
        _rect(slide, x, y, w, 0.07, R.tokens.kpi_bar)
        _styled_box(slide, x, y + 0.35, w, 1.0, k.value, value_ts.model_copy(update={"background": None}), size=_size_or(value_ts, "kpi_value", t.kpi_value),
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _styled_box(slide, x, y + 1.45, w, 0.45, k.label.upper(), label_ts.model_copy(update={"background": None}), size=_size_or(label_ts, "kpi_label", t.kpi_label), align=PP_ALIGN.CENTER)
        if k.note:
            _text_box(slide, x, y + 1.9, w, 0.4, k.note, size=max(t.kpi_label - 2, theme.SLIDE_MIN_BODY_PT), colour=theme.INK_FAINT, align=PP_ALIGN.CENTER)


def _timeline(slide, s: PlannedSlide, plan: DeckPlan) -> None:
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    t = plan.type
    m = theme.SLIDE_MARGIN_IN
    n = max(1, len(s.steps))
    left = m
    width = theme.SLIDE_W_IN - 2 * m
    line_y = theme.SLIDE_BODY_TOP_IN + 1.6
    step_w = width / n
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(left), Inches(line_y), Inches(left + width), Inches(line_y))
    connector.line.color.rgb = _rgb(theme.BORDER)
    connector.line.width = Inches(0.06)
    for i, (label, text) in enumerate(s.steps):
        cx = left + step_w * i + step_w / 2
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - 0.16), Inches(line_y - 0.16), Inches(0.32), Inches(0.32))
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb(theme.ACCENT)
        dot.line.color.rgb = _rgb(theme.WHITE)
        dot.shadow.inherit = False
        _text_box(slide, left + step_w * i, line_y - 1.0, step_w, 0.6, label, size=t.body, bold=True, colour=theme.NAVY, align=PP_ALIGN.CENTER)
        _text_box(slide, left + step_w * i + 0.05, line_y + 0.35, step_w - 0.1, 2.0, text, size=max(t.body - 2, theme.SLIDE_MIN_BODY_PT),
                  colour=theme.INK_MUTED, align=PP_ALIGN.CENTER)


def _columns(slide, s: PlannedSlide, plan: DeckPlan, *, cards: bool) -> None:
    t = plan.type
    m = theme.SLIDE_MARGIN_IN
    gap = 0.4
    w = (theme.SLIDE_W_IN - 2 * m - gap) / 2
    y = theme.SLIDE_BODY_TOP_IN
    h = theme.SLIDE_BODY_H_IN
    for i, (title, items, default) in enumerate(((s.left_title, s.left, "Option A"), (s.right_title, s.right, "Option B"))):
        x = m + i * (w + gap)
        if cards:
            _rect(slide, x, y, w, h - 0.3, theme.SURFACE)
            _rect(slide, x, y, w, 0.06, theme.ACCENT if i == 0 else theme.TEAL)
            title = title or default
        inner_x = x + (0.25 if cards else 0)
        inner_w = w - (0.5 if cards else 0)
        top = y + (0.25 if cards else 0)
        if title:
            _text_box(slide, inner_x, top, inner_w, 0.5, title, size=t.body + 2, bold=True, colour=theme.BOARDROOM)
            if not cards:
                _rect(slide, inner_x, top + 0.5, inner_w, 0.015, theme.BORDER)
            top += 0.65
        _bullets(slide, inner_x, top, inner_w, h - (top - y) - 0.3, items, t.body)


def _sources_slide(prs, s: PlannedSlide, plan: DeckPlan, total: int) -> None:
    from pptx.util import Inches, Pt

    t = plan.type
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
    _background(slide)
    _slide_title(slide, s, plan)
    m = theme.SLIDE_MARGIN_IN
    box = slide.shapes.add_textbox(Inches(m), Inches(theme.SLIDE_BODY_TOP_IN), Inches(theme.SLIDE_W_IN - 2 * m), Inches(theme.SLIDE_BODY_H_IN))
    tf = box.text_frame
    tf.word_wrap = True
    size = Pt(max(t.body - 3, theme.SLIDE_MIN_BODY_PT))
    for i, c in enumerate(plan.spec.sources):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = f"{i + 1}.  {c.title}"
        run.font.size = size
        run.font.name = theme.OFFICE_SANS
        run.font.color.rgb = _rgb(theme.INK)
        if c.url:
            r2 = p.add_run()
            r2.text = f"  —  {c.url}"
            r2.font.size = size
            r2.font.name = theme.OFFICE_SANS
            r2.font.color.rgb = _rgb(theme.ACCENT)
    _footer(slide, s, plan, total)


def _background(slide) -> None:
    R = _style()
    if R.page.background:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(R.page.background)


def _content_slide(prs, s: PlannedSlide, plan: DeckPlan, total: int) -> None:
    t = plan.type
    m = theme.SLIDE_MARGIN_IN
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
    _background(slide)
    _slide_title(slide, s, plan)
    x, y, w, h = m, theme.SLIDE_BODY_TOP_IN, theme.SLIDE_W_IN - 2 * m, theme.SLIDE_BODY_H_IN
    numbered = plan.template_id == "training"
    if s.layout == "chart" and s.chart is not None:
        if s.bullets:
            _bullets(slide, x, y, 4.0, h, s.bullets, t.body, numbered=numbered)
            _native_chart(slide, x + 4.35, y, w - 4.35, h - 0.2, s.chart, t.body)
        else:
            _native_chart(slide, x + 0.5, y, w - 1.0, h - 0.2, s.chart, t.body)
        if s.chart.caption:
            cap = _style().element("caption")
            _styled_box(slide, x, y + h - 0.35, w, 0.35, s.chart.caption, cap, size=_size_or(cap, "caption", theme.SLIDE_MIN_BODY_PT))
    elif s.layout == "table" and s.table is not None:
        _native_table(slide, x, y, w, h - 0.4, s.table, t.body, s.number)
        if s.table.caption:
            cap = _style().element("caption")
            _styled_box(slide, x, y + h - 0.35, w, 0.35, s.table.caption, cap, size=_size_or(cap, "caption", theme.SLIDE_MIN_BODY_PT))
    elif s.layout == "kpis":
        _kpi_boxes(slide, s, plan)
    elif s.layout == "timeline":
        _timeline(slide, s, plan)
    elif s.layout == "two_column":
        _columns(slide, s, plan, cards=False)
    elif s.layout == "comparison":
        _columns(slide, s, plan, cards=True)
    else:
        if s.bullets:
            _bullets(slide, x, y, w, h, s.bullets, t.body, numbered=numbered)
    _footer(slide, s, plan, total)
    _notes(slide, s)


def render_pptx(spec: S.PresentationSpec, out_path: str | Path, *, plan: Optional[DeckPlan] = None,
                resolved: Optional[ST.ResolvedStyle] = None) -> Path:
    """Write the deck to `out_path` and return it. Every layout decision —
    bullets trimmed, table rows cut — is the plan's, with its warnings; this
    writer makes none of its own, so the .pptx cannot differ from the preview."""

    global _R
    plan = plan or plan_deck(spec)
    _R = resolved or ST.resolve(spec)
    try:
        return _render(spec, out_path, plan)
    finally:
        _R = None


def _render(spec: S.PresentationSpec, out_path: str | Path, plan: DeckPlan) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(theme.SLIDE_W_IN)
    prs.slide_height = Inches(theme.SLIDE_H_IN)
    total = len(plan.slides)
    for s in plan.slides:
        if s.layout in ("title", "section", "closing"):
            _band_slide(prs, s, plan)
        elif s.is_sources_slide:
            _sources_slide(prs, s, plan, total)
        else:
            _content_slide(prs, s, plan, total)
    props = prs.core_properties
    props.title = spec.title
    props.subject = theme.core_property(spec.subtitle or spec.purpose)
    props.author = spec.author or "TechSara Local AI"
    props.comments = "Generated by TechSara Local AI Artifact Studio"
    out = Path(out_path)
    prs.save(str(out))
    return out


__all__ = ["render_pptx"]
