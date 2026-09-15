"""Page images and workbook grids for the in-app viewer — the `preview` stage.

Pages come from `preview.pdf` through pypdfium2, the same rasteriser
core/pdf.py uses for uploaded PDFs (RENDER_SCALE there is a fixed 2.0; here
the scale is derived from the WIDTH the API asked for — 240 px for a card
thumbnail, 1400 px for a page — so a thumbnail is not a downscaled 1400 px
render). PNG bytes are returned; the API caches them under `previews/`.

The workbook preview is a GRID, not an image: `sheet_grid()` opens the .xlsx
read-only with `data_only=False`, so a formula cell is returned as its TEXT
(`=SUM(B2:B41)`) and is never evaluated by this code — the browser shows the
formula, exactly as the API contract says. Hidden sheets are excluded; bounds
are enforced so a 10,000-row sheet does not become a 10,000-row JSON.

STYLES. `grid_for()` and `sheet_grid()` also return what a cell LOOKS like
— `header_styles`, `cell_styles` ({"r:c": {fill, color, bold, italic,
underline}}) and `display` (values formatted by their number format, totals
computed) — from the cell styles and the conditional formats the XLSX
writer produced, evaluated here in Python (see styled_window). A workbook
over 8 MB is served without them.

`grid_for()` is the one shape the `GET …/grid` route serves for an xlsx AND
a csv (CONTRACT-2 §11): `sheets`, `sheet`, `columns`, `rows`, `total_rows`,
`total_columns`, `truncated`, `formulas_as_text: true` — an xlsx through
`sheet_grid` (trimmed to the sheet's real last column, paged by offset), a
csv through render/csv.read_csv_grid (one pass, one page in memory).
"""
from __future__ import annotations

import datetime as _dt
import functools
import io
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .. import types as T


def _lock():
    """The ONE lock every pypdfium2 call in the process takes — owned by
    core/pdf.py, whose uploaded-PDF readers run in the same worker threads.
    PDFium is not thread-safe; two threads inside it crash the process
    (reproduced 2026-09-11: SIGSEGV at 80 concurrent page renders, SIGABRT
    at 16). The worker subprocess has its own process and its own lock."""
    from ...core.pdf import PDFIUM_LOCK

    return PDFIUM_LOCK


def page_count(pdf_path: str | Path) -> int:
    import pypdfium2 as pdfium  # lazy: arm64 wheel, no system deps

    with _lock():
        pdf = pdfium.PdfDocument(str(pdf_path))
        try:
            return len(pdf)
        finally:
            pdf.close()


def rasterise_page(pdf_path: str | Path, page: int, width: int) -> bytes:
    """PNG bytes of 1-based `page`, `width` pixels wide (height follows the
    page's aspect). Raises IndexError past the last page, ValueError for a
    width outside PREVIEW_WIDTHS' range."""
    import pypdfium2 as pdfium  # lazy

    if width < 32 or width > max(T.PREVIEW_WIDTHS) * 2:
        raise ValueError(f"preview width {width} is out of range")
    with _lock():
        pdf = pdfium.PdfDocument(str(pdf_path))
        try:
            n = len(pdf)
            if page < 1 or page > n:
                raise IndexError(f"page {page} of {n}")
            pg = pdf[page - 1]
            try:
                page_w, _ = pg.get_size()  # points
                scale = width / page_w if page_w > 0 else 1.0
                bitmap = pg.render(scale=scale)
                pil = bitmap.to_pil().convert("RGB")
            finally:
                pg.close()
        finally:
            pdf.close()
    if pil.width != width:  # rounding: make the promised width exact
        from PIL import Image

        pil = pil.resize((width, max(1, round(pil.height * width / pil.width))), Image.LANCZOS)
    buf = io.BytesIO()
    pil.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
            return str(value)
        return value
    if isinstance(value, _dt.datetime):
        # openpyxl reads a date cell back as a datetime at midnight; the
        # grid shows the date a person typed, not "2026-08-03T00:00:00".
        if value.hour == value.minute == value.second == value.microsecond == 0 and value.tzinfo is None:
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    return str(value)


def sheet_grid(xlsx_path: str | Path, sheet: Optional[str], max_rows: int, max_cols: int, *, style_offset: int = 0) -> Dict[str, Any]:
    """The viewer's grid: `{sheets: [{name, rows, cols}], sheet: {name,
    columns, rows, truncated, formulas}}`. Formulas as text, never
    evaluated; hidden sheets excluded; `sheet` None means the first."""
    from openpyxl import load_workbook  # lazy

    max_rows = max(1, min(int(max_rows), T.MAX_ROWS_PER_SHEET))
    max_cols = max(1, min(int(max_cols), T.MAX_COLUMNS_PER_SHEET))
    wb = load_workbook(str(xlsx_path), read_only=True, data_only=False)
    try:
        visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
        listing: List[Dict[str, Any]] = [
            {"name": ws.title, "rows": int(ws.max_row or 0), "cols": int(ws.max_column or 0)} for ws in visible
        ]
        if not visible:
            return {"sheets": [], "sheet": None}
        chosen = visible[0]
        if sheet:
            for ws in visible:
                if ws.title == sheet:
                    chosen = ws
                    break
            else:
                raise KeyError(sheet)
        # Trimmed to the sheet's real last column: iter_rows pads every row
        # to max_col, and a three-column sheet used to come back as sixty
        # columns, fifty-seven of them blank.
        real_cols = int(chosen.max_column or 0)
        width = max(1, min(max_cols, real_cols)) if real_cols else max_cols
        columns: List[str] = []
        rows: List[List[Any]] = []
        formulas: Dict[str, str] = {}
        truncated = False
        for r_index, row in enumerate(chosen.iter_rows(min_row=1, max_col=width, values_only=False), start=1):
            if r_index == 1:
                columns = ["" if c.value is None else str(c.value) for c in row]
                continue
            if len(rows) >= max_rows:
                truncated = True
                break
            out_row: List[Any] = []
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and c.data_type == "f":
                    formulas[c.coordinate] = v
                    out_row.append(v)
                else:
                    out_row.append(_json_value(v))
            rows.append(out_row)
        cols_truncated = int(chosen.max_column or 0) > max_cols
        chosen_title = chosen.title
        out = {
            "sheets": listing,
            "sheet": {
                "name": chosen_title,
                "columns": columns,
                "rows": rows,
                "truncated": bool(truncated or cols_truncated),
                "formulas": formulas,
            },
        }
    finally:
        wb.close()
    try:
        # Only the rows the caller shows are evaluated (grid_for pages from
        # `style_offset`): the CF evaluation is per cell, and a last page of
        # a 5,000-row sheet must not evaluate the 4,500 rows before it.
        style_offset = max(0, min(int(style_offset), len(rows)))
        styled = (styled_grid(xlsx_path, chosen_title, offset=style_offset, limit=len(rows) - style_offset, width=max(1, len(columns)))
                  if len(rows) > style_offset else {"header_styles": [], "cell_styles": {}, "display": []})
    except Exception:  # noqa: BLE001 — styles are a courtesy; the grid still answers
        styled = None
    if styled is not None:
        out["sheet"].update(styled)
    return out


def grid_for(path: str | Path, fmt: str, *, sheet: Optional[str] = "", offset: int = 0, limit: int = 500,
             max_cols: int = 60, title: str = "") -> Dict[str, Any]:
    """CONTRACT-2 §11: the grid dict for one file. `sheets` lists the
    workbook's visible sheet names (for a csv, the one name the file has:
    `title`, else its stem); `sheet` is the one shown; `rows` is the page
    from `offset` (0-based data rows) of at most `limit`; `total_rows` and
    `total_columns` are the whole sheet's; `truncated` says a row or a
    column lies outside the page. Formulas are text, never evaluated.
    Raises KeyError for a sheet the workbook does not have."""
    offset = max(0, int(offset))
    limit = max(1, min(int(limit), T.MAX_ROWS_PER_SHEET))
    max_cols = max(1, min(int(max_cols), T.MAX_COLUMNS_PER_SHEET))
    if fmt == "csv":
        from .csv import read_csv_grid  # lazy: keeps this module's imports light

        page = read_csv_grid(path, offset=offset, limit=limit, max_cols=max_cols)
        name = (title or "").strip() or Path(path).stem
        return {
            "sheets": [name], "sheet": name, "columns": list(page["columns"]), "rows": list(page["rows"]),
            "total_rows": int(page["total_rows"]), "total_columns": int(page["total_columns"]),
            "truncated": bool(page["truncated"]), "formulas_as_text": True,
        }
    if fmt != "xlsx":
        raise ValueError(f"no grid for {fmt}")
    # sheet_grid reads from the top: the page is sliced out of offset+limit
    # rows, which is one read and at most the page plus its offset in
    # memory (a 10,000-row sheet paged at 500 costs 10,000 rows once, on
    # the last page — acceptable for a viewer that pages forward).
    raw = sheet_grid(path, sheet or None, max_rows=offset + limit, max_cols=max_cols, style_offset=offset)
    styled_page: Dict[str, Any] = {}
    chosen_raw = raw.get("sheet") or {}
    if chosen_raw.get("display") is not None:
        # sheet_grid evaluated the styles from `offset` on: display and the
        # cell_styles keys are already relative to the page.
        styled_page = {
            "header_styles": chosen_raw.get("header_styles") or [],
            "display": list(chosen_raw.get("display") or [])[:limit],
            "cell_styles": {k: v for k, v in (chosen_raw.get("cell_styles") or {}).items() if int(k.split(":")[0]) < limit},
        }
    chosen = raw.get("sheet") or {}
    listing = raw.get("sheets") or []
    names = [str(s.get("name")) for s in listing]
    all_rows = list(chosen.get("rows") or [])
    page_rows = all_rows[offset: offset + limit]
    facts = next((s for s in listing if str(s.get("name")) == str(chosen.get("name"))), {})
    total_rows = max(0, int(facts.get("rows") or 0) - 1) if facts else len(all_rows)
    total_columns = int(facts.get("cols") or len(chosen.get("columns") or []))
    return {
        "sheets": names, "sheet": str(chosen.get("name") or ""), "columns": list(chosen.get("columns") or []),
        "rows": page_rows, "total_rows": total_rows, "total_columns": total_columns,
        "truncated": total_rows > offset + len(page_rows) or total_columns > max_cols,
        "formulas_as_text": True,
        **styled_page,
    }




# ---------------------------------------------------------- styled grid --
#
# WHY IN PYTHON. openpyxl reads conditional formats but never evaluates them,
# and a browser cannot either; yet the colours a person asked for ("Status
# red for Blocked", banding, a yellow row) live in those rules, because Excel
# paints a CF fill above a cell fill (render/xlsx.sheet_cf_plan). So the grid
# evaluates the rules the XLSX writer produces — expressions over the row's
# own cells (TRIM/OR/AND/NOT/ISNUMBER/SEARCH/LEN/MOD/ROW/TODAY, comparisons,
# string and number literals) and colour scales — in priority order with
# stopIfTrue, exactly as Excel resolves them, and returns a per-cell style.
# A rule outside that grammar is skipped (no colour), never guessed. Totals
# written as SUBTOTAL/SUM/AVERAGE/COUNTA/MIN/MAX over a same-column range are
# computed for display; the formula text is still what `rows` carries.

_GRID_STYLE_MAX_BYTES = 8 * 1024 * 1024
_HEX6 = re.compile(r"^[0-9A-Fa-f]{6}$")


def _rgb_of(color: Any) -> Optional[str]:
    """'#RRGGBB' from an openpyxl Color with an rgb value, else None (theme
    and indexed colours are not ours and are not guessed)."""
    try:
        rgb = getattr(color, "rgb", None)
    except Exception:  # noqa: BLE001 — openpyxl raises on some descriptor states
        return None
    if isinstance(rgb, str) and len(rgb) in (6, 8):
        tail = rgb[-6:]
        if _HEX6.match(tail):
            return "#" + tail.upper()
    return None


class _FormulaError(ValueError):
    pass


_TOKEN_RE = re.compile(
    r'\s*(?:(?P<str>"(?:[^"]|"")*")|(?P<num>\d+(?:\.\d+)?)|(?P<range>\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)|'
    r"(?P<ref>\$?[A-Z]{1,3}\$?\d+)|(?P<func>[A-Z][A-Z0-9.]*)\s*\(|(?P<bool>TRUE|FALSE)\b|(?P<op><>|>=|<=|[=<>+\-*/(),&]))",
)


def _tokens(text: str) -> Tuple[Tuple[str, str], ...]:
    """The formula's tokens, cached: a CF formula is the same text for every
    cell of its range, so a 500-row page tokenises it once, not per cell."""
    return _tokens_cached(text)


@functools.lru_cache(maxsize=512)
def _tokens_cached(text: str) -> Tuple[Tuple[str, str], ...]:
    if len(text) > 4096:
        raise _FormulaError("formula too long")
    out: List[Tuple[str, str]] = []
    pos = 0
    text = text.strip()
    if text.startswith("="):
        text = text[1:]
    while pos < len(text):
        m = _TOKEN_RE.match(text, pos)
        if not m or m.end() == pos:
            raise _FormulaError(f"unreadable at {pos}")
        kind = m.lastgroup or ""
        out.append((kind, m.group(kind)))
        pos = m.end()
        if kind == "func":
            out.append(("op", "("))
    return tuple(out)


def _col_num(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def _split_ref(ref: str) -> Tuple[int, int, bool, bool]:
    m = re.match(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$", ref)
    if not m:
        raise _FormulaError("bad ref")
    return _col_num(m.group(2)), int(m.group(4)), bool(m.group(1)), bool(m.group(3))


class _Evaluator:
    """Evaluate one CF/total formula at a cell. `value_at(col, row)` returns
    the cell's value (1-based); `anchor` is the range's top-left the formula
    is relative to."""

    def __init__(self, formula: str, value_at, anchor: Tuple[int, int], at: Tuple[int, int]):
        self.toks = _tokens(formula)
        self.i = 0
        self.value_at = value_at
        self.dc = at[0] - anchor[0]
        self.dr = at[1] - anchor[1]
        self.at = at

    def peek(self) -> Tuple[str, str]:
        return self.toks[self.i] if self.i < len(self.toks) else ("end", "")

    def take(self) -> Tuple[str, str]:
        tok = self.peek()
        self.i += 1
        return tok

    def run(self) -> Any:
        v = self.compare()
        if self.peek()[0] != "end":
            raise _FormulaError("trailing tokens")
        return v

    def compare(self) -> Any:
        left = self.additive()
        while self.peek() in (("op", "="), ("op", "<>"), ("op", "<"), ("op", ">"), ("op", "<="), ("op", ">=")):
            op = self.take()[1]
            right = self.additive()
            left = _cmp(left, op, right)
        return left

    def additive(self) -> Any:
        left = self.term()
        while self.peek() in (("op", "+"), ("op", "-"), ("op", "&")):
            op = self.take()[1]
            right = self.term()
            if op == "&":
                left = f"{_text(left)}{_text(right)}"
            else:
                a, b = _num(left), _num(right)
                if a is None or b is None:
                    raise _FormulaError("#VALUE!")
                left = a + b if op == "+" else a - b
        return left

    def term(self) -> Any:
        left = self.unary()
        while self.peek() in (("op", "*"), ("op", "/")):
            op = self.take()[1]
            right = self.unary()
            a, b = _num(left), _num(right)
            if a is None or b is None or (op == "/" and b == 0):
                raise _FormulaError("#VALUE!")
            left = a * b if op == "*" else a / b
        return left

    def unary(self) -> Any:
        if self.peek() == ("op", "-"):
            self.take()
            v = _num(self.unary())
            if v is None:
                raise _FormulaError("#VALUE!")
            return -v
        return self.primary()

    def cell(self, ref: str) -> Any:
        c, r, abs_c, abs_r = _split_ref(ref)
        return self.value_at(c if abs_c else c + self.dc, r if abs_r else r + self.dr)

    def primary(self) -> Any:
        kind, text = self.take()
        if kind == "str":
            return text[1:-1].replace('""', '"')
        if kind == "num":
            return float(text)
        if kind == "bool":
            return text == "TRUE"
        if kind == "ref":
            return self.cell(text)
        if kind == "range":
            a, b = text.split(":")
            c1, r1, _, _ = _split_ref(a)
            c2, r2, _, _ = _split_ref(b)
            if (abs(c2 - c1) + 1) * (abs(r2 - r1) + 1) > 100_000:
                # Our totals span one column of at most 5,000 rows; a range
                # this large is not ours and would be a denial of service.
                raise _FormulaError("range too large")
            return [self.value_at(c, r) for r in range(min(r1, r2), max(r1, r2) + 1) for c in range(min(c1, c2), max(c1, c2) + 1)]
        if kind == "op" and text == "(":
            v = self.compare()
            if self.take() != ("op", ")"):
                raise _FormulaError("missing )")
            return v
        if kind == "func":
            self.take()  # the "(" _tokens inserted
            args: List[Any] = []
            if self.peek() != ("op", ")"):
                while True:
                    args.append(self.compare())
                    if self.peek() == ("op", ","):
                        self.take()
                        continue
                    break
            if self.take() != ("op", ")"):
                raise _FormulaError("missing )")
            return _call(text, args, self.at)
        raise _FormulaError(f"unexpected {text!r}")


def _num(v: Any) -> Optional[float]:
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, _dt.datetime):
        return (v - _dt.datetime(1899, 12, 30)).total_seconds() / 86400.0
    if isinstance(v, _dt.date):
        return float((v - _dt.date(1899, 12, 30)).days)
    return None


def _text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def _cmp(a: Any, op: str, b: Any) -> bool:
    na, nb = _num(a), _num(b)
    if na is not None and nb is not None and not isinstance(a, str) and not isinstance(b, str):
        x, y = na, nb
    elif isinstance(a, str) or isinstance(b, str) or a is None or b is None:
        if (na is not None and not isinstance(a, (str, type(None)))) or (nb is not None and not isinstance(b, (str, type(None)))):
            # A number compared with text: Excel orders every number below
            # every string, and they are never equal.
            if op in ("=",):
                return False
            if op == "<>":
                return True
            x, y = (0, 1) if na is not None and not isinstance(a, str) else (1, 0)
        else:
            x, y = _text(a).casefold(), _text(b).casefold()
    else:
        x, y = _text(a), _text(b)
    return {"=": x == y, "<>": x != y, "<": x < y, ">": x > y, "<=": x <= y, ">=": x >= y}[op]


def _flat(args: List[Any]) -> List[Any]:
    out: List[Any] = []
    for a in args:
        out.extend(a if isinstance(a, list) else [a])
    return out


def _call(name: str, args: List[Any], at: Tuple[int, int]) -> Any:
    if name == "AND":
        return all(bool(_num(a) if not isinstance(a, str) else a) for a in _flat(args))
    if name == "OR":
        return any(bool(_num(a) if not isinstance(a, str) else a) for a in _flat(args))
    if name == "NOT":
        return not bool(args[0])
    if name == "TRIM":
        return " ".join(_text(args[0]).split())
    if name == "ISNUMBER":
        return _num(args[0]) is not None and not isinstance(args[0], (str, bool))
    if name == "LEN":
        return float(len(_text(args[0])))
    if name == "SEARCH":
        needle = _text(args[0]).replace("~*", "\x00").replace("~?", "\x01").replace("~~", "~").casefold()
        needle = needle.replace("\x00", "*").replace("\x01", "?")
        hay = _text(args[1]).casefold()
        idx = hay.find(needle)
        if idx < 0:
            raise _FormulaError("#VALUE!")
        return float(idx + 1)
    if name == "MOD":
        a, b = _num(args[0]), _num(args[1])
        if a is None or not b:
            raise _FormulaError("#DIV/0!")
        return a % b
    if name == "ROW":
        return float(at[1])
    if name == "TODAY":
        return _dt.datetime.combine(_dt.date.today(), _dt.time())
    if name in ("SUBTOTAL", "SUM", "AVERAGE", "COUNTA", "MIN", "MAX", "COUNT"):
        fn = name
        values = args
        if name == "SUBTOTAL":
            code = int(_num(args[0]) or 0)
            fn = {101: "AVERAGE", 1: "AVERAGE", 102: "COUNT", 2: "COUNT", 103: "COUNTA", 3: "COUNTA", 104: "MAX", 4: "MAX", 105: "MIN", 5: "MIN", 109: "SUM", 9: "SUM"}.get(code, "")
            values = args[1:]
        flat = _flat(values)
        nums = [_num(v) for v in flat if not isinstance(v, (str, bool)) and _num(v) is not None]
        if fn == "SUM":
            return float(sum(nums))
        if fn == "AVERAGE":
            if not nums:
                raise _FormulaError("#DIV/0!")
            return float(sum(nums) / len(nums))
        if fn == "COUNT":
            return float(len(nums))
        if fn == "COUNTA":
            return float(sum(1 for v in flat if v not in (None, "")))
        if fn == "MIN":
            return float(min(nums)) if nums else 0.0
        if fn == "MAX":
            return float(max(nums)) if nums else 0.0
    raise _FormulaError(f"unsupported function {name}")


def evaluate_formula(formula: str, value_at, *, anchor: Tuple[int, int] = (1, 1), at: Tuple[int, int] = (1, 1)) -> Any:
    """Evaluate `formula` (CF grammar and the totals functions) at cell
    `at` (1-based col, row) relative to `anchor`. Raises ValueError for a
    formula outside the grammar or an Excel error value."""
    return _Evaluator(formula, value_at, anchor, at).run()


def _mix(a: str, b: str, t: float) -> str:
    ra, ga, ba = (int(a[i:i + 2], 16) for i in (1, 3, 5))
    rb, gb, bb = (int(b[i:i + 2], 16) for i in (1, 3, 5))
    return "#%02X%02X%02X" % (round(ra + (rb - ra) * t), round(ga + (gb - ga) * t), round(ba + (bb - ba) * t))


_NUMBER_FORMAT_CURRENCY = re.compile(r'"([₹$€£])"')


def display_value(value: Any, number_format: str) -> str:
    """A cell value as its number format shows it: grouping, decimals,
    percent, currency (INR in lakh/crore grouping), dd-mmm-yyyy or ISO
    dates. Text and unknown formats come back as written."""
    from . import theme

    if value is None:
        return ""
    fmt = number_format or "General"
    if isinstance(value, (_dt.date, _dt.datetime)):
        moment = value if isinstance(value, _dt.datetime) else _dt.datetime.combine(value, _dt.time())
        if "mmm" in fmt:
            return moment.strftime("%d-%b-%Y" + (" %H:%M" if "hh" in fmt else ""))
        return moment.strftime("%Y-%m-%d" + (" %H:%M" if "hh" in fmt else ""))
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    number = float(value)
    if fmt in ("General", "@"):
        return f"{int(number)}" if number == int(number) and abs(number) < 1e15 else f"{number:g}"
    first = fmt.split(";")[-1] if fmt.startswith("[>=") else fmt.split(";")[0]
    if "%" in first:
        decimals = len(first.split(".")[1].rstrip("%")) if "." in first else 0
        return f"{number * 100:,.{decimals}f}%"
    decimals = len(re.sub(r"[^0]", "", first.split(".")[1])) if "." in first else 0
    symbol_m = _NUMBER_FORMAT_CURRENCY.search(fmt)
    symbol = symbol_m.group(1) if symbol_m else ""
    if symbol == "₹":
        whole, _, frac = f"{abs(number):.{decimals}f}".partition(".")
        text = theme.indian_grouping(whole) + (f".{frac}" if frac else "")
    else:
        text = f"{abs(number):,.{decimals}f}"
    sign = "-" if number < 0 and float(text.replace(",", "") or 0) != 0 else ""
    return f"{sign}{symbol}{text}"


def styled_window(ws, first_row: int, last_row: int, width: int) -> Dict[str, Any]:
    """Per-cell styles and display strings for rows first_row..last_row
    (1-based sheet rows; row 1 is the header) of a worksheet loaded with
    styles. Returns {header_styles: [..], cell_styles: {"r:c": {...}},
    display: [[...]]} with r, c 0-based within the data window."""
    values: Dict[Tuple[int, int], Any] = {}
    max_row = int(ws.max_row or 0)

    def value_at(col: int, row: int) -> Any:
        key = (col, row)
        if key not in values:
            if row < 1 or col < 1 or row > max_row:
                values[key] = None
            else:
                values[key] = ws.cell(row=row, column=col).value
        return values[key]

    def static(cell) -> Dict[str, Any]:
        out: Dict[str, Any] = {}
        fill = cell.fill
        if fill is not None and getattr(fill, "fill_type", None) == "solid":
            colour = _rgb_of(fill.fgColor)
            if colour and colour not in ("#FFFFFF", "#000000"):
                out["fill"] = colour
        font = cell.font
        if font is not None:
            colour = _rgb_of(font.color)
            if colour and colour not in ("#000000",):
                out["color"] = colour
            if font.bold:
                out["bold"] = True
            if font.italic:
                out["italic"] = True
            if font.underline:
                out["underline"] = True
        return out

    # Conditional formats, highest priority first.
    rules: List[Tuple[int, Any, Any]] = []
    try:
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                rules.append((int(rule.priority or 10_000), cf.sqref, rule))
    except Exception:  # noqa: BLE001 — a workbook we did not write: no CF colours
        rules = []
    rules.sort(key=lambda r: r[0])
    parsed_ranges = []
    for priority, sqref, rule in rules:
        boxes = []
        for rng in str(sqref).split():
            parts = rng.split(":")
            try:
                c1, r1, _, _ = _split_ref(parts[0])
                c2, r2, _, _ = _split_ref(parts[-1])
            except _FormulaError:
                continue
            boxes.append((min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2)))
        parsed_ranges.append((boxes, rule))

    scale_cache: Dict[int, Tuple[float, float, float]] = {}

    def scale_colour(idx: int, boxes, rule, number: float) -> Optional[str]:
        cs = rule.colorScale
        if cs is None:
            return None
        colours = ["#" + (_rgb_of(c) or "#FFFFFF")[-6:] for c in cs.color]
        if idx not in scale_cache:
            nums = []
            for c1, r1, c2, r2 in boxes:
                for r in range(r1, min(r2, max_row) + 1):
                    for c in range(c1, c2 + 1):
                        v = value_at(c, r)
                        if isinstance(v, (int, float)) and not isinstance(v, bool):
                            nums.append(float(v))
            if not nums:
                scale_cache[idx] = (0.0, 0.0, 0.0)
            else:
                nums.sort()
                scale_cache[idx] = (nums[0], nums[len(nums) // 2] if len(nums) % 2 else (nums[len(nums) // 2 - 1] + nums[len(nums) // 2]) / 2, nums[-1])
        lo, mid, hi = scale_cache[idx]
        if hi == lo:
            return colours[-1]
        if len(colours) == 3:
            if number <= mid:
                return _mix(colours[0], colours[1], 0 if mid == lo else (number - lo) / (mid - lo))
            return _mix(colours[1], colours[2], 0 if hi == mid else (number - mid) / (hi - mid))
        return _mix(colours[0], colours[-1], (number - lo) / (hi - lo))

    header_styles = [static(ws.cell(row=1, column=c + 1)) for c in range(width)]
    cell_styles: Dict[str, Dict[str, Any]] = {}
    display: List[List[str]] = []
    for r in range(first_row, last_row + 1):
        row_out: List[str] = []
        for c in range(1, width + 1):
            cell = ws.cell(row=r, column=c)
            style = static(cell)
            decided: set = set()
            for idx, (boxes, rule) in enumerate(parsed_ranges):
                inside = next(((c1, r1) for c1, r1, c2, r2 in boxes if c1 <= c <= c2 and r1 <= r <= r2), None)
                if inside is None:
                    continue
                hit = False
                if rule.type == "colorScale":
                    v = value_at(c, r)
                    if isinstance(v, (int, float)) and not isinstance(v, bool) and "fill" not in decided:
                        colour = scale_colour(idx, boxes, rule, float(v))
                        if colour:
                            style["fill"] = colour
                            decided.add("fill")
                    continue
                if rule.type != "expression" or not rule.formula:
                    continue
                try:
                    hit = bool(evaluate_formula(rule.formula[0], value_at, anchor=inside, at=(c, r)))
                except (ValueError, IndexError, KeyError, TypeError):
                    hit = False
                if not hit:
                    continue
                dxf = rule.dxf
                if dxf is not None:
                    if dxf.fill is not None and "fill" not in decided:
                        colour = _rgb_of(dxf.fill.bgColor) or _rgb_of(dxf.fill.fgColor)
                        if colour:
                            style["fill"] = colour
                            decided.add("fill")
                    if dxf.font is not None:
                        colour = _rgb_of(dxf.font.color)
                        if colour and "color" not in decided:
                            style["color"] = colour
                            decided.add("color")
                        for attr in ("bold", "italic", "underline"):
                            if getattr(dxf.font, attr) and attr not in decided:
                                style[attr] = True
                                decided.add(attr)
                if rule.stopIfTrue:
                    break
            if style:
                cell_styles[f"{r - first_row}:{c - 1}"] = style
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and cell.data_type == "f":
                try:
                    computed = evaluate_formula(v, value_at, at=(c, r))
                    row_out.append(display_value(computed, cell.number_format))
                except (ValueError, IndexError, KeyError, TypeError):
                    row_out.append(v)
            else:
                row_out.append(display_value(v, cell.number_format))
        display.append(row_out)
    return {"header_styles": header_styles, "cell_styles": cell_styles, "display": display}


def styled_grid(xlsx_path: str | Path, sheet: Optional[str], *, offset: int, limit: int, width: int) -> Optional[Dict[str, Any]]:
    """styled_window for one sheet of a workbook file, or None when the file
    is too large to load with styles (the plain grid still answers)."""
    from openpyxl import load_workbook  # lazy

    p = Path(xlsx_path)
    if p.stat().st_size > _GRID_STYLE_MAX_BYTES:
        return None
    wb = load_workbook(str(p), read_only=False, data_only=False)
    try:
        visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
        if not visible:
            return None
        chosen = visible[0]
        if sheet:
            chosen = next((ws for ws in visible if ws.title == sheet), chosen)
        first = 2 + max(0, int(offset))
        last = min(int(chosen.max_row or 1), first + max(1, int(limit)) - 1)
        if last < first:
            return {"header_styles": [], "cell_styles": {}, "display": []}
        return styled_window(chosen, first, last, max(1, int(width)))
    finally:
        wb.close()


def rasterise_image(image_path: str | Path, width: int) -> bytes:
    """PNG bytes of a standalone chart image (preview_kind 'image'), `width`
    pixels wide: a PNG is downscaled with PIL; an SVG is drawn through
    WeasyPrint (which never fetches: the assets-only fetcher) and PDFium."""
    from PIL import Image

    if width < 32 or width > max(T.PREVIEW_WIDTHS) * 2:
        raise ValueError(f"preview width {width} is out of range")
    p = Path(image_path)
    if p.suffix.lower() == ".svg":
        from .pdf import make_url_fetcher

        from weasyprint import HTML  # lazy: native libs

        html = f'<html><body style="margin:0"><img src="{p.name}" style="width:100%"></body></html>'
        pdf_bytes = HTML(string=html, base_url=str(p.parent.resolve()) + "/", url_fetcher=make_url_fetcher(p.parent)).write_pdf()
        tmp = p.parent / (p.stem + ".preview-tmp.pdf")
        tmp.write_bytes(pdf_bytes)
        try:
            return rasterise_page(tmp, 1, width)
        finally:
            tmp.unlink(missing_ok=True)
    with Image.open(p) as im:
        im = im.convert("RGB")
        if im.width != width:
            im = im.resize((width, max(1, round(im.height * width / im.width))), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, format="PNG", optimize=True)
        return buf.getvalue()


__all__ = ["page_count", "rasterise_page", "rasterise_image", "sheet_grid", "grid_for", "styled_grid", "styled_window", "evaluate_formula", "display_value"]
