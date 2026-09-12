"""Reopen every rendered file and say what it is — the `validate` stage.

WHY. The legacy report path advertised a file on `is_file()` alone
(CURRENT_STATE.md: "no %PDF magic, page count, zip integrity, DOCX/XLSX
reopen"). A truncated PDF or a zip that lost its central directory reaches
the person as a card whose download opens nothing. Here each output is
reopened with the library that would read it — pypdfium2, python-docx,
python-pptx, openpyxl — and its structural facts (pages, slides, sheets,
size, sha256) are recorded in validation.json.

WHAT IS REFUSED. An Office file with a `vbaProject.bin` part (a macro) or a
relationship whose TargetMode is External (a link that resolves outside the
package — the classic remote-template phone-home) fails validation; neither
can come from python-docx/pptx/openpyxl writing our spec, so their presence
means the file is not the one we wrote. A PDF without the `%PDF` header,
with zero pages, over types.MAX_PAGES, or with a page whose media box is
empty fails too. A validation failure is a RenderError('validation_failure')
upstream, never a card.

ROW COUNTS (CONTRACT-2 §11). "Exactly 500 rows" is checked here, by the
process that rendered: a CSV's data rows are counted by csv.reader
(render/csv.validate_csv) and compared with the spec sheet it came from;
an XLSX's rows are counted per sheet — the rows between the header and the
totals row, inside the sheet's own columns, so a chart's data block beside
the table is not a row — and compared with the spec. A mismatch is the
sentence "the CSV has 499 data rows; 500 were required". `validate_all`
takes the spec for that, and keys its files by the render report's
`role:format:sheet_slug` (a bare format is still accepted).
"""
from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .. import types as T

_EXTERNAL_TARGET_RE = re.compile(rb'TargetMode\s*=\s*"External"', re.IGNORECASE)


class ValidationFailed(ValueError):
    """The file is not one we can release. The message is safe to show."""


def sha256_of(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _check_zip(path: Path) -> None:
    """zipfile integrity, no macro part, no external relationship target."""
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ValidationFailed("the file is not a valid Office package") from exc
    with zf:
        bad = zf.testzip()
        if bad is not None:
            raise ValidationFailed("the Office package has a corrupt part")
        names = zf.namelist()
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            raise ValidationFailed("the file carries a macro project, which is never produced here")
        for name in names:
            if name.endswith(".rels"):
                if _EXTERNAL_TARGET_RE.search(zf.read(name)):
                    raise ValidationFailed("the file links to an external target, which is never produced here")


def validate_pdf(path: str | Path, *, max_pages: int = T.MAX_PAGES) -> Dict[str, object]:
    import pypdfium2 as pdfium  # lazy

    p = Path(path)
    with open(p, "rb") as fh:
        head = fh.read(5)
    if head != b"%PDF-":
        raise ValidationFailed("the PDF does not start with a PDF header")
    try:
        pdf = pdfium.PdfDocument(str(p))
    except Exception as exc:  # pypdfium2 raises its own error type on a broken file
        raise ValidationFailed("the PDF could not be opened") from exc
    try:
        pages = len(pdf)
        if pages <= 0:
            raise ValidationFailed("the PDF has no pages")
        if pages > max_pages:
            raise ValidationFailed(f"the PDF has {pages} pages; the ceiling is {max_pages}")
        for i in range(pages):
            page = pdf[i]
            try:
                w, h = page.get_size()
            finally:
                page.close()
            if w <= 0 or h <= 0:
                raise ValidationFailed(f"page {i + 1} of the PDF has an empty media box")
    finally:
        pdf.close()
    return {"ok": True, "pages": pages, "size": p.stat().st_size}


def validate_docx(path: str | Path) -> Dict[str, object]:
    from docx import Document  # lazy

    p = Path(path)
    _check_zip(p)
    try:
        document = Document(str(p))
    except Exception as exc:
        raise ValidationFailed("the DOCX could not be opened by Word's library") from exc
    paragraphs = len(document.paragraphs)
    if paragraphs == 0 and not document.tables:
        raise ValidationFailed("the DOCX is empty")
    return {"ok": True, "paragraphs": paragraphs, "tables": len(document.tables), "size": p.stat().st_size}


def validate_pptx(path: str | Path) -> Dict[str, object]:
    from pptx import Presentation  # lazy

    p = Path(path)
    _check_zip(p)
    try:
        prs = Presentation(str(p))
    except Exception as exc:
        raise ValidationFailed("the PPTX could not be opened by PowerPoint's library") from exc
    slides = len(prs.slides)
    if slides <= 0:
        raise ValidationFailed("the PPTX has no slides")
    if slides > T.MAX_SLIDES + 1:  # + the sources slide
        raise ValidationFailed(f"the PPTX has {slides} slides; the ceiling is {T.MAX_SLIDES}")
    return {"ok": True, "slides": slides, "size": p.stat().st_size}


def _data_rows_of(ws, n_cols: int) -> int:
    """The data rows of one worksheet as the XLSX writer lays them out:
    row 1 is the header, the data follows, and a totals row (the first row
    with a formula cell inside the sheet's columns) ends it. Only the first
    `n_cols` columns are read, so a chart's data block written beside the
    table (xlsx._draw_charts) is never counted. Interior blank rows count:
    a pasted row whose every cell was blank is still a row."""
    last_data = 1
    for r_index, row in enumerate(ws.iter_rows(min_row=2, max_col=max(1, n_cols)), start=2):
        values = [c.value for c in row]
        if any(isinstance(c.value, str) and c.value.startswith("=") and c.data_type == "f" for c in row):
            break  # the totals row
        if any(v is not None for v in values):
            last_data = r_index
    return max(0, last_data - 1)


def validate_xlsx(path: str | Path, *, spec: Any = None) -> Dict[str, object]:
    """Reopen the workbook; with the spec, count each spec sheet's data
    rows and refuse a count that differs from the sheet's rows."""
    from openpyxl import load_workbook  # lazy

    p = Path(path)
    _check_zip(p)
    try:
        wb = load_workbook(str(p), read_only=True, data_only=False)
    except Exception as exc:
        raise ValidationFailed("the XLSX could not be opened by Excel's library") from exc
    facts: Dict[str, object] = {"ok": True, "size": p.stat().st_size}
    try:
        names = list(wb.sheetnames)
        if not names:
            raise ValidationFailed("the XLSX has no sheets")
        facts["sheets"] = len(names)
        facts["sheet_names"] = names
        body = getattr(spec, "body", spec)
        sheets = getattr(body, "sheets", None)
        if sheets:
            from .xlsx import sheet_titles

            counts: Dict[str, int] = {}
            for sheet, title in zip(sheets, sheet_titles(body)):
                if title not in wb.sheetnames:
                    raise ValidationFailed(f"the XLSX has no sheet named {title!r}")
                counted = _data_rows_of(wb[title], len(sheet.columns))
                expected = min(len(sheet.rows), T.MAX_ROWS_PER_SHEET)
                if counted != expected:
                    raise ValidationFailed(f"sheet {sheet.name!r} of the XLSX has {counted:,} data rows; {expected:,} were required")
                counts[sheet.name] = counted
            facts["sheet_rows"] = counts
            facts["rows"] = sum(counts.values())
            facts["columns"] = max(len(sh.columns) for sh in sheets)
    finally:
        wb.close()
    return facts


def validate_csv(path: str | Path, *, expected_rows: Optional[int] = None, expected_columns: Optional[int] = None) -> Dict[str, object]:
    """render/csv.validate_csv, under the validators' common shape."""
    from .csv import validate_csv as _validate_csv  # lazy: the module imports this one

    return dict(_validate_csv(path, expected_rows=expected_rows, expected_columns=expected_columns))


_VALIDATORS = {"pdf": validate_pdf, "docx": validate_docx, "pptx": validate_pptx, "xlsx": validate_xlsx, "csv": validate_csv}


def validate_file(path: str | Path, fmt: str, **facts_kw: Any) -> Dict[str, object]:
    """Reopen `path` as `fmt`; return the facts, or raise ValidationFailed.
    Keyword arguments go to the format's validator (`spec` for xlsx,
    `expected_rows`/`expected_columns` for csv, `max_pages` for pdf)."""
    if fmt not in _VALIDATORS:
        raise ValidationFailed(f"no validator for {fmt}")
    p = Path(path)
    if not p.is_file() or p.stat().st_size == 0:
        raise ValidationFailed(f"the {fmt} file was not written")
    if p.stat().st_size > T.MAX_FILE_BYTES:
        raise ValidationFailed(f"the {fmt} file is larger than {T.MAX_FILE_BYTES // (1024 * 1024)} MB")
    facts = _VALIDATORS[fmt](p, **facts_kw)
    facts["sha256"] = sha256_of(p)
    facts["format"] = fmt
    facts["filename"] = p.name
    return facts


def split_key(key: str) -> Tuple[str, str, str]:
    """`role:format:sheet_slug` → (role, format, sheet_slug); a bare format
    is ("", format, "")."""
    parts = str(key).split(":")
    if len(parts) == 1:
        return "", parts[0], ""
    if len(parts) == 2:
        return parts[0], parts[1], ""
    return parts[0], parts[1], ":".join(parts[2:])


def _sheet_for(spec: Any, sheet_slug: str) -> Any:
    """The spec sheet a per-sheet CSV came from: by slug of its name, or
    the only sheet when the file carries no part."""
    body = getattr(spec, "body", spec)
    sheets: List[Any] = list(getattr(body, "sheets", None) or [])
    if not sheets:
        return None
    if not sheet_slug:
        return sheets[0] if len(sheets) == 1 else None
    for sh in sheets:
        if T.slug_for(sh.name, fallback="part") == sheet_slug:
            return sh
    return None


def validate_all(files: Dict[str, str], preview_pdf: Optional[str] = None, *, spec: Any = None,
                 max_pages: int = T.MAX_PAGES) -> Dict[str, object]:
    """{key: path} (+ the preview) → the validation.json dict, `files` keyed
    as given. A key is the render report's `role:format:sheet_slug` or a
    bare format. With `spec` (a WorkbookSpec or its ArtifactSpec) a CSV is
    held to its sheet's row and column counts and an XLSX to every sheet's
    (CONTRACT-2 §11); `rows` and `columns` are returned per file. Raises
    on the first failure; the caller turns it into a RenderError."""
    out: Dict[str, object] = {"files": {}, "renderer_version": T.RENDERER_VERSION, "template_version": T.TEMPLATE_VERSION}
    body = getattr(spec, "body", spec) if spec is not None else None
    sheets = list(getattr(body, "sheets", None) or []) if body is not None else []
    for key, path in files.items():
        _role, fmt, sheet_slug = split_key(key)
        kw: Dict[str, Any] = {}
        if fmt == "csv" and sheets:
            sheet = _sheet_for(body, sheet_slug)
            if sheet is not None:
                kw = {"expected_rows": min(len(sheet.rows), T.MAX_ROWS_PER_SHEET), "expected_columns": len(sheet.columns)}
        elif fmt == "xlsx" and sheets:
            kw = {"spec": body}
        elif fmt == "pdf":
            kw = {"max_pages": max_pages}
        facts = validate_file(path, fmt, **kw)
        if sheets and fmt in ("docx", "pdf"):
            # The tabular document carries every sheet's rows.
            facts.setdefault("rows", sum(min(len(sh.rows), T.MAX_ROWS_PER_SHEET) for sh in sheets))
            facts.setdefault("columns", max(len(sh.columns) for sh in sheets))
        out["files"][key] = facts
    if preview_pdf:
        out["preview"] = validate_file(preview_pdf, "pdf", max_pages=max_pages)
    return out


__all__ = [
    "ValidationFailed", "sha256_of", "validate_file", "validate_all", "validate_pdf", "validate_docx", "validate_pptx",
    "validate_xlsx", "validate_csv", "split_key",
]
