"""Read the chart fixtures as DataTables, the shape material_in hands the
chart code. Independent of the app's readers on purpose: openpyxl cached
values, csv.reader, python-docx table cells."""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, List

HERE = Path(__file__).resolve().parent

FILES = {
    "sales": "sales_daily.csv", "tickets": "tickets.xlsx", "employees": "employees.csv", "projects": "projects.csv",
    "cashflow": "cashflow.csv", "funnel": "funnel.csv", "dates_dmy": "dates_dmy.csv", "dates_mdy": "dates_mdy.csv",
    "dates_ambiguous": "dates_ambiguous.csv", "units": "units.docx",
    "prices": "prices.csv", "targets": "targets.csv", "defects": "defects.csv",
}


def ground_truth() -> dict:
    return json.loads((HERE / "ground_truth.json").read_text(encoding="utf-8"))


def _num(cell: str) -> Any:
    try:
        f = float(cell.replace(",", ""))
        return int(f) if f.is_integer() else f
    except ValueError:
        return cell


def table(name: str, table_id: str = "") -> Any:
    from app.artifacts.compose import DataTable

    fname = FILES[name]
    path = HERE / fname
    tid = table_id or f"upload_{name}"
    if fname.endswith(".csv"):
        with open(path, newline="", encoding="utf-8") as fh:
            rows = list(csv.reader(fh))
        return DataTable(id=tid, title=fname, columns=rows[0], rows=[[_num(c) for c in r] for r in rows[1:]])
    if fname.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return DataTable(id=tid, title=fname, columns=[str(c) for c in rows[0]], rows=rows[1:])
    if fname.endswith(".docx"):
        from docx import Document

        doc = Document(str(path))
        t = doc.tables[0]
        rows: List[List[Any]] = [[c.text for c in r.cells] for r in t.rows]
        return DataTable(id=tid, title=fname, columns=rows[0], rows=[[_num(c) for c in r] for r in rows[1:]])
    raise KeyError(name)


def all_tables() -> list:
    return [table(n) for n in FILES]
