"""Build the synthetic chart fixtures and their ground truth.

Every file here is invented (no user data). The ground truth is computed in
plain Python from the same rows, independently of app/artifacts/chart_data.py
(which uses pandas/duckdb), so a test comparing the two checks real values.

    python tests/fixtures/charts/make_fixtures.py   # rewrites the files

Deterministic: a fixed seed, fixed dates, no clock.
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import random
from collections import OrderedDict, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
SEED = 20260915
MONTH_ABBR = ("", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")

#: (month, rows, total amount) — the sums the acceptance criteria name.
SALES_MONTHS = ((1, 23, 64514), (2, 22, 71319), (3, 23, 74708), (4, 22, 58936), (5, 23, 63092), (6, 22, 77350), (7, 15, 23072))
REGIONS = ("North", "South", "East", "West")
PRODUCTS = ("Alpha", "Bravo", "Charlie", "Delta")
STATUS_COUNTS = (("Resolved", 24), ("Open", 14), ("In Progress", 11), ("Closed", 8), ("On Hold", 3))
PRIORITIES = ("Low", "Medium", "High", "Critical")
OWNERS = ("Asha", "Ravi", "Meera", "John", "Priya")
CATEGORIES = ("Billing", "Login", "Performance", "Reports", "Integrations")


def _alloc(total: int, n: int, rng: random.Random) -> list:
    weights = [rng.uniform(0.4, 1.6) for _ in range(n)]
    s = sum(weights)
    amounts = [int(total * w / s) for w in weights]
    for i in range(total - sum(amounts)):
        amounts[i % n] += 1
    assert sum(amounts) == total
    return amounts


def sales(rng: random.Random):
    rows = []
    for month, n, total in SALES_MONTHS:
        days_in = 31 if month in (1, 3, 5, 7) else (28 if month == 2 else 30)
        days = sorted(rng.sample(range(1, days_in + 1), n))
        for day, amount in zip(days, _alloc(total, n, rng)):
            rows.append({
                "Date": dt.date(2026, month, day).isoformat(),
                "Region": rng.choice(REGIONS),
                "Product": rng.choice(PRODUCTS),
                "Units": rng.randint(1, 40),
                "Amount": amount,
            })
    return rows


def tickets(rng: random.Random):
    statuses = [s for s, c in STATUS_COUNTS for _ in range(c)]
    rng.shuffle(statuses)
    rows = []
    for i, status in enumerate(statuses, start=1):
        created = dt.datetime(2026, 1, 5) + dt.timedelta(days=rng.randint(0, 170), hours=rng.randint(8, 18))
        due = created.date() + dt.timedelta(days=rng.randint(3, 30))
        rows.append({
            "Ticket": f"T-{i:03d}",
            "Title": f"Synthetic issue {i}",
            "Status": status,
            "Priority": rng.choice(PRIORITIES),
            "Owner": rng.choice(OWNERS),
            "Category": rng.choice(CATEGORIES),
            "Created": created,
            "Due": due,
            "Score": rng.randint(1, 5),
            "Hours": round(rng.uniform(0.5, 24.0), 1),
        })
    return rows


def employees(rng: random.Random):
    depts = ("Engineering", "Sales", "Support", "Finance")
    rows = []
    for i in range(1, 81):
        exp = rng.randint(0, 20)
        dept = rng.choice(depts)
        salary = 300000 + exp * 85000 + rng.randint(-60000, 60000) + (150000 if dept == "Engineering" else 0)
        rows.append({"Employee": f"E{i:03d}", "Department": dept, "Experience": exp, "Salary": salary,
                     "Rating": round(rng.uniform(2.0, 5.0), 1), "Age": 22 + exp + rng.randint(0, 6)})
    return rows


def projects():
    return [
        {"Task": "Discovery", "Owner": "Asha", "Start": "05-01-2026", "End": "16-01-2026"},
        {"Task": "Design", "Owner": "Ravi", "Start": "19-01-2026", "End": "13-02-2026"},
        {"Task": "Build", "Owner": "Meera", "Start": "16-02-2026", "End": "24-04-2026"},
        {"Task": "Testing", "Owner": "John", "Start": "20-04-2026", "End": "22-05-2026"},
        {"Task": "Launch", "Owner": "Priya", "Start": "25-05-2026", "End": "29-05-2026"},
    ]


def cashflow():
    return [
        {"Item": "Opening balance", "Amount": 500000}, {"Item": "Sales", "Amount": 320000},
        {"Item": "Services", "Amount": 85000}, {"Item": "Salaries", "Amount": -410000},
        {"Item": "Rent", "Amount": -60000}, {"Item": "Marketing", "Amount": -45000}, {"Item": "Tax", "Amount": -38000},
    ]


def funnel():
    return [{"Stage": "Visitors", "Count": 12000}, {"Stage": "Sign-ups", "Count": 3100}, {"Stage": "Trials", "Count": 1240},
            {"Stage": "Paid", "Count": 410}, {"Stage": "Renewed", "Count": 290}]


def dated(rng: random.Random, order: str, ambiguous: bool):
    rows = []
    start = dt.date(2026, 1, 1)
    for i in range(40):
        d = start + dt.timedelta(days=i * 4 + rng.randint(0, 2))
        if ambiguous:
            d = dt.date(2026, (i % 6) + 1, (i % 12) + 1)
        a, b = (d.day, d.month) if order == "dmy" else (d.month, d.day)
        rows.append({"Date": f"{a:02d}-{b:02d}-{d.year}", "Amount": rng.randint(100, 999), "_iso": d.isoformat()})
    return rows


def write_csv(path: Path, rows, fields):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def write_tickets_xlsx(path: Path, rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    fields = list(rows[0].keys())
    ws.append(fields)
    for r in rows:
        ws.append([r[f] for f in fields])
    wb.properties.creator = "fixtures"
    wb.save(path)


def write_units_docx(path: Path, units):
    from docx import Document

    doc = Document()
    doc.core_properties.author = "fixtures"
    doc.add_heading("Units shipped by quarter", level=1)
    doc.add_paragraph("Synthetic figures for the chart tests.")
    table = doc.add_table(rows=1, cols=5)
    for i, h in enumerate(("Product", "Q1", "Q2", "Q3", "Q4")):
        table.rows[0].cells[i].text = h
    for name, qs in units.items():
        cells = table.add_row().cells
        cells[0].text = name
        for i, q in enumerate(qs, start=1):
            cells[i].text = f"{q:,}"
    doc.save(path)


def write_headcount_pdf(path: Path, headcount):
    from weasyprint import HTML

    rows = "".join(f"<tr><td>{d}</td><td>{h}</td><td>{a}</td></tr>" for d, h, a in headcount)
    html = ("<html><body style='font-family: DejaVu Sans'><h1>Headcount by department</h1>"
            "<table border='1' cellpadding='6'><tr><th>Department</th><th>Headcount</th><th>Attrition %</th></tr>"
            f"{rows}</table></body></html>")
    HTML(string=html).write_pdf(path)


def main() -> None:
    rng = random.Random(SEED)
    s_rows = sales(rng)
    t_rows = tickets(rng)
    e_rows = employees(rng)
    dmy = dated(rng, "dmy", ambiguous=False)
    mdy = dated(rng, "mdy", ambiguous=False)
    amb = dated(rng, "dmy", ambiguous=True)
    units = OrderedDict((("Alpha", (1200, 1350, 1100, 1500)), ("Bravo", (800, 950, 1020, 1100)), ("Charlie", (400, 380, 520, 610)),
                         ("Delta", (1500, 1420, 1390, 1600)), ("Echo", (230, 310, 290, 350))))
    headcount = (("Engineering", 142, 8.5), ("Sales", 96, 14.2), ("Support", 64, 11.0), ("Finance", 28, 6.1), ("HR", 12, 4.0))

    write_csv(HERE / "sales_daily.csv", s_rows, ["Date", "Region", "Product", "Units", "Amount"])
    write_tickets_xlsx(HERE / "tickets.xlsx", t_rows)
    write_csv(HERE / "employees.csv", e_rows, ["Employee", "Department", "Experience", "Salary", "Rating", "Age"])
    write_csv(HERE / "projects.csv", projects(), ["Task", "Owner", "Start", "End"])
    write_csv(HERE / "cashflow.csv", cashflow(), ["Item", "Amount"])
    write_csv(HERE / "funnel.csv", funnel(), ["Stage", "Count"])
    write_csv(HERE / "dates_dmy.csv", dmy, ["Date", "Amount"])
    write_csv(HERE / "dates_mdy.csv", mdy, ["Date", "Amount"])
    write_csv(HERE / "dates_ambiguous.csv", amb, ["Date", "Amount"])
    write_units_docx(HERE / "units.docx", units)
    write_headcount_pdf(HERE / "headcount.pdf", headcount)

    gt: dict = {}
    by_month = defaultdict(int)
    by_region = defaultdict(int)
    by_product = defaultdict(int)
    units_by_product = defaultdict(int)
    count_by_month = defaultdict(int)
    q_region = defaultdict(int)
    region_n = defaultdict(int)
    month_region = defaultdict(int)
    for r in s_rows:
        d = dt.date.fromisoformat(r["Date"])
        m = f"{MONTH_ABBR[d.month]} 2026"
        by_month[m] += r["Amount"]
        count_by_month[m] += 1
        by_region[r["Region"]] += r["Amount"]
        region_n[r["Region"]] += 1
        by_product[r["Product"]] += r["Amount"]
        units_by_product[r["Product"]] += r["Units"]
        month_region[f"{m}|{r['Region']}"] += r["Amount"]
        if d.month <= 6:
            q_region[f"Q{(d.month - 1) // 3 + 1} 2026|{r['Region']}"] += r["Amount"]
    gt["sales"] = {
        "rows": len(s_rows), "amount_by_month": dict(by_month), "count_by_month": dict(count_by_month),
        "amount_by_region": dict(by_region), "amount_by_product": dict(by_product), "units_by_product": dict(units_by_product),
        "amount_by_quarter_region_h1": dict(q_region), "amount_by_month_region": dict(month_region),
        "avg_amount_by_region": {k: by_region[k] / region_n[k] for k in by_region}, "total_amount": sum(by_region.values()),
    }
    status = defaultdict(int)
    prio = defaultdict(int)
    owner_hours = defaultdict(float)
    avg_score_owner = defaultdict(list)
    status_priority = defaultdict(int)
    tickets_by_month = defaultdict(int)
    for r in t_rows:
        status[r["Status"]] += 1
        prio[r["Priority"]] += 1
        owner_hours[r["Owner"]] += r["Hours"]
        avg_score_owner[r["Owner"]].append(r["Score"])
        status_priority[f"{r['Status']}|{r['Priority']}"] += 1
        tickets_by_month[f"{MONTH_ABBR[r['Created'].month]} 2026"] += 1
    gt["tickets"] = {
        "rows": len(t_rows), "status_counts": dict(status), "priority_counts": dict(prio),
        "hours_by_owner": dict(owner_hours),
        "avg_score_by_owner": {k: sum(v) / len(v) for k, v in avg_score_owner.items()},
        "status_priority_counts": dict(status_priority), "tickets_by_month": dict(tickets_by_month),
        "open_by_priority": {p: sum(1 for r in t_rows if r["Status"] == "Open" and r["Priority"] == p)
                             for p in PRIORITIES if any(r["Status"] == "Open" and r["Priority"] == p for r in t_rows)},
    }
    n = len(e_rows)
    xs = [r["Experience"] for r in e_rows]
    ys = [r["Salary"] for r in e_rows]
    mx, my = sum(xs) / n, sum(ys) / n
    slope = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sum((x - mx) ** 2 for x in xs)
    intercept = my - slope * mx
    dept_n = defaultdict(int)
    dept_salary = defaultdict(int)
    for r in e_rows:
        dept_n[r["Department"]] += 1
        dept_salary[r["Department"]] += r["Salary"]
    gt["employees"] = {"rows": n, "trend_salary_on_experience": {"slope": slope, "intercept": intercept},
                       "count_by_department": dict(dept_n), "avg_salary_by_department": {k: dept_salary[k] / dept_n[k] for k in dept_n},
                       "salaries": ys, "experience": xs}
    gt["projects"] = {r["Task"]: (dt.datetime.strptime(r["End"], "%d-%m-%Y") - dt.datetime.strptime(r["Start"], "%d-%m-%Y")).days for r in projects()}
    gt["cashflow"] = {r["Item"]: r["Amount"] for r in cashflow()}
    gt["funnel"] = {r["Stage"]: r["Count"] for r in funnel()}
    gt["units"] = {k: list(v) for k, v in units.items()}
    gt["headcount"] = {d: h for d, h, _ in headcount}
    for name, rows in (("dates_dmy", dmy), ("dates_mdy", mdy), ("dates_ambiguous", amb)):
        agg = defaultdict(int)
        for r in rows:
            d = dt.date.fromisoformat(r["_iso"])
            agg[f"{MONTH_ABBR[d.month]} 2026"] += r["Amount"]
        gt[name] = {"amount_by_month": dict(agg)}
    (HERE / "ground_truth.json").write_text(json.dumps(gt, indent=1, sort_keys=True, default=str), encoding="utf-8")


if __name__ == "__main__":
    main()
