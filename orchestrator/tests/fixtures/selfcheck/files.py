"""Real rendered files and FILE-LEVEL defect injection for the self-check tests.

Every defect here is made by editing the produced file (the OOXML parts
with lxml, the workbook/deck through openpyxl/python-pptx, the PDF through
pdfium's object API) — never by changing the spec and re-rendering — so the
inspector is tested on what a person downloads. Each case has a TWIN: the
same edit with the value the checklist item expects, which must pass.

All content is invented.
"""
from __future__ import annotations

import ctypes
import io
import os
import shutil
import zipfile
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

from lxml import etree

from app.artifacts import spec as S
from app.artifacts.render import render_version

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W = "{%s}" % W


def doc_spec(title: str = "Vendor Access Review", orientation: str = "portrait") -> S.ArtifactSpec:
    return S.parse_body("document", {
        "title": title, "subtitle": "Quarterly control check", "orientation": orientation,
        "blocks": [
            {"type": "heading", "level": 1, "text": "Scope"},
            {"type": "paragraph", "text": "This review covers vendor accounts with standing access to production systems."},
            {"type": "heading", "level": 1, "text": "Findings"},
            {"type": "paragraph", "text": "Four accounts had keys older than the rotation window."},
            {"type": "table", "table": {"columns": ["Area", "Status", "Count"], "rows": [["Logins", "Open", 4], ["Keys", "Closed", 9], ["Tokens", "Open", 2]], "numeric_columns": [2]}},
            {"type": "heading", "level": 2, "text": "Detail"},
            {"type": "bullets", "items": ["Rotate keys", "Remove idle vendors"]},
            {"type": "heading", "level": 1, "text": "Recommendations"},
            {"type": "paragraph", "text": "Adopt a ninety day rotation and review access monthly."},
        ]})


def workbook_spec(rows: int = 12) -> S.ArtifactSpec:
    statuses = ["Open", "Closed", "In Progress"]
    return S.parse_body("workbook", {"title": "Ticket Tracker", "sheets": [{
        "name": "Tickets",
        "columns": [{"name": "ID"}, {"name": "Status"}, {"name": "Owner"}, {"name": "Amount", "type": "number"}],
        "rows": [[f"T-{i + 1}", statuses[i % 3], f"Agent {i % 4}", float(10 * (i + 1))] for i in range(rows)],
        "totals": [{"column": "Amount"}],
        "charts": [{"type": "bar", "title": "Amount by ticket", "categories": ["T-1", "T-2", "T-3"], "series": [{"name": "Amount", "values": [10, 20, 30]}]}],
    }]})


def deck_spec() -> S.ArtifactSpec:
    return S.parse_body("presentation", {"title": "Q3 Results", "slides": [
        {"layout": "title", "title": "Q3 Results"},
        {"layout": "chart", "title": "Revenue", "chart": {"type": "line", "title": "Revenue", "categories": ["Jul", "Aug", "Sep"], "series": [{"name": "Revenue", "values": [3, 5, 4]}]}},
        {"layout": "table", "title": "Detail", "table": {"columns": ["Region", "Units"], "rows": [["North", 120], ["South", 95]]}},
    ]})


def synthetic_audit_markdown(target_chars: int = 40_000) -> str:
    """An invented audit answer of about `target_chars`: H1, H2 sections,
    paragraphs, bullet lists and a findings table per section."""
    areas = ["Identity", "Network", "Endpoints", "Backups", "Logging", "Vendors", "Change control", "Incident response"]
    parts = ["# Security Posture Audit", "", "An invented review for tests; no real organisation.", ""]
    n = 0
    while sum(len(p) + 1 for p in parts) < target_chars:
        area = areas[n % len(areas)]
        parts += [f"## {n + 1}. {area} review", ""]
        parts += [f"The {area.lower()} controls were sampled across {12 + n} systems during week {n % 52 + 1}. "
                  "Evidence was collected from configuration exports and interviews with the owners of each system. " * 3, ""]
        parts += [f"- Control {n}-A is documented", f"- Control {n}-B lacks an owner", ""]
        parts += ["| Control | Status | Severity | Owner |", "|---|---|---|---|"]
        for k in range(6):
            parts.append(f"| {area[:3].upper()}-{n:02d}{k} | {['Open', 'Closed', 'In Progress'][k % 3]} | {['High', 'Medium', 'Low'][k % 3]} | Team {chr(65 + (n + k) % 6)} |")
        parts.append("")
        n += 1
    return "\n".join(parts)


def markdown_to_spec(md: str) -> S.ArtifactSpec:
    """A test-only markdown -> DocumentSpec (headings, paragraphs, bullets,
    pipe tables), standing in for the intent track's md_import."""
    import re as _re

    lines = md.splitlines()
    title = "Document"
    blocks: List[dict] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        m = _re.match(r"^(#{1,6})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            if level == 1 and title == "Document":
                title = m.group(2).strip()
            else:
                blocks.append({"type": "heading", "level": min(3, max(1, level - 1)), "text": m.group(2).strip()})
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not _re.fullmatch(r"\|[-:| ]+\|", lines[i].strip()):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append({"type": "table", "table": {"columns": rows[0], "rows": rows[1:]}})
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append({"type": "bullets", "items": items})
            continue
        if line.strip():
            blocks.append({"type": "paragraph", "text": line.strip()[:6000]})
        i += 1
    return S.parse_body("document", {"title": title[:120], "blocks": blocks[:400]})


def render(spec: S.ArtifactSpec, formats: List[str], out_dir: Path, slug: str = "sample", version: int = 1) -> Dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    report = render_version(spec, formats, str(out_dir), title_slug=slug, version=version)
    data = report.to_json()
    return {f["filename"]: out_dir / f["filename"] for f in data["files"]}


# ----------------------------------------------------------- OOXML edits --


def rewrite_part(path: Path, part: str, fn: Callable[[etree._Element], None]) -> None:
    """Rewrite one XML part of a zip package in place."""
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == part:
                root = etree.fromstring(data)
                fn(root)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(info, data)
    os.replace(tmp, path)


def add_part(path: Path, part: str, data: bytes) -> None:
    with zipfile.ZipFile(path, "a") as z:
        z.writestr(part, data)


def _style_el(root, name: str):
    for st in root.findall(_W + "style"):
        n = st.find(_W + "name")
        if n is not None and n.get(_W + "val", "").lower() == name.lower():
            return st
    raise KeyError(name)


def _rpr(el):
    rpr = el.find(_W + "rPr")
    if rpr is None:
        rpr = etree.SubElement(el, _W + "rPr")
    return rpr


def _set_child(parent, tag: str, **attrs):
    el = parent.find(_W + tag)
    if el is None:
        el = etree.SubElement(parent, _W + tag)
    for k, v in attrs.items():
        el.set(_W + k, v)
    return el


def docx_style(path: Path, style: str, *, color: Optional[str] = None, size_pt: Optional[float] = None, font: Optional[str] = None,
               bold: Optional[bool] = None, italic: Optional[bool] = None, underline: Optional[bool] = None, spacing: Optional[int] = None) -> None:
    def fn(root):
        rpr = _rpr(_style_el(root, style))
        if color is not None:
            _set_child(rpr, "color", val=color.lstrip("#"))
        if size_pt is not None:
            _set_child(rpr, "sz", val=str(int(size_pt * 2)))
        if font is not None:
            _set_child(rpr, "rFonts", ascii=font, hAnsi=font, cs=font, eastAsia=font)
        if bold is not None:
            _set_child(rpr, "b", val="1" if bold else "0")
        if italic is not None:
            _set_child(rpr, "i", val="1" if italic else "0")
        if underline is not None:
            _set_child(rpr, "u", val="single" if underline else "none")
        if spacing is not None:
            _set_child(rpr, "spacing", val=str(spacing))
    rewrite_part(path, "word/styles.xml", fn)


def docx_table_header(path: Path, *, fill: Optional[str] = None, color: Optional[str] = None) -> None:
    def fn(root):
        for tbl in root.iter(_W + "tbl"):
            tr = tbl.find(_W + "tr")
            for tc in tr.findall(_W + "tc"):
                if fill is not None:
                    tcpr = tc.find(_W + "tcPr")
                    if tcpr is None:
                        tcpr = etree.SubElement(tc, _W + "tcPr")
                        tc.insert(0, tcpr)
                    _set_child(tcpr, "shd", val="clear", color="auto", fill=fill.lstrip("#"))
                if color is not None:
                    for r in tc.iter(_W + "r"):
                        _set_child(_rpr(r), "color", val=color.lstrip("#"))
    rewrite_part(path, "word/document.xml", fn)


def docx_section(path: Path, *, landscape: Optional[bool] = None, size: Optional[Tuple[int, int]] = None, margins_twips: Optional[int] = None) -> None:
    def fn(root):
        for sect in root.iter(_W + "sectPr"):
            pgsz = sect.find(_W + "pgSz")
            w, h = int(pgsz.get(_W + "w")), int(pgsz.get(_W + "h"))
            if size is not None:
                w, h = size
            if landscape is not None:
                short, long_ = sorted((w, h))
                w, h = (long_, short) if landscape else (short, long_)
                if landscape:
                    pgsz.set(_W + "orient", "landscape")
                elif _W + "orient" in pgsz.attrib:
                    del pgsz.attrib[_W + "orient"]
            pgsz.set(_W + "w", str(w))
            pgsz.set(_W + "h", str(h))
            if margins_twips is not None:
                pgmar = sect.find(_W + "pgMar")
                for k in ("top", "bottom", "left", "right"):
                    pgmar.set(_W + k, str(margins_twips))
    rewrite_part(path, "word/document.xml", fn)


def docx_strip_page_fields(path: Path) -> None:
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if n.startswith("word/footer") or n.startswith("word/header")]
    for part in parts:
        def fn(root):
            for el in list(root.iter(_W + "instrText")):
                el.text = " DATE "
            for el in list(root.iter(_W + "fldSimple")):
                el.set(_W + "instr", "DATE")
        rewrite_part(path, part, fn)


def docx_external_link(path: Path, target: str) -> None:
    def fn(root):
        rel = etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
        rel.set("Id", "rIdInjected9")
        rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink")
        rel.set("Target", target)
        rel.set("TargetMode", "External")
    rewrite_part(path, "word/_rels/document.xml.rels", fn)


def docx_delete_paragraph_text(path: Path, contains: str) -> None:
    def fn(root):
        for p in root.iter(_W + "p"):
            text = "".join(t.text or "" for t in p.iter(_W + "t"))
            if contains in text:
                for t in p.iter(_W + "t"):
                    t.text = ""
    rewrite_part(path, "word/document.xml", fn)


def docx_replace_text(path: Path, old: str, new: str) -> None:
    def fn(root):
        for t in root.iter(_W + "t"):
            if t.text and old in t.text:
                t.text = t.text.replace(old, new)
    rewrite_part(path, "word/document.xml", fn)


# ------------------------------------------------------- workbook / deck --


def xlsx_edit(path: Path, fn: Callable) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(path))
    fn(wb)
    wb.save(str(path))


def pptx_edit(path: Path, fn: Callable) -> None:
    from pptx import Presentation

    prs = Presentation(str(path))
    fn(prs)
    prs.save(str(path))


# --------------------------------------------------------------------- PDF --


def pdf_edit_text(path: Path, match: Callable[[str], bool], *, rgb: Tuple[int, int, int]) -> int:
    """Recolour every text object whose text satisfies `match`, regenerate
    the page content and save the PDF in place. Returns objects changed."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as R

    pdf = pdfium.PdfDocument(str(path))
    changed = 0
    for i in range(len(pdf)):
        page = pdf[i]
        tp = page.get_textpage()
        n = R.FPDFPage_CountObjects(page.raw)
        for j in range(n):
            obj = R.FPDFPage_GetObject(page.raw, j)
            if R.FPDFPageObj_GetType(obj) != R.FPDF_PAGEOBJ_TEXT:
                continue
            ln = R.FPDFTextObj_GetText(obj, tp.raw, None, 0)
            buf = ctypes.create_string_buffer(ln * 2)
            R.FPDFTextObj_GetText(obj, tp.raw, ctypes.cast(buf, ctypes.POINTER(ctypes.c_ushort)), ln)
            text = buf.raw.decode("utf-16-le", "ignore").rstrip("\x00")
            if match(text):
                R.FPDFPageObj_SetFillColor(obj, rgb[0], rgb[1], rgb[2], 255)
                changed += 1
        R.FPDFPage_GenerateContent(page.raw)
    out = io.BytesIO()
    pdf.save(out)
    pdf.close()
    Path(path).write_bytes(out.getvalue())
    return changed


def pdf_swap_mediabox(path: Path) -> None:
    import pypdfium2 as pdfium
    import pypdfium2.raw as R

    pdf = pdfium.PdfDocument(str(path))
    for i in range(len(pdf)):
        page = pdf[i]
        l, b, r, t = page.get_mediabox()
        R.FPDFPage_SetMediaBox(page.raw, l, b, l + (t - b), b + (r - l))
    out = io.BytesIO()
    pdf.save(out)
    pdf.close()
    Path(path).write_bytes(out.getvalue())


def copy_files(files: Dict[str, Path], dest: Path) -> Dict[str, Path]:
    dest.mkdir(parents=True, exist_ok=True)
    out = {}
    for name, p in files.items():
        q = dest / name
        shutil.copyfile(p, q)
        out[name] = q
    return out
