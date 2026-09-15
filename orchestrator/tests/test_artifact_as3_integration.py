"""AS3 integration: the seams between the five tracks (intent-capability,
styling-engine, charts, prompt-edits, agentic-selfcheck) once they run in
one tree. Each test pins a defect the combined tree had before the
integration commit, or a wire the tracks left for it."""
from __future__ import annotations

import asyncio
import json
import re
import zipfile
from pathlib import Path

import pytest

from app.artifacts import chart_data as CD
from app.artifacts import chart_spec as CS
from app.artifacts import edits as E
from app.artifacts import formats as F
from app.artifacts import pipeline, selfcheck as SC, store
from app.artifacts import spec as S
from app.artifacts import style as ST
from app.artifacts import types as T
from app.artifacts.render import render_version
from app.config import settings
from tests.fixtures.charts import loader
from tests.test_artifact_jobs import _accept, _run, isolated, owner  # noqa: F401 — fixtures by name


@pytest.fixture(scope="module")
def tables():
    return [loader.table(n) for n in loader.FILES]


# ------------------------------------------------------------- Chart v2 --


def test_the_spec_chart_is_chart_v2_and_old_literal_charts_still_load():
    assert S.Chart is CS.Chart
    old = S.parse_body("document", {"title": "Old", "blocks": [
        {"type": "chart", "chart": {"type": "bar", "title": "T", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}}]})
    assert old.body.blocks[0].chart.is_literal
    new = S.parse_body("document", {"title": "New", "blocks": [
        {"type": "chart", "chart": {"type": "heatmap", "title": "H", "data": {"table_id": "upload1", "x": "Status", "group_by": "Priority", "agg": "count"}}}]})
    assert new.body.blocks[0].chart.data.table_id == "upload1"


@pytest.mark.parametrize("kind", ["document", "presentation", "workbook"])
def test_the_guided_schema_never_offers_a_plotted_number(kind):
    schema = S.schema_for(kind)
    chart = schema["$defs"]["Chart"]
    assert not {"categories", "series", "extra", "provenance"} & set(chart["properties"])
    assert set(chart["required"]) >= {"type", "title", "data"}
    assert set(schema["$defs"]["ChartStyle"]["properties"]) <= set(CS.GUIDED_STYLE_FIELDS)


def test_v2_charts_render_natively_in_xlsx_and_pptx_and_as_images_in_docx_pdf(tmp_path, tables):
    charts = [
        {"type": "stacked_bar", "title": "By region", "data": {"table_id": "upload_sales", "x": "Product", "group_by": "Region", "y": ["Amount"]}},
        {"type": "heatmap", "title": "Heat", "data": {"table_id": "upload_tickets", "x": "Status", "group_by": "Priority", "agg": "count"}},
    ]
    doc = S.parse_body("document", {"title": "Doc", "blocks": [{"type": "chart", "chart": c} for c in charts]})
    doc, _ = CD.resolve_spec(doc, tables)
    for sub in ("d", "p", "w"):
        (tmp_path / sub).mkdir()
    rep = render_version(doc, ["docx", "pdf"], str(tmp_path / "d"), title_slug="doc", version=1)
    with zipfile.ZipFile(tmp_path / "d" / "doc-v1.docx") as z:
        assert len([n for n in z.namelist() if n.startswith("word/media/")]) == 2
    assert {f.format for f in rep.files} == {"docx", "pdf"}

    deck = S.parse_body("presentation", {"title": "Deck", "slides": [{"layout": "chart", "title": c["title"], "chart": c} for c in charts]})
    deck, _ = CD.resolve_spec(deck, tables)
    render_version(deck, ["pptx", "pdf"], str(tmp_path / "p"), title_slug="deck", version=1)
    with zipfile.ZipFile(tmp_path / "p" / "deck-v1.pptx") as z:
        names = z.namelist()
    assert len([n for n in names if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)]) == 1, "stacked bar native"
    assert any(n.startswith("ppt/media/") for n in names), "heatmap is a picture"

    t = next(t for t in tables if t.id == "upload_tickets")
    wb = S.parse_body("workbook", {"title": "WB", "sheets": [{
        "name": "Tickets", "columns": [{"name": c} for c in t.columns],
        "rows": [[str(c) if hasattr(c, "isoformat") else c for c in r] for r in t.rows],
        "charts": [{"type": "donut", "title": "Status", "data": {"table_id": "", "x": "Status", "agg": "count"}}]}]})
    wb, _ = CD.resolve_spec(wb, tables)
    render_version(wb, ["xlsx"], str(tmp_path / "w"), title_slug="wb", version=1)
    from openpyxl import load_workbook

    book = load_workbook(tmp_path / "w" / "wb-v1.xlsx")
    assert book.sheetnames[-1] == "Chart data"
    rows = [r for r in book["Chart data"].iter_rows(values_only=True) if r[0] is not None]
    counts = {}
    for r in t.rows:
        counts[r[t.columns.index("Status")]] = counts.get(r[t.columns.index("Status")], 0) + 1
    assert {r[0]: r[1] for r in rows[2:]} == counts, "values recomputed by code"
    with zipfile.ZipFile(tmp_path / "w" / "wb-v1.xlsx") as z:
        assert "doughnutChart" in z.read("xl/charts/chart1.xml").decode()


def test_standalone_chart_images_are_live_and_never_named_like_scratch(tmp_path, tables):
    spec = S.parse_body("document", {"title": "Img", "blocks": [
        {"type": "chart", "chart": {"type": "pie", "title": "Status", "data": {"table_id": "upload_tickets", "x": "Status", "agg": "count"}}}]})
    spec, _ = CD.resolve_spec(spec, tables)
    (tmp_path / "i").mkdir()
    (tmp_path / "j").mkdir()
    rep = render_version(spec, ["png"], str(tmp_path / "i"), title_slug="img", version=2)
    assert [f.filename for f in rep.files] == ["img-v2-chart-1.png"] and rep.preview_kind == "image"
    rep = render_version(spec, ["docx", "svg"], str(tmp_path / "j"), title_slug="img", version=2)
    names = {f.filename for f in rep.files}
    assert "img-v2-chart-1.svg" in names and not set(rep.chart_files) & names, "publish deletes chart_files as scratch"


@pytest.mark.parametrize("text, kind, formats", [
    ("make a bar chart of sales by month as a png", "document", ["png"]),
    ("give me a pie chart of this as svg", "document", ["svg"]),
    ("bar chart bana do png me", "document", ["png"]),
    ("a word report with a line chart, plus the chart as png", "document", ["docx", "pdf", "png"]),
    ("excel sheet with a pie chart of status and the chart as png", "workbook", ["xlsx", "png"]),
])
def test_chart_image_formats_are_read_next_to_chart_words(text, kind, formats):
    d = F.decide(text)
    assert (d.kind, d.formats) == (kind, formats), d


@pytest.mark.parametrize("text", ["draw a logo png for my bakery", "a pdf report on png compression"])
def test_a_png_without_chart_words_is_not_a_chart_image(text):
    assert not set(F.decide(text).formats) & set(T.IMAGE_FORMATS)


def test_chart_style_fonts_are_canonical_and_a_bare_three_letter_word_is_no_colour():
    assert CS.allowed_font("calibri") == "Calibri" and CS.allowed_font("GEORGIA") == "Georgia"
    assert CS.resolve_color("bed") is None and CS.resolve_color("#bed") == "#BBEEDD" and CS.resolve_color("navy")


# -------------------------------------------------------------- styling --


class _Ctx:
    def __init__(self, instruction, kind, formats, effort="fast"):
        self.instruction, self.kind, self.formats, self.effort, self.operation = instruction, kind, formats, effort, "create"
        self.warnings = []

    def warn(self, text):
        self.warnings.append(text)


def test_a_create_applies_the_requested_styling_to_the_spec(monkeypatch):
    from app.engines import artifact as engine

    async def no_llm(*a, **k):  # every phrase here is read by the parser
        raise AssertionError("no style model call expected")

    monkeypatch.setattr(ST, "extract_patch_llm", no_llm)
    doc = S.parse_body("document", {"title": "Retention", "blocks": [{"type": "heading", "level": 1, "text": "Scope"}, {"type": "paragraph", "text": "x"}]})
    ctx = _Ctx("Word report on data retention with purple headings and a dark green title in Georgia 28pt", "document", ["docx"])
    out = asyncio.run(engine._apply_requested_style(ctx, doc))
    R = ST.resolve(out)
    assert R.element("heading", level=1).color == "#6D5AE6"
    title = R.element("title")
    assert (title.color, title.font_family, title.size_pt) == ("#1E6B34", "Georgia", 28.0)


def test_a_styled_csv_request_also_makes_the_excel_file():
    d = F.decide("csv of 20 orders with a blue header row")
    assert d.formats == ["csv", "xlsx"] and "formatting is in the Excel" in d.data_only_note


# ---------------------------------------------------------------- edits --


@pytest.mark.parametrize("text, undo", [
    ("undo", True), ("undo that please", True), ("पिछला बदलाव हटा दो", True), ("go back to the previous version", True),
    ("restore version 1", False), ("make headings blue like the previous version", False),
    ("make it look like the last version but with a red title", False),
])
def test_an_undo_is_the_whole_request_never_an_undo_word_inside_an_edit(text, undo):
    assert E.undo_signal(text) is undo


def test_restore_version_is_a_restore_not_an_undo():
    from tests.test_artifact_edits import nine_section_doc

    plan = E.preplan("restore version 1", nine_section_doc())
    assert [o.op for o in plan.ops] == ["restore_version"]


@pytest.mark.parametrize("text", ["delete rows where status is Closed", "add a column for owner", "make it landscape"])
def test_no_empty_or_duplicate_style_op_rides_on_other_edits(text):
    from tests.test_artifact_edits import four_sheet_workbook, nine_section_doc

    parent = nine_section_doc() if "landscape" in text else four_sheet_workbook()
    plan = E.preplan(text, parent)
    assert plan is None or "set_style" not in [o.op for o in plan.ops], plan


# ------------------------------------------------------------ selfcheck --


def test_a_literal_chart_is_not_a_failed_values_check_and_a_sheet_chart_is_recomputed(tables):
    literal = S.parse_body("document", {"title": "Old", "blocks": [
        {"type": "chart", "chart": {"type": "bar", "title": "T", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}}]})
    assert SC.chart_values_check(literal, [], []) in (None, (True, []))
    wb = S.parse_body("workbook", {"title": "WB", "sheets": [{
        "name": "T", "columns": [{"name": "Status"}, {"name": "Hours", "type": "number"}],
        "rows": [["Open", 2], ["Closed", 3], ["Open", 5]],
        "charts": [{"type": "bar", "title": "H", "data": {"table_id": "", "x": "Status", "y": ["Hours"], "agg": "sum"}}]}]})
    wb, _ = CD.resolve_spec(wb, [])
    assert SC.chart_values_check(wb, [], []) in (None, (True, [])), "no native file observed: nothing failed"
    bad = wb.model_copy(deep=True)
    bad.body.sheets[0].charts[0].series[0].values[0] = 99.0
    ok, diffs = SC.chart_values_check(bad, [], [])
    assert not ok and diffs


def test_a_create_job_publishes_the_requested_style_in_the_file(owner, monkeypatch, tmp_path):  # noqa: F811
    """End to end through the pipeline with the REAL engine composer (the
    model call stubbed): the DOCX carries the requested heading colour and
    the self-check finds it met."""
    from app.artifacts import compose as C
    from app.engines import artifact as engine

    async def render(work_dir, spec, formats, title_slug, version, effort, **kw):
        report = await asyncio.to_thread(render_version, spec, formats, work_dir, title_slug=title_slug, version=version, effort=effort)
        return report.to_json()

    doc = S.parse_body("document", {"title": "Data Retention", "blocks": [
        {"type": "heading", "level": 1, "text": "Scope"}, {"type": "paragraph", "text": "What is kept and for how long."},
        {"type": "heading", "level": 1, "text": "Rules"}, {"type": "paragraph", "text": "Seven years for invoices."}]})

    async def fake_compose(req, progress=None):
        return C.ComposeResult(spec=doc) if hasattr(C, "ComposeResult") else type("R", (), {"spec": doc, "warnings": [], "corrections": 0, "transform": {}})()

    monkeypatch.setattr(pipeline, "_render_in_subprocess", render)
    monkeypatch.setattr(C, "compose", fake_compose)
    monkeypatch.setattr(settings, "artifact_selfcheck", True)
    pipeline.set_composer(engine.compose_for_pipeline)
    row = _accept(owner, instruction="Word report on data retention with purple headings", formats=["docx"], format_reason="explicit: word",
                  template_id="generic", material={"history_text": "user: Word report on data retention with purple headings"})
    fresh = _run(row["id"])
    assert fresh["status"] in ("completed", "completed_with_warnings"), fresh
    vdir = Path(store.version_dir(owner, row["artifact_id"], 1))
    spec = json.loads((vdir / "spec.json").read_text())
    assert any(r["target"]["kind"] == "heading" and r["style"].get("color") == "#6D5AE6" for r in spec["document"]["style"]["rules"])
    report = json.loads((vdir / SC.SELFCHECK_NAME).read_text())
    assert not any("headings" in u for u in report["unmet"]) and report["outcome"] in ("clean", "repaired"), report


def test_an_export_of_the_answer_publishes_through_the_real_composer(owner, monkeypatch):  # noqa: F811
    """The import payload is the DocumentSpec BODY (md_import); the composer
    must wrap it in the envelope. Before the integration fix every export
    failed "The content could not be written" (live run 2026-09-15)."""
    from app.artifacts import md_import
    from app.engines import artifact as engine

    async def render(work_dir, spec, formats, title_slug, version, effort, **kw):
        report = await asyncio.to_thread(render_version, spec, formats, work_dir, title_slug=title_slug, version=version, effort=effort)
        return report.to_json()

    monkeypatch.setattr(pipeline, "_render_in_subprocess", render)
    pipeline.set_composer(engine.compose_for_pipeline)
    doc, notes = md_import.markdown_to_document("# Audit\n\n## Findings\n\n| Area | Status |\n|---|---|\n| Keys | Open |\n\n## Next steps\n- Rotate keys\n")
    instruction = "give it in docs in a classy format with dark blue headings"
    row = _accept(owner, instruction=instruction, formats=["docx"], format_reason=f"{engine.IMPORT_REASON}: previous answer · explicit: docx",
                  template_id="generic", material={"history_text": "", "previous_answer": "# Audit"})
    engine._attach_payload(owner, row["artifact_id"], int(row["version"]), "import",
                           {"spec": doc.model_dump(mode="json", by_alias=True, exclude_none=True), "notes": list(notes)})
    fresh = _run(row["id"])
    assert fresh["status"] in ("completed", "completed_with_warnings"), fresh
    spec = json.loads((Path(store.version_dir(owner, row["artifact_id"], 1)) / "spec.json").read_text())
    assert [b["text"] for b in spec["document"]["blocks"] if b["type"] == "heading"][:2] == ["Findings", "Next steps"] or spec["document"]["title"] == "Audit"
    assert any(r["target"]["kind"] == "heading" and r["style"].get("color") == "#1F3864" for r in spec["document"]["style"]["rules"])


@pytest.mark.parametrize("formats", [["png"], ["svg"], ["docx", "png"]])
def test_a_chart_image_version_publishes_one_file_per_chart(owner, monkeypatch, tables, formats):  # noqa: F811
    """Live run 2026-09-15: every png/svg-only job failed publication with
    "the png file is not named <title>-v1.png" — the pipeline keyed files by
    (role, format) and a chart image is one file per chart."""
    from tests.test_artifact_jobs import _composer

    async def render(work_dir, spec, formats_, title_slug, version, effort, **kw):
        report = await asyncio.to_thread(render_version, spec, formats_, work_dir, title_slug=title_slug, version=version, effort=effort)
        return report.to_json()

    spec = S.parse_body("document", {"title": "Ticket Mix", "blocks": [
        {"type": "chart", "chart": {"type": "pie", "title": "Status", "data": {"table_id": "upload_tickets", "x": "Status", "agg": "count"}}},
        {"type": "chart", "chart": {"type": "bar", "title": "Owners", "data": {"table_id": "upload_tickets", "x": "Owner", "agg": "count"}}}]})
    spec, _ = CD.resolve_spec(spec, tables)
    monkeypatch.setattr(pipeline, "_render_in_subprocess", render)
    pipeline.set_composer(_composer(spec))
    row = _accept(owner, instruction="pie chart of status as an image", formats=formats, format_reason="explicit: chart image", template_id="generic")
    fresh = _run(row["id"])
    assert fresh["status"] in ("completed", "completed_with_warnings"), fresh
    from app.artifacts import db as adb

    version = adb.get_version(row["artifact_id"], 1, owner)
    names = sorted(f["filename"] for f in version["files"])
    images = [n for n in names if n.rsplit(".", 1)[-1] in T.IMAGE_FORMATS]
    assert images == sorted(f"ticket-mix-v1-chart-{n}.{formats[-1]}" for n in (1, 2)), names
    assert len({f["file_id"] for f in version["files"]}) == len(version["files"])
    vdir = Path(store.version_dir(owner, row["artifact_id"], 1))
    assert all((vdir / n).is_file() for n in names), "published, not swept as scratch"


# ------------------------------------------------- live-run 2026-09-15 fixes --


@pytest.mark.parametrize("text, expected", [
    ("make Status red where Blocked and green where Done", {"cond:status:eq:blocked:background": "#C62828", "cond:status:eq:done:background": "#3F8F4F"}),
    ("Priority orange where Critical", {"cond:priority:eq:critical:background": "#E07B00"}),
    ("Growth % column: negative in red, positive in green", {"cond:growth %:lt:0.0:background": "#C62828", "cond:growth %:gt:0.0:background": "#3F8F4F"}),
    ("Overdue wale red me", {"cond:status:eq:overdue:background": "#C62828"}),
    ("put a colour scale on the Score column", {"scale:score": True}),
    ("the Risks section paragraphs in italic", {"rule:paragraph|section=risks:italic": True}),
    ("headings in Arial 16", {"rule:heading:font_family": "Arial", "rule:heading:size_pt": 16.0}),
])
def test_the_style_parser_reads_the_live_phrasings(text, expected):
    kind = "document" if ("section" in text or "headings" in text) else "workbook"
    patch, _unparsed = ST.parse_style_request(text, kind)
    got = ST.patch_fields(patch)
    assert all(got.get(k) == v for k, v in expected.items()), got


@pytest.mark.parametrize("text", ["make the header red where possible", "make the title red for the launch", "use Calibri 2 times", "a report on the red team exercise"])
def test_the_new_readings_do_not_invent_conditions_or_sizes(text):
    got = ST.patch_fields(ST.parse_style_request(text, "workbook")[0])
    assert not any(k.startswith("cond:") for k in got) and not any(k.endswith("size_pt") or k == "base_size_pt" for k in got), got


def test_the_style_model_can_say_conditions_and_scales_and_never_leaks_prose():
    patch, notes = ST.patch_from_llm_object({
        "rules": [], "conditional": [{"column": "Status", "op": "eq", "value": "Blocked", "background": "red"}, {"column": "Growth %", "op": "lt", "value": "0", "color": "#C62828"}],
        "scales": [{"column": "Score"}],
        "not_understood": ["The request implies conditional formatting based on cell content. It cannot be expressed.", "wavy borders"]}, "workbook")
    got = ST.patch_fields(patch)
    assert got["cond:status:eq:blocked:background"] == "#C62828" and got["cond:growth %:lt:0.0:color"] == "#C62828" and got["scale:score"] is True
    assert notes == ["Not applied: 'wavy borders'."]


def test_a_requested_header_colour_beats_the_models_column_highlight(tmp_path):
    from openpyxl import load_workbook

    wb = S.parse_body("workbook", {"title": "Costs", "sheets": [{
        "name": "Costs", "columns": [{"name": "Item"}, {"name": "Q1", "type": "number"}],
        "rows": [["Rent", 1200], ["Power", 340]], "style": {"highlight": [{"column": "Item", "color": "blue"}]}}]})
    wb, _f, _n, _u = ST.apply_request(wb, "excel with a yellow header row", formats=["xlsx"])
    (tmp_path / "w").mkdir()
    render_version(wb, ["xlsx"], str(tmp_path / "w"), title_slug="costs", version=1)
    ws = load_workbook(tmp_path / "w" / "costs-v1.xlsx").worksheets[0]
    assert {c.fill.fgColor.rgb[-6:] for c in ws[1]} == {"FFD54F"}


def test_the_models_legacy_sheet_style_gives_way_to_the_house_and_the_request():
    from app.engines import artifact as engine

    wb = S.parse_body("workbook", {"title": "V", "sheets": [{
        "name": "Vendors", "columns": [{"name": "Vendor"}, {"name": "Owner"}], "rows": [["A", "B"]],
        "style": {"highlight": [{"column": "Vendor", "color": "blue"}], "header_fill": "light"}}]})
    engine._house_sheet_style(wb, "excel sheet of vendors, Vendor column bold", styled=True)
    assert wb.body.sheets[0].style.highlight == [] and wb.body.sheets[0].style.header_fill == "dark"
    wb2 = S.parse_body("workbook", {"title": "V", "sheets": [{
        "name": "Vendors", "columns": [{"name": "Vendor"}], "rows": [["A"]], "style": {"header_fill": "light"}}]})
    engine._house_sheet_style(wb2, "a plain sheet with a light header", styled=False)
    assert wb2.body.sheets[0].style.header_fill == "light"


def test_the_binding_repair_runs_before_the_numbers_are_computed(tables):
    from app.engines import artifact as engine

    spec = S.parse_body("document", {"title": "H", "blocks": [{"type": "chart", "chart": {
        "type": "heatmap", "title": "Ticket Count by Status and Priority", "data": {"table_id": "upload_tickets", "x": "Status", "agg": "count"}}}]})
    warned = []
    out = asyncio.run(engine._post_process(spec, tables, warned.append, "heatmap of ticket count by Status and Priority in a pdf"))
    chart = out.body.blocks[0].chart
    assert chart.data.group_by == "Priority" and chart.series and any("Priority" in w for w in warned)


def test_a_documents_chart_binding_is_an_enum_of_real_columns(tables):
    schema = S.schema_for("document", tables=tables[:1])
    binding = schema["$defs"]["Binding"]["properties"]
    assert binding["table_id"]["enum"] == [tables[0].id] and "Month" not in binding["x"]["enum"]
    assert "enum" not in json.dumps(S.schema_for("workbook", tables=tables[:1])["$defs"]["Binding"]["properties"]["x"])


@pytest.mark.parametrize("text, create", [
    ("scatter of Salary vs Experience with a trend line in a pdf", True),
    ("pie of tickets by status", True),
    ("the pie of my dreams", False),
    ("funnel of the marketing process explained", False),
])
def test_a_chart_type_named_as_a_noun_is_a_chart_request(text, create):
    from app.artifacts import intent as I

    assert (I.decide(text, upload_formats=["csv"]).action == "create") is create


def test_a_conditionally_filled_column_drops_the_red_negative_number_format(tmp_path):
    from openpyxl import load_workbook

    wb = S.parse_body("workbook", {"title": "Growth", "sheets": [{
        "name": "Sales", "columns": [{"name": "Region"}, {"name": "Growth %", "type": "number"}], "rows": [["North", 20.8], ["East", -13.3]]}]})
    wb, _f, _n, _u = ST.apply_request(wb, "show negative Growth % in red", formats=["xlsx"])
    (tmp_path / "w").mkdir()
    render_version(wb, ["xlsx"], str(tmp_path / "w"), title_slug="growth", version=1)
    ws = load_workbook(tmp_path / "w" / "growth-v1.xlsx").worksheets[0]
    assert "[Red]" not in ws["B3"].number_format and ws["B3"].value == -13.3
