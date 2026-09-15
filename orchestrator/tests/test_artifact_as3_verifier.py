"""AS3 adversarial verification of the integrated layer (2026-09-15).

Each test failed on the integrated tree before its fix:

- super-linear text readers on the event loop: the checklist extractor
  (10 s on a 4,000-character instruction), the composer's style-clause strip
  (quadratic), the style parser's clause split on a whitespace run (cubic),
  the typed-figures reader over a 20 KB message (11 s);
- typed aggregates that still reached a published workbook: a figure in a
  column the model declared `date`, a KPI line "Total Amount: 1,047,638";
- an edit that added a typed column to a computed sheet dropped the sheet;
- ordinary chat in a conversation holding a file edited that file ("give me
  bullet points on climate change"), and a story's plot made a chart file;
- a totals row added by a style rule came back on every render.

The data is written here; no model is called.
"""
from __future__ import annotations

import random
import time

import pytest

from app.artifacts import chart_data as CD
from app.artifacts import compose as C
from app.artifacts import derived as D
from app.artifacts import intent as I
from app.artifacts import requirements as RQ
from app.artifacts import spec as S
from app.artifacts import style as ST
from tests.fixtures.charts import loader


def _elapsed(fn, *args, **kwargs) -> float:
    t0 = time.perf_counter()
    fn(*args, **kwargs)
    return time.perf_counter() - t0


# ------------------------------------------------------ bounded readers --


def test_the_checklist_extractor_is_linear_on_a_long_clause_of_repeated_parts():
    text = "row 3 A1:B2 #FF00FF " * 200  # 4,000 characters: the decision cap
    small = RQ.extract_rules("row 3 A1:B2 #FF00FF " * 2, kind="workbook", operation="create")
    assert _elapsed(RQ.extract_rules, text, kind="workbook", operation="create") < 2.0
    assert [i.key() for i in RQ.extract_rules(text, kind="workbook", operation="create")] == [i.key() for i in small]


def test_the_composer_strips_style_clauses_exactly_as_the_pattern_did_and_in_linear_time():
    rng = random.Random(7)
    words = ("bold italic font 12 pt 12pt with colours colour colored color for of in red text dark blue navy report, "
             "section; findings. background colour fill\n Georgia a the and highlighted headings header 3 points scheme").split(" ")
    seps = [" ", " ", ", ", ". ", "; ", "\n", "  ", ".", ",", "\n\n", "\t"]
    for _ in range(5000):
        text = "".join(rng.choice(words) + rng.choice(seps) for _ in range(rng.randint(0, 16)))
        assert C.strip_style_clauses(text) == C._STYLE_CLAUSE_RE.sub(" ", text), repr(text)
    assert _elapsed(C.strip_style_clauses, "1 " * 4000) < 1.0
    assert _elapsed(C.requested_sections, "a " * 4000) < 1.0


def test_the_style_parser_reads_a_long_whitespace_run_in_linear_time():
    assert _elapsed(ST.parse_style_request_with_notes, " " * 2000, "workbook") < 2.0
    patch, unparsed = ST.parse_style_request("make the header row bold   and   the title navy", "workbook")
    assert not patch.is_empty() and unparsed == []


def test_the_typed_figures_reader_is_bounded_on_a_long_message_and_still_reads_a_series():
    assert _elapsed(CD.parse_prompt_data, "Jan " * 5000) < 2.0
    assert _elapsed(CD.parse_prompt_data, "make it " * 2500) < 2.0
    table = CD.parse_prompt_data("sales by region: north 120; south 95; east 70")
    assert table is not None and [r[1] for r in table.rows] == [120, 95, 70]


# ------------------------------------------------------ typed aggregates --


@pytest.fixture()
def sales():
    return loader.table("sales", "upload1")


def _figures(spec) -> list:
    return [c for sh in spec.body.sheets for r in sh.rows for c in r if isinstance(c, (int, float)) or isinstance(c, str)]


def test_a_figure_in_a_column_declared_date_is_not_published(sales):
    spec = S.parse_body("workbook", {"title": "T", "sheets": [{
        "name": "Summary", "columns": [{"name": "Region", "type": "text"}, {"name": "Revenue", "type": "date"}],
        "rows": [["North", 123456], ["South", 98765]]}]})
    out, notes, report = D.enforce(spec, [sales])
    assert 123456 not in _figures(out) and 98765 not in _figures(out)
    assert report["dropped"] == ["Summary"]


def test_a_kpi_line_with_a_typed_total_is_not_published(sales):
    spec = S.parse_body("workbook", {"title": "T", "sheets": [{
        "name": "KPIs", "columns": [{"name": "KPI", "type": "text"}],
        "rows": [["Total Amount: 1,047,638"], ["Average daily amount: 234,567"]]}]})
    out, notes, report = D.enforce(spec, [sales])
    assert not any("1,047,638" in str(c) for c in _figures(out))
    assert report["dropped"] == ["KPIs"]
    # A text sheet with no aggregate figure is left alone.
    plain = S.parse_body("workbook", {"title": "T", "sheets": [{
        "name": "Notes", "columns": [{"name": "Note", "type": "text"}], "rows": [["Regions are as recorded in the upload"]]}]})
    assert D.enforce(plain, [sales])[1] == []


def test_an_edit_that_types_a_new_column_on_a_computed_sheet_keeps_the_sheet(sales):
    draft = S.parse_body("workbook", {"title": "T", "sheets": [{
        "name": "Monthly", "columns": [{"name": "Month", "type": "text"}, {"name": "Total Amount", "type": "number"}],
        "rows": [["Jan-2026", 1], ["Feb-2026", 1]]}]})
    parent, _notes, _report = D.enforce(draft, [sales])
    computed = parent.body.sheets[0]
    assert computed.computed_from == "upload1"
    total = sum(r[1] for r in computed.rows)
    edited = computed.model_dump()
    edited["computed_from"] = None
    edited["columns"] = edited["columns"] + [{"name": "Share %", "type": "percent"}]
    edited["rows"] = [r + [round(r[1] / total, 4)] for r in edited["rows"]]
    child = S.parse_body("workbook", {"title": "T", "sheets": [edited]})
    out, notes, report = D.enforce(child, [sales], parent=parent)
    assert [s.name for s in out.body.sheets] == ["Monthly"], notes
    sheet = out.body.sheets[0]
    assert [c.name for c in sheet.columns] == ["Month", "Total Amount"] and sheet.rows == computed.rows
    assert sheet.computed_from == "upload1"
    assert any("'Share %'" in n and "left out" in n for n in notes)


# ------------------------------------------------------------- the gate --


@pytest.mark.parametrize("text", [
    "give me bullet points on climate change", "highlight the main takeaways from our discussion",
    "bold claim: AI will replace jobs, discuss", "summarize the key points in italics", "write the answer in bold",
    "highlight the key risks in this contract",
])
def test_ordinary_chat_in_a_conversation_with_a_file_is_not_an_edit_of_that_file(text):
    d = I.decide(text, has_artifacts=True, artifact_hints=["Sales Report"], has_assistant_answer=True)
    assert d.action == "none", (text, d.rule)


@pytest.mark.parametrize("text", ["make it bold", "make the headings dark blue", "highlight the overdue rows in red",
                                  "can you make this bold", "शीर्षकों को गहरा नीला कर दो"])
def test_a_style_edit_that_points_at_the_file_is_still_an_edit(text):
    d = I.decide(text, has_artifacts=True, artifact_hints=["Sales Report"], has_assistant_answer=True)
    assert d.action == "edit", (text, d.rule)


def test_right_after_a_file_card_a_bare_style_request_is_an_edit():
    d = I.decide("highlight the key risks", has_artifacts=True, has_assistant_answer=True, last_turn_is_artifact=True)
    assert d.action == "edit"


@pytest.mark.parametrize("text", ["plot of the movie Inception", "explain the graph of y=x^2", "what is the plot of the novel Dune"])
def test_a_story_plot_or_a_graph_to_explain_is_not_a_chart_file(text):
    assert I.decide(text).action == "none", text


@pytest.mark.parametrize("text", ["Histogram of hours spent per ticket.", "Box plot of hours by priority.", "plot of sales by month",
                                  "Heatmap of tickets by status and priority."])
def test_a_noun_first_chart_request_is_still_a_chart(text):
    assert I.decide(text).action == "create", text


# ------------------------------------------------------------ totals row --


def _tickets_with_total_style() -> S.ArtifactSpec:
    spec = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="Tickets", sheets=[
        S.Sheet(name="Tickets", columns=[S.Column(name="Ticket"), S.Column(name="Amount", type="number")], rows=[["A", 10], ["B", 20]]),
    ]))
    spec, _f, _notes, _u = ST.apply_request(spec, "bold total row", formats=["xlsx"])
    assert [t.column for t in spec.body.sheets[0].totals] == [1]
    return spec


def test_a_render_does_not_bring_back_a_totals_row_an_edit_removed():
    spec = _tickets_with_total_style()
    spec.body.sheets[0].totals = []  # an edit removed it; the style rule stays
    _, notes = ST.normalize_spec_style(spec, add_totals=False)
    assert spec.body.sheets[0].totals == [] and not any("had no totals row" in n for n in notes)


def test_rendering_draws_the_spec_as_saved(tmp_path):
    from openpyxl import load_workbook

    from app.artifacts.render import render_version

    spec = _tickets_with_total_style()
    spec.body.sheets[0].totals = []
    report = render_version(spec, ["xlsx"], str(tmp_path), title_slug="t", version=2)
    path = next(p for p in report.paths.values() if str(p).endswith(".xlsx"))
    ws = load_workbook(path)["Tickets"]
    assert ws.max_row == 3, [[c.value for c in r] for r in ws.iter_rows()]
