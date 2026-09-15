"""Native XLSX/PPTX charts (render/chart_native.py): the numbers in the FILE
(the Chart data sheet, the PPTX c:numCache) equal ground truth; styling lands
in the chart XML; image fallbacks are declared; label text is never a formula."""
from __future__ import annotations

import re
import zipfile

import pytest

pytest.importorskip("openpyxl")
pytest.importorskip("pptx")

from app.artifacts import chart_data as CD  # noqa: E402
from app.artifacts import chart_spec as CS  # noqa: E402
from app.artifacts.render import chart_native as N  # noqa: E402
from app.artifacts.render import validate as V  # noqa: E402
from tests.fixtures.charts import loader  # noqa: E402

GT = loader.ground_truth()

PIE = {"type": "pie", "title": "Status", "data": {"table_id": "upload_tickets", "x": "Status", "agg": "count"}}
LINE = {"type": "line", "title": "Monthly", "data": {"table_id": "upload_sales", "x": "Date", "y": ["Amount"]}}
STACK = {"type": "stacked_bar", "title": "H1", "data": {"table_id": "upload_sales", "x": "Date", "date_bucket": "quarter", "group_by": "Region",
                                                         "y": ["Amount"], "filters": [{"column": "Date", "op": "lt", "value": "2026-07-01"}]}}


@pytest.fixture(scope="module")
def tables():
    return [loader.table(n) for n in loader.FILES]


def resolved(raw, tables):
    c, _, msg = CD.resolve_chart(CS.Chart.model_validate(raw), tables)
    assert c is not None, msg
    return c


def workbook_with(charts, tmp_path, name="w.xlsx"):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Report"
    results = [N.add_xlsx_chart(ws, c, f"D{2 + 20 * i}") for i, c in enumerate(charts)]
    wb.create_sheet("Later")  # a sheet added after the charts
    N.keep_chart_data_last(wb)
    path = tmp_path / name
    wb.save(path)
    return path, results


def deck_with(charts, tmp_path, name="d.pptx"):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    results = []
    for c in charts:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        results.append(N.add_pptx_chart(slide, c, (0.5, 0.5, 12.0, 6.5)))
    path = tmp_path / name
    prs.save(path)
    return path, results


def chart_xmls(path, prefix):
    with zipfile.ZipFile(path) as zf:
        names = sorted((n for n in zf.namelist() if re.fullmatch(rf"{prefix}/charts/chart\d+\.xml", n)), key=lambda n: int(re.findall(r"\d+", n)[-1]))
        return [zf.read(n).decode("utf-8") for n in names]


def unprefixed(xml):
    """openpyxl writes the chart namespace as the default one; add the c:
    prefix to bare element names so both writers read the same."""
    return re.sub(r"<(/?)(?![a-z]+:)([A-Za-z]+)", r"<\1c:\2", xml)


def test_xlsx_chart_data_sheet_holds_exact_values(tmp_path, tables):
    charts = [resolved(r, tables) for r in (PIE, LINE, STACK)]
    path, results = workbook_with(charts, tmp_path)
    assert all(r["mode"] == "native" for r in results)
    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb.sheetnames[-1] == "Chart data" and wb["Chart data"].sheet_state == "visible"
    assert V.chart_data_mismatches(path, charts) == []
    rows = [r for r in wb["Chart data"].iter_rows(values_only=True)]
    # Pie block: title, header, five statuses.
    assert rows[0][0] == "Status" and dict((r[0], r[1]) for r in rows[2:7]) == GT["tickets"]["status_counts"]
    facts = V.validate_file(path, "xlsx", expected_charts=3)
    assert facts["charts"] == 3


def test_xlsx_tampered_value_is_detected(tmp_path, tables):
    charts = [resolved(LINE, tables)]
    path, _ = workbook_with(charts, tmp_path)
    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb["Chart data"]["B3"] = 64515
    wb.save(path)
    diffs = V.chart_data_mismatches(path, charts)
    assert diffs and "64515" in diffs[0]
    with pytest.raises(V.ValidationFailed, match="native charts"):
        V.validate_file(path, "xlsx", expected_charts=2)


def test_xlsx_styling_lands_in_the_chart_xml(tmp_path, tables):
    bar = resolved({"type": "bar", "title": "Revenue", "data": {"table_id": "upload_sales", "x": "Region", "y": ["Amount"]},
                    "style": {"color": "#8E44AD", "legend_position": "top", "y_min": 0, "y_max": 200000, "title": {"size_pt": 16}, "font_family": "Georgia"}}, tables)
    pie = resolved({**PIE, "style": {"category_colors": {"Resolved": "green", "Open": "red"}}}, tables)
    scatter = resolved({"type": "scatter", "title": "Pay", "data": {"table_id": "upload_employees", "x": "Experience", "y": ["Salary"], "trendline": True}}, tables)
    combo = resolved({"type": "combo", "title": "Units", "y2_label": "Q4", "data": {"table_id": "upload_units", "x": "Product", "y": ["Q1", "Q2"], "y2": ["Q4"]}}, tables)
    pct = resolved({"type": "percent_stacked_bar", "title": "Mix", "data": {"table_id": "upload_tickets", "x": "Owner", "group_by": "Priority", "agg": "count"}}, tables)
    path, _ = workbook_with([bar, pie, scatter, combo, pct], tmp_path)
    x_bar, x_pie, x_scatter, x_combo, x_pct = (unprefixed(x) for x in chart_xmls(path, "xl"))
    assert '<a:srgbClr val="8E44AD"' in x_bar
    assert 'formatCode="#,##0" sourceLinked="0"' in x_bar
    assert '<c:min val="0"' in x_bar and '<c:max val="200000"' in x_bar
    assert 'sz="1600"' in x_bar and 'typeface="Georgia"' in x_bar
    assert '<c:dLblPos val="outEnd"' in x_bar
    assert "3F8F4F" in x_pie and "C62828" in x_pie and "<c:dPt>" in x_pie
    assert "<c:trendline>" in x_scatter and '<c:trendlineType val="linear"' in x_scatter
    assert "<c:barChart>" in x_combo and "<c:lineChart>" in x_combo and '<c:crosses val="max"' in x_combo
    assert '<c:grouping val="percentStacked"' in x_pct and 'formatCode="0%"' in x_pct
    # One legend for a single series is none; legend_position top when asked (and more than one series).
    top = resolved({**STACK, "style": {"legend_position": "top"}}, tables)
    path2, _ = workbook_with([top], tmp_path, "w2.xlsx")
    assert '<c:legendPos val="t"' in unprefixed(chart_xmls(path2, "xl")[0])


def test_xlsx_image_fallback_is_declared_and_counted(tmp_path, tables):
    heat = resolved({"type": "heatmap", "title": "Heat", "data": {"table_id": "upload_tickets", "x": "Status", "group_by": "Priority", "agg": "count"}}, tables)
    pie = resolved(PIE, tables)
    path, results = workbook_with([heat, pie], tmp_path)
    assert [r["mode"] for r in results] == ["image", "native"] and "picture" in results[0]["note"]
    native, images = N.expected_native_counts([heat, pie], "xlsx")
    assert (native, images) == (1, 1)
    assert V.validate_file(path, "xlsx", expected_charts=native)["charts"] == 1
    with zipfile.ZipFile(path) as zf:
        assert any(n.startswith("xl/media/") for n in zf.namelist())


def test_xlsx_formula_leads_are_text(tmp_path):
    from app.artifacts.compose import DataTable
    from openpyxl import load_workbook

    t = DataTable(id="paste1", title="p", columns=["Label", "N"], rows=[["=HYPERLINK(\"http://x\")", 3], ["+cmd", 4], ["ok", 5]])
    c, _, _ = CD.resolve_chart(CS.Chart.model_validate({"type": "bar", "title": "=SUM(A1)", "data": {"table_id": "paste1", "x": "Label", "y": ["N"]}}), [t])
    path, _ = workbook_with([c], tmp_path)
    wb = load_workbook(path)
    ws = wb["Chart data"]
    cells = [ws.cell(row=r, column=1) for r in range(1, 6)]
    assert sum(1 for c in cells if isinstance(c.value, str) and c.value[:1] in "=+") == 3  # title + two labels
    for cell in cells:
        if isinstance(cell.value, str) and cell.value[:1] in "=+":
            assert cell.data_type == "s" and cell.quotePrefix, cell.coordinate
    assert V.validate_file(path, "xlsx")["ok"]


def test_pptx_cached_values_equal_ground_truth(tmp_path, tables):
    charts = [resolved(r, tables) for r in (PIE, LINE, STACK)]
    path, results = deck_with(charts, tmp_path)
    assert all(r["mode"] == "native" for r in results)
    xmls = chart_xmls(path, "ppt")

    def cached(xml):
        out = []
        for ser in re.findall(r"<c:ser>(.*?)</c:ser>", xml, re.S):
            name = re.search(r"<c:tx>.*?<c:v>(.*?)</c:v>", ser, re.S).group(1)
            cats = re.findall(r"<c:pt idx=\"\d+\"><c:v>(.*?)</c:v>", re.search(r"<c:cat>(.*?)</c:cat>", ser, re.S).group(1))
            vals = [float(v) for v in re.findall(r"<c:pt idx=\"\d+\"><c:v>(.*?)</c:v>", re.search(r"<c:val>(.*?)</c:val>", ser, re.S).group(1))]
            out.append((name, dict(zip(cats, vals))))
        return out

    (pie_series,) = cached(xmls[0])
    assert pie_series[1] == GT["tickets"]["status_counts"]
    (line_series,) = cached(xmls[1])
    assert line_series[1] == GT["sales"]["amount_by_month"]
    stacked = {f"{cat}|{name}": v for name, vals in cached(xmls[2]) for cat, v in vals.items()}
    assert stacked == GT["sales"]["amount_by_quarter_region_h1"]
    assert V.validate_file(path, "pptx", expected_charts=3)["charts"] == 3


def test_pptx_styling_combo_trendline_and_image(tmp_path, tables):
    bar = resolved({"type": "bar", "title": "Revenue", "data": {"table_id": "upload_sales", "x": "Region", "y": ["Amount"]}, "style": {"color": "#8E44AD", "title": {"size_pt": 20, "color": "dark blue"}}}, tables)
    combo = resolved({"type": "combo", "title": "Units", "data": {"table_id": "upload_units", "x": "Product", "y": ["Q1", "Q2"], "y2": ["Q4"]}}, tables)
    scatter = resolved({"type": "scatter", "title": "Pay", "data": {"table_id": "upload_employees", "x": "Experience", "y": ["Salary"], "trendline": True}}, tables)
    gantt = resolved({"type": "gantt", "title": "Plan", "data": {"table_id": "upload_projects", "label": "Task", "start": "Start", "end": "End"}}, tables)
    path, results = deck_with([bar, combo, scatter, gantt], tmp_path)
    assert [r["mode"] for r in results] == ["native", "native", "native", "image"]
    x_bar, x_combo, x_scatter = chart_xmls(path, "ppt")
    assert 'val="8E44AD"' in x_bar and 'sz="2000"' in x_bar and 'val="1F3864"' in x_bar
    assert "<c:lineChart>" in x_combo and "<c:barChart>" in x_combo and '<c:crosses val="max"/>' in x_combo
    assert x_combo.count("<c:valAx>") == 2
    assert "<c:trendline>" in x_scatter
    from pptx import Presentation

    prs = Presentation(str(path))
    assert any(sh.shape_type == 13 for sh in prs.slides[3].shapes)  # PICTURE
    assert V.validate_file(path, "pptx", expected_charts=3)["charts"] == 3


def test_pptx_labels_drop_formula_leads(tmp_path):
    from app.artifacts.compose import DataTable

    t = DataTable(id="paste1", title="p", columns=["Label", "N"], rows=[["=cmd|' /C calc'!A0", 3], ["@SUM(1)", 4]])
    c, _, _ = CD.resolve_chart(CS.Chart.model_validate({"type": "bar", "title": "x", "data": {"table_id": "paste1", "x": "Label", "y": ["N"]}}), [t])
    path, _ = deck_with([c], tmp_path)
    xml = chart_xmls(path, "ppt")[0]
    assert "<c:v>=cmd" not in xml and "<c:v>@SUM" not in xml and "cmd|" in xml
