"""B1 (AS3 fix round 2026-09-15): a workbook never publishes an aggregate the
model typed. Live case: "plot total Amount by month as a line in an excel
sheet" over sales_daily.csv published a code-computed chart beside a sheet
the model typed with 63,668 in every month. Every expected figure below is
recomputed here with the csv module, independently of chart_data."""
from __future__ import annotations

import asyncio
import csv
import datetime as dt
from collections import OrderedDict
from pathlib import Path

import pytest

from app.artifacts import chart_data as CD
from app.artifacts import compose as C
from app.artifacts import derived as D
from app.artifacts import spec as S
from app.artifacts import types as T
from tests.fixtures.charts import loader

SALES_CSV = Path(loader.__file__).resolve().parent / "sales_daily.csv"
TYPED = 63668  # the one figure the model typed for every month


def _monthly_sums() -> "OrderedDict[str, float]":
    out: "OrderedDict[str, float]" = OrderedDict()
    with open(SALES_CSV, newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            d = dt.date.fromisoformat(row["Date"])
            key = d.strftime("%b %Y")
            out[key] = out.get(key, 0.0) + float(row["Amount"])
    return out


@pytest.fixture()
def sales():
    return loader.table("sales", "upload1")


@pytest.fixture()
def tickets():
    return loader.table("tickets", "upload1")


def _typed_monthly_sheet(with_chart: bool = True, **extra):
    months = list(_monthly_sums())
    sheet = {
        "name": "Monthly Sales",
        "columns": [{"name": "Month", "type": "text"}, {"name": "Total Amount", "type": "number"}],
        "rows": [[m.replace(" ", "-"), TYPED] for m in months],
        "totals": [{"column": 1, "fn": "sum", "label": "Total"}],
        "style": {"highlight": [{"column": "Total Amount", "color": "blue"}]},
        **extra,
    }
    if with_chart:
        sheet["charts"] = [{"type": "line", "title": "Total Amount by Month",
                            "data": {"table_id": "upload1", "x": "Date", "y": ["Amount"], "agg": "sum", "date_bucket": "month"}}]
    return S.parse_body("workbook", {"title": "Monthly Sales Analysis", "sheets": [sheet]})


def _all_figures(spec) -> list:
    return [c for sh in spec.body.sheets for r in sh.rows for c in r[1:] if isinstance(c, (int, float))]


# ------------------------------------------------------------ the live case --


def test_the_live_case_the_typed_monthly_table_becomes_the_computed_one(sales):
    spec, _ = CD.resolve_spec(_typed_monthly_sheet(), [sales])
    assert TYPED in _all_figures(spec), "the draft carries the model's figure"
    out, notes, report = D.enforce(spec, [sales])
    sheet = out.body.sheets[0]
    truth = _monthly_sums()
    assert [r[0] for r in sheet.rows] == list(truth)
    assert [r[1] for r in sheet.rows] == pytest.approx(list(truth.values()))
    assert TYPED not in _all_figures(out)
    assert [c.name for c in sheet.columns] == ["Month", "Total Amount"]
    assert sheet.computed_from == "upload1" and sheet.rows_are_code_computed
    assert [(t.column, t.fn) for t in sheet.totals] == [(1, "sum")]
    assert sheet.charts and sheet.charts[0].series[0].values == pytest.approx(list(truth.values()))
    assert report["computed"] == ["Monthly Sales"]
    assert any("computed from sales_daily.csv" in n and "not used" in n for n in notes)
    assert not S.unsupported_figures(out, C.material_text(C.ComposeRequest(kind="workbook", formats=["xlsx"], template_id="generic",
                                                                                 effort="fast", material=C.Material(instruction="x", tables=[sales]))))


def test_without_a_chart_the_columns_are_read_as_the_binding(sales):
    out, notes, report = D.enforce(_typed_monthly_sheet(with_chart=False), [sales])
    truth = _monthly_sums()
    sheet = out.body.sheets[0]
    assert [r[0] for r in sheet.rows] == list(truth)
    assert [r[1] for r in sheet.rows] == pytest.approx(list(truth.values()))
    assert report["computed"] == ["Monthly Sales"] and sheet.computed_from == "upload1"


def test_a_count_and_an_average_column_are_computed_and_a_sum_over_averages_is_not_totalled(tickets):
    status = {}
    hours = {}
    for r in tickets.rows:
        status[r[2]] = status.get(r[2], 0) + 1
        hours.setdefault(r[2], []).append(float(r[9]))
    spec = S.parse_body("workbook", {"title": "Tickets", "sheets": [{
        "name": "By status",
        "columns": [{"name": "Status"}, {"name": "Number of tickets", "type": "integer"}, {"name": "Average Hours", "type": "number"}],
        "rows": [[s, 1, 99.5] for s in status],
        "totals": [{"column": "Number of tickets", "fn": "sum"}, {"column": "Average Hours", "fn": "sum"}],
    }]})
    out, notes, _ = D.enforce(spec, [tickets])
    sheet = out.body.sheets[0]
    got = {r[0]: (r[1], r[2]) for r in sheet.rows}
    assert set(got) == set(status)
    for s, n in status.items():
        assert got[s][0] == n
        assert got[s][1] == pytest.approx(sum(hours[s]) / len(hours[s]))
    assert [(t.column, t.fn) for t in sheet.totals] == [(1, "sum")]
    assert any("average values" in n for n in notes)


def test_rows_copied_from_the_data_are_left_exactly_as_they_are(sales):
    rows = [list(r) for r in sales.rows[:5]]
    spec = S.parse_body("workbook", {"title": "Sample", "sheets": [{
        "name": "First days", "columns": [{"name": c} for c in sales.columns], "rows": rows}]})
    out, notes, report = D.enforce(spec, [sales])
    assert out is spec and notes == [] and report["computed"] == report["dropped"] == []


def test_a_typed_total_row_under_copied_rows_becomes_a_totals_formula(sales):
    rows = [[r[0], r[1], r[2], r[3], r[4]] for r in sales.rows[:4]]
    spec = S.parse_body("workbook", {"title": "Sample", "sheets": [{
        "name": "First days",
        "columns": [{"name": "Date", "type": "date"}, {"name": "Region"}, {"name": "Product"}, {"name": "Units", "type": "integer"}, {"name": "Amount", "type": "number"}],
        "rows": rows + [["Total", None, None, 999, 123456]]}]})
    out, notes, report = D.enforce(spec, [sales])
    sheet = out.body.sheets[0]
    assert sheet.rows == rows
    assert sorted((t.column, t.fn) for t in sheet.totals) == [(3, "sum"), (4, "sum")]
    assert report["totals"] == ["First days"] and 123456 not in _all_figures(out)


def test_figures_nothing_can_compute_are_left_out_and_the_source_is_copied(sales):
    spec = S.parse_body("workbook", {"title": "KPIs", "sheets": [{
        "name": "Dashboard", "columns": [{"name": "Metric"}, {"name": "Value", "type": "number"}],
        "rows": [["Revenue growth", 12.5], ["Best region share", 41.0]]}]})
    out, notes, report = D.enforce(spec, [sales])
    assert report["dropped"] == ["Dashboard"] and report["copied_source"] == "upload1"
    sheet = out.body.sheets[0]
    assert sheet.rows_from == "upload1" and len(sheet.rows) == len(sales.rows) and sheet.rows[0] == list(sales.rows[0])
    assert 12.5 not in _all_figures(out) and 41.0 not in _all_figures(out)
    assert any("left out" in n for n in notes)


def test_a_headline_dashboard_is_computed_from_the_whole_table(sales):
    amounts = [float(r[4]) for r in sales.rows]
    units = [float(r[3]) for r in sales.rows]
    regions = {r[1] for r in sales.rows}
    spec = S.parse_body("workbook", {"title": "Dashboard", "sheets": [
        {"name": "Dashboard", "columns": [{"name": "Metric"}, {"name": "Value", "type": "number"}],
         "rows": [["Total Amount", 400000], ["Average Units", 20], ["Number of orders", 150], ["Number of regions", 5], ["Highest Amount", 9999],
                  ["Average Amount per Month", 174281.5], ["Revenue growth", 12.5]],
         "totals": [{"column": "Value", "fn": "sum"}]}]})
    out, notes, report = D.enforce(spec, [sales])
    sheet = out.body.sheets[0]
    monthly = list(_monthly_sums().values())
    assert report["computed"] == ["Dashboard"] and sheet.computed_from == "upload1"
    assert [r[0] for r in sheet.rows] == ["Total Amount", "Average Units", "Number of orders", "Number of regions", "Highest Amount", "Average Amount per Month"]
    assert [r[1] for r in sheet.rows] == pytest.approx([sum(amounts), sum(units) / len(units), len(sales.rows), len(regions), max(amounts),
                                                         sum(monthly) / len(monthly)])
    assert sheet.totals == [] and any("would add up" in n for n in notes)
    assert any("'Revenue growth' was left out" in n for n in notes)


def test_the_live_shaped_workbook_of_2026_09_15_every_figure_from_the_data(sales):
    """The shape the live engine wrote for F11 (a Dashboard of headline
    figures, a monthly summary with two measures, invented numbers in both;
    the numbers below are this test's own)."""
    daily = {}
    months = {}
    units_by_month = {}
    with open(SALES_CSV, newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            d = dt.date.fromisoformat(row["Date"])
            daily[d] = daily.get(d, 0.0) + float(row["Amount"])
            months[d.strftime("%b %Y")] = months.get(d.strftime("%b %Y"), 0.0) + float(row["Amount"])
            units_by_month[d.strftime("%b %Y")] = units_by_month.get(d.strftime("%b %Y"), 0.0) + float(row["Units"])
    chart = {"type": "line", "title": "Total Amount by Month", "data": {"table_id": "upload1", "x": "Date", "y": ["Amount"], "agg": "sum", "date_bucket": "month"}}
    spec = S.parse_body("workbook", {"title": "Sales Analysis Workbook", "sheets": [
        {"name": "Dashboard", "columns": [{"name": "Metric"}, {"name": "Value", "type": "number"}],
         "rows": [["Total Amount (Jan-Jul 2026)", "1,111,111"], ["Total Units Sold", "2,222"], ["Average Daily Amount", "3,333"],
                  ["Peak Month", "March"], ["Peak Month Amount", "444,444"], ["Total Amount (North)", "5,555"], ["Source", "sales_daily.csv"]],
         "charts": [chart]},
        {"name": "Monthly Sales Summary", "columns": [{"name": "Month"}, {"name": "Total Amount", "type": "number"}, {"name": "Total Units", "type": "integer"}],
         "rows": [[m, 100000 + i, 500 + i] for i, m in enumerate(["January 2026", "February 2026", "March 2026", "April 2026", "May 2026", "June 2026", "July 2026"])],
         "totals": [{"column": "Total Amount", "fn": "sum"}, {"column": "Total Units", "fn": "sum"}], "charts": [chart]},
    ]})
    spec, _ = CD.resolve_spec(spec, [sales])
    out, notes, report = D.enforce(spec, [sales])
    dash, summary = out.body.sheets
    assert report["computed"] == ["Dashboard", "Monthly Sales Summary"]
    got = {r[0]: r[1] for r in dash.rows}
    assert got == pytest.approx({
        "Total Amount (Jan-Jul 2026)": sum(months.values()), "Total Units Sold": sum(units_by_month.values()),
        "Average Daily Amount": sum(daily.values()) / len(daily), "Peak Month Amount": max(months.values()),
    } | {"Source": got.get("Source")})
    assert got["Source"] == "sales_daily.csv" and "Peak Month" not in got and "Total Amount (North)" not in got
    assert dash.charts and dash.charts[0].series[0].values == pytest.approx(list(months.values()))
    assert [c.name for c in summary.columns] == ["Month", "Total Amount", "Total Units"]
    assert [r[0] for r in summary.rows] == list(months)
    assert [r[1] for r in summary.rows] == pytest.approx(list(months.values()))
    assert [r[2] for r in summary.rows] == pytest.approx(list(units_by_month.values()))
    assert sorted((t.column, t.fn) for t in summary.totals) == [(1, "sum"), (2, "sum")]
    assert not {1111111, 2222, 3333, 444444, 5555} & {int(v) for v in _all_figures(out)}


def test_the_copied_source_carries_dates_as_text_cells_a_sheet_can_hold(tickets):
    spec = S.parse_body("workbook", {"title": "KPIs", "sheets": [{
        "name": "Summary", "columns": [{"name": "Metric"}, {"name": "Value", "type": "number"}], "rows": [["Backlog health", 87.5]]}]})
    out, _, report = D.enforce(spec, [tickets])
    sheet = out.body.sheets[0]
    assert report["copied_source"] == "upload1" and len(sheet.rows) == len(tickets.rows)
    assert sheet.rows[0][6] == "2026-01-27T14:00:00".replace("T", " ") and sheet.rows[0][8] == 5


def test_a_revision_draft_with_an_uncomputed_chart_keeps_the_chart_when_the_table_cannot_be_computed(sales):
    spec = S.parse_body("workbook", {"title": "Mix", "sheets": [
        {"name": "Mix", "columns": [{"name": "Segment"}, {"name": "Share", "type": "percent"}], "rows": [["Big", 70], ["Small", 30]],
         "charts": [{"type": "pie", "title": "Amount by region", "data": {"table_id": "upload1", "x": "Region", "y": ["Amount"], "agg": "sum"}}]}]})
    assert not spec.body.sheets[0].charts[0].series
    out, _, report = D.enforce(spec, [sales])
    mix = out.body.sheets[0]
    regions = {}
    for r in sales.rows:
        regions[r[1]] = regions.get(r[1], 0.0) + float(r[4])
    assert report["dropped"] == ["Mix"] and 70 not in _all_figures(out)
    chart = mix.charts[0]
    assert dict(zip(chart.categories, chart.series[0].values)) == pytest.approx(regions)


def test_a_typed_table_beside_a_computed_chart_it_does_not_match_keeps_the_chart_only(sales):
    spec = S.parse_body("workbook", {"title": "Mix", "sheets": [
        {"name": "Data", "columns": [{"name": c} for c in sales.columns], "rows": [], "rows_from": "upload1"},
        {"name": "Mix", "columns": [{"name": "Segment"}, {"name": "Share", "type": "percent"}],
         "rows": [["Big", 70], ["Small", 30]],
         "charts": [{"type": "pie", "title": "Amount by region", "data": {"table_id": "upload1", "x": "Region", "y": ["Amount"], "agg": "sum"}}]},
    ]})
    spec.body.sheets[0].rows = [list(r) for r in sales.rows]
    spec, _ = CD.resolve_spec(spec, [sales])
    out, notes, report = D.enforce(spec, [sales])
    mix = out.body.sheets[1]
    assert report["dropped"] == ["Mix"]
    assert [c.name for c in mix.columns] == ["Note"] and 70 not in _all_figures(out)
    assert mix.charts and mix.charts[0].provenance.table_id == "upload1"


def test_a_model_cannot_certify_its_own_figures_with_the_code_marker(sales):
    raw = {"title": "Monthly", "sheets": [{"name": "Monthly Sales", "computed_from": "upload1",
                                           "columns": [{"name": "Month"}, {"name": "Total Amount", "type": "number"}],
                                           "rows": [["Jan", TYPED]]}]}
    req = C.ComposeRequest(kind="workbook", formats=["xlsx"], template_id="generic", effort="fast",
                           material=C.Material(instruction="x", tables=[sales]))
    C._fill_code_made_rows(raw, req, [])
    assert "computed_from" not in raw["sheets"][0]
    # And a spec that still carries it (a stored or hand-made one) is checked again.
    spec = S.parse_body("workbook", {**raw, "sheets": [{**raw["sheets"][0], "computed_from": "upload1", "rows": [["Jan 2026", TYPED]]}]})
    out, _, report = D.enforce(spec, [sales])
    assert report["computed"] == ["Monthly Sales"] and TYPED not in _all_figures(out)


def test_an_edit_that_kept_a_computed_sheet_keeps_its_marker(sales):
    parent, _, _ = D.enforce(CD.resolve_spec(_typed_monthly_sheet(), [sales])[0], [sales])
    echoed = parent.body.model_dump()
    echoed["sheets"][0].pop("computed_from")
    child = S.parse_body("workbook", echoed)
    out, notes, report = D.enforce(child, [sales], parent=parent)
    assert out.body.sheets[0].computed_from == "upload1" and notes == [] and report["computed"] == []


def test_an_edit_that_reorders_rows_the_edited_version_already_had_is_not_a_typed_aggregate(sales):
    legacy = S.parse_body("workbook", {"title": "Old", "sheets": [{
        "name": "Summary", "columns": [{"name": "Metric"}, {"name": "Value", "type": "number"}],
        "rows": [["Backlog", 1200.5], ["Churn", 3.25]]}]})
    child = S.parse_body("workbook", {"title": "Old", "sheets": [{
        "name": "Summary", "columns": [{"name": "Metric"}, {"name": "Value", "type": "number"}],
        "rows": [["Churn", 3.25], ["Backlog", 1200.5]]}]})
    out, notes, _ = D.enforce(child, [sales], parent=legacy)
    assert out is child and notes == []
    # Without the edited version the same rows are figures nothing computed.
    _, notes, report = D.enforce(child, [sales])
    assert report["dropped"] == ["Summary"]


def test_no_data_tables_means_nothing_to_compute_from(sales):
    spec = _typed_monthly_sheet(with_chart=False)
    out, notes, _ = D.enforce(spec, [])
    assert out is spec and notes == []


# ------------------------------------------------------ through the composer --


class _Ctx:
    def __init__(self, material):
        self.material = material
        self.instruction = "plot total Amount by month as a line in an excel sheet"
        self.kind, self.formats, self.template_id, self.effort = "workbook", ["xlsx"], "generic", "fast"
        self.operation = "create"
        self.parent_spec = None
        self.budget = T.EFFORT_BUDGETS["fast"]
        self.warnings = []
        self.transforms = []
        self.job = {}

    async def progress_stage(self, *a):
        pass

    async def progress(self, *a):
        pass

    def warn(self, text):
        if text not in self.warnings:
            self.warnings.append(text)

    def record_transform(self, t):
        self.transforms.append(t)


def test_the_composer_publishes_computed_figures_and_no_stale_figures_warning(sales, monkeypatch, tmp_path):
    from app.engines import artifact as engine
    from app.artifacts.render import render_version
    from openpyxl import load_workbook

    draft = _typed_monthly_sheet()

    async def fake_compose(req, *, progress=None):
        figures = S.unsupported_figures(draft, C.material_text(req))
        assert "63668" in figures
        return C.ComposeResult(spec=draft, warnings=[C.FIGURES_WARNING + ", ".join(figures)])

    monkeypatch.setattr(C, "compose", fake_compose)
    material = {"history_text": "", "notes": [], "tables": [{"id": "upload1", "title": sales.title, "columns": sales.columns, "rows": sales.rows}]}
    ctx = _Ctx(material)
    spec = asyncio.run(engine.compose_for_pipeline(ctx))
    truth = _monthly_sums()
    sheet = spec.body.sheets[0]
    assert [r[1] for r in sheet.rows] == pytest.approx(list(truth.values()))
    assert not any(w.startswith(C.FIGURES_WARNING) for w in ctx.warnings), ctx.warnings
    assert any("computed from sales_daily.csv" in w for w in ctx.warnings)

    report = render_version(spec, ["xlsx"], str(tmp_path), title_slug="monthly", version=1)
    path = next(Path(tmp_path, f.filename) for f in report.files if f.format == "xlsx")
    ws = load_workbook(path).worksheets[0]
    cells = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_row=1 + len(truth), max_col=2)]
    assert [c[1] for c in cells] == pytest.approx(list(truth.values()))
    assert all(v != TYPED for row in ws.iter_rows(values_only=True) for v in row)


def test_a_typed_table_the_composer_cannot_compute_keeps_the_figures_warning_off_the_left_out_numbers(sales, monkeypatch):
    from app.engines import artifact as engine

    draft = S.parse_body("workbook", {"title": "KPIs", "sheets": [
        {"name": "Data", "columns": [{"name": c} for c in sales.columns], "rows": [list(r) for r in sales.rows], "rows_from": "upload1"},
        {"name": "Dashboard", "columns": [{"name": "Metric"}, {"name": "Value", "type": "number"}], "rows": [["Growth", 12345.5]]}]})

    async def fake_compose(req, *, progress=None):
        return C.ComposeResult(spec=draft, warnings=["sheet 'Dashboard': 1 row was typed by the model, not copied from the pasted table", C.FIGURES_WARNING + "12,345.5"],
                               transform={"rows": 150, "typed_rows": 1, "typed_sheets": ["Dashboard"]})

    monkeypatch.setattr(C, "compose", fake_compose)
    ctx = _Ctx({"history_text": "", "notes": [], "tables": [{"id": "upload1", "title": sales.title, "columns": sales.columns, "rows": sales.rows}]})
    spec = asyncio.run(engine.compose_for_pipeline(ctx))
    assert [sh.name for sh in spec.body.sheets] == ["Data"]
    assert not any(w.startswith(C.FIGURES_WARNING) for w in ctx.warnings)
    assert any("Dashboard sheet was left out" in w for w in ctx.warnings)
    assert not any("typed by the model" in w for w in ctx.warnings), "a note about a sheet that is no longer in the file"
    assert ctx.transforms and "typed_rows" not in ctx.transforms[-1]


def test_a_revision_is_held_to_the_same_rule(sales, monkeypatch):
    typed = _typed_monthly_sheet(with_chart=False).body.model_dump(mode="json", exclude_none=True)

    async def fake_once(req, budget, *, outline_json, extra=""):
        return typed

    async def no_rewrite(req, spec):
        return 0, [], {}

    monkeypatch.setattr(C, "_compose_once", fake_once)
    monkeypatch.setattr(C, "_rewrite_columns", no_rewrite)
    req = C.ComposeRequest(kind="workbook", formats=["xlsx"], template_id="generic", effort="max",
                           material=C.Material(instruction="x", tables=[sales]))
    start = S.parse_body("workbook", {"title": "M", "sheets": [{"name": "Data", "columns": [{"name": c} for c in sales.columns], "rows": [list(r) for r in sales.rows], "rows_from": "upload1"}]})
    fixed = asyncio.run(C.revise(req, start, [{"where": "x", "problem": "y", "fix": "z"}]))
    assert TYPED not in _all_figures(fixed)
    assert fixed.body.sheets[0].computed_from == "upload1"
