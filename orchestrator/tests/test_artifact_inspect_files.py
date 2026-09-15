"""artifacts/inspect_files.py against REAL produced files, and the seeded
defect corpus (tests/fixtures/selfcheck/seeded.py): 61 defects injected by
editing the file bytes/XML, each with a correct twin. Offline, no database.
"""
from __future__ import annotations

import pytest

from app.artifacts import inspect_files as I
from app.artifacts import requirements as RQ
from app.artifacts import selfcheck as SC
from app.artifacts import style as ST
from tests.fixtures.selfcheck import files as F
from tests.fixtures.selfcheck import seeded

pytest.importorskip("pypdfium2")
pytest.importorskip("weasyprint")


@pytest.fixture(scope="module")
def bases(tmp_path_factory):
    root = tmp_path_factory.mktemp("selfcheck-bases")
    out = {
        "doc": (F.render(F.doc_spec(), ["docx", "pdf"], root / "doc", slug="vendor-access-review"), F.doc_spec()),
        "wb": (F.render(F.workbook_spec(), ["xlsx", "csv"], root / "wb", slug="ticket-tracker"), F.workbook_spec()),
        "deck": (F.render(F.deck_spec(), ["pptx", "pdf"], root / "deck", slug="q3-results"), F.deck_spec()),
    }
    svg_dir = root / "svg"
    svg_dir.mkdir()
    (svg_dir / "chart-v1.svg").write_bytes(seeded.SVG_CLEAN)
    out["svg"] = ({"chart-v1.svg": svg_dir / "chart-v1.svg"}, None)
    return out


def _values(obs, fmt, target, prop):
    return [o.value for o in obs if o.format == fmt and o.target == target and o.property == prop]


# ------------------------------------------------------------ real files --


def test_docx_observations_resolve_styles_geometry_fields_and_text(bases):
    files, spec = bases["doc"]
    obs = I.inspect({n: p for n, p in files.items() if n.endswith(".docx")}, spec)
    assert _values(obs, "docx", "page", "orientation") == ["portrait"]
    assert _values(obs, "docx", "page", "page_size") == ["A4"]
    assert _values(obs, "docx", "page", "page_numbers") == [True]
    assert _values(obs, "docx", "title", "size_pt") and _values(obs, "docx", "title", "bold") == [True]
    # Heading colour resolved through the style chain (no run-level colour in the file).
    # AS3 integration: the house colours come from the styling track's theme.
    house = ST.resolve(None).tokens
    assert set(_values(obs, "docx", "heading1", "color")) == {I.norm_hex(house.h1)}
    assert _values(obs, "docx", "document", "headings") == [["Scope", "Findings", "Detail", "Recommendations"]]
    assert "Tokens" in _values(obs, "docx", "document", "table_cells")[0]
    assert set(_values(obs, "docx", "table_header", "background")) == {house.header_fill}
    assert _values(obs, "docx", "security", "unsafe_external_targets") == [[]]


def test_pdf_observations_read_fill_colours_fonts_sizes_and_filled_rects(bases):
    files, spec = bases["doc"]
    obs = I.inspect({n: p for n, p in files.items() if n.endswith(".pdf")}, spec)
    house = ST.resolve(None).tokens  # AS3 integration: the styling track's theme
    assert _values(obs, "pdf", "title", "color") == [house.title]
    assert _values(obs, "pdf", "title", "size_pt") == [float(ST.resolve(None).element("title").size_pt)], "size is font size x the text matrix, in points"
    assert set(_values(obs, "pdf", "table_header", "background")) == {house.header_fill}
    assert set(_values(obs, "pdf", "table_header", "color")) == {"#FFFFFF"}
    assert set(_values(obs, "pdf", "heading1", "background")) == {"#FFFFFF"}, "a rule line under a heading is not its background"
    assert _values(obs, "pdf", "page", "page_numbers") == [True]
    # The house body font is Calibri: drawn with Carlito where it is installed
    # (the orchestrator images), Liberation Sans on a host without it.
    from app.artifacts.render import theme

    body = theme.resolve_font(ST.font_face("Calibri")).family or "Liberation Sans"
    assert body in _values(obs, "pdf", "document", "fonts")[0]


def test_xlsx_observations_see_formulas_quote_prefix_charts_and_their_cells(bases):
    files, spec = bases["wb"]
    obs = I.inspect(files, spec)
    assert _values(obs, "xlsx", "sheet", "columns") == [["ID", "Status", "Owner", "Amount"]]
    assert _values(obs, "xlsx", "sheet", "row_count") == [12]
    assert _values(obs, "csv", "sheet", "row_count") == [12]
    assert _values(obs, "xlsx", "sheet", "freeze_panes") == ["A2"]
    assert _values(obs, "xlsx", "chart", "type") == ["bar"]
    values = _values(obs, "xlsx", "chart", "values")[0]
    assert values["series"] == [[10, 20, 30]]
    assert _values(obs, "xlsx", "security", "unexpected_formulas") == [[]]


def test_pptx_observations_find_slide_titles_drawn_as_text_boxes(bases):
    files, spec = bases["deck"]
    obs = I.inspect({n: p for n, p in files.items() if n.endswith(".pptx")}, spec)
    assert set(_values(obs, "pptx", "document", "headings")[0]) == {"Q3 Results", "Revenue", "Detail"}
    assert _values(obs, "pptx", "chart", "type") == ["line"]
    assert _values(obs, "pptx", "chart", "values")[0]["series"] == [[3.0, 5.0, 4.0]]


def test_an_unreadable_file_is_an_observation_not_an_exception(tmp_path):
    bad = tmp_path / "broken-v1.docx"
    bad.write_bytes(b"not a zip")
    obs = I.inspect({bad.name: bad})
    assert [(o.target, o.property, o.value) for o in obs] == [("file", "readable", False)]


def test_spec_only_properties_are_unverifiable_never_pass(bases):
    """PPTX carries no page-number evidence and a DOCX chart image carries
    no chart type: both must come back unverifiable."""
    files, spec = bases["deck"]
    obs = I.inspect({n: p for n, p in files.items() if n.endswith(".pptx")}, spec)
    checklist = RQ.Checklist(items=[
        RQ.ChecklistItem("c1", "layout", "page", "page_numbers", True),
        RQ.ChecklistItem("c2", "style", "slide_title", "background", "#1F3864"),
    ])
    results = SC.evaluate(checklist, obs, SC.EvalContext(spec=spec))
    assert [r.result for r in results] == ["unverifiable", "unverifiable"]
    dfiles, dspec = bases["doc"]
    dobs = I.inspect({n: p for n, p in dfiles.items() if n.endswith(".docx")}, dspec)
    chart_type = SC.evaluate(RQ.Checklist(items=[RQ.ChecklistItem("c3", "chart", "chart", "type", "bar")]), dobs, SC.EvalContext(spec=dspec))
    assert chart_type[0].result == "unverifiable"


# ------------------------------------------------------- seeded corpus --

_FMT_FOR_PREFIX = {"d": ("docx",), "x": ("xlsx",), "p": ("pptx",), "f": ("pdf",), "s": ("svg",)}


def _run_case(case, bases, tmp_path, which):
    files, spec = bases[case.base]
    exts = _FMT_FOR_PREFIX[case.id[0]]
    if case.id == "x20":
        exts = ("csv",)
    chosen = {n: p for n, p in files.items() if n.rsplit(".", 1)[-1] in exts}
    work = F.copy_files(chosen, tmp_path / f"{case.id}-{which}")
    (case.twin if which == "twin" else case.defect)(work)
    obs = I.inspect(work, spec)
    structure = {}
    if spec is not None and spec.kind == "document":
        body = spec.body
        structure = {"headings": [b.text for b in body.blocks if b.type == "heading"],
                     "cells": [c for b in body.blocks if b.type == "table" for c in [*b.table.columns, *[str(v) for r in b.table.rows for v in r]]]}
    ectx = SC.EvalContext(spec=spec, source_structure=structure)
    ectx.chart_values_ok = SC.chart_values_check(spec, obs)
    return SC.evaluate(RQ.Checklist(items=[case.item]), obs, ectx)[0]


CASES = seeded.cases()


def test_the_corpus_has_at_least_sixty_file_level_defects():
    assert len(CASES) >= 60
    assert len({c.id for c in CASES}) == len(CASES)


@pytest.mark.parametrize("case", CASES, ids=[c.id for c in CASES])
def test_seeded_defect_is_detected_and_its_twin_passes(case, bases, tmp_path):
    twin = _run_case(case, bases, tmp_path, "twin")
    defect = _run_case(case, bases, tmp_path, "defect")
    assert twin.result == "pass", (case.id, "twin", twin.evidence, twin.by_format)
    assert defect.result == "fail", (case.id, "defect", defect.evidence, defect.by_format)


# ------------------------------------ B2 self-check accuracy (2026-09-15) --
#
# Each case is a false "unmet" (or a false should-item failure) the live
# integration run produced on files that were correct. Every test renders
# the file for real, and each holds a twin whose file really lacks the
# property and must still fail.


def _scoped_doc(style_request=None, bullets_in_risks=False):
    blocks = [
        {"type": "heading", "level": 1, "text": "Background"},
        {"type": "paragraph", "text": "Staff work from home three days a week on company laptops."},
        {"type": "paragraph", "text": "The helpdesk supports remote staff during office hours."},
        {"type": "heading", "level": 1, "text": "Risks"},
        {"type": "paragraph", "text": "Home routers are rarely patched and often keep their factory passwords."},
        {"type": "paragraph", "text": "Shared family devices can expose company sessions."},
    ]
    if bullets_in_risks:
        blocks.append({"type": "bullets", "items": ["Unpatched home routers", "Shared devices"]})
    blocks += [
        {"type": "heading", "level": 1, "text": "Controls"},
        {"type": "paragraph", "text": "Enforce a VPN and full disk encryption on every laptop."},
        {"type": "paragraph", "text": "Run a monthly phishing drill for all staff."},
    ]
    spec = F.S.parse_body("document", {"title": "Remote Work Security", "blocks": blocks})
    if style_request:
        patch, unparsed = ST.parse_style_request(style_request, "document")
        assert not unparsed, unparsed
        spec = spec.model_copy(update={"document": spec.body.model_copy(update={"style": ST.merge(spec.body.style, patch)})})
    return spec


def _check(instruction, files, spec, *, kind="document"):
    items = RQ.extract_rules(instruction, kind=kind, operation="create")
    for n, it in enumerate(items):
        it.id = f"c{n:02d}"
    return SC.evaluate(RQ.Checklist(items=items), I.inspect(files, spec), SC.EvalContext(spec=spec, instruction=instruction))


@pytest.mark.parametrize("instruction", [
    "make it a docx with an underlined title and the Risks section paragraphs in italic",
    "paragraphs in the Risks section italic",
    "Risks section ke paragraphs italic kar do",
])
def test_a_style_scoped_to_one_section_is_read_with_its_section(instruction):
    items = [i for i in RQ.extract_rules(instruction, kind="document", operation="create") if i.category == "style" and i.property == "italic"]
    assert len(items) == 1 and items[0].target == "paragraph" and items[0].must
    assert items[0].locator.get("section") == "Risks", items[0].locator
    assert "Risks section" in RQ.describe(items[0])


def test_section_words_that_name_no_section_stay_global():
    for instruction, target in (("make section headings dark blue", "heading"), ("headings in each section blue", "heading"), ("body text italic", "paragraph")):
        items = [i for i in RQ.extract_rules(instruction, kind="document", operation="create") if i.category == "style"]
        assert items and all(i.target == target and not i.locator.get("section") and i.locator.get("section_index") is None for i in items), (instruction, [i.locator for i in items])
    two = [i for i in RQ.extract_rules("make the Controls section bold and the Risks section italic", kind="document", operation="create") if i.category == "style"]
    assert {(i.property, i.locator.get("section")) for i in two} == {("bold", "Controls"), ("italic", "Risks")}
    assert [i.locator.get("section_index") for i in RQ.extract_rules("make section 2 text bold", kind="document", operation="create") if i.category == "style"] == [2]


def test_risks_only_italic_is_met_in_docx_and_pdf_and_its_twin_is_not(tmp_path):
    """Live E6: the Risks paragraphs were italic and nothing else was; the
    check read it as global body italic ("docx: 8 of 9 show False")."""
    instruction = "make it a docx with an underlined title and the Risks section paragraphs in italic"
    spec = _scoped_doc("underlined title and the Risks section paragraphs in italic", bullets_in_risks=True)
    files = F.render(spec, ["docx", "pdf"], tmp_path / "styled", slug="remote-work")
    got = {(r.item.target, r.item.property): r for r in _check(instruction, files, spec)}
    italic = got[("paragraph", "italic")]
    assert italic.result == "pass" and italic.by_format == {"docx": "pass", "pdf": "pass"}, (italic.evidence, italic.by_format)
    # The PDF draws text-decoration underline as a stroked rule inside the glyph box.
    assert got[("title", "underline")].by_format == {"docx": "pass", "pdf": "pass"}, got[("title", "underline")].evidence

    # Twin: the same document without the style really is not italic in Risks.
    plain = _scoped_doc()
    plain_files = F.render(plain, ["docx", "pdf"], tmp_path / "plain", slug="remote-work")
    twin = {(r.item.target, r.item.property): r for r in _check(instruction, plain_files, plain)}
    assert twin[("paragraph", "italic")].by_format == {"docx": "fail", "pdf": "fail"}, twin[("paragraph", "italic")].evidence
    assert twin[("title", "underline")].by_format == {"docx": "fail", "pdf": "fail"}

    # A GLOBAL request on the Risks-only file is still honestly unmet.
    global_ = {(r.item.target, r.item.property): r for r in _check("body text italic", files, spec)}
    assert global_[("paragraph", "italic")].result == "fail"


def test_a_scoped_style_on_a_section_the_file_does_not_have_is_not_a_pass(tmp_path):
    spec = _scoped_doc("the Risks section paragraphs in italic")
    files = F.render(spec, ["docx"], tmp_path / "d", slug="remote-work")
    r = next(r for r in _check("the Appendix section paragraphs in italic", files, spec) if r.item.property == "italic")
    assert r.result == "unverifiable" and any("Appendix" in e for e in r.evidence), (r.result, r.evidence)


def test_a_landscape_pdf_footer_page_number_is_seen(tmp_path):
    """Live C2: after "make it landscape" every PDF version failed "page
    numbers": the footer band was 7.5 % of a 595 pt page height."""
    spec = F.doc_spec(orientation="landscape")
    files = F.render(spec, ["pdf"], tmp_path / "land", slug="vendor-access-review")
    obs = I.inspect(files, spec)
    assert set(_values(obs, "pdf", "page", "orientation")) == {"landscape"}
    assert _values(obs, "pdf", "page", "page_numbers") == [True]
    item = RQ.ChecklistItem("h", "house_style", "page", "page_numbers", True)
    assert SC.evaluate(RQ.Checklist(items=[item]), obs, SC.EvalContext(spec=spec))[0].result == "pass"
    # Twin: a landscape PDF whose footer text is gone has no page numbers.
    import ctypes

    import pypdfium2 as pdfium
    import pypdfium2.raw as R

    pdf_path = next(p for n, p in files.items() if n.endswith(".pdf"))
    doc = pdfium.PdfDocument(str(pdf_path))
    removed = 0
    for index in range(len(doc)):
        page = doc[index]
        textpage = page.get_textpage()
        for j in reversed(range(R.FPDFPage_CountObjects(page.raw))):
            obj = R.FPDFPage_GetObject(page.raw, j)
            if R.FPDFPageObj_GetType(obj) != R.FPDF_PAGEOBJ_TEXT:
                continue
            n = R.FPDFTextObj_GetText(obj, textpage.raw, None, 0)
            buf = ctypes.create_string_buffer(max(1, n) * 2)
            R.FPDFTextObj_GetText(obj, textpage.raw, ctypes.cast(buf, ctypes.POINTER(ctypes.c_ushort)), n)
            if buf.raw.decode("utf-16-le", "ignore").rstrip("\x00").startswith("Page "):
                R.FPDFPage_RemoveObject(page.raw, obj)
                R.FPDFPageObj_Destroy(obj)
                removed += 1
        R.FPDFPage_GenerateContent(page.raw)
    stripped = tmp_path / "stripped.pdf"
    doc.save(str(stripped))
    doc.close()
    assert removed == len(_values(obs, "pdf", "page", "orientation"))
    assert _values(I.inspect({"stripped.pdf": stripped}, spec), "pdf", "page", "page_numbers") == [False]


def _workbook_with_notes_and_native_chart():
    return F.S.parse_body("workbook", {"title": "Vendor Spend", "assumptions": ["Spend is in INR, excluding tax."], "sheets": [{
        "name": "Vendors",
        "columns": [{"name": "Vendor"}, {"name": "Owner"}, {"name": "Spend", "type": "number"}],
        "rows": [[f"Vendor {i}", f"Owner {i % 3}", float(1000 + 250 * i)] for i in range(1, 9)],
        "charts": [{"type": "area", "title": "Spend by vendor", "categories": [f"Vendor {i}" for i in range(1, 9)],
                    "series": [{"name": "Spend", "values": [float(1000 + 250 * i) for i in range(1, 9)]}]}],
    }]})


def test_a_filled_header_row_is_judged_on_data_sheets_not_notes_or_chart_data(tmp_path):
    """Live: "a filled header row" failed on 8 of 8 workbooks that had a
    Notes or a Chart data sheet; their bold A1 label was read as a header."""
    import openpyxl
    from openpyxl.styles import PatternFill

    spec = _workbook_with_notes_and_native_chart()
    files = F.render(spec, ["xlsx"], tmp_path / "wb", slug="vendor-spend")
    xlsx = next(p for n, p in files.items() if n.endswith(".xlsx"))
    names = openpyxl.load_workbook(str(xlsx)).sheetnames
    assert "Notes" in names and "Chart data" in names, names
    item = RQ.ChecklistItem("h", "house_style", "table_header", "fill_present", True, must=False)
    obs = I.inspect(files, spec)
    assert {o.locator["sheet"] for o in obs if o.target == "table_header"} == {"Vendors"}
    assert SC.evaluate(RQ.Checklist(items=[item]), obs, SC.EvalContext(spec=spec))[0].result == "pass"
    assert _values(I.inspect(files, None), "xlsx", "sheet", "row_count") == [8], "no spec: the renderer's own sheets are still recognised by name"

    # Twin: a data sheet whose header row really has no fill still fails.
    def unfill(wb):
        for cell in wb["Vendors"][1]:
            cell.fill = PatternFill(fill_type=None)

    F.xlsx_edit(xlsx, unfill)
    assert SC.evaluate(RQ.Checklist(items=[item]), I.inspect(files, spec), SC.EvalContext(spec=spec))[0].result == "fail"


def test_a_spec_sheet_called_notes_is_still_a_table(tmp_path):
    spec = F.S.parse_body("workbook", {"title": "Meeting Log", "sheets": [{
        "name": "Notes", "columns": [{"name": "Date"}, {"name": "Topic"}], "rows": [["2026-09-01", "Budget"], ["2026-09-08", "Hiring"]]}]})
    files = F.render(spec, ["xlsx"], tmp_path / "wb", slug="meeting-log")
    obs = I.inspect(files, spec)
    assert [o.locator["text"] for o in obs if o.target == "table_header" and o.property == "background"] == ["Date", "Topic"]
    assert _values(obs, "xlsx", "sheet", "row_count") == [2]
