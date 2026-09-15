"""material_in: what a file turn is made from, gathered by code.

Uploads are authored in the test (a 60-row status workbook with a formula
cell carrying a cached value, a 30,000-character Word file, CSV, markdown).
"""
from __future__ import annotations

import asyncio
import base64
import io
import re
import time
import zipfile
from collections import Counter
from pathlib import Path

import pytest

from app import db
from app.artifacts import intent as I
from app.artifacts import material_in as M
from app.artifacts import spec as S

STATUS_COUNTS = {"Resolved": 24, "Open": 14, "In Progress": 11, "Closed": 8, "On Hold": 3}


def _tickets_xlsx(path: Path, *, rows: int = 60) -> None:
    """60 tickets; Status counts 24/14/11/8/3; a second sheet with a formula
    whose CACHED value (42) is written into the XML the way Excel saves it."""
    from openpyxl import Workbook

    statuses = [s for s, n in STATUS_COUNTS.items() for _ in range(n)]
    assert len(statuses) == rows
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    ws.append(["ID", "Title", "Status", "Hours"])
    for i, st in enumerate(statuses):
        ws.append([f"T-{i + 1:03d}", f"Ticket {i + 1}", st, (i % 7) + 1])
    calc = wb.create_sheet("Calc")
    calc.append(["Metric", "Value"])
    calc.append(["Answer", "=6*7"])
    wb.save(path)
    # openpyxl writes a formula without a cached value; add <v>42</v> as Excel would.
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet2.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<f>6\*7</f>(<v></v>|<v/>)?", "<f>6*7</f><v>42</v>", text)
                assert "<v>42</v>" in text
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def test_an_uploaded_workbook_is_read_as_cached_values_on_every_sheet(tmp_path):
    src = tmp_path / "tickets.xlsx"
    _tickets_xlsx(src)
    tables, notes = M.read_xlsx(str(src), name="tickets.xlsx")
    assert [t.id for t in tables] == ["upload1", "upload2"]
    tickets = tables[0]
    assert tickets.columns == ["ID", "Title", "Status", "Hours"] and len(tickets.rows) == 60
    assert dict(Counter(r[2] for r in tickets.rows)) == STATUS_COUNTS
    calc = tables[1]
    assert calc.rows == [["Answer", 42]], "the formula's cached value, never the formula"
    assert not any(isinstance(c, str) and c.startswith("=") for t in tables for r in t.rows for c in r)


def test_gather_reads_a_same_turn_workbook_upload(tmp_path):
    src = tmp_path / "tickets.xlsx"
    _tickets_xlsx(src)
    b64 = base64.b64encode(src.read_bytes()).decode()
    g = asyncio.run(M.gather(history=[], pdf_uploads=[("tickets.xlsx", b64)], conversation_id="conv-mat-1", text="pie of status"))
    assert g.upload_names == ["tickets.xlsx"]
    assert len(g.upload_tables[0].rows) == 60
    assert dict(Counter(r[2] for r in g.upload_tables[0].rows)) == STATUS_COUNTS


def _long_answer() -> str:
    return ("# Onboarding Audit\n\n## Findings\n| # | Area | Severity |\n|---|---|---|\n| 1 | Access | High |\n| 2 | Training | Medium |\n\n"
            + "\n".join(f"- Observation {i} with enough words to be a real answer." for i in range(12)))


def test_the_previous_answer_is_the_substantial_one_even_after_pleasantries():
    history = [{"role": "user", "content": "audit onboarding"}, {"role": "assistant", "content": _long_answer()},
               {"role": "user", "content": "thanks"}, {"role": "assistant", "content": "You're welcome!"}]
    g = asyncio.run(M.gather(history=history, text="give it in docs", save_documents=False))
    assert g.previous_answer_turn_index == 1
    assert g.previous_answer_md.startswith("# Onboarding Audit\n\n## Findings\n|"), "markdown and newlines kept"
    assert [t.id for t in g.answer_tables] == ["answer1"]
    assert g.answer_tables[0].source_id == "assistant_answer"
    assert g.answer_tables[0].rows == [["1", "Access", "High"], ["2", "Training", "Medium"]]


def test_a_file_card_is_never_the_previous_answer():
    history = [{"role": "user", "content": "make a pdf"}, {"role": "assistant", "content": "Created **Audit** as PDF."}]
    md, idx, _ = M.previous_answer(history)
    assert md == "" and idx is None


def test_a_huge_previous_answer_is_clipped_with_a_note():
    history = [{"role": "assistant", "content": "# Big\n\n" + "x " * 100_000}]
    md, idx, notes = M.previous_answer(history)
    assert idx == 0 and len(md) == M.PREVIOUS_ANSWER_MAX_CHARS and notes


def _big_docx(path: Path) -> tuple:
    import docx

    d = docx.Document()
    d.add_paragraph("Supplier Handbook", style="Title")
    heads, cells, size, n = [], [], 0, 0
    while size < 30_000:
        n += 1
        h = f"Chapter {n} obligations"
        d.add_heading(h, level=1)
        heads.append(h)
        body = f"Chapter {n} explains the supplier obligations, the review cadence and the escalation route. " * 6
        d.add_paragraph(body)
        size += len(body)
        t = d.add_table(rows=3, cols=3)
        for r in range(3):
            for c in range(3):
                t.cell(r, c).text = f"Ch{n}-R{r}C{c}"
                cells.append(f"Ch{n}-R{r}C{c}")
    d.save(str(path))
    return heads, cells


def test_a_30k_docx_upload_with_a_hinglish_request_is_imported_whole_and_remembered(tmp_path):
    from app.artifacts.render import render_version

    src = tmp_path / "handbook.docx"
    heads, cells = _big_docx(src)
    text = "is file ko professional docx me bana do"
    intent = I.decide(text, upload_formats=["docx"])
    assert intent.action == "create" and intent.target == "upload" and intent.formats == ["docx"]
    b64 = base64.b64encode(src.read_bytes()).decode()
    g = asyncio.run(M.gather(history=[], pdf_uploads=[("handbook.docx", b64)], conversation_id="conv-mat-docx", intent=intent, text=text))
    doc = g.upload_docs[0]["spec_or_text"]
    assert isinstance(doc, S.DocumentSpec) and doc.title == "Supplier Handbook"
    work = tmp_path / "v1"
    work.mkdir()
    render_version(S.ArtifactSpec(kind="document", document=doc), ["docx"], str(work), title_slug="handbook", version=1)
    import docx

    made = docx.Document(str(next(work.glob("*.docx"))))
    blob = " ".join([p.text for p in made.paragraphs] + [c.text for t in made.tables for r in t.rows for c in r.cells])
    assert all(h in blob for h in heads) and all(c in blob for c in cells)
    stored = db.get_documents("conv-mat-docx")
    assert [d["filename"] for d in stored] == ["handbook.docx"] and "Chapter 1 obligations" in stored[0]["text"]
    assert len([t for t in g.upload_tables if t.id.startswith("upload")]) == len(heads)


def test_earlier_turn_documents_are_read_in_full_when_the_target_is_the_upload():
    long = "# Policy\n\n" + ("Clause text that goes on. " * 3000)
    db.save_document("conv-mat-earlier", "policy.md", long, 0)
    intent = I.ArtifactIntent("create", target="upload")
    g = asyncio.run(M.gather(history=[], conversation_id="conv-mat-earlier", intent=intent, text="make the policy I uploaded a pdf"))
    assert g.upload_docs and g.upload_docs[0]["name"] == "policy.md"
    assert len(g.uploads_text) >= len(long)


def test_csv_markdown_and_pdf_uploads():
    csv_b64 = base64.b64encode("Region,Sales\nNorth,120\nSouth,95\n".encode()).decode()
    md_b64 = base64.b64encode("# Notes\n\n- one\n- two\n".encode()).decode()
    g = asyncio.run(M.gather(history=[], pdf_uploads=[("sales.csv", csv_b64), ("notes.md", md_b64)], save_documents=False))
    assert g.upload_tables[0].columns == ["Region", "Sales"] and g.upload_tables[0].rows == [["North", "120"], ["South", "95"]]
    kinds = {d["name"]: d["kind"] for d in g.upload_docs}
    assert kinds == {"sales.csv": "csv", "notes.md": "md"}
    assert isinstance(next(d for d in g.upload_docs if d["name"] == "notes.md")["spec_or_text"], S.DocumentSpec)
    pages = ["Name   Units   Price\nAlpha   10   2.50\nBeta   12   3.00\nGamma   7   9.10\n"]
    tables = M.pdf_tables_approximate(pages, name="price.pdf")
    assert tables and "(approximate)" in tables[0].title and tables[0].rows[0] == ["Alpha", "10", "2.50"]


def test_a_slow_reader_hits_the_deadline_and_leaves_a_note(monkeypatch):
    def slow(*a, **k):
        time.sleep(0.5)
        return {"name": "x.txt", "kind": "txt", "doc": None, "text": "late", "tables": [], "notes": [], "pages": 0}

    monkeypatch.setattr(M, "_read_one_upload", slow)
    b64 = base64.b64encode(b"hello").decode()
    g = asyncio.run(M.gather(history=[], pdf_uploads=[("x.txt", b64)], save_documents=False, deadline_s=0.1))
    assert any("longer than" in n for n in g.notes)


def test_a_large_workbook_is_parsed_off_the_event_loop(tmp_path):
    """The CPU work runs in a thread: a heartbeat on the loop keeps beating
    (max gap < 250 ms: the thread shares the GIL, and 123 ms was seen
    once under a parallel suite) while a 50,000-row workbook is read. Inline
    on the loop the gap is the whole parse (seconds)."""
    from openpyxl import Workbook

    src = tmp_path / "big.xlsx"
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Data")
    ws.append(["A", "B", "C"])
    for i in range(50_000):
        ws.append([i, f"row {i}", i * 1.5])
    wb.save(src)
    b64 = base64.b64encode(src.read_bytes()).decode()

    async def run():
        gaps = []
        stop = asyncio.Event()

        async def beat():
            last = time.perf_counter()
            while not stop.is_set():
                await asyncio.sleep(0.01)
                now = time.perf_counter()
                gaps.append(now - last)
                last = now

        task = asyncio.create_task(beat())
        g = await M.gather(history=[], pdf_uploads=[("big.xlsx", b64)], save_documents=False)
        stop.set()
        await task
        return g, max(gaps)

    g, worst = asyncio.run(run())
    assert len(g.upload_tables[0].rows) == 50_000
    assert worst < 0.25, worst


def test_the_material_dict_is_json_safe():
    import json

    g = asyncio.run(M.gather(history=[{"role": "assistant", "content": _long_answer()}], save_documents=False,
                             pdf_uploads=[("n.md", base64.b64encode(b"# N\n\nbody").decode())]))
    json.dumps(M.to_material_dict(g))


# --------------------------------------------------------------------------
# Verifier 2026-09-15.


def test_verifier_a_forged_sheet_dimension_does_not_pad_every_row(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Status"])
    for i in range(20000):
        ws.append([f"r{i}", "Open"])
    plain = tmp_path / "plain.xlsx"
    wb.save(plain)
    forged = tmp_path / "forged.xlsx"
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(forged, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"/>', b'<dimension ref="A1:XFD20001"/>', data)
            zout.writestr(info, data)
    t0 = time.perf_counter()
    tables, _ = M.read_xlsx(str(forged), name="forged.xlsx")
    assert time.perf_counter() - t0 < 2.0
    assert tables[0].columns == ["Name", "Status"] and len(tables[0].rows) == 20000


def test_verifier_a_banner_title_row_is_not_the_header_but_a_one_column_sheet_keeps_its_header(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    one = wb.active
    one.title = "OneCol"
    for v in ("Name", "Alice", "Bob"):
        one.append([v])
    banner = wb.create_sheet("Banner")
    banner.merge_cells("A1:C1")
    banner["A1"] = "Q3 Sales"
    banner.append(["Region", "Units", "Price"])
    banner.append(["North", 1, 2])
    banner.append(["South", 3, 4, "extra"])
    path = tmp_path / "shapes.xlsx"
    wb.save(path)
    tables, _ = M.read_xlsx(str(path), name="shapes.xlsx")
    assert tables[0].columns == ["Name"] and tables[0].rows == [["Alice"], ["Bob"]]
    assert tables[1].columns == ["Region", "Units", "Price", "Column 4"] and "Q3 Sales" in tables[1].title
    assert tables[1].rows == [["North", 1, 2, None], ["South", 3, 4, "extra"]]
