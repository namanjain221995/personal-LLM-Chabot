"""md_import: markdown and uploaded Word files → DocumentSpec, by code.

The acceptance case: a synthetic ~40,000-character audit answer exported to
DOCX keeps every heading line and every table cell (checked on the PRODUCED
file), with no model call; unsafe links and remote images never reach the
file. All documents here are authored for the test.
"""
from __future__ import annotations

import random
import re
import zipfile
from pathlib import Path

import pytest

from app.artifacts import md_import
from app.artifacts import spec as S
from app.artifacts import types as T


def audit_markdown(target_chars: int = 40_000, seed: int = 7) -> str:
    rng = random.Random(seed)
    areas = ["Access management", "Change control", "Backup and restore", "Vendor risk", "Logging", "Incident response",
             "Data retention", "Endpoint security", "Network segmentation", "Training", "Physical security", "Payroll controls"]
    owners = ["IT Ops", "Security", "Finance", "HR", "Facilities", "Procurement"]
    out = ["# Internal Controls Audit Report — FY2026", "",
           "**Prepared for:** the Audit Committee  ", "_Scope:_ twelve control areas across three offices.", "",
           "## 1. Executive summary", "",
           "The audit found **14 findings**. See [the framework](https://example.org/framework), "
           "not [this](javascript:alert(1)) or [that](file:///etc/passwd) or [data](data:text/html;base64,PHNjcmlwdD4=).", "",
           "![org chart](https://tracker.example.net/pixel.png)", "",
           "<script>alert('x')</script> <b>Bold HTML</b>", "",
           "> Management accepted every finding.", ""]
    section = 2
    while sum(len(x) + 1 for x in out) < target_chars:
        area = areas[(section - 2) % len(areas)]
        out += [f"## {section}. {area}", "", f"### {section}.1 Observations", ""]
        for _ in range(3):
            out.append(" ".join(f"Control {area.lower()} step {rng.randint(1, 99)} was tested on {rng.randint(5, 60)} samples." for _ in range(4)))
            out.append("")
        out += [f"### {section}.2 Findings", "", "| # | Finding | Severity | Owner | Due | Cost (₹) |", "|---|---|:---:|---|---|---:|"]
        for k in range(rng.randint(4, 9)):
            out.append(f"| {section}.{k + 1} | {area} gap {rng.randint(100, 999)} in `system-{rng.randint(1, 9)}` | "
                       f"{rng.choice(['High', 'Medium', 'Low'])} | {rng.choice(owners)} | 2026-{rng.randint(10, 12)}-{rng.randint(10, 28)} | "
                       f"{rng.randint(1, 90)},{rng.randint(100, 999)} |")
        out += ["", f"### {section}.3 Recommendations", ""]
        for k in range(4):
            out.append(f"{k + 1}. Fix {area.lower()} item {k + 1}.")
            if k == 1:
                out.append(f"   - Sub-step: review {area.lower()} evidence monthly.")
        out += ["", "- Keep evidence for 7 years", "- Report progress *monthly*", ""]
        if section == 4:
            out += ["```sql", "SELECT user_id FROM access_review WHERE stale;", "```", "", "```mermaid", "flowchart TD", "A-->B", "```", ""]
        section += 1
    out += ["## Conclusion", "", "Fixing access reviews first gives the largest reduction in risk.", ""]
    return "\n".join(out)


def _norm(s: str) -> str:
    return " ".join(s.split())


def _expected(md: str):
    heads, cells = [], []
    for line in md.split("\n"):
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            heads.append(_norm(md_import._inline(m.group(2), [])))
        if "|" in line and not md_import._TABLE_SEP_RE.match(line):
            cells += [_norm(md_import._inline(c, [])) for c in md_import._cells(line) if c.strip()]
    return heads, cells


def _docx_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts += [c.text for c in r.cells]
    return _norm("\n".join(parts))


def test_the_40k_audit_export_is_faithful_in_the_produced_docx(tmp_path, monkeypatch):
    from app import llm
    from app.artifacts.render import render_version

    async def no_model(*a, **k):  # the importer must never call the model
        raise AssertionError("model called")

    monkeypatch.setattr(llm, "json_completion", no_model)
    md = audit_markdown()
    assert len(md) >= 40_000
    doc, notes = md_import.markdown_to_document(md)
    spec = S.ArtifactSpec(kind="document", document=doc)
    work = tmp_path / "v1"
    work.mkdir()
    render_version(spec, ["docx"], str(work), title_slug="audit", version=1)
    produced = next(work.glob("*.docx"))
    blob = _docx_text(produced)
    heads, cells = _expected(md)
    assert len(heads) > 50 and len(cells) > 500
    assert [h for h in heads if h not in blob] == []
    assert [c for c in cells if c not in blob] == []
    with zipfile.ZipFile(produced) as z:
        rels = "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".rels"))
        body = z.read("word/document.xml").decode("utf-8", "replace")
    assert "javascript:" not in body + rels and "file:///" not in body + rels and "data:text/html" not in body + rels
    assert "tracker.example.net" not in rels and "<script>" not in body
    assert re.findall(r'TargetMode="External"', rels) == []
    assert doc.title == "Internal Controls Audit Report — FY2026"
    assert any("mermaid" in n for n in notes)


def test_the_export_html_never_carries_unsafe_urls_or_remote_images(tmp_path):
    from app.artifacts.render import html as H

    doc, _ = md_import.markdown_to_document(audit_markdown(8_000))
    spec = S.ArtifactSpec(kind="document", document=doc)
    text = H.document_html(spec.document)  # the HTML WeasyPrint turns into the PDF
    assert "javascript:" not in text and "file:///" not in text and "tracker.example.net" not in text
    assert "<script" not in text.lower().replace("<script>", "") and "<script>" not in text


def test_inline_markup_links_and_images():
    notes = []
    assert md_import._inline("**bold** and _it_ and `code` and ~~gone~~", notes) == "bold and it and code and gone"
    assert md_import._inline("[site](https://a.example/x)", notes) == "site (https://a.example/x)"
    assert md_import._inline("[x](javascript:alert(1))", notes) == "x"
    assert md_import._inline("![alt text](http://evil.example/p.png)", notes) == "alt text"
    assert md_import._inline("snake_case_name stays", notes) == "snake_case_name stays"
    assert md_import._inline("<img src=x onerror=alert(1)>hi", notes) == "hi"


def test_numbered_list_with_a_nested_item_keeps_its_numbering():
    doc, _ = md_import.markdown_to_document("# T\n\n1. one\n2. two\n   - nested\n3. three\n")
    nums = [b for b in doc.blocks if isinstance(b, S.Numbered)]
    assert len(nums) == 1 and nums[0].items == ["one", "two – nested", "three"]


def test_long_items_paragraphs_and_lists_are_split_not_clipped():
    long_item = "word " * 400  # 2,000 characters, over the 600 ceiling
    many = "\n".join(f"- item {i}" for i in range(95))
    para = ("A sentence of the long paragraph. " * 400).strip()  # ~13,600 characters
    doc, _ = md_import.markdown_to_document(f"# T\n\n- {long_item}\n\n{many}\n\n{para}\n")
    text = " ".join(" ".join(b.items) if isinstance(b, S.Bullets) else getattr(b, "text", "") for b in doc.blocks)
    assert text.count("word") == 400
    assert all(f"item {i}" in text for i in range(95))
    assert _norm(para) in _norm(" ".join(b.text for b in doc.blocks if isinstance(b, S.Paragraph)).replace("\n", " "))
    assert all(len(b.items) <= 40 for b in doc.blocks if isinstance(b, S.Bullets))


def test_wide_and_long_tables_are_split_with_every_cell_kept():
    header = "| " + " | ".join(f"C{i}" for i in range(20)) + " |"
    sep = "|" + "---|" * 20
    rows = ["| " + " | ".join(f"r{r}c{c}" for c in range(20)) + " |" for r in range(450)]
    doc, notes = md_import.markdown_to_document("# T\n\n" + "\n".join([header, sep, *rows]))
    tables = [b.table for b in doc.blocks if isinstance(b, S.TableBlock)]
    assert all(len(t.columns) <= T.MAX_TABLE_COLUMNS and len(t.rows) <= T.MAX_TABLE_ROWS for t in tables)
    seen = {str(c) for t in tables for r in t.rows for c in r if c is not None}
    assert all(f"r{r}c{c}" in seen for r in range(450) for c in range(20))
    assert any("split" in n for n in notes)


def test_numeric_columns_are_inferred_by_code():
    doc, _ = md_import.markdown_to_document("# T\n\n| Name | Amount | Share |\n|---|---|---|\n| A | ₹1,20,000 | 12% |\n| B | 3,400.50 | 7.5% |\n| C | — | n/a |\n")
    table = next(b.table for b in doc.blocks if isinstance(b, S.TableBlock))
    assert table.numeric_columns == [1, 2]


def test_title_rules():
    doc, _ = md_import.markdown_to_document("Plain first line. More text here.\n\nSecond paragraph.", title_hint="")
    assert doc.title == "Plain first line."
    doc, _ = md_import.markdown_to_document("## Report\n\nBody", title_hint="Hint Title")
    assert doc.title == "Hint Title"
    doc, _ = md_import.markdown_to_document("# Same\n\n## Same\n\nBody")
    assert doc.title == "Same" and not (isinstance(doc.blocks[0], S.Heading) and doc.blocks[0].text == "Same")
    with pytest.raises(ValueError):
        md_import.markdown_to_document("   \n  ")


def _synthetic_docx(path: Path, target_chars: int = 30_000) -> tuple:
    """A Word file with a title, headings, lists, tables, a hyperlink and an
    INCLUDEPICTURE field pointing at a remote image."""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rng = random.Random(3)
    d = docx.Document()
    d.add_paragraph("Vendor Risk Review 2026", style="Title")
    heads, cells, size = [], [], 0
    n = 0
    while size < target_chars:
        n += 1
        h = f"Section {n}: Supplier group {rng.randint(10, 99)}"
        d.add_heading(h, level=1 if n % 3 else 2)
        heads.append(h)
        for _ in range(3):
            p = f"Supplier review {n} paragraph with {rng.randint(1, 999)} observations about contracts and renewals. " * 3
            d.add_paragraph(p)
            size += len(p)
        d.add_paragraph(f"Bullet for section {n}", style="List Bullet")
        t = d.add_table(rows=4, cols=4)
        for r in range(4):
            for c in range(4):
                v = f"S{n}R{r}C{c}"
                t.cell(r, c).text = v
                cells.append(v)
    # A hyperlink relationship and a field code: neither may reach the import.
    part = d.part
    r_id = part.relate_to("https://exfil.example/steal", RT.HYPERLINK, is_external=True)
    p = d.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "visible link text"
    run.append(t)
    link.append(run)
    p._p.append(link)
    fld = d.add_paragraph()._p
    r1 = OxmlElement("w:r"); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r1.append(fc); fld.append(r1)
    r2 = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.text = ' INCLUDEPICTURE "https://tracker.example.net/x.png" '; r2.append(it); fld.append(r2)
    r3 = OxmlElement("w:r"); fc3 = OxmlElement("w:fldChar"); fc3.set(qn("w:fldCharType"), "end"); r3.append(fc3); fld.append(r3)
    d.save(str(path))
    return heads, cells


def test_docx_upload_import_keeps_every_heading_and_cell_and_nothing_else(tmp_path):
    from app.artifacts.render import render_version

    src = tmp_path / "upload.docx"
    heads, cells = _synthetic_docx(src)
    doc, _ = md_import.docx_to_document(str(src))
    assert doc.title == "Vendor Risk Review 2026"
    spec = S.ArtifactSpec(kind="document", document=doc)
    text = _norm(S.text_of(spec))
    assert all(h in text for h in heads) and all(c in text for c in cells)
    assert "visible link text" in text
    assert "exfil.example" not in text and "tracker.example.net" not in text and "INCLUDEPICTURE" not in text
    work = tmp_path / "v1"
    work.mkdir()
    render_version(spec, ["docx"], str(work), title_slug="vendor-risk", version=1)
    produced = next(work.glob("*.docx"))
    blob = _docx_text(produced)
    assert all(h in blob for h in heads) and all(c in blob for c in cells)
    with zipfile.ZipFile(produced) as z:
        everything = "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith((".xml", ".rels")))
    assert "exfil.example" not in everything
    assert "tracker.example.net" not in everything
    assert "INCLUDEPICTURE" not in everything  # the renderer's own PAGE/NUMPAGES fields are fine


def test_a_zip_bomb_docx_is_refused_before_python_docx_opens_it(tmp_path, monkeypatch):
    from app.core import archive

    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"<" + b"a" * 5_000_000 + b">")
    monkeypatch.setattr(archive, "_limits", lambda: (1_000_000, 100, 50))
    with pytest.raises(archive.ArchiveError):
        md_import.docx_to_document(str(bomb))


# --------------------------------------------------------------------------
# Verifier 2026-09-15.


def test_verifier_unmatched_emphasis_markers_are_linear():
    import time

    for blob in ("*a " * 30000, "_a " * 30000, "**a " * 20000, "~~a " * 20000):
        t0 = time.perf_counter()
        md_import._inline(blob, [])
        md_import.markdown_to_document(blob)
        assert time.perf_counter() - t0 < 1.0, blob[:8]


def test_verifier_script_and_style_content_is_not_text():
    assert md_import._inline("<script>alert(1)</script><b>Bold</b> <style>p{}</style>ok", []) == "Bold ok"


def test_verifier_docx_numbers_and_merged_cells_import_faithfully(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Policy", 0)
    d.add_paragraph("1. A sentence that starts with a number.")
    t = d.add_table(rows=3, cols=3)
    for i in range(3):
        for j in range(3):
            t.cell(i, j).text = f"c{i}{j}"
    t.cell(1, 0).merge(t.cell(1, 1))
    t.cell(1, 0).text = "merged"
    path = tmp_path / "m.docx"
    d.save(path)
    doc, _ = md_import.docx_to_document(str(path))
    paras = [b.text for b in doc.blocks if isinstance(b, S.Paragraph)]
    assert "1. A sentence that starts with a number." in paras and not any("\\" in p for p in paras)
    table = next(b.table for b in doc.blocks if isinstance(b, S.TableBlock))
    assert table.columns == ["c00", "c01", "c02"]
    assert table.rows[0] == ["merged", None, "c12"], "cells after a merged cell stay under their own heading"
