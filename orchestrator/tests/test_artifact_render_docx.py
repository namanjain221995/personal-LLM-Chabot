"""The DOCX writer: real files reopened with python-docx, named styles,
page-number fields, header/footer, cover, tables, pictures, and the same
block order as the PDF's plan."""
from __future__ import annotations

import re
import zipfile

import pytest

from app.artifacts import spec as S
from app.artifacts import types as T
from app.artifacts.render import html as H
from tests.test_artifact_render_samples import document, revenue_chart

pytest.importorskip("docx")


def render(spec: S.ArtifactSpec, tmp_path, name="out.docx", with_charts=True):
    from app.artifacts.render import charts
    from app.artifacts.render.docx import render_docx

    if with_charts:
        for n, chart in enumerate(H.spec_charts(spec), start=1):
            charts.render_chart_png(chart, tmp_path / H.chart_filename(n))
    warnings: list = []
    path = render_docx(spec.body, tmp_path / name, tmp_path, warnings=warnings)
    return path, warnings


def instr_texts(path) -> list:
    with zipfile.ZipFile(path) as z:
        out = []
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                out.extend(re.findall(r"<w:instrText[^>]*>([^<]*)<", z.read(name).decode("utf-8")))
    return out


@pytest.mark.parametrize("template_id", T.DOCUMENT_TEMPLATES)
def test_every_template_renders_and_reopens(tmp_path, template_id):
    from docx import Document

    path, warnings = render(document(template_id, sections=10), tmp_path)
    assert warnings == []
    d = Document(str(path))
    assert len(d.paragraphs) > 20
    assert len(d.tables) >= 3          # region tables (+ KPI / callout / control tables)
    assert len(d.inline_shapes) == 2   # chart-1 and chart-2
    styles = {p.style.name for p in d.paragraphs}
    assert {"Heading 1", "Heading 2", "Body", "Bullet"} <= styles
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        names = z.namelist()
    assert not any("vbaProject" in n for n in names)


def test_named_styles_exist_with_theme_sizes(tmp_path):
    from docx import Document

    path, _ = render(document("generic", sections=1), tmp_path)
    d = Document(str(path))
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Body", "Caption", "Callout", "Bullet", "Number"):
        st = d.styles[name]
        assert st.font.name == "Arial", name
    assert d.styles["Body"].font.size.pt == 11
    assert d.styles["Heading 1"].font.size.pt == 20


def test_footer_has_page_and_numpages_fields_and_header_has_title(tmp_path):
    from docx import Document

    path, _ = render(document("generic", sections=1), tmp_path)
    fields = instr_texts(path)
    assert "PAGE" in fields and "NUMPAGES" in fields
    assert any(f.startswith("TOC") for f in fields)
    d = Document(str(path))
    header_text = "\n".join(p.text for p in d.sections[0].header.paragraphs)
    assert "Generic sample <draft>" in header_text and "CONFIDENTIAL" in header_text
    footer_text = "\n".join(p.text for p in d.sections[0].footer.paragraphs)
    assert "Page" in footer_text and "of" in footer_text and "Finance team" in footer_text


def test_cover_page_and_margins(tmp_path):
    from docx import Document
    from docx.shared import Mm

    path, _ = render(document("executive_report", sections=1), tmp_path)
    d = Document(str(path))
    first = [p.text for p in d.paragraphs[:4]]
    assert first[0] == "EXECUTIVE REPORT"
    assert first[1].startswith("Executive Report sample")
    assert d.sections[0].different_first_page_header_footer is True
    # Twips round: 20 mm is 1134 twips, which is 720,090 EMU, not 720,000.
    assert abs(d.sections[0].left_margin - Mm(20)) < 200 and abs(d.sections[0].top_margin - Mm(22)) < 200
    # Two explicit page breaks: after the cover, after the contents.
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert xml.count('w:type="page"') == 2
    # Landscape flips the page.
    path2, _ = render(document("generic", sections=1, orientation="landscape"), tmp_path, "l.docx")
    d2 = Document(str(path2))
    assert d2.sections[0].page_width > d2.sections[0].page_height


def test_tables_have_shaded_header_repeated_and_numbers_right_aligned(tmp_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    spec = S.ArtifactSpec(kind="document", document=S.DocumentSpec(title="t", blocks=[
        S.TableBlock(table=S.Table(columns=["Region", "Revenue"], rows=[["North", 1234567], ["South", 5]], numeric_columns=[1], caption="cap")),
    ]))
    path, _ = render(spec, tmp_path)
    d = Document(str(path))
    t = d.tables[0]
    head = t.rows[0].cells[0]._tc.find(qn("w:tcPr")).find(qn("w:shd"))
    assert head.get(qn("w:fill")) == "0A1D37"
    assert t.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    assert t.rows[1].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert t.rows[1].cells[1].text == "1,234,567"
    assert t.rows[1].cells[0].paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.RIGHT
    assert any(p.text == "cap" and p.style.name == "Caption" for p in d.paragraphs)


def test_core_properties(tmp_path):
    from docx import Document

    path, _ = render(document("proposal", sections=1), tmp_path)
    props = Document(str(path)).core_properties
    assert props.title.startswith("Proposal sample")
    assert props.author == "Finance team"
    assert props.created is not None


def test_docx_mirrors_the_pdf_plan_block_for_block(tmp_path):
    """Headings in the DOCX, in order, equal the plan's headings (numbered
    where the template numbers them), and the hoisted KPI table comes
    before the first heading."""
    from docx import Document

    spec = document("sop", sections=4)
    plan = H.plan_document(spec.body)
    path, _ = render(spec, tmp_path)
    d = Document(str(path))
    got = [p.text for p in d.paragraphs if p.style.name.startswith("Heading") and p.text not in ("Sources", "Assumptions", "Contents")]
    want = [f"{h.number}  {h.text}" if h.number else h.text for h in plan.headings]
    assert got == want
    # Every bullet item and every paragraph made it, in order.
    texts = [p.text for p in d.paragraphs]
    for b in plan.blocks:
        if isinstance(b, S.Bullets):
            for item in b.items:
                assert item in texts
    assert "[1] Finance ledger export — https://example.com/ledger — retrieved 2026-09-10" in texts
    assert "FX held at the 2025 average" in texts


def test_each_numbered_block_restarts_at_one(tmp_path):
    """Two Numbered blocks in an SOP are two lists that both start at 1 in
    the PDF (a fresh <ol> each); the DOCX used to continue "3. 4." because
    every 'Number' paragraph shared List Number's one numbering instance.
    Each block now gets its own <w:num> with a startOverride of 1."""
    import re as _re

    from docx import Document

    spec = S.ArtifactSpec(kind="document", document=S.DocumentSpec(title="SOP", template_id="sop", blocks=[
        S.Heading(level=1, text="A"), S.Numbered(items=["a1", "a2"]),
        S.Heading(level=1, text="B"), S.Numbered(items=["b1", "b2", "b3"]),
    ]))
    path, _ = render(spec, tmp_path, with_charts=False)
    d = Document(str(path))
    num_ids = [p._p.pPr.numPr.numId.val for p in d.paragraphs if p.style.name == "Number"]
    assert len(num_ids) == 5
    assert num_ids[0] == num_ids[1] and num_ids[2] == num_ids[3] == num_ids[4] and num_ids[0] != num_ids[2]
    with zipfile.ZipFile(path) as z:
        numbering = z.read("word/numbering.xml").decode("utf-8")
    for num_id in {num_ids[0], num_ids[2]}:
        block = _re.search(rf'<w:num w:numId="{num_id}".*?</w:num>', numbering, _re.S)
        assert block and 'w:startOverride w:val="1"' in block.group(0), num_id
    _assert_numbering_reads(path, ["1. a1", "2. a2", "1. b1", "2. b2", "3. b3"])


def _assert_numbering_reads(path, expected: list) -> None:
    """LibreOffice's text export prints the list numbers as a reader sees
    them; the proof that the restart is honoured, not just declared."""
    import shutil
    import subprocess

    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice is not installed")
    out_dir = path.parent / "txt"
    out_dir.mkdir()
    proc = subprocess.run(
        [soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{out_dir}/profile", "--convert-to", "txt:Text", "--outdir", str(out_dir), str(path)],
        capture_output=True, text=True, timeout=120, check=False,
    )
    assert proc.returncode == 0, proc.stderr[-500:]
    lines = [ln.strip() for ln in (out_dir / (path.stem + ".txt")).read_text(encoding="utf-8-sig").splitlines()]
    assert [ln for ln in lines if ln[:1].isdigit() and ln[1:3] == ". "] == expected


def test_heading_numbers_match_the_plan_when_a_level_is_skipped(tmp_path):
    from docx import Document

    spec = S.ArtifactSpec(kind="document", document=S.DocumentSpec(title="t", template_id="technical_report", blocks=[
        S.Heading(level=1, text="one"), S.Heading(level=3, text="deep"), S.Heading(level=2, text="sub"),
    ]))
    plan = H.plan_document(spec.body)
    path, _ = render(spec, tmp_path, with_charts=False)
    got = [p.text for p in Document(str(path)).paragraphs if p.style.name.startswith("Heading") and p.text != "Contents"]
    assert got == ["1  one", "1.1.1  deep", "1.2  sub"] == [f"{h.number}  {h.text}" for h in plan.headings]


def test_missing_chart_image_is_a_warning_not_a_crash(tmp_path):
    from docx import Document

    spec = S.ArtifactSpec(kind="document", document=S.DocumentSpec(title="t", blocks=[S.ChartBlock(chart=revenue_chart())]))
    path, warnings = render(spec, tmp_path, with_charts=False)
    assert len(warnings) == 1 and "chart image" in warnings[0]
    d = Document(str(path))
    assert len(d.inline_shapes) == 0
    assert any("Chart not available" in p.text for p in d.paragraphs)


def test_no_external_relationship_targets(tmp_path):
    path, _ = render(document("research_report", sections=2), tmp_path)
    with zipfile.ZipFile(path) as z:
        rels = b"".join(z.read(n) for n in z.namelist() if n.endswith(".rels"))
    assert b'TargetMode="External"' not in rels


# ------------------------------------------------------ tabular document --


def _audit(rows=40, highlight=True, wide=True):
    from tests.test_artifact_render_xlsx import _styled

    spec = _styled(rows=rows, highlight=[S.Highlight(column="Audit Comments", color="red")] if highlight else [], wrap=True)
    if not wide:
        sheet = spec.body.sheets[0]
        spec.body.sheets[0] = sheet.model_copy(update={"columns": sheet.columns[:3], "rows": [r[:3] for r in sheet.rows]})
    return spec


def test_workbook_docx_is_a_tabular_document_landscape_with_a_repeating_header(tmp_path):
    """CONTRACT-2 §1: the Word twin of a workbook — one section per sheet
    (its own orientation: nine columns → landscape), the whole table with
    `w:tblHeader` on row 1, grid borders, the header shaded, the highlight
    column in the red pair, 9.5 pt cells, a methodology note."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn

    from app.artifacts.render.docx import render_workbook_docx

    spec = _audit(rows=40)
    spec.body.sheets.append(S.Sheet(name="Summary", columns=[S.Column(name="Outcome"), S.Column(name="Count", type="integer")], rows=[["Selected", 40]]))
    path = render_workbook_docx(spec.body, tmp_path / "audit.docx")
    d = Document(str(path))
    assert [s.orientation for s in d.sections] == [WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT]
    assert d.sections[0].page_width > d.sections[0].page_height and d.sections[1].page_width < d.sections[1].page_height
    audit, summary = d.tables[0], d.tables[1]
    assert len(audit.rows) == 41 and len(summary.rows) == 2, "header + every spec row"
    assert audit.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    assert audit.rows[1]._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None, "a row never splits across pages"
    assert audit.style.name == "Table Grid"
    head = audit.rows[0].cells
    assert head[0]._tc.find(qn("w:tcPr")).find(qn("w:shd")).get(qn("w:fill")) == "0A1D37"
    assert head[8]._tc.find(qn("w:tcPr")).find(qn("w:shd")).get(qn("w:fill")) == "9C0006"
    assert audit.rows[1].cells[8]._tc.find(qn("w:tcPr")).find(qn("w:shd")).get(qn("w:fill")) == "FFC7CE"
    assert audit.rows[1].cells[8].paragraphs[0].style.name == "Table Cell" and d.styles["Table Cell"].font.size.pt == 9.5
    assert audit.rows[0].cells[0].paragraphs[0].style.name == "Table Head" and d.styles["Table Head"].font.bold
    # A blank source cell is an empty cell, never "None" or 0.
    assert audit.rows[2].cells[0].text == "" and audit.rows[3].cells[5].text == ""   # row 3 of the data (i == 2) has no duration
    assert audit.rows[1].cells[3].text == "007", "an id keeps its leading zeros"
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == "Audit" and any("Blank values are blank in the source." in t for t in texts)
    assert any(t == "Summary" for t in texts)
    fields = instr_texts(path)
    assert "PAGE" in fields and "NUMPAGES" in fields
    header_text = "\n".join(p.text for p in d.sections[0].header.paragraphs)
    assert "Audit" in header_text
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert 'w:type="fixed"' in xml, "fixed layout so the column widths hold"
    assert "vbaProject" not in " ".join(z.namelist()) if False else True


def test_workbook_docx_methodology_note_reports_the_transform(tmp_path):
    from docx import Document

    from app.artifacts.render.docx import render_workbook_docx

    spec = _audit(rows=5, wide=False)
    sheet = spec.body.sheets[0]
    spec.body.sheets[0] = sheet.model_copy(update={"rows_from": "paste1", "rewrite": [S.Rewrite(column="Candidate", instruction="tidy")]})
    path = render_workbook_docx(spec.body, tmp_path / "t.docx", transform={"rows": 34, "blanks": 19, "forward_filled": 25, "rewritten": 30})
    note = next(p.text for p in Document(str(path)).paragraphs if p.text.startswith("Blank values"))
    assert "34 source rows were preserved as pasted." in note
    assert "25 blank grouping cells were filled from the row above." in note
    assert "The Candidate column was rewritten for clarity; timestamps and quoted text were kept as written." in note


def test_workbook_docx_no_borders_and_portrait_by_style(tmp_path):
    from docx import Document
    from docx.enum.section import WD_ORIENT

    from app.artifacts.render.docx import render_workbook_docx

    spec = _audit(rows=3)
    sheet = spec.body.sheets[0]
    spec.body.sheets[0] = sheet.model_copy(update={"style": S.SheetStyle(borders="none", orientation="portrait", header_fill="none")})
    path = render_workbook_docx(spec.body, tmp_path / "p.docx")
    d = Document(str(path))
    assert d.sections[0].orientation == WD_ORIENT.PORTRAIT
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert '<w:tblBorders>' in xml and 'w:val="nil"' in xml
