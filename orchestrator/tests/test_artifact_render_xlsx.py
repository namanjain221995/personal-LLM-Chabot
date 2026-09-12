"""The XLSX writer and the grid preview: typed formats, real formulas written
by code, formula-injection neutralisation, freeze/filter, native charts, the
dashboard template — every assertion on a file reopened with openpyxl."""
from __future__ import annotations

import datetime as dt
import zipfile

import pytest

from app.artifacts import spec as S
from app.artifacts import types as T
from app.artifacts.render import preview
from tests.test_artifact_render_samples import workbook

pytest.importorskip("openpyxl")


def render(spec: S.ArtifactSpec, tmp_path, name="wb.xlsx"):
    from app.artifacts.render.xlsx import render_xlsx

    warnings: list = []
    path = render_xlsx(spec.body, tmp_path / name, warnings=warnings)
    return path, warnings


@pytest.mark.parametrize("template_id", T.WORKBOOK_TEMPLATES)
def test_every_template_renders_and_reopens(tmp_path, template_id):
    from openpyxl import load_workbook

    path, warnings = render(workbook(template_id), tmp_path)
    wb = load_workbook(str(path))
    names = wb.sheetnames
    expected = ["Pipeline", "Regions", "Notes   draft", "Notes"]
    if template_id == "dashboard":
        expected = ["Dashboard"] + expected
    assert names == expected
    assert warnings == []
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        assert not any("vbaProject" in n for n in z.namelist())
        rels = b"".join(z.read(n) for n in z.namelist() if n.endswith(".rels"))
    assert b'TargetMode="External"' not in rels


def test_formulas_are_written_by_code_from_the_row_count(tmp_path):
    from openpyxl import load_workbook

    path, _ = render(workbook(rows=40), tmp_path)
    ws = load_workbook(str(path))["Pipeline"]
    assert ws["C42"].value == "=SUM(C2:C41)" and ws["C42"].data_type == "f"
    assert ws["D42"].value == "=SUM(D2:D41)"
    assert ws["E42"].value == "=AVERAGE(E2:E41)"
    assert ws["A42"].value == "Total"          # the label of the first total
    regions = load_workbook(str(path))["Regions"]
    assert regions["B5"].value == "=SUM(B2:B4)" and regions["C5"].value == "=SUM(C2:C4)"


def test_formula_injection_is_neutralised_to_text(tmp_path):
    from openpyxl import load_workbook

    path, _ = render(workbook(), tmp_path)
    ws = load_workbook(str(path))["Pipeline"]
    hostile = {
        "F2": '=HYPERLINK("http://evil.example/x")',
        "F3": "+cmd|' /C calc'!A0",
        "F4": "-2+3",
        "F5": "@SUM(A1:A9)",
    }
    for ref, text in hostile.items():
        cell = ws[ref]
        assert cell.data_type == "s", ref
        assert cell.value == text, ref
        assert cell.quotePrefix is True, ref
    # The spec strips leading whitespace, so a TAB or CR lead never reaches
    # the writer through a validated spec; the writer guards anyway (below).
    assert ws["F6"].value == "Tab lead" and ws["F6"].quotePrefix is False
    assert ws["F8"].value == "plain text" and ws["F8"].quotePrefix is False
    # In the XML the cell is an inline string, never an <f> element.
    with zipfile.ZipFile(path) as z:
        sheet_xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>HYPERLINK" not in sheet_xml and "<f>SUM(A1:A9)" not in sheet_xml
    assert sheet_xml.count("<f>") == 3   # exactly the three totals


def test_writer_guards_tab_and_cr_leads_directly():
    from openpyxl import Workbook

    from app.artifacts.render.xlsx import _write_text, is_formula_like

    ws = Workbook().active
    for i, text in enumerate(("\tTab lead", "\rCR lead", "=1+1", "safe"), start=1):
        _write_text(ws.cell(row=i, column=1), text)
    assert [ws.cell(row=i, column=1).quotePrefix for i in range(1, 5)] == [True, True, True, False]
    assert all(ws.cell(row=i, column=1).data_type == "s" for i in range(1, 5))
    assert is_formula_like("\tx") and is_formula_like("\rx") and not is_formula_like("x=1")


def test_number_formats_dates_and_percent_scale(tmp_path):
    from openpyxl import load_workbook

    path, warnings = render(workbook(), tmp_path)
    ws = load_workbook(str(path))["Pipeline"]
    assert ws["B2"].value == dt.datetime(2026, 2, 2) and ws["B2"].number_format == "yyyy-mm-dd"
    assert ws["C2"].value == 1000.5 and ws["C2"].number_format == "#,##0.00"
    assert ws["D2"].value == 3 and ws["D2"].number_format == "#,##0"
    assert ws["E2"].value == 0.01 and ws["E2"].number_format == "0.0%"
    assert ws["A2"].number_format == "@"
    # Whole-number percentages are scaled down and warned about.
    whole = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", sheets=[
        S.Sheet(name="p", columns=[S.Column(name="k"), S.Column(name="rate", type="percent")], rows=[["a", 12.5], ["b", 40]]),
    ]))
    path2, warnings2 = render(whole, tmp_path, "pct.xlsx")
    ws2 = load_workbook(str(path2))["p"]
    assert ws2["B2"].value == pytest.approx(0.125) and ws2["B3"].value == pytest.approx(0.4)
    assert len(warnings2) == 1 and "12.5%" in warnings2[0]
    # A value that cannot be read as the column's type stays as text.
    odd = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", sheets=[
        S.Sheet(name="p", columns=[S.Column(name="n", type="number"), S.Column(name="d", type="date")], rows=[["n/a", "next week"], ["1,234", "2026-13-40"]]),
    ]))
    path3, _ = render(odd, tmp_path, "odd.xlsx")
    ws3 = load_workbook(str(path3))["p"]
    assert ws3["A2"].value == "n/a" and ws3["A2"].data_type == "s"
    assert ws3["A3"].value == 1234.0
    assert ws3["B2"].value == "next week" and ws3["B3"].value == "2026-13-40"


def test_freeze_autofilter_widths_and_header_style(tmp_path):
    from openpyxl import load_workbook

    path, _ = render(workbook(rows=10), tmp_path)
    wb = load_workbook(str(path))
    ws = wb["Pipeline"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:F11"
    assert ws.column_dimensions["C"].width == 14 and ws.column_dimensions["F"].width == 30
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("0A1D37")
    notes = wb["Notes   draft"]
    assert notes.freeze_panes is None and notes.auto_filter.ref is None


def test_native_charts_reference_sheet_cells(tmp_path):
    from openpyxl import load_workbook

    path, _ = render(workbook(), tmp_path)
    wb = load_workbook(str(path))
    regions = wb["Regions"]
    assert len(regions._charts) == 2
    line, pie = regions._charts
    assert line.tagname == "lineChart" and pie.tagname == "pieChart"
    refs = [s.val.numRef.f for s in line.series]
    assert refs == ["'Regions'!$B$2:$B$4", "'Regions'!$C$2:$C$4"]
    assert [s.tx.strRef.f for s in line.series] == ["'Regions'!B1", "'Regions'!C1"]
    pipeline = wb["Pipeline"]
    assert len(pipeline._charts) == 1 and pipeline._charts[0].tagname == "barChart"
    # Its categories are the first five deals and its series is the Amount
    # column, so it charts the sheet's own cells.
    assert pipeline._charts[0].series[0].val.numRef.f == "'Pipeline'!$C$2:$C$6"
    # A chart whose numbers are not in the sheet gets an auxiliary data block
    # to the right and still references real cells — never literal values.
    foreign = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", sheets=[
        S.Sheet(name="p", columns=[S.Column(name="k"), S.Column(name="v", type="number")], rows=[["a", 1]],
                charts=[S.Chart(type="bar", title="Elsewhere", categories=["x", "y"], series=[S.Series(name="other", values=[3, 4])])]),
    ]))
    path2, _ = render(foreign, tmp_path, "foreign.xlsx")
    ws2 = load_workbook(str(path2))["p"]
    assert ws2._charts[0].series[0].val.numRef.f == "'p'!$O$2:$O$3"
    assert ws2["N1"].value == "Elsewhere" and ws2["N2"].value == "x" and ws2["O2"].value == 3.0


def test_dashboard_template_puts_a_summary_sheet_first(tmp_path):
    from openpyxl import load_workbook

    path, _ = render(workbook("dashboard"), tmp_path)
    wb = load_workbook(str(path))
    assert wb.sheetnames[0] == "Dashboard"
    d = wb["Dashboard"]
    assert d["B2"].value == "Sales tracker"
    kpis = {d.cell(row=5, column=c).value: d.cell(row=6, column=c).value for c in range(2, 13, 2) if d.cell(row=5, column=c).value}
    assert kpis["Amount (sum)"] == "='Pipeline'!C42"
    assert kpis["Q1 (sum)"] == "='Regions'!B5"
    assert len(d._charts) == 3   # the Pipeline chart and the two Regions charts, all over sheet columns


def test_dashboard_title_and_purpose_are_text_never_formulas(tmp_path):
    """The Dashboard sheet's B2/B3 were the only two cells written without
    _write_text (review 2026-09-11): a title of =HYPERLINK(...) was a live
    formula. CONTRACT §8: every cell that begins with = + - @ is neutralised."""
    from openpyxl import load_workbook

    hostile = workbook("dashboard").body.model_copy(update={
        "title": '=HYPERLINK("http://evil.example/t","click")', "purpose": "+cmd|' /C calc'!A0",
    })
    path, _ = render(S.ArtifactSpec(kind="workbook", workbook=hostile), tmp_path)
    d = load_workbook(str(path))["Dashboard"]
    for ref, text in (("B2", hostile.title), ("B3", hostile.purpose)):
        assert d[ref].data_type == "s" and d[ref].value == text and d[ref].quotePrefix is True, ref
    with zipfile.ZipFile(path) as z:
        sheet_xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")   # sheet1 is the Dashboard (index 0)
    assert "<f>HYPERLINK" not in sheet_xml and "HYPERLINK" in sheet_xml
    # The only formulas on the dashboard are the KPI cells this code writes
    # (the XML stores a formula without its leading "=").
    import re as _re

    assert _re.findall(r"<f>([^<]*)</f>", sheet_xml) == ["'Pipeline'!C42", "'Pipeline'!D42", "'Pipeline'!E42", "'Regions'!B5", "'Regions'!C5"]


def test_dashboard_kpi_formula_quotes_an_apostrophe_in_the_sheet_name(tmp_path):
    """"Q1's data" is a legal sheet name; ='Q1's data'!B3 is not a legal
    formula (LibreOffice: Err:509). openpyxl doubles the apostrophe."""
    from openpyxl import load_workbook

    sheet = S.Sheet(name="Q1's data", columns=[S.Column(name="k"), S.Column(name="v", type="number")], rows=[["a", 2], ["b", 3]], totals=[S.Total(column=1)])
    spec = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", template_id="dashboard", sheets=[sheet]))
    path, _ = render(spec, tmp_path, "apos.xlsx")
    wb = load_workbook(str(path))
    assert wb.sheetnames == ["Dashboard", "Q1's data"]
    assert wb["Dashboard"]["B6"].value == "='Q1''s data'!B4"
    # The same doubling openpyxl uses for its own chart references.
    assert wb["Dashboard"]._charts == [] and wb["Q1's data"]["B4"].value == "=SUM(B2:B3)"
    _assert_kpi_evaluates(path, "Dashboard", "B6", 5.0)


def _assert_kpi_evaluates(path, sheet: str, ref: str, expected: float) -> None:
    """Recalculate with LibreOffice headless when it is installed (it is on
    the dev box and in the :cpu image), and read the cached value back:
    the proof the formula is one a spreadsheet application accepts."""
    import shutil
    import subprocess

    from openpyxl import load_workbook

    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice is not installed")
    out_dir = path.parent / "recalc"
    out_dir.mkdir()
    proc = subprocess.run(
        [soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{out_dir}/profile", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
        capture_output=True, text=True, timeout=120, check=False,
    )
    assert proc.returncode == 0, proc.stderr[-500:]
    recalculated = load_workbook(str(out_dir / path.name), data_only=True)
    assert recalculated[sheet][ref].value == pytest.approx(expected)


def test_row_ceiling_is_enforced_with_a_warning(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    import app.artifacts.render.xlsx as X

    monkeypatch.setattr(X.T, "MAX_ROWS_PER_SHEET", 10)
    path, warnings = render(workbook(rows=25), tmp_path)
    assert any("10-row ceiling" in w for w in warnings)
    assert load_workbook(str(path))["Pipeline"].max_row == 12   # header + 10 rows + totals


# --------------------------------------------------------------- preview --


def test_sheet_grid_returns_formulas_as_text_and_honours_bounds(tmp_path):
    path, _ = render(workbook("dashboard", rows=30), tmp_path)
    grid = preview.sheet_grid(path, None, max_rows=200, max_cols=50)
    assert [s["name"] for s in grid["sheets"]] == ["Dashboard", "Pipeline", "Regions", "Notes   draft", "Notes"]
    assert grid["sheet"]["name"] == "Dashboard"
    assert "='Pipeline'!C32" in grid["sheet"]["formulas"].values()
    pipe = preview.sheet_grid(path, "Pipeline", max_rows=5, max_cols=3)
    assert pipe["sheet"]["columns"] == ["Deal", "Close date", "Amount"]
    assert len(pipe["sheet"]["rows"]) == 5 and pipe["sheet"]["truncated"] is True
    assert pipe["sheet"]["rows"][0][1] == "2026-02-02"   # a date cell reads back as midnight; the grid shows the date
    full = preview.sheet_grid(path, "Pipeline", max_rows=200, max_cols=50)
    assert full["sheet"]["formulas"]["C32"] == "=SUM(C2:C31)"
    assert full["sheet"]["truncated"] is False
    # The hostile note is data, as text.
    assert full["sheet"]["rows"][0][5] == '=HYPERLINK("http://evil.example/x")'
    assert "F2" not in full["sheet"]["formulas"]
    with pytest.raises(KeyError):
        preview.sheet_grid(path, "No such sheet", max_rows=10, max_cols=10)


def test_sheet_grid_skips_hidden_sheets(tmp_path):
    from openpyxl import load_workbook

    path, _ = render(workbook(), tmp_path)
    wb = load_workbook(str(path))
    wb["Notes"].sheet_state = "hidden"
    wb.save(str(path))
    names = [s["name"] for s in preview.sheet_grid(path, None, 10, 10)["sheets"]]
    assert "Notes" not in names and "Pipeline" in names


# ---------------------------------------------------------------- style --


def _styled(rows=6, **style):
    """A nine-column audit-shaped sheet with a style (CONTRACT-2 §4)."""
    columns = [S.Column(name="Host"), S.Column(name="Candidate"), S.Column(name="Date", type="date"), S.Column(name="Session ID"),
               S.Column(name="Meeting ID"), S.Column(name="Duration", type="integer"), S.Column(name="Ratio", type="number"),
               S.Column(name="Outcome"), S.Column(name="Audit Comments")]
    data = []
    for i in range(rows):
        data.append([
            "Ravi Sharma" if i % 3 == 0 else None, f"Cand {i}", "2026-08-03" if i % 2 == 0 else None, "007" if i == 0 else f"S-{1041 + i}",
            f"MTG-{77812 + i}", 42 if i != 2 else None, 0.81, "Selected", "a long comment " * 8,
        ])
    return S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="Audit", sheets=[
        S.Sheet(name="Audit", columns=columns, rows=data, style=S.SheetStyle(**style) if style else None),
    ]))


def test_default_style_is_thin_black_borders_bold_navy_header_and_top_aligned_cells(tmp_path):
    """A sheet with no style gets SheetStyle's defaults (CONTRACT-2 §4)."""
    from openpyxl import load_workbook

    from app.artifacts.render.xlsx import BORDER_COLOUR

    path, _ = render(_styled(), tmp_path)
    ws = load_workbook(str(path))["Audit"]
    for ref in ("A1", "I1", "A2", "I2", "E7"):
        b = ws[ref].border
        assert (b.left.style, b.right.style, b.top.style, b.bottom.style) == ("thin",) * 4, ref
        assert b.left.color.rgb.endswith(BORDER_COLOUR) and b.top.color.rgb.endswith(BORDER_COLOUR), ref
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("0A1D37") and ws["A1"].font.color.rgb.endswith("FFFFFF")
    assert ws["I2"].alignment.vertical == "top" and ws["I2"].alignment.wrap_text in (None, False)
    # Blanks are blank cells, never 0 or "".
    assert ws["A3"].value is None and ws["C3"].value is None and ws["F4"].value is None
    # An id with leading zeros is text in a text column; a date column holds dates.
    assert ws["D2"].value == "007" and ws["D2"].data_type == "s"
    assert ws["C2"].value == dt.datetime(2026, 8, 3) and ws["C2"].number_format == "yyyy-mm-dd"


def test_leading_zero_ids_are_never_coerced_even_in_a_numeric_column(tmp_path):
    from openpyxl import load_workbook

    spec = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", sheets=[
        S.Sheet(name="p", columns=[S.Column(name="code", type="integer"), S.Column(name="n", type="integer")], rows=[["007", "7"], ["00123", "0"], ["0", "0.5"]]),
    ]))
    path, _ = render(spec, tmp_path, "zeros.xlsx")
    ws = load_workbook(str(path))["p"]
    assert ws["A2"].value == "007" and ws["A2"].data_type == "s"
    assert ws["A3"].value == "00123" and ws["A3"].data_type == "s"
    assert ws["A4"].value == 0 and ws["B2"].value == 7 and ws["B3"].value == 0 and ws["B4"].value == 0.5


def test_highlight_column_uses_readable_red_pairs_and_wrap_top_aligns(tmp_path):
    """CONTRACT-2 §4 SheetStyle: the highlighted column's header is white
    on dark red (9C0006) and its cells dark red on light red (FFC7CE) —
    the pairs pinned exactly; wrap sets wrap_text and top alignment."""
    from openpyxl import load_workbook

    from app.artifacts.render.xlsx import HIGHLIGHT_COLOURS

    assert HIGHLIGHT_COLOURS["red"] == ("9C0006", "FFFFFF", "FFC7CE", "9C0006")
    path, _ = render(_styled(highlight=[S.Highlight(column="audit comments", color="red")], wrap=True), tmp_path, "hl.xlsx")
    ws = load_workbook(str(path))["Audit"]
    assert ws["I1"].fill.fgColor.rgb.endswith("9C0006") and ws["I1"].font.color.rgb.endswith("FFFFFF") and ws["I1"].font.bold
    assert ws["I2"].fill.fgColor.rgb.endswith("FFC7CE") and ws["I2"].font.color.rgb.endswith("9C0006")
    assert ws["H2"].fill.fill_type is None, "only the named column is highlighted"
    assert ws["I2"].alignment.wrap_text is True and ws["I2"].alignment.vertical == "top"
    assert ws["A1"].alignment.wrap_text is True
    assert ws.column_dimensions["I"].width <= 45, "a wrapped comment column is a paragraph, not a strip"
    # Every other colour has its pair too.
    for colour, (h_fill, _, c_fill, c_font) in HIGHLIGHT_COLOURS.items():
        p, _ = render(_styled(highlight=[S.Highlight(column="Outcome", color=colour)]), tmp_path, f"{colour}.xlsx")
        w = load_workbook(str(p))["Audit"]
        assert w["H1"].fill.fgColor.rgb.endswith(h_fill) and w["H2"].fill.fgColor.rgb.endswith(c_fill) and w["H2"].font.color.rgb.endswith(c_font)


def test_header_fill_light_none_and_no_borders(tmp_path):
    from openpyxl import load_workbook

    path, _ = render(_styled(header_fill="light", borders="none", header_bold=False), tmp_path, "light.xlsx")
    ws = load_workbook(str(path))["Audit"]
    assert ws["A1"].fill.fgColor.rgb.endswith("ECECEC") and ws["A1"].font.color.rgb.endswith("0D0D0D")
    assert not ws["A1"].font.bold
    assert ws["A2"].border.left is None or ws["A2"].border.left.style is None
    assert ws["A2"].border.top is None or ws["A2"].border.top.style is None
    assert ws["A1"].border.bottom.style == "thin", "no-borders keeps a rule under the header"
    path, _ = render(_styled(header_fill="none"), tmp_path, "none.xlsx")
    ws = load_workbook(str(path))["Audit"]
    assert ws["A1"].fill.fill_type is None and ws["A1"].font.bold and ws["A1"].font.color.rgb.endswith("0D0D0D")


def test_is_landscape_follows_the_style_and_the_six_column_rule():
    from app.artifacts.render.xlsx import is_landscape

    nine = _styled().body.sheets[0]
    assert is_landscape(nine) is True                         # auto, 9 columns
    assert is_landscape(nine.model_copy(update={"style": S.SheetStyle(orientation="portrait")})) is False
    three = workbook().body.sheets[1]                          # Regions: 3 columns
    assert is_landscape(three) is False
    assert is_landscape(three.model_copy(update={"style": S.SheetStyle(orientation="landscape")})) is True


def test_a_chart_over_a_long_sheet_aggregates_to_top_20_and_other(tmp_path):
    """Charts over > MAX_CHART_POINTS rows aggregate rather than fail: the
    series summed per category, the top 20 by value, then "Other" — from
    a data block beside the sheet the chart still references."""
    from openpyxl import load_workbook

    from app.artifacts.render.xlsx import CHART_TOP_CATEGORIES, aggregate_chart

    rows = [[f"Region {i % 30}", 10 + i] for i in range(500)]
    # As the model writes it: the label column's header for categories and
    # a series named after the numeric column (spec.Sheet fills the cells).
    sheet = S.Sheet.model_validate({
        "name": "Sales", "columns": [{"name": "Region"}, {"name": "Amount", "type": "number"}], "rows": rows,
        "charts": [{"type": "bar", "title": "Amount by region", "categories": ["Region"], "series": [{"name": "Amount"}]}],
    })
    spec = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", template_id="dashboard", sheets=[sheet]))
    chart = spec.body.sheets[0].charts[0]
    assert len(chart.categories) == T.MAX_CHART_POINTS, "the spec keeps the chart under its cap"
    cats, values = aggregate_chart(spec.body.sheets[0], chart, 0, [1])
    assert len(cats) == CHART_TOP_CATEGORIES + 1 and cats[-1] == "Other"
    assert values[0][-1] == sum(v for r, v in rows if r not in cats[:-1])
    assert values[0][0] == max(values[0][:-1]), "top categories first"
    path, warnings = render(spec, tmp_path, "agg.xlsx")
    assert warnings == []
    wb = load_workbook(str(path))
    ws = wb["Sales"]
    assert len(ws._charts) == 1
    ref = ws._charts[0].series[0].val.numRef.f
    assert ref == f"'Sales'!$O$2:$O${1 + CHART_TOP_CATEGORIES + 1}", ref
    assert ws["N1"].value == "Amount by region" and ws["N22"].value == "Other" and ws["O22"].value == values[0][-1]
    # The dashboard draws the same aggregate from the same block.
    assert len(wb["Dashboard"]._charts) == 1 and wb["Dashboard"]._charts[0].series[0].val.numRef.f == ref
