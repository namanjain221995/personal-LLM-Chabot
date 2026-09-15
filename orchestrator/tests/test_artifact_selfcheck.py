"""The artifact self-check (artifacts/selfcheck.py) end to end through the job
runner: real database (the test DSN), real renders (render_version in a
thread), a stub composer. Models are stubs; nothing leaves the process.
"""
from __future__ import annotations

import asyncio
import hashlib
import json
import os
from pathlib import Path

import pytest

from app import metrics
from app.artifacts import inspect_files as I, pipeline, requirements as RQ, selfcheck as SC, store
from app.artifacts import spec as S
from app.artifacts.render import render_version
from app.config import Settings, settings
from tests.fixtures.selfcheck import files as F
from tests.test_artifact_jobs import _accept, _composer, _run, isolated, owner  # noqa: F401 — fixtures by name

pytest.importorskip("weasyprint")


@pytest.fixture(autouse=True)
def real_render(monkeypatch):
    async def render(work_dir, spec, formats, title_slug, version, effort, **kw):
        report = await asyncio.to_thread(render_version, spec, formats, work_dir, title_slug=title_slug, version=version, effort=effort)
        return report.to_json()

    monkeypatch.setattr(pipeline, "_render_in_subprocess", render)
    monkeypatch.setattr(settings, "artifact_selfcheck", True)
    monkeypatch.setattr(settings, "artifact_selfcheck_repair", True)
    monkeypatch.setattr(settings, "artifact_selfcheck_budget_fast_s", 20.0)
    monkeypatch.setattr(settings, "artifact_selfcheck_budget_think_s", 60.0)
    monkeypatch.setattr(settings, "artifact_selfcheck_budget_max_s", 90.0)
    SC.set_content_repairer(None)
    yield
    SC.set_content_repairer(None)


def _report(owner_id: int, row: dict, version: int = 1) -> dict:
    return store.read_json(os.path.join(store.version_dir(owner_id, row["artifact_id"], version), SC.SELFCHECK_NAME))


def _files(owner_id: int, row: dict, version: int = 1, ext: str = "docx") -> Path:
    d = Path(store.version_dir(owner_id, row["artifact_id"], version))
    return next(p for p in d.iterdir() if p.suffix == "." + ext and p.name != "preview.pdf")


def _docx_orientation(path: Path) -> str:
    import docx

    section = docx.Document(str(path)).sections[0]
    return "landscape" if section.page_width > section.page_height else "portrait"


# ------------------------------------------------------------- the loop --


def test_a_landscape_request_the_composer_missed_is_repaired_by_code(owner):
    pipeline.set_composer(_composer(F.doc_spec(orientation="portrait")))
    row = _accept(owner, instruction="Make a Word report on vendor access, landscape", formats=["docx"], format_reason="explicit: word")
    fresh = _run(row["id"])
    assert fresh["status"] == "completed", fresh
    report = _report(owner, row)
    assert report["outcome"] == "repaired" and report["repair"]["accepted"] is True
    assert report["model_calls"] == 0, "Fast: code repair only"
    assert _docx_orientation(_files(owner, row)) == "landscape", "read back with python-docx, not the inspector"
    assert json.loads(Path(store.version_dir(owner, row["artifact_id"], 1), "spec.json").read_text())["document"]["orientation"] == "landscape"


def test_an_unmet_style_is_repaired_by_code_once_the_styling_engine_is_merged(owner):
    """AS3 integration: with edits + style merged, the code repair applies
    the requested heading colour and the version publishes clean."""
    pipeline.set_composer(_composer(F.doc_spec()))
    row = _accept(owner, instruction="Make a PDF report with purple headings", formats=["pdf"], format_reason="explicit: pdf")
    fresh = _run(row["id"])
    report = _report(owner, row)
    assert fresh["status"] in ("completed", "completed_with_warnings") and report["outcome"] == "repaired", report
    assert report["repair"]["accepted"] and "code" in report["repair"]["kinds"]
    assert not any("headings" in line for line in report["unmet"])


def test_an_unmet_style_is_named_and_the_version_is_completed_with_warnings(owner, monkeypatch):
    # Repair off, so the unmet item is what publishes (the repaired path is the test above).
    monkeypatch.setattr(settings, "artifact_selfcheck_repair", False)
    pipeline.set_composer(_composer(F.doc_spec()))
    row = _accept(owner, instruction="Make a PDF report with purple headings", formats=["pdf"], format_reason="explicit: pdf")
    fresh = _run(row["id"])
    assert fresh["status"] == "completed_with_warnings"
    report = _report(owner, row)
    assert report["outcome"] == "unmet"
    assert any("headings" in line and "#6D5AE6" in line for line in report["unmet"]), report["unmet"]
    assert any("headings text #6D5AE6" == c for c in report["false_claim_guard"]["not_claimable"])
    from app.artifacts import db as adb

    version = adb.get_version(row["artifact_id"], 1, owner)
    assert any("not met: headings" in w for w in version["warnings"])


_BREAKS = {
    "rename_scope": lambda b: [x.update(text="Overview") for x in b if x.get("type") == "heading" and x["text"] == "Scope"],
    "rename_findings": lambda b: [x.update(text="Results") for x in b if x.get("type") == "heading" and x["text"] == "Findings"],
    "rename_recommendations": lambda b: [x.update(text="Next steps") for x in b if x.get("type") == "heading" and x["text"] == "Recommendations"],
    "drop_tokens_row": lambda b: [x["table"]["rows"].pop() for x in b if x.get("type") == "table"],
    "rename_a_cell": lambda b: [x["table"]["rows"][0].__setitem__(0, "Log") for x in b if x.get("type") == "table"],
    "change_a_number": lambda b: [x["table"]["rows"][1].__setitem__(2, 8) for x in b if x.get("type") == "table"],
    "rename_a_column": lambda b: [x["table"]["columns"].__setitem__(0, "Zone") for x in b if x.get("type") == "table"],
    "drop_detail_heading": lambda b: b.__delitem__(next(i for i, x in enumerate(b) if x.get("type") == "heading" and x["text"] == "Detail")),
    "drop_the_table": lambda b: b.__delitem__(next(i for i, x in enumerate(b) if x.get("type") == "table")),
    "blank_status_cell": lambda b: [x["table"]["rows"][2].__setitem__(1, "Unknown") for x in b if x.get("type") == "table"],
}

_EXPORT_SOURCE = """# Vendor Access Review

## Scope
This review covers vendor accounts with standing access to production systems.

## Findings
Four accounts had keys older than the rotation window.

| Area | Status | Count |
|---|---|---|
| Logins | Open | 4 |
| Keys | Closed | 9 |
| Tokens | Open | 2 |

### Detail
- Rotate keys
- Remove idle vendors

## Recommendations
Adopt a ninety day rotation and review access monthly.
"""


@pytest.mark.parametrize("name", sorted(_BREAKS))
def test_a_repair_that_fixes_one_item_and_breaks_another_is_rejected_byte_identical(owner, monkeypatch, name):
    pipeline.set_composer(_composer(F.doc_spec(orientation="portrait")))
    snapshot = {}
    real_try = pipeline._try_revision

    async def spying_try(runner, ctx, stages, revised, label, *, accept=None):
        snapshot.update({n: hashlib.sha256(Path(ctx.work_dir, n).read_bytes()).hexdigest() for n in os.listdir(ctx.work_dir)
                         if n.endswith((".docx", ".pdf", ".json")) and Path(ctx.work_dir, n).is_file()})
        return await real_try(runner, ctx, stages, revised, label, accept=accept)

    async def breaking_plan(results, spec, **kw):
        data = spec.model_dump(mode="json")
        data["document"]["orientation"] = "landscape"  # fixes the failing must
        _BREAKS[name](data["document"]["blocks"])  # and breaks something that passed
        return SC.RepairPlan(spec=S.load(data), item_ids=["x"], kinds=["code"])

    monkeypatch.setattr(pipeline, "_try_revision", spying_try)
    monkeypatch.setattr(SC, "plan_repair", breaking_plan)
    row = _accept(owner, instruction="convert your last answer to a word doc in landscape with sections: Scope, Findings, Recommendations",
                  formats=["docx", "pdf"], format_reason="explicit: word", material={"previous_answer": _EXPORT_SOURCE, "history_text": ""})
    fresh = _run(row["id"])
    report = _report(owner, row)
    assert report["repair"]["attempted"] and report["repair"]["accepted"] is False, report["repair"]
    assert report["repair"]["reason"] in ("regressed", "not_improved"), report["repair"]
    published = Path(store.version_dir(owner, row["artifact_id"], 1))
    for filename, digest in snapshot.items():
        if filename.endswith((".docx", ".pdf")) and filename != "preview.pdf":
            assert hashlib.sha256((published / filename).read_bytes()).hexdigest() == digest, f"{filename} changed after a rejected repair"
    assert _docx_orientation(_files(owner, row)) == "portrait"
    assert fresh["status"] == "completed_with_warnings", "the orientation is still unmet and says so"
    assert metrics.render().count("artifact_selfcheck_repair_rejected_total") >= 1


def test_fast_makes_no_model_call_and_think_makes_at_most_two(owner, monkeypatch):
    calls = {"proposer": 0, "repair": 0}

    async def proposer(instruction, summary, kind):
        calls["proposer"] += 1
        return [{"category": "content", "target": "section:budget", "property": "present", "expected": "true"}]

    async def repairer(ccx, spec, issues):
        calls["repair"] += 1
        data = spec.model_dump(mode="json")
        data["document"]["blocks"].append({"type": "heading", "level": 1, "text": "Budget"})
        data["document"]["blocks"].append({"type": "paragraph", "text": "The budget is twelve thousand."})
        return S.load(data)

    monkeypatch.setattr(RQ, "_default_proposer", proposer)
    SC.set_content_repairer(repairer)
    pipeline.set_composer(_composer(F.doc_spec()))
    instruction = "Word document covering Scope, Findings and Budget"
    fast = _accept(owner, instruction=instruction, formats=["docx"], effort="fast", generation_id="g-fast")
    _run(fast["id"])
    assert calls == {"proposer": 0, "repair": 0}
    assert _report(owner, fast)["outcome"] == "unmet", "Fast cannot write a missing section"
    think = _accept(owner, instruction=instruction, formats=["docx"], effort="think", generation_id="g-think")
    fresh = _run(think["id"])
    report = _report(owner, think)
    assert calls["proposer"] == 1 and calls["repair"] == 1
    assert report["model_calls"] <= 2
    assert report["outcome"] == "repaired" and fresh["status"] == "completed", report


def test_a_selfcheck_exception_never_fails_the_job(owner, monkeypatch):
    async def boom(*a, **k):
        raise RuntimeError("checker bug")

    monkeypatch.setattr(RQ, "build", boom)
    pipeline.set_composer(_composer(F.doc_spec()))
    row = _accept(owner, instruction="Make a Word report, landscape", formats=["docx"])
    fresh = _run(row["id"])
    assert fresh["status"] == "completed"
    assert _report(owner, row)["outcome"] == "error"


def test_the_switch_is_a_real_settings_field_and_off_means_no_report(owner, monkeypatch):
    monkeypatch.setenv("ARTIFACT_SELFCHECK", "false")
    monkeypatch.setenv("ARTIFACT_SELFCHECK_REPAIR", "false")
    monkeypatch.setenv("ARTIFACT_SELFCHECK_BUDGET_THINK_S", "33")
    fresh_settings = Settings()
    assert fresh_settings.artifact_selfcheck is False and fresh_settings.artifact_selfcheck_repair is False
    assert fresh_settings.artifact_selfcheck_budget_think_s == 33.0
    monkeypatch.setattr(settings, "artifact_selfcheck", False)
    pipeline.set_composer(_composer(F.doc_spec()))
    row = _accept(owner, instruction="Make a Word report, landscape", formats=["docx"])
    assert _run(row["id"])["status"] == "completed"
    assert not os.path.exists(os.path.join(store.version_dir(owner, row["artifact_id"], 1), SC.SELFCHECK_NAME))


def test_no_budget_left_for_a_rerender_skips_the_repair(owner, monkeypatch):
    monkeypatch.setattr(settings, "artifact_selfcheck_budget_fast_s", 0.05)
    pipeline.set_composer(_composer(F.doc_spec(orientation="portrait")))
    row = _accept(owner, instruction="Make a Word report on vendor access, landscape", formats=["docx"])
    fresh = _run(row["id"])
    report = _report(owner, row)
    assert report["outcome"] == "skipped_budget", report
    assert _docx_orientation(_files(owner, row)) == "portrait"
    assert report["repair"]["attempted"] is False and fresh["status"] == "completed_with_warnings"


def test_a_restore_is_never_rechecked(monkeypatch):
    async def must_not_run(*a, **k):
        raise AssertionError("a restore is a byte copy of a checked version")

    monkeypatch.setattr(SC, "run_hook", must_not_run)

    class Runner:
        deferred = None
        job_id = "j"

    class Ctx:
        job = {"operation": "edit", "progress": {"edit": {"restore_version": 1}}}

    assert asyncio.run(pipeline._selfcheck(Runner(), Ctx(), {}, {})) is None


# ------------------------------------------------------------ faithfulness --


def test_the_production_shape_export_is_faithful_and_honest_about_style(owner):
    md = F.synthetic_audit_markdown(40_000)
    assert len(md) >= 40_000
    spec = F.markdown_to_spec(md)
    pipeline.set_composer(_composer(spec))
    instruction = ("just give it in docs in a standard and classy format, provide a dox file, "
                   "landscape, headings dark blue and Georgia body font")
    row = _accept(owner, instruction=instruction, formats=["docx"], format_reason="explicit: docx",
                  material={"previous_answer": md, "history_text": ""})
    fresh = _run(row["id"])
    report = _report(owner, row)
    by = {(i["category"], i["property"]): i for i in report["items"]}
    assert by[("faithfulness", "headings_covered")]["result"] == "pass"
    assert by[("faithfulness", "table_cells_covered")]["result"] == "pass"
    assert by[("layout", "orientation")]["result"] == "pass", "repaired by code"
    assert _docx_orientation(_files(owner, row)) == "landscape"
    heading_colour = next(i for i in report["items"] if i["category"] == "style" and i["target"] == "heading")
    assert heading_colour["result"] == "pass", "the default navy headings ARE dark blue (hue family + shade)"
    body_font = next(i for i in report["items"] if i["category"] == "style" and i["property"] == "font_family")
    assert body_font["result"] == "fail" and any("Georgia" in line for line in report["unmet"]), report["unmet"]
    assert fresh["status"] == "completed_with_warnings"
    assert not any(i["category"] == "content" for i in report["items"]), "no 'White Bold Text' / 'Landscape' sections"


def test_a_dropped_table_row_in_an_export_is_caught(owner):
    md = F.synthetic_audit_markdown(6_000)
    spec = F.markdown_to_spec(md)
    data = spec.model_dump(mode="json")
    table = next(b for b in data["document"]["blocks"] if b["type"] == "table")
    dropped = table["table"]["rows"].pop()
    pipeline.set_composer(_composer(S.load(data)))
    row = _accept(owner, instruction="give it in a word file", formats=["docx"], material={"previous_answer": md, "history_text": ""})
    _run(row["id"])
    report = _report(owner, row)
    cells = next(i for i in report["items"] if i["property"] == "table_cells_covered")
    assert cells["result"] == "fail" and report["outcome"] == "unmet"
    assert any(dropped[0] in e for e in cells["evidence"]) or cells["evidence"]


# ------------------------------------------------------------ preservation --


def test_edit_preservation_is_checked_on_the_files_of_both_versions(owner, monkeypatch):
    """v2 = orientation only: every section's FILE text equals v1's (pass).
    v3 = the same spec, but the produced DOCX differs in a section the spec
    did not change (a renderer-level change the spec guard cannot see):
    preservation fails, and the version says so."""
    base = F.doc_spec()
    pipeline.set_composer(_composer(base))
    v1 = _accept(owner, instruction="Make a Word report on vendor access", formats=["docx"], generation_id="g1")
    assert _run(v1["id"])["status"] == "completed"
    data = base.model_dump(mode="json")
    data["document"]["orientation"] = "landscape"
    edited = S.load(data)
    pipeline.set_composer(_composer(edited))
    v2 = _accept(owner, operation="edit", instruction="make it landscape", formats=["docx"], parent=(v1["artifact_id"], 1), generation_id="g2")
    assert _run(v2["id"])["status"] == "completed"
    preserve = next(i for i in _report(owner, v2, 2)["items"] if i["category"] == "preservation")
    assert preserve["result"] == "pass", preserve

    async def drifting_render(work_dir, spec, formats, title_slug, version, effort, **kw):
        report = await asyncio.to_thread(render_version, spec, formats, work_dir, title_slug=title_slug, version=version, effort=effort)
        data = report.to_json()
        for f in data["files"]:
            if f["filename"].endswith(".docx"):
                F.docx_replace_text(Path(work_dir, f["filename"]), "Four accounts", "Five accounts")
                body = Path(work_dir, f["filename"]).read_bytes()
                f["size"], f["sha256"] = len(body), hashlib.sha256(body).hexdigest()
        return data

    monkeypatch.setattr(pipeline, "_render_in_subprocess", drifting_render)
    v3 = _accept(owner, operation="edit", instruction="make it landscape", formats=["docx"], parent=(v1["artifact_id"], 2), generation_id="g3")
    fresh = _run(v3["id"])
    report = _report(owner, v3, 3)
    preserve = next(i for i in report["items"] if i["category"] == "preservation")
    assert preserve["result"] == "fail" and "Findings" in preserve["evidence"][0], preserve
    assert fresh["status"] == "completed_with_warnings"


def test_preservation_uses_the_edit_plans_touched_set_when_present():
    before = F.doc_spec()
    data = before.model_dump(mode="json")
    for b in data["document"]["blocks"]:
        if b["type"] == "paragraph" and b["text"].startswith("Adopt"):
            b["text"] = "Adopt a thirty day rotation."
    after = S.load(data)
    item = RQ.ChecklistItem("c1", "preservation", "untouched", "unchanged", True)
    outside = SC.evaluate(RQ.Checklist(items=[item]), [], SC.EvalContext(spec=after, parent_spec=before, touched={"Scope"}))
    assert outside[0].result == "fail" and "Recommendations" in outside[0].evidence[0]
    inside = SC.evaluate(RQ.Checklist(items=[item]), [], SC.EvalContext(spec=after, parent_spec=before, touched={"Recommendations"}))
    assert inside[0].result == "unverifiable", "no parent files: nothing file-level to compare"


# ------------------------------------------------------------ units --


def _res(id_, result, must=True, category="style"):
    return SC.ItemResult(RQ.ChecklistItem(id_, category, "t", "p", True, must=must), result)


@pytest.mark.parametrize("before,after,ok,reason", [
    (["fail", "pass"], ["pass", "pass"], True, ""),
    (["fail", "pass"], ["pass", "fail"], False, "regressed"),
    (["fail", "fail"], ["fail", "fail"], False, "not_improved"),
    (["fail", "fail"], ["pass", "fail"], True, ""),
    (["fail", "pass", "pass"], ["pass", "pass", "fail"], False, "regressed"),
    (["fail", "unverifiable"], ["fail", "pass"], False, "not_improved"),
    (["fail", "pass"], ["unverifiable", "pass"], True, ""),
    (["fail", "fail", "pass"], ["pass", "pass", "fail"], False, "regressed"),
    (["fail", "pass"], ["fail", "fail"], False, "regressed"),
    (["fail", "fail"], ["pass", "unverifiable"], True, ""),
])
def test_the_strict_acceptance_rule(before, after, ok, reason):
    b = [_res(f"c{i}", r) for i, r in enumerate(before)]
    a = [_res(f"c{i}", r) for i, r in enumerate(after)]
    assert SC.accept_repair(b, a, operation="create") == (ok, reason)


def test_a_should_item_that_regresses_also_rejects_and_preservation_rejects_edits():
    b = [_res("c0", "fail"), _res("c1", "pass", must=False)]
    a = [_res("c0", "pass"), _res("c1", "fail", must=False)]
    assert SC.accept_repair(b, a, operation="create") == (False, "regressed")
    b = [_res("c0", "fail"), _res("c1", "unverifiable", category="preservation")]
    a = [_res("c0", "pass"), _res("c1", "fail", category="preservation")]
    assert SC.accept_repair(b, a, operation="edit") == (False, "preservation")


def test_named_colours_match_by_hue_family_and_hex_exactly():
    named = {"color_name": "red", "shade": ""}
    assert SC.colour_matches("#FFC7CE", "#C62828", named)[0], "a red highlight is a light red fill"
    assert not SC.colour_matches("#2F6FB2", "#C62828", named)[0]
    assert SC.colour_matches("#0A1D37", "#1F3864", {"color_name": "dark blue", "shade": "dark"})[0]
    assert not SC.colour_matches("#DCE6F2", "#1F3864", {"color_name": "dark blue", "shade": "dark"})[0]
    assert not SC.colour_matches("#0A1D37", "#1F3864", {})[0], "a typed hex is exact"
    assert SC.colour_matches("#1F3865", "#1F3864", {})[0]


def test_font_substitutes_pass_with_a_note():
    assert SC.font_matches("Carlito", "Calibri") == (True, "Calibri was set in its substitute Carlito")
    assert SC.font_matches("Liberation Sans", "Georgia")[0] is False


def test_contested_items_are_phrased_as_a_reading_and_never_claimed():
    item = RQ.ChecklistItem("c1", "style", "title", "color", "#1F3864", must=False, contested=True,
                            phrase="headings dark blue", note="the request can also be read as heading color #1F3864")
    report = SC.summarize([SC.ItemResult(item, "contested")])
    assert report.unmet == ["I read 'headings dark blue' as title text #1F3864 (it could also mean heading color #1F3864)"]
    assert report.false_claim_guard["claimable"] == []


def test_carry_code_owned_puts_back_style_and_chart_bindings(monkeypatch):
    class Fake:
        def __init__(self, data):
            self.data = data
            self.kind = "document"

        def model_dump(self, **kw):
            return json.loads(json.dumps(self.data))

    original = Fake({"kind": "document", "document": {"style": {"rules": [1]}, "blocks": [{"type": "chart", "chart": {"data": {"table_id": "upload1"}, "title": "A"}}]}})
    revised = Fake({"kind": "document", "document": {"blocks": [{"type": "chart", "chart": {"title": "A (fixed)"}}]}})
    monkeypatch.setattr(S, "load", lambda data: Fake(data))
    carried = pipeline.carry_code_owned(original, revised)
    assert carried.data["document"]["style"] == {"rules": [1]}
    assert carried.data["document"]["blocks"][0]["chart"] == {"title": "A (fixed)", "data": {"table_id": "upload1"}}
    same = pipeline.carry_code_owned(F.doc_spec(), F.doc_spec(title="Other"))
    assert same.title == "Other", "no code-owned fields: the revision stands as returned"


def test_metrics_and_trace_event_are_recorded(owner):
    from app.core import tracing

    events = []

    class Recorder(tracing.TraceRecorder):
        async def _persist(self, fn, *args, **kwargs):
            events.append((fn.__name__, args))

    pipeline.set_composer(_composer(F.doc_spec()))
    row = _accept(owner, instruction="Make a Word report, landscape", formats=["docx"])

    async def scenario():
        rec = Recorder("trace-test")
        token = rec.activate()
        try:
            assert await pipeline.ensure_running(row["id"])
            done = await pipeline.wait_for(row["id"])
            await asyncio.sleep(0.05)
            return done
        finally:
            tracing.TraceRecorder.deactivate(token)

    asyncio.run(scenario())
    text = metrics.render()
    for name in ("artifact_selfcheck_jobs_total", "artifact_selfcheck_items_total", "artifact_selfcheck_seconds"):
        assert name in text, name
    assert any(args and args[2] == "artifact_selfcheck" for name, args in events if name == "append_query_trace_event"), events


def test_the_report_is_published_but_never_downloadable():
    assert SC.SELFCHECK_NAME in store.PUBLISHED_NAMES and SC.SELFCHECK_NAME in store.NOT_DOWNLOADABLE_NAMES
    with pytest.raises(store.PathRefused):
        store.resolve_version_file(1, "a" * 32, 1, "json", SC.SELFCHECK_NAME)


# ------------------------------------------- verifier cases (2026-09-15) --


def _doc(blocks_override=None, title="Vendor Access Review"):
    data = F.doc_spec(title=title).model_dump(mode="json")
    if blocks_override is not None:
        data["document"]["blocks"] = blocks_override(data["document"]["blocks"])
    return S.load(data)


def _preserve(parent, child, instruction):
    item = RQ.ChecklistItem("p", "preservation", "untouched", "unchanged", True)
    ectx = SC.EvalContext(spec=child, parent_spec=parent, instruction=instruction, parent_observations=None)
    return SC._eval_preservation(item, [], ectx)


def _rewrite_scope(blocks):
    out = [dict(b) for b in blocks]
    for b in out:
        if b["type"] == "paragraph" and b["text"].startswith("This review"):
            b["text"] = "A rewritten scope nobody asked for."
    return out


def _edit_recommendations(blocks):
    out = [dict(b) for b in blocks]
    for b in out:
        if b["type"] == "paragraph" and b["text"].startswith("Adopt"):
            b["text"] = "Adopt a thirty day rotation."
    return out


def test_preservation_fails_when_an_unnamed_section_is_rewritten_without_a_touched_set():
    parent = F.doc_spec()
    both = _doc(lambda b: _edit_recommendations(_rewrite_scope(b)))
    r = _preserve(parent, both, "update the Recommendations section to a thirty day rotation")
    assert r.result == "fail" and "Scope" in r.evidence[0], r.evidence
    only_named = _doc(_edit_recommendations)
    assert _preserve(parent, only_named, "update the Recommendations section to a thirty day rotation").result != "fail"
    assert _preserve(parent, only_named, "update section 3").result != "fail", "section 3 = the third H1"
    assert _preserve(parent, only_named, "change the rotation period").result != "fail", "one unnamed change is the edit itself"


def test_preservation_fails_when_a_section_disappears_and_is_silent_for_whole_document_edits():
    parent = F.doc_spec()

    def drop_findings(blocks):
        out, skip = [], False
        for b in blocks:
            if b["type"] == "heading" and b["level"] == 1:
                skip = b["text"] == "Findings"
            if not skip:
                out.append(b)
        return _edit_recommendations(out)

    dropped = _doc(drop_findings)
    assert _preserve(parent, dropped, "update the Recommendations section").result == "fail"
    assert _preserve(parent, _doc(lambda b: _edit_recommendations(_rewrite_scope(b))), "make the whole thing shorter").result == "unverifiable"


def test_an_unshaded_colour_name_is_not_met_by_near_black():
    assert SC.colour_matches("#0A1D37", "#2F6FB2", {"color_name": "blue", "shade": ""})[0] is False
    assert SC.colour_matches("#0A1D37", "#1F3864", {"color_name": "dark blue", "shade": "dark"})[0] is True
    assert SC.colour_matches("#2E75B6", "#2F6FB2", {"color_name": "blue", "shade": ""})[0] is True


def test_a_blank_source_cell_drawn_as_a_gap_matches_a_zero_in_the_spec():
    assert SC._series_equal([[5.0, None, 7.0]], [[5.0, 0.0, 7.0]])
    assert not SC._series_equal([[5.0, None, 7.0]], [[5.0, 6.0, 7.0]])


def test_the_default_content_repairer_passes_the_jobs_tables_and_row_count(monkeypatch):
    from app.artifacts import compose as C

    seen = {}

    async def fake_revise(req, spec, issues):
        seen["material"] = req.material
        return spec

    monkeypatch.setattr(C, "revise", fake_revise)

    class CCX:
        instruction, kind, formats, template_id, effort = "make a report with a Risks section", "document", ["docx"], "generic", "think"
        material = {"history_text": "h", "tables": [{"id": "t1", "title": "Sales", "columns": ["A", "B"], "rows": [["x", 1]], "extra": "ignored"}],
                    "sources": [], "row_count": 30, "transform": {"rows": 1}, "notes": ["n"]}

    asyncio.run(SC._default_content_repairer(CCX(), F.doc_spec(), [{"where": "Risks", "problem": "missing", "fix": "add"}]))
    m = seen["material"]
    assert [t.id for t in m.tables] == ["t1"] and m.tables[0].rows == [["x", 1]]
    assert m.row_count == 30 and m.transform == {"rows": 1}


def test_an_interrupted_revision_is_rolled_back_before_the_stages_run(tmp_path):
    work = tmp_path / "v1.tmp"
    kept = work / pipeline.REVISION_KEPT_DIR
    (kept / "previews").mkdir(parents=True)
    (work / "previews").mkdir()
    (kept / "spec.json").write_text("GOOD SPEC")
    (kept / "report-v1.docx").write_bytes(b"GOOD DOCX")
    (kept / "previews" / "p1.png").write_bytes(b"GOOD PNG")
    (work / "spec.json").write_text("CANDIDATE SPEC")
    (work / "report-v1.docx").write_bytes(b"HALF BUILT")
    (work / "previews" / "p1.png").write_bytes(b"CANDIDATE PNG")
    (work / store.MATERIAL_NAME).write_text("{}")
    assert pipeline.recover_interrupted_revision(str(work)) is True
    assert (work / "spec.json").read_text() == "GOOD SPEC"
    assert (work / "report-v1.docx").read_bytes() == b"GOOD DOCX"
    assert (work / "previews" / "p1.png").read_bytes() == b"GOOD PNG"
    assert (work / store.MATERIAL_NAME).exists() and not kept.exists()
    assert pipeline.recover_interrupted_revision(str(work)) is False


def test_an_edit_that_rewrites_an_unnamed_section_publishes_with_a_warning(owner):
    base = F.doc_spec()
    pipeline.set_composer(_composer(base))
    v1 = _accept(owner, instruction="Make a Word report on vendor access", formats=["docx"], generation_id="g1")
    assert _run(v1["id"])["status"] == "completed"
    pipeline.set_composer(_composer(_doc(lambda b: _edit_recommendations(_rewrite_scope(b)))))
    v2 = _accept(owner, operation="edit", instruction="update the Recommendations section to a thirty day rotation", formats=["docx"],
                 parent=(v1["artifact_id"], 1), generation_id="g2")
    fresh = _run(v2["id"])
    report = _report(owner, v2, 2)
    preserve = next(i for i in report["items"] if i["category"] == "preservation")
    assert preserve["result"] == "fail", preserve
    assert fresh["status"] == "completed_with_warnings"
    assert any("the rest of the document unchanged" in line for line in report["unmet"]), report["unmet"]


# ------------------------------- the 2026-09-16 chart / format incident --

#: "visualise this table on pie chart" — the table the assistant itself had
#: written one turn earlier. Six states, an estimated "~67" cell, and a
#: summary row whose RANK cell says "Total" and whose state cell says
#: "Distinct States".
_INCIDENT_TABLE = {
    "id": "answer1", "title": "Table from the assistant's earlier answer",
    "columns": ["Rank", "State", "Count (Approx)", "% of Total Records"],
    "rows": [[1, "Texas", 21, "2.57%"], [2, "Missouri", 11, "1.35%"], [3, "Illinois", 9, "1.10%"],
             [4, "California", 8, "0.98%"], [5, "New Jersey", 7, "0.86%"],
             ["—", "Other 25 States", "~67", "8.20%"], ["Total", "Distinct States", 30, "3.67%"]],
}


def _equal_slice_pie_spec():
    """The spec the platform published: seven slices, every one of them 1."""
    from app.artifacts import chart_spec as CS

    chart = CS.Chart(
        type="pie", title="Records by state", data=CS.Binding(table_id="answer1", x="State", agg="count"),
        categories=["Texas", "Missouri", "Illinois", "California", "New Jersey", "Other 25 States", "Distinct States"],
        series=[CS.Series(name="Count", values=[1.0] * 7)],
        provenance=CS.Provenance(table_id="answer1", table_title="Table from the assistant's earlier answer", agg="count"),
    )
    return S.parse_body("document", {
        "title": "Records by state",
        "blocks": [{"type": "paragraph", "text": "The share of records held by each state."},
                   {"type": "chart", "chart": chart.model_dump(mode="python")}],
    })


def test_rerunning_the_same_binding_cannot_see_the_wrong_pie(owner):
    """THE TAUTOLOGY, stated as a test. chart_values_check verifies a chart by
    calling chart_data.recompute_matches, which re-runs the SAME binding: a
    wrong binding agrees with itself. The audit is a different reading."""
    from app.artifacts import chart_audit as CA
    from app.artifacts import chart_data as CD
    from app.artifacts import chart_spec as CS

    table = {"id": "answer1", "title": "States", "columns": ["State", "Count", "Population"],
             "rows": [["Texas", 21, 30], ["Missouri", 11, 6], ["Illinois", 9, 12], ["California", 8, 39], ["New Jersey", 7, 9]]}
    resolved, _notes, msg = CD.resolve_chart(CS.Chart(type="pie", title="Records by state",
                                                      data=CS.Binding(table_id="answer1", x="State", agg="count")), [table])
    assert msg == "" and [s.values for s in resolved.series] == [[1.0] * 5], "five states, five equal slices"
    assert CD.recompute_matches(resolved, [table]) == (True, []), "the same binding, the same answer"
    spec = S.parse_body("document", {"title": "States", "blocks": [{"type": "chart", "chart": resolved.model_dump(mode="python")}]})
    assert SC.chart_values_check(spec, [], [table]) is None, "the old check has nothing to say about it"
    audit = CA.audit_spec(spec, [table])
    assert audit.codes() == ["binding"] and "bound to the wrong column" in audit.findings[0].message
    # And a re-bind does NOT fix this one — the table offers two measures, so
    # chart_data leaves the row count standing. The person is told instead.
    fixed, _ = CD.resolve_spec(spec, [table])
    assert CA.audit_spec(fixed, [table]).codes() == ["binding"]


def test_a_pie_of_equal_slices_fails_the_check_and_is_rebound_by_code(owner):
    pipeline.set_composer(_composer(_equal_slice_pie_spec()))
    row = _accept(owner, instruction="visualise this table on pie chart", formats=["docx"], format_reason="explicit: word",
                  material={"tables": [_INCIDENT_TABLE]})
    fresh = _run(row["id"])
    assert fresh["status"] == "completed", fresh
    report = _report(owner, row)
    assert report["outcome"] == "repaired" and report["repair"]["accepted"] is True, report
    assert report["repair"]["kinds"] == ["code"] and report["model_calls"] == 0, "no model repairs a number"
    published = json.loads(Path(store.version_dir(owner, row["artifact_id"], 1), "spec.json").read_text())
    chart = published["document"]["blocks"][1]["chart"]
    assert chart["categories"] == ["Other 25 States", "Texas", "Missouri", "Illinois", "California", "New Jersey"]
    assert chart["series"][0]["values"] == [67.0, 21.0, 11.0, 9.0, 8.0, 7.0], "the Count column, not a row count"
    assert "Distinct States" not in chart["categories"], "the summary row is not a slice"
    results = {i["property"]: i["result"] for i in report["items"] if i["category"] == "chart"}
    assert results["binding_plausible"] == "pass" and results["no_summary_category"] == "pass"
    assert results["values_recomputed"] == "pass", "the published numbers survive a regrouping by hand"


def test_an_unrepairable_wrong_pie_is_said_plainly_instead(owner, monkeypatch):
    # Repair off, so what the person is told is what publishes.
    monkeypatch.setattr(settings, "artifact_selfcheck_repair", False)
    pipeline.set_composer(_composer(_equal_slice_pie_spec()))
    row = _accept(owner, instruction="visualise this table on pie chart", formats=["docx"], format_reason="explicit: word",
                  material={"tables": [_INCIDENT_TABLE]})
    fresh = _run(row["id"])
    assert fresh["status"] == "completed_with_warnings"
    report = _report(owner, row)
    assert report["outcome"] == "unmet"
    unmet = " | ".join(report["unmet"])
    assert "all 7 slices are 1" in unmet and "Count (Approx)" in unmet, report["unmet"]
    assert "Distinct States" in unmet and "counted twice" in unmet, report["unmet"]
    results = {(i["property"], i["result"]) for i in report["items"] if i["category"] == "chart"}
    assert ("binding_plausible", "fail") in results and ("no_summary_category", "fail") in results, sorted(results)
    from app.artifacts import db as adb

    assert any("all 7 slices are 1" in w for w in adb.get_version(row["artifact_id"], 1, owner)["warnings"])


def test_a_file_nobody_asked_for_is_reported_never_repaired(owner, monkeypatch):
    """Production 2026-09-16 delivered a Word file AND a PDF for a request
    that named neither. _validate_files only ever complains that a SELECTED
    format is missing, so every check passed."""
    async def render(work_dir, spec, formats, title_slug, version, effort, **kw):
        report = await asyncio.to_thread(render_version, spec, list(formats) + ["pdf"], work_dir,
                                         title_slug=title_slug, version=version, effort=effort)
        return report.to_json()

    monkeypatch.setattr(pipeline, "_render_in_subprocess", render)
    pipeline.set_composer(_composer(F.doc_spec()))
    row = _accept(owner, instruction="Make a Word report on vendor access", formats=["docx"], format_reason="explicit: word")
    fresh = _run(row["id"])
    assert fresh["status"] == "completed_with_warnings", fresh
    report = _report(owner, row)
    assert {f["format"] for f in store.read_json(os.path.join(store.version_dir(owner, row["artifact_id"], 1), "validation.json"))["files"]} == {"docx", "pdf"}
    assert any("a PDF, which the request did not ask for" in line for line in report["unmet"]), report["unmet"]
    assert report["repair"].get("accepted") is not True, "a repair can add a format, never withdraw a delivered file"


def test_a_delivered_format_the_request_named_is_not_an_extra(owner):
    pipeline.set_composer(_composer(F.doc_spec()))
    row = _accept(owner, instruction="Make a Word report on vendor access and a PDF", formats=["docx", "pdf"],
                  format_reason="explicit: word, pdf")
    fresh = _run(row["id"])
    assert fresh["status"] == "completed", fresh
    report = _report(owner, row)
    assert not any("did not ask for" in line for line in report["unmet"]), report["unmet"]


# ------------------------------------------------------------- honesty --


def test_a_must_that_could_not_be_verified_is_not_a_success():
    item = RQ.ChecklistItem("c1", "chart", "chart", "values_recomputed", True, must=True)
    report = SC.summarize([SC.ItemResult(item, "unverifiable", ["the chart's binding cannot be regrouped by a second reading"])])
    assert report.unconfirmed == ["chart values that match the table regrouped by hand"]
    assert report.unmet == ["not confirmed: chart values that match the table regrouped by hand — "
                            "the chart's binding cannot be regrouped by a second reading"]
    assert report.false_claim_guard["claimable"] == []


def test_the_checker_failing_to_open_a_file_is_not_an_unconfirmed_requirement():
    """_eval_format already says why: the validate stage reopened every file
    with the format's own library, so a reader that cannot parse it is a fact
    about the reader. It must not become a warning the person reads."""
    item = RQ.ChecklistItem("c1", "format", "file", "format", "pdf", must=True)
    unreadable = [I.Observation("file", "readable", False, "pdf", {"name": "x.pdf"})]
    results = SC.evaluate(RQ.Checklist(items=[item]), unreadable, SC.EvalContext())
    assert [r.result for r in results] == ["unverifiable"] and results[0].unreadable is True
    assert SC.summarize(results).unmet == []


def test_a_should_item_that_cannot_be_verified_stays_silent():
    item = RQ.ChecklistItem("c1", "style", "heading", "color", "#1F3864", must=False)
    report = SC.summarize([SC.ItemResult(item, "unverifiable")])
    assert report.unconfirmed == [] and report.unmet == []
