"""artifacts/style.py: the style model, the request parser, the resolver,
the declared support matrix, the contrast policy and the injection guards.

Every figure the style guide states is recomputed here from the WCAG
formula; every injection case is asserted on the produced string or file.
"""
from __future__ import annotations

import asyncio
import itertools
import json
import math
import random
import re
import time
import zipfile

import pytest
from pydantic import ValidationError

from app.artifacts import spec as S
from app.artifacts import style as ST
from tests.fixtures import style_phrases as F
from tests.test_artifact_render_samples import deck, document, workbook

# ------------------------------------------------------------------ model --


def test_colour_names_resolve_in_every_language_form():
    assert ST.resolve_color("dark blue") == ST.resolve_color("navy") == "#1F3864"
    assert ST.resolve_color("#1f3864") == ST.resolve_color("1F3864") == "#1F3864"
    assert ST.resolve_color("#abc") == "#AABBCC"
    # Hinglish, Hindi, Gujarati: the same values (style guide §1).
    assert ST.resolve_color("neela") == ST.resolve_color("नीला") == ST.resolve_color("વાદળી") == "#2F6FB2"
    assert ST.resolve_color("gehra neela") == ST.resolve_color("गहरा नीला") == ST.resolve_color("ઘેરો વાદળી") == "#1F3864"
    assert ST.resolve_color("lal") == ST.resolve_color("लाल") == ST.resolve_color("લાલ") == "#C62828"
    assert ST.resolve_color("hara") == ST.resolve_color("हरा") == ST.resolve_color("લીલો") == "#3F8F4F"
    assert ST.resolve_color("peela") == ST.resolve_color("पीला") == ST.resolve_color("પીળો") == "#FFD54F"
    assert ST.resolve_color("safed") == ST.resolve_color("सफ़ेद") == ST.resolve_color("સફેદ") == "#FFFFFF"
    for bad in ("", "blurple", "#12345", "rgb(1,2,3)", "url(x)", "red; background:url(x)", None, 12):
        assert ST.resolve_color(bad) is None, bad


def test_text_style_validates_colours_fonts_and_sizes():
    ts = ST.TextStyle(color="dark blue", background="#fff", font_family="georgia", size_pt=12)
    assert (ts.color, ts.background, ts.font_family, ts.size_pt) == ("#1F3864", "#FFFFFF", "Georgia", 12.0)
    with pytest.raises(ValidationError):
        ST.TextStyle(font_family="Georgia; } @import url(x)")
    with pytest.raises(ValidationError):
        ST.TextStyle(font_family="Comic Neue Evil")
    with pytest.raises(ValidationError):
        ST.TextStyle(color="red;background:url(http://x)")
    for size in (5, 73, float("nan"), float("inf")):
        with pytest.raises(ValidationError):
            ST.TextStyle(size_pt=size)


def test_targets_carry_only_their_locators():
    ST.StyleTarget(kind="heading", level=2, text="Scope")
    ST.StyleTarget(kind="cell_range", a1="b2:d4", sheet="Data")
    assert ST.StyleTarget(kind="cell_range", a1=" $B$2 : D4 ").a1 == "B2:D4"
    for bad in (dict(kind="paragraph", index=3), dict(kind="title", text="x"), dict(kind="column"), dict(kind="row"),
                dict(kind="cell_range", a1="B2:D"), dict(kind="cell_range", a1="=HYPERLINK(1)"), dict(kind="row", index=1, sheet="a[b]")):
        with pytest.raises(ValidationError):
            ST.StyleTarget(**bad)


def test_condition_values_are_bounded_data():
    ST.CondRule(column="Status", op="eq", value="x" * 100, style=ST.TextStyle(background="red"))
    with pytest.raises(ValidationError):
        ST.CondRule(column="Status", op="eq", value="x" * 101, style=ST.TextStyle(background="red"))
    with pytest.raises(ValidationError):
        ST.CondRule(column="Score", op="gt", value=float("inf"), style=ST.TextStyle(background="red"))
    with pytest.raises(ValidationError):
        ST.CondRule(column="Score", op="gt", value="high", style=ST.TextStyle(background="red"))
    with pytest.raises(ValidationError):
        ST.CondRule(column="Status", op="eq", value="Done", style=ST.TextStyle())
    with pytest.raises(ValidationError):
        ST.StyleSpec(rules=[ST.StyleRule(target=ST.StyleTarget(kind="title"), style=ST.TextStyle(bold=True))] * 51)
    # Header/footer text is plain text: control characters become spaces and
    # it is cut at 120 characters.
    assert ST.HeaderFooter(footer_text="x" * 130 + "\r\n").footer_text == "x" * 120
    assert ST.HeaderFooter(header_text="a\x00b\tc").header_text == "a b c"


def test_merge_combines_rules_on_the_same_target_and_replaces_scalars():
    base = ST.merge(None, ST.parse_style_request("headings dark blue, landscape")[0])
    later = ST.merge(base, ST.parse_style_request("headings bold, portrait, footer 'Q3'")[0])
    heading = [r for r in later.rules if r.target.kind == "heading"]
    assert len(heading) == 1 and heading[0].style.color == "#1F3864" and heading[0].style.bold is True
    assert later.page.orientation == "portrait" and later.header_footer.footer_text == "Q3"
    again = ST.merge(later, ST.StylePatch(rules=[ST.StyleRule(target=ST.StyleTarget(kind="heading"), style=ST.TextStyle(color="red"))]))
    assert [r.style.color for r in again.rules if r.target.kind == "heading"] == ["#C62828"]
    cond = ST.CondRule(column="Status", value="Fail", style=ST.TextStyle(background="red"))
    twice = ST.merge(ST.merge(None, ST.StylePatch(conditional=[cond])), ST.StylePatch(conditional=[cond.model_copy(update={"style": ST.TextStyle(background="orange")})]))
    assert len(twice.conditional) == 1 and twice.conditional[0].style.background == "#E07B00"


# ----------------------------------------------------------------- parser --


def _fields(text, kind):
    patch, unparsed = ST.parse_style_request(text, kind)
    return ST.patch_fields(patch), unparsed


def test_parse_style_request_accuracy_on_80_authored_phrases():
    """>= 90% field-level accuracy on 25 English, 20 Hinglish, 15 Hindi,
    15 Gujarati and 5 typo-heavy phrases; the residual model call is needed
    on <= 20% of them."""
    result = F.score(F.PHRASES, _fields)
    assert result["n"] == 80
    assert result["accuracy"] >= 0.90, result["misses"]
    assert result["llm_rate"] <= 0.20
    for group in (F.ENGLISH, F.HINGLISH, F.HINDI, F.GUJARATI):
        assert F.score(group, _fields)["accuracy"] >= 0.85


def test_parse_heldout_is_reported_not_tuned():
    """The held-out phrases were written with the set above and never used
    to change the parser; the floor here is what it scored when frozen."""
    result = F.score(F.HELDOUT, _fields)
    assert result["n"] == 20 and result["accuracy"] >= 0.85


def test_production_shape_request_parses_to_the_style_it_asks_for():
    text = ("just give it in docs in a standard and classy format, headings dark blue, body font Georgia 12pt, "
            "table header dark green with white bold text, landscape. provide a dox file")
    fields, unparsed = _fields(text, "document")
    assert unparsed == []
    assert fields == {
        "preset": "classic", "page.orientation": "landscape", "fonts.body": "Georgia", "base_size_pt": 12.0,
        "rule:heading:color": "#1F3864", "rule:table_header:background": "#1E6B34", "rule:table_header:color": "#FFFFFF",
        "rule:table_header:bold": True,
    }
    # Content clauses are not style: nothing about "docs" or "dox file".
    assert ST.parse_style_request("give me a summary of the audit")[0].is_empty()


def test_parser_notes_clamped_sizes_and_text_safe_substitutions():
    patch, _, notes = ST.parse_style_request_with_notes("title 100pt, orange headings", "document")
    assert any("72" in n for n in notes) and any("#B35F00" in n for n in notes)
    title = next(r for r in patch.rules if r.target.kind == "title")
    assert title.style.size_pt == 72
    # A user hex code is used as given, even when it fails contrast.
    patch, _, notes = ST.parse_style_request_with_notes("headings #E07B00", "document")
    assert patch.rules[0].style.color == "#E07B00" and not notes


def test_unreadable_style_phrases_are_returned_for_the_model_and_content_is_ignored():
    patch, unparsed = ST.parse_style_request("make the headings pop with a sort of vintage colour vibe", "document")
    assert unparsed and patch.rules == []
    patch, unparsed = ST.parse_style_request("summarise section 3 and add a risks table", "document")
    assert patch.is_empty() and unparsed == []


# --------------------------------------------------------- model fallback --


def test_extract_patch_llm_validates_every_rule_and_makes_one_call():
    calls = []

    async def fake(messages, **kw):
        calls.append(kw)
        return json.dumps({
            "rules": [
                {"target": {"kind": "heading", "level": 2}, "style": {"color": "burgundy-ish"}},
                {"target": {"kind": "heading"}, "style": {"color": "maroon", "italic": True}},
                {"target": {"kind": "title"}, "style": {"font_family": "Georgia; } @import url(x)"}},
                {"target": {"kind": "nonsense"}, "style": {"bold": True}},
            ],
            "orientation": "landscape", "body_font": "Wingdings Evil", "not_understood": ["vintage vibe"],
        })

    patch, notes = asyncio.run(ST.extract_patch_llm(["headings maroon italic, vintage vibe"], "document", completion=fake))
    assert len(calls) == 1 and calls[0]["max_tokens"] == 300 and calls[0]["thinking"] is False
    assert [(r.target.kind, r.style.color, r.style.italic) for r in patch.rules] == [("heading", "#7B1E1E", True)]
    assert patch.page.orientation == "landscape" and patch.fonts is None
    assert len(notes) == 5  # three bad rules, the font, the phrase not understood
    assert asyncio.run(ST.extract_patch_llm([], "document", completion=fake)) == (ST.StylePatch(), [])
    assert len(calls) == 1, "no phrases, no call"


def test_extract_patch_llm_times_out_to_nothing():
    async def slow(messages, **kw):
        await asyncio.sleep(2)
        return "{}"

    started = time.perf_counter()
    patch, notes = asyncio.run(ST.extract_patch_llm(["something"], "document", completion=slow, timeout_s=0.2))
    assert patch.is_empty() and notes and time.perf_counter() - started < 1.5


# --------------------------------------------------------------- resolver --


def test_contrast_ratio_matches_the_wcag_figures_in_the_style_guide():
    cr = ST.contrast_ratio
    assert round(cr("#1F2937", "#FFFFFF"), 1) == 14.7
    assert round(cr("#6B7280", "#F3F6FA"), 2) == 4.46   # caption on band: below 4.5, so never used there
    assert round(cr("#0E9D9A", "#FFFFFF"), 2) == 3.33   # teal as text fails
    assert round(cr("#E07B00", "#FFFFFF"), 1) == 3.0    # orange with white fails
    assert round(cr("#B7791F", "#FFFFFF"), 2) == 3.64   # amber as text fails
    assert round(cr("#FFFFFF", "#1F3864"), 1) == 11.6


def _system_pairs(tokens: ST.Tokens):
    """Every text/background pair the SYSTEM chooses in a preset."""
    return {
        "body on white": (tokens.ink, "#FFFFFF"), "body on band": (tokens.ink, tokens.band), "title": (tokens.title, "#FFFFFF"),
        "h1": (tokens.h1, "#FFFFFF"), "h2": (tokens.h2, "#FFFFFF"), "h3": (tokens.h3, "#FFFFFF"), "muted on white": (tokens.muted, "#FFFFFF"),
        "muted on band": (tokens.muted, tokens.band), "caption on white": (tokens.caption, "#FFFFFF"),
        "header": (tokens.header_text, tokens.header_fill), "totals": (tokens.total_text, tokens.total_fill),
        "kpi value on band": (tokens.primary, tokens.band), "link": (tokens.accent_text, "#FFFFFF"),
        "cover title": (ST.readable_on(tokens.cover_band), tokens.cover_band),
        **{f"status {k}": (text, fill) for k, (fill, text) in ST.STATUS_PAIRS.items()},
        **{f"score {c}": (tokens.ink, c) for c in ST.SCORE_SCALE},
    }


@pytest.mark.parametrize("preset", sorted(ST.PRESETS))
def test_every_system_chosen_pair_reaches_4_5_in_every_preset(preset):
    tokens = ST.resolve(S.ArtifactSpec(kind="document", document=S.DocumentSpec(title="t", blocks=[S.Paragraph(text="x")], style=ST.StyleSpec(preset=preset)))).tokens
    low = {name: round(ST.contrast_ratio(*pair), 2) for name, pair in _system_pairs(tokens).items() if ST.contrast_ratio(*pair) < 4.5}
    assert low == {}


def test_data_label_colour_is_the_better_of_white_and_ink_and_ink_outside_bars_passes():
    d = ST.resolve(None).chart_defaults
    for fill in ST.CHART_PALETTE:
        chosen = d.label_color_for(fill)
        other = ST.INK if chosen == ST.WHITE else ST.WHITE
        assert ST.contrast_ratio(chosen, fill) >= ST.contrast_ratio(other, fill)
    # Labels outside the bar sit on white in ink (style guide §6).
    assert ST.contrast_ratio(ST.INK, "#FFFFFF") >= 4.5


def _deutan_lab(hex_value: str):
    m = [[0.367322, 0.860646, -0.227968], [0.280085, 0.672501, 0.047413], [-0.011820, 0.042940, 0.968881]]
    h = hex_value.lstrip("#")
    c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    lin = [x / 12.92 if x <= 0.04045 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    rgb = [max(0.0, min(1.0, sum(m[i][j] * lin[j] for j in range(3)))) for i in range(3)]
    X = 0.4124 * rgb[0] + 0.3576 * rgb[1] + 0.1805 * rgb[2]
    Y = 0.2126 * rgb[0] + 0.7152 * rgb[1] + 0.0722 * rgb[2]
    Z = 0.0193 * rgb[0] + 0.1192 * rgb[1] + 0.9505 * rgb[2]
    f = lambda t: t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116  # noqa: E731
    fx, fy, fz = f(X / 0.9505), f(Y), f(Z / 1.089)
    return (116 * fy - 16, 500 * (fx - fy), 200 * (fy - fz))


def test_chart_palette_first_five_are_distinct_under_deuteranopia():
    for a, b in itertools.combinations(ST.CHART_PALETTE[:5], 2):
        assert math.dist(_deutan_lab(a), _deutan_lab(b)) >= 25, (a, b)


def test_system_text_flips_on_a_user_fill_and_explicit_pairs_are_honoured_with_one_warning():
    body = S.DocumentSpec(title="t", blocks=[S.TableBlock(table=S.Table(columns=["A"], rows=[["x"]]))])
    body.style = ST.StyleSpec(rules=[
        ST.StyleRule(target=ST.StyleTarget(kind="table_header"), style=ST.TextStyle(background="light yellow")),
        ST.StyleRule(target=ST.StyleTarget(kind="caption"), style=ST.TextStyle(color="yellow")),
    ])
    R = ST.resolve(body)
    header = R.element("table_header")
    assert header.background == "#FFF1C7" and header.color == ST.INK, "white on light yellow flips to ink"
    caption = R.element("caption")
    assert caption.color == "#FFD54F", "an explicit user colour is never changed"
    R.element("caption")
    assert len([w for w in R.warnings if "hard to read" in w]) == 1


def test_resolve_defaults_are_techsara_classic():
    R = ST.resolve(None)
    assert R.tokens.header_fill == "#1F3864" and R.tokens.band == "#F3F6FA" and R.tokens.grid == "#E5E9F0"
    assert R.body_face.name == "Calibri" and "Carlito" in R.body_face.css_stack
    h1, h2, h3 = (R.element("heading", level=n) for n in (1, 2, 3))
    assert (h1.size_pt, h1.color, h2.size_pt, h2.color, h3.size_pt, h3.color) == (18, "#1F3864", 14, "#2E5597", 12, "#1F2937")
    assert R.element("title").size_pt == 28 and R.element("caption").italic is True


def test_specific_rules_only_touch_their_element():
    body = S.DocumentSpec(title="t", blocks=[S.Heading(level=1, text="Scope"), S.Heading(level=1, text="3. Recommendations")])
    body.style = ST.merge(None, ST.parse_style_request("make the Recommendations heading red")[0])
    R = ST.resolve(body)
    assert R.element("heading", level=1, text="Scope").color == "#1F3864"
    assert R.element("heading", level=1, text="3. Recommendations").color == "#C62828"
    assert R.element("heading", generic_only=True, level=1).color == "#1F3864"


def test_normalize_drops_rules_whose_targets_do_not_exist_with_a_note():
    spec = document("generic", sections=1)
    spec.body.style = ST.merge(None, ST.parse_style_request("make the Appendix heading red, Owner column bold, slide titles green, headings navy")[0])
    _, notes = ST.normalize_spec_style(spec)
    kinds = [(r.target.kind, r.target.text or r.target.name) for r in spec.body.style.rules]
    assert kinds == [("heading", None)]
    assert len(notes) == 3 and all(n.startswith("Not applied:") for n in notes)


def test_a_status_condition_binds_to_the_sheet_s_real_status_column():
    spec = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", sheets=[
        S.Sheet(name="Tasks", columns=[S.Column(name="Task"), S.Column(name="Current state")], rows=[["a", "Blocked"]]),
    ]))
    spec.body.style = ST.merge(None, ST.parse_style_request("Blocked in red", "workbook")[0])
    ST.normalize_spec_style(spec)
    assert [c.column for c in spec.body.style.conditional] == ["Current state"]


# --------------------------------------------------------- support matrix --


def test_support_matrix_is_declared_per_format_and_kind():
    docx = ST.support_matrix("docx", "document")
    assert docx[("heading", "color")] == "supported" and docx[("cell_range", "background")] == "unsupported"
    assert docx[("page", "background")] == "unsupported"
    assert ST.support_matrix("pptx", "presentation")[("page", "background")] == "supported"
    assert ST.support_matrix("xlsx", "workbook")[("cell_range", "background")] == "supported"
    assert all(v == "unsupported" for v in ST.support_matrix("csv", "workbook").values())
    assert ST.support_matrix("xlsx", "document") == {}


def test_warnings_for_emit_exactly_one_sentence_per_unsupported_group():
    spec = deck("generic")
    spec.body.style = ST.merge(None, ST.parse_style_request("slide text right aligned on a yellow background, slide background light grey", "presentation")[0])
    spec.body.style.rules.append(ST.StyleRule(target=ST.StyleTarget(kind="slide_body"), style=ST.TextStyle(background="yellow", align="right")))
    warnings = ST.warnings_for(spec, ["pptx", "pdf"])
    assert len([w for w in warnings if "slide text" in w and "PowerPoint" in w]) == 1
    assert len([w for w in warnings if "slide text" in w and "PDF" in w]) == 1
    assert not any("page" in w for w in warnings), "a slide background is supported in both"
    doc = document("generic", sections=1)
    doc.body.style = ST.StyleSpec(page=ST.PageStyle(background="light yellow"))
    assert ST.warnings_for(doc, ["docx", "pdf"]) == [
        "The Word file has no page background colour (Word does not print one and it floods a PDF with ink); element backgrounds were kept.",
        "The PDF has no page background colour (Word does not print one and it floods a PDF with ink); element backgrounds were kept.",
    ]
    wb = workbook()
    wb.body.style = ST.merge(None, ST.parse_style_request("blue header", "workbook")[0])
    assert ST.warnings_for(wb, ["csv", "xlsx"]) == [ST.CSV_STYLE_SENTENCE]


def test_csv_with_styling_also_delivers_a_styled_xlsx(tmp_path):
    from openpyxl import load_workbook

    from app.artifacts.render import render_version

    spec = workbook(rows=6)
    spec, formats, notes, unparsed = ST.apply_request(spec, "blue header, bold totals", formats=["csv"])
    assert formats == ["csv", "xlsx"] and notes == [ST.CSV_STYLE_SENTENCE] and unparsed == []
    report = render_version(spec, formats, str(tmp_path), title_slug="styled", version=1)
    assert sorted({f.format for f in report.files}) == ["csv", "xlsx"]
    xlsx = next(f for f in report.files if f.format == "xlsx")
    ws = load_workbook(str(tmp_path / xlsx.filename))["Pipeline"]
    assert "#" + ws["A1"].fill.fgColor.rgb[-6:] == ST.resolve_color("blue")
    assert ST.CSV_STYLE_SENTENCE in report.warnings
    # Unstyled: the CSV alone, no sentence.
    assert ST.formats_with_style(["csv"], kind="workbook", styled=False) == (["csv"], [])


# ----------------------------------------------------------------- schema --


#: The guided-JSON schema lengths before styling existed (measured on the
#: base tree, cb0b4e3): styling adds 0 characters to guided decoding.
#: AS3 integration: plus the charts track's GUIDED Chart v2 (binding, the
#: guided ChartStyle subset, filters) — still no styling characters.
#: 2026-09-16: +1826 for the six advanced chart types (their names in the
#: `type` enum and the bullet chart's `target` column). The number is pinned
#: so that anything ELSE creeping into guided decoding is caught here.
BASE_SCHEMA_LENGTHS = {"document": 14972, "presentation": 12872, "workbook": 21356}


@pytest.mark.parametrize("kind", sorted(BASE_SCHEMA_LENGTHS))
def test_style_is_not_in_the_guided_schema(kind):
    schema = json.dumps(S.schema_for(kind))
    assert len(schema) == BASE_SCHEMA_LENGTHS[kind]
    assert "style" not in S.schema_for(kind)["properties"]
    assert "StyleSpec" not in schema and "ColumnFormat" not in schema and "TextStyle" not in schema


def test_parse_body_drops_model_written_style_but_load_keeps_it():
    body = {"title": "t", "blocks": [{"type": "paragraph", "text": "x"}], "style": {"preset": "boardroom"}}
    assert S.parse_body("document", body).body.style is None
    wb = {"title": "w", "sheets": [{"name": "s", "columns": [{"name": "a", "type": "currency", "format": {"currency": "INR"}, "align": "center"}], "rows": [[1]]}]}
    col = S.parse_body("workbook", wb).body.sheets[0].columns[0]
    assert col.format is None and col.align is None
    stored = S.ArtifactSpec(kind="document", document=S.DocumentSpec(title="t", blocks=[S.Paragraph(text="x")], style=ST.StyleSpec(preset="boardroom")))
    again = S.load(stored.model_dump(mode="json", exclude_none=True))
    assert again.body.style.preset == "boardroom"


def test_style_numbers_are_not_unsupported_figures():
    spec = workbook(rows=3)
    before = S.unsupported_figures(spec, "")
    spec.body.style = ST.merge(None, ST.parse_style_request("score above 80000.5 green, title 13pt", "workbook")[0])
    assert S.unsupported_figures(spec, "") == before


# -------------------------------------------------------------- injection --


def _formula_literals(formula: str):
    return re.findall(r'"(?:[^"]|"")*"', formula)


def test_condition_value_cannot_escape_its_formula_literal(tmp_path):
    from openpyxl import load_workbook

    from app.artifacts.render.xlsx import render_xlsx

    hostile = '"),HYPERLINK("http://x'
    spec = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", sheets=[
        S.Sheet(name="Data", columns=[S.Column(name="Status"), S.Column(name="Note")], rows=[["Open", hostile], ["Done", "fine"]]),
    ]))
    spec.body.style = ST.StyleSpec(conditional=[
        ST.CondRule(column="Note", op="eq", value=hostile, style=ST.TextStyle(background="red")),
        ST.CondRule(column="Note", op="contains", value='*?~"x', style=ST.TextStyle(background="yellow")),
    ])
    path = render_xlsx(spec.body, tmp_path / "inj.xlsx")
    ws = load_workbook(str(path))["Data"]
    formulas = [r.formula[0] for cf in ws.conditional_formatting for r in cf.rules if r.formula]
    eq = next(f for f in formulas if "TRIM($B2)" in f)
    assert _formula_literals(eq) == ['"""),HYPERLINK(""http://x"'] and eq == 'TRIM($B2)="""),HYPERLINK(""http://x"'
    contains = next(f for f in formulas if "SEARCH" in f)
    assert _formula_literals(contains) == ['"~*~?~~""x"']
    # Outside the literals only code-made tokens remain.
    for f in formulas:
        outside = re.sub(r'"(?:[^"]|"")*"', "", f)
        assert "HYPERLINK" not in outside and "http" not in outside, f
    assert ST.xlsx_formula_literal("a" * 300) == '"' + "a" * 100 + '"'


def test_footer_text_ampersands_are_doubled_in_the_xlsx_header_footer(tmp_path):
    from app.artifacts.render.xlsx import render_xlsx

    spec = workbook(rows=3)
    spec.body.style = ST.StyleSpec(header_footer=ST.HeaderFooter(footer_text="&F secret", header_text="R&D &A"))
    path = render_xlsx(spec.body, tmp_path / "hf.xlsx")
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    footer = re.search(r"<oddFooter>(.*?)</oddFooter>", xml).group(1)
    header = re.search(r"<oddHeader>(.*?)</oddHeader>", xml).group(1)
    assert "&amp;&amp;F secret" in footer and "&amp;P of &amp;N" in footer
    assert "R&amp;&amp;D &amp;&amp;A" in header
    assert ST.xlsx_header_footer_text("&F secret") == "&&F secret"


def test_generated_css_never_contains_url_or_import_for_random_specs():
    from app.artifacts.render import html as H

    rng = random.Random(7)
    fonts = list(ST.FONT_ALLOWLIST.values())
    colours = list(ST.COLOR_NAMES)
    kinds = ["title", "subtitle", "heading", "paragraph", "bullet", "caption", "table_header", "table_body", "kpi_value", "callout", "header_footer"]
    for i in range(20):
        rules = []
        for _ in range(rng.randint(1, 8)):
            kind = rng.choice(kinds)
            target = ST.StyleTarget(kind=kind, text="url(x) } @import 'y'" if kind == "heading" and rng.random() < 0.5 else None) if kind == "heading" else ST.StyleTarget(kind=kind)
            rules.append(ST.StyleRule(target=target, style=ST.TextStyle(
                font_family=rng.choice(fonts).name, size_pt=rng.uniform(6, 72), bold=rng.random() < 0.5, italic=rng.random() < 0.5,
                underline=rng.random() < 0.5, color=rng.choice(colours), background=rng.choice(colours), align=rng.choice(["left", "center", "right", "justify"]))))
        spec = document("generic", sections=2, title=f"t{i} url(http://evil) @import x")
        spec.body.blocks.insert(0, S.Heading(level=1, text="url(x) } @import 'y'"))
        spec.body.style = ST.StyleSpec(preset=rng.choice(sorted(ST.PRESETS)), rules=rules, header_footer=ST.HeaderFooter(footer_text="url(z) @import"))
        out = H.document_html(spec.body)
        css = re.search(r"<style>(.*)</style>", out, re.S).group(1)
        assert "url(" not in css and "@import" not in css, i
        assert "evil" not in css and "'y'" not in css


def test_a_huge_cell_range_is_clipped_to_the_used_range_quickly(tmp_path):
    from openpyxl import load_workbook

    from app.artifacts.render.xlsx import render_xlsx

    spec = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="w", sheets=[
        S.Sheet(name="Data", columns=[S.Column(name=f"c{j}") for j in range(5)], rows=[[f"r{i}c{j}" for j in range(5)] for i in range(50)]),
    ]))
    spec.body.style = ST.StyleSpec(rules=[ST.StyleRule(target=ST.StyleTarget(kind="cell_range", a1="A1:XFD1048576"), style=ST.TextStyle(background="light green"))])
    assert ST.clip_a1("A1:XFD1048576", 5, 51) == (1, 1, 5, 51)
    started = time.perf_counter()
    path = render_xlsx(spec.body, tmp_path / "huge.xlsx")
    assert time.perf_counter() - started < 2.0
    ws = load_workbook(str(path))["Data"]
    ranges = [str(cf.sqref) for cf in ws.conditional_formatting for r in cf.rules if r.formula == ["TRUE"]]
    assert ranges == ["A2:E51"]
    assert ws["E51"].fill.fgColor.rgb.endswith("E3F2E6") and ws.max_column == 5


# ------------------------------------------------------------ formatting --


def test_format_cell_shares_the_xlsx_display_rules():
    from app.artifacts.render import theme

    col = lambda t, **f: S.Column(name="x", type=t, format=S.ColumnFormat(**f) if f else None)  # noqa: E731
    assert theme.format_cell(1234567, col("integer")) == "1,234,567"
    assert theme.format_cell(1234.5, col("number")) == "1,234.50" and theme.format_cell(1234.0, col("number")) == "1,234"
    assert theme.format_cell(0.125, col("percent")) == "12.5%" and theme.format_cell(12.5, col("percent"), percent_scale=0.01) == "12.5%"
    assert theme.format_cell(12000000.5, col("currency", currency="INR")) == "₹1,20,00,000.50"
    assert theme.format_cell(-1500, col("currency", currency="USD")) == "-$1,500.00"
    assert theme.format_cell("2026-08-03", col("date")) == "03-Aug-2026" and theme.format_cell("2026-08-03", col("date", date_style="iso")) == "2026-08-03"
    assert theme.format_cell("007", col("integer")) == "007" and theme.format_cell("n/a", col("number")) == "n/a"
    assert theme.format_cell(2024, "text") == "2024"


# ------------------------------------------------ verifier regressions --


def _doc_with_sections() -> S.ArtifactSpec:
    blocks = [S.Heading(level=1, text="Section 1: Overview"), S.Paragraph(text="Overview text."),
              S.Heading(level=1, text="Risks"), S.Paragraph(text="Risk text."), S.Heading(level=2, text="Detail"), S.Paragraph(text="Detail text.")]
    return S.ArtifactSpec(kind="document", document=S.DocumentSpec(title="Review", blocks=blocks))


@pytest.mark.parametrize("kind,first,second,target_kind", [
    ("workbook", "row 5 yellow", "remove the yellow from row 5", "row"),
    ("document", "headings red", "remove the red colour from headings", "heading"),
    ("document", "table header dark green", "remove the background from the table header", "table_header"),
    ("document", "headings red", "हेडिंग से लाल रंग हटाओ", "heading"),
])
def test_a_removal_takes_the_colour_off_and_never_adds_it(kind, first, second, target_kind):
    style = ST.merge(None, ST.parse_style_request(first, kind)[0])
    patch, unparsed = ST.parse_style_request(second, kind)
    assert not patch.rules and patch.clear and not unparsed, ST.patch_fields(patch)
    merged = ST.merge(style, patch)
    assert not [r for r in merged.rules if r.target.kind == target_kind and (r.style.color or r.style.background)]


def test_a_removal_keeps_the_other_properties_of_the_rule():
    style = ST.merge(None, ST.parse_style_request("row 5 yellow and bold", "workbook")[0])
    merged = ST.merge(style, ST.parse_style_request("row 5 no fill", "workbook")[0])
    (rule,) = merged.rules
    assert rule.style.bold is True and rule.style.background is None


def test_removing_a_column_colour_drops_its_conditional_colours():
    style = ST.merge(None, ST.parse_style_request("Status column red for Blocked", "workbook")[0])
    assert style.conditional
    merged = ST.merge(style, ST.parse_style_request("Status column: remove the red", "workbook")[0])
    assert not merged.conditional


@pytest.mark.parametrize("text", ["don't make the headings blue", "do not make the header blue", "headings ko blue mat karo"])
def test_a_negated_request_adds_nothing(text):
    patch, unparsed = ST.parse_style_request(text, "document")
    assert patch.is_empty() and not unparsed, ST.patch_fields(patch)


def test_remove_all_formatting_resets_to_the_house_style():
    style = ST.merge(None, ST.parse_style_request("headings red, landscape, footer 'X'", "document")[0])
    patch, _ = ST.parse_style_request("remove all formatting", "document")
    assert patch.reset
    assert ST.merge(style, patch) == ST.StyleSpec()


@pytest.mark.parametrize("text,level,prop,value", [
    ("heading 1 size 20", 1, "size_pt", 20.0),
    ("make heading 2 blue", 2, "color", "#2F6FB2"),
    ("level 3 headings italic", 3, "italic", True),
])
def test_a_heading_number_is_a_level_never_a_size(text, level, prop, value):
    patch, _ = ST.parse_style_request(text, "document")
    (rule,) = patch.rules
    assert rule.target.level == level and getattr(rule.style, prop) == value


def test_headings_12pt_is_still_a_size():
    (rule,) = ST.parse_style_request("headings 12pt", "document")[0].rules
    assert rule.target.level is None and rule.style.size_pt == 12


@pytest.mark.parametrize("text", [
    "write a report on the red team exercise and the blue ocean strategy",
    "make a pdf about the white house press briefing",
    "a sheet of the top 10 black friday deals",
    "write about green energy in India, 5 pages",
])
def test_colour_words_inside_content_are_not_style(text):
    patch, unparsed, notes = ST.parse_style_request_with_notes(text, "document")
    assert patch.is_empty() and not unparsed and not notes


def test_a_trailing_colour_in_a_content_clause_is_still_a_style_phrase():
    _, unparsed = ST.parse_style_request("write a report on sales in blue", "document")
    assert unparsed


def test_an_unknown_font_containing_sans_is_not_silently_arial():
    patch, unparsed = ST.parse_style_request("title font Comic Sans MS", "document")
    assert not any(r.style.font_family for r in patch.rules) and unparsed
    (rule,) = ST.parse_style_request("title font sans-serif", "document")[0].rules
    assert rule.style.font_family == "Arial"


def test_paragraphs_of_one_section_stay_in_that_section():
    spec = _doc_with_sections()
    patch, _ = ST.parse_style_request("paragraphs in the Risks section italic", "document")
    (rule,) = patch.rules
    assert rule.target.kind == "paragraph" and rule.target.section == "Risks"
    spec.document.style = ST.merge(None, ST.parse_style_request("paragraphs in the Budget section italic", "document")[0])
    _, notes = ST.normalize_spec_style(spec)
    assert not spec.document.style.rules and any("Budget" in n for n in notes)


def test_a_spreadsheet_column_letter_maps_to_the_column_position():
    spec = S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="T", sheets=[S.Sheet(
        name="S", columns=[S.Column(name="Name"), S.Column(name="Owner")], rows=[["a", "b"]])]))
    spec.workbook.style = ST.merge(None, ST.parse_style_request("column B yellow", "workbook")[0])
    _, notes = ST.normalize_spec_style(spec)
    (rule,) = spec.workbook.style.rules
    assert rule.target.index == 2 and rule.target.name is None and not notes
