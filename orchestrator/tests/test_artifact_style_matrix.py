"""The declared support matrix, proven on PRODUCED files.

For every (format, kind) the style module declares, and every target kind
with a "supported" property, a spec gets ONE rule that sets every supported
property to a value no default has (a serif family, 15 pt, the opposite
weight and slant, underline, a burgundy text colour, a mint background,
centred); the file is rendered through render_version and read back by
tests/artifact_file_readers.py — python-docx, python-pptx, openpyxl and
PDFium page objects, never the renderer's own style objects. A declared
cell that the file does not show fails with its (format, kind, target,
property) name.
"""
from __future__ import annotations

from typing import Dict, List, Tuple

import pytest

from app.artifacts import spec as S
from app.artifacts import style as ST
from app.artifacts.render import render_version, theme
from tests import artifact_file_readers as FR

FONT = "DejaVu Serif"
SIZE = 15.0
COLOR = "#8B1E3F"
BACKGROUND = "#D9F2E6"

pytest.importorskip("docx")
pytest.importorskip("pptx")
pytest.importorskip("openpyxl")


def _document() -> S.ArtifactSpec:
    return S.ArtifactSpec(kind="document", document=S.DocumentSpec(
        title="MKTITLE Report", subtitle="MKSUB line", author="Audit team", date="2026-09-15",
        blocks=[
            S.Heading(level=1, text="MKHEAD Scope"),
            S.Paragraph(text="MKPARA The review covered the controls in scope."),
            S.Bullets(items=["MKBULLET first point", "second point"]),
            S.TableBlock(table=S.Table(columns=["MKCOL", "Count"], rows=[["MKCELL", 1], ["MKROW2", 2], ["third", 3]], numeric_columns=[1], caption="MKCAP table")),
            S.KPIRow(items=[S.KPI(label="MKKPIL", value="MKKPIV")]),
            S.Callout(kind="note", title="Note", text="MKCALL call-out text"),
        ]))


def _workbook() -> S.ArtifactSpec:
    return S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="MKTITLE Book", purpose="MKSUB line", sheets=[
        S.Sheet(name="Data", columns=[S.Column(name="MKCOL"), S.Column(name="Count", type="integer")],
                rows=[["MKCELL", 1], ["MKROW2", 2], ["third", 3]], totals=[S.Total(column="Count", label="MKTOTAL")], notes="MKCAP note"),
    ]))


def _deck() -> S.ArtifactSpec:
    chart = S.Chart(type="bar", title="MKCHART", categories=["A", "B"], series=[S.Series(name="One", values=[1, 2]), S.Series(name="Two", values=[2, 1])])
    return S.ArtifactSpec(kind="presentation", presentation=S.PresentationSpec(title="MKDECK", slides=[
        S.Slide(layout="title", title="MKTITLE deck", subtitle="MKSUB line"),
        S.Slide(layout="bullets", title="MKSLIDE findings", bullets=["MKBODY first point", "second point"]),
        S.Slide(layout="table", title="Table slide", table=S.Table(columns=["MKCOL", "Count"], rows=[["MKCELL", 1], ["MKROW2", 2]], numeric_columns=[1], caption="MKCAP table")),
        S.Slide(layout="kpis", title="KPI slide", kpis=[S.KPI(label="MKKPIL", value="MKKPIV")]),
        S.Slide(layout="chart", title="Chart slide", chart=chart),
    ]))


#: target kind → (the rule's target, the marker whose element shows it).
DOC_TARGETS: Dict[str, Tuple[ST.StyleTarget, str]] = {
    "title": (ST.StyleTarget(kind="title"), "MKTITLE"), "subtitle": (ST.StyleTarget(kind="subtitle"), "MKSUB"),
    "heading": (ST.StyleTarget(kind="heading"), "MKHEAD"), "paragraph": (ST.StyleTarget(kind="paragraph"), "MKPARA"),
    "bullet": (ST.StyleTarget(kind="bullet"), "MKBULLET"), "caption": (ST.StyleTarget(kind="caption"), "MKCAP"),
    "table_header": (ST.StyleTarget(kind="table_header"), "MKCOL"), "table_body": (ST.StyleTarget(kind="table_body"), "MKCELL"),
    "table": (ST.StyleTarget(kind="table"), "MKCELL"), "column": (ST.StyleTarget(kind="column", name="MKCOL"), "MKCELL"),
    "row": (ST.StyleTarget(kind="row", index=2), "MKROW2"), "kpi_value": (ST.StyleTarget(kind="kpi_value"), "MKKPIV"),
    "kpi_label": (ST.StyleTarget(kind="kpi_label"), "MKKPIL"), "callout": (ST.StyleTarget(kind="callout"), "MKCALL"),
    "header_footer": (ST.StyleTarget(kind="header_footer"), "MKFOOT"),
}
WB_TARGETS: Dict[str, Tuple[ST.StyleTarget, str]] = {
    "title": (ST.StyleTarget(kind="title"), "MKTITLE"), "subtitle": (ST.StyleTarget(kind="subtitle"), "MKSUB"),
    "caption": (ST.StyleTarget(kind="caption"), "MKCAP"), "table_header": (ST.StyleTarget(kind="table_header"), "MKCOL"),
    "table_body": (ST.StyleTarget(kind="table_body"), "MKCELL"), "table_total": (ST.StyleTarget(kind="table_total"), "MKTOTAL"),
    "table": (ST.StyleTarget(kind="table"), "MKCELL"), "column": (ST.StyleTarget(kind="column", name="MKCOL"), "MKCELL"),
    "row": (ST.StyleTarget(kind="row", index=3), "MKROW2"), "cell_range": (ST.StyleTarget(kind="cell_range", a1="A2"), "MKCELL"),
    "header_footer": (ST.StyleTarget(kind="header_footer"), "MKHF"),
}
DECK_TARGETS: Dict[str, Tuple[ST.StyleTarget, str]] = {
    "title": (ST.StyleTarget(kind="title"), "MKTITLE"), "subtitle": (ST.StyleTarget(kind="subtitle"), "MKSUB"),
    "slide_title": (ST.StyleTarget(kind="slide_title"), "MKSLIDE"), "slide_body": (ST.StyleTarget(kind="slide_body"), "MKBODY"),
    "bullet": (ST.StyleTarget(kind="bullet"), "MKBODY"), "table_header": (ST.StyleTarget(kind="table_header"), "MKCOL"),
    "table_body": (ST.StyleTarget(kind="table_body"), "MKCELL"), "table": (ST.StyleTarget(kind="table"), "MKCELL"),
    "column": (ST.StyleTarget(kind="column", name="MKCOL"), "MKCELL"), "row": (ST.StyleTarget(kind="row", index=2), "MKROW2"),
    "kpi_value": (ST.StyleTarget(kind="kpi_value"), "MKKPIV"), "kpi_label": (ST.StyleTarget(kind="kpi_label"), "MKKPIL"),
    "caption": (ST.StyleTarget(kind="caption"), "MKCAP"), "header_footer": (ST.StyleTarget(kind="header_footer"), "MKFOOT"),
}
CHART_TARGETS = ("chart_title", "chart_axis", "chart_legend", "chart_labels")


def _supported(fmt: str, kind: str, target: str) -> List[str]:
    m = ST.support_matrix(fmt, kind)
    return [p for p in ST._TEXT_PROPS if m.get((target, p)) == "supported"]


def _rule_style(target: str, props: List[str]) -> ST.TextStyle:
    base = ST.resolve(None).base(target if target not in ("column", "row", "cell_range", "table") else "table_body")
    values = {"font_family": FONT, "size_pt": SIZE, "bold": not base.bold, "italic": not base.italic, "underline": True,
              "color": COLOR, "background": BACKGROUND, "align": "center"}
    return ST.TextStyle(**{p: values[p] for p in props})


def _expect(prop: str, style: ST.TextStyle):
    return {"font_family": FONT, "size_pt": SIZE, "bold": style.bold, "italic": style.italic, "underline": True,
            "color": COLOR, "background": BACKGROUND, "align": "center"}[prop]


def _check(observed: Dict, props: List[str], style: ST.TextStyle, where: str, *, pdf: bool = False) -> List[str]:
    failures = []
    if observed is None:
        return [f"{where}: element not found"]
    for prop in props:
        want = _expect(prop, style)
        if prop == "font_family":
            got = observed["font"] or ""
            ok = FONT.replace(" ", "") in got.replace(" ", "").replace("-", "") if pdf else got == FONT
            if pdf and not theme.font_installed(FONT):
                continue
        elif prop == "size_pt":
            got = observed["size"]
            ok = got is not None and abs(float(got) - SIZE) <= 0.3
        else:
            got = observed[{"color": "color", "background": "background", "align": "align"}.get(prop, prop)]
            ok = got == want
        if not ok:
            failures.append(f"{where}.{prop}: want {want!r}, got {got!r}")
    return failures


def _render(spec: S.ArtifactSpec, formats: List[str], tmp_path, name: str):
    out = tmp_path / name
    out.mkdir()
    report = render_version(spec, formats, str(out), title_slug="m", version=1)
    return {f.format: str(out / f.filename) for f in report.files}, report


@pytest.mark.parametrize("target", sorted(DOC_TARGETS))
def test_document_docx_and_pdf_show_every_supported_property(tmp_path, target):
    from docx import Document

    docx_props, pdf_props = _supported("docx", "document", target), _supported("pdf", "document", target)
    props = sorted(set(docx_props) | set(pdf_props))
    assert props, target
    rule_target, marker = DOC_TARGETS[target]
    style = _rule_style(target, props)
    spec = _document()
    spec.body.style = ST.StyleSpec(rules=[ST.StyleRule(target=rule_target, style=style)], header_footer=ST.HeaderFooter(footer_text="MKFOOT footer"))
    files, _ = _render(spec, ["docx", "pdf"], tmp_path, target)
    failures = _check(FR.docx_observe(Document(files["docx"]), marker), docx_props, style, f"docx/document/{target}")
    failures += _check(FR.pdf_observe(files["pdf"], marker), pdf_props, style, f"pdf/document/{target}", pdf=True)
    assert not failures, "\n" + "\n".join(failures)


@pytest.mark.parametrize("target", sorted(WB_TARGETS))
def test_workbook_xlsx_docx_and_pdf_show_every_supported_property(tmp_path, target):
    from docx import Document

    by_fmt = {fmt: _supported(fmt, "workbook", target) for fmt in ("xlsx", "docx", "pdf")}
    props = sorted({p for v in by_fmt.values() for p in v})
    if not props:
        pytest.skip(f"no format declares {target}")
    rule_target, marker = WB_TARGETS[target]
    style = _rule_style(target, props)
    spec = _workbook()
    spec.body.style = ST.StyleSpec(rules=[ST.StyleRule(target=rule_target, style=style)],
                                   header_footer=ST.HeaderFooter(header_text="MKHF header"), banded=False)
    files, _ = _render(spec, ["xlsx", "docx", "pdf"], tmp_path, target)
    failures = []
    if by_fmt["xlsx"]:
        failures += _check(FR.xlsx_observe(files["xlsx"], marker), by_fmt["xlsx"], style, f"xlsx/workbook/{target}")
    if by_fmt["docx"]:
        failures += _check(FR.docx_observe(Document(files["docx"]), marker), by_fmt["docx"], style, f"docx/workbook/{target}")
    if by_fmt["pdf"]:
        failures += _check(FR.pdf_observe(files["pdf"], marker), by_fmt["pdf"], style, f"pdf/workbook/{target}", pdf=True)
    assert not failures, "\n" + "\n".join(failures)


@pytest.mark.parametrize("target", sorted(DECK_TARGETS))
def test_presentation_pptx_and_pdf_show_every_supported_property(tmp_path, target):
    from pptx import Presentation

    pptx_props, pdf_props = _supported("pptx", "presentation", target), _supported("pdf", "presentation", target)
    props = sorted(set(pptx_props) | set(pdf_props))
    rule_target, marker = DECK_TARGETS[target]
    style = _rule_style(target, props)
    spec = _deck()
    spec.body.style = ST.StyleSpec(rules=[ST.StyleRule(target=rule_target, style=style)], header_footer=ST.HeaderFooter(footer_text="MKFOOT deck"))
    files, _ = _render(spec, ["pptx", "pdf"], tmp_path, target)
    failures = _check(FR.pptx_observe(Presentation(files["pptx"]), marker), pptx_props, style, f"pptx/presentation/{target}")
    failures += _check(FR.pdf_observe(files["pdf"], marker), pdf_props, style, f"pdf/presentation/{target}", pdf=True)
    assert not failures, "\n" + "\n".join(failures)


def test_pptx_native_charts_show_their_supported_properties(tmp_path):
    from pptx import Presentation

    rules = []
    styles = {}
    for target in CHART_TARGETS:
        props = _supported("pptx", "presentation", target)
        styles[target] = (props, _rule_style(target, props))
        rules.append(ST.StyleRule(target=ST.StyleTarget(kind=target), style=styles[target][1]))
    spec = _deck()
    spec.body.style = ST.StyleSpec(rules=rules)
    files, _ = _render(spec, ["pptx"], tmp_path, "charts")
    facts = FR.pptx_chart_fonts(Presentation(files["pptx"]))
    failures = []
    for target, (props, style) in styles.items():
        observed = facts.get(target)
        failures += _check({**observed, "background": None, "align": None} if observed else None, props, style, f"pptx/{target}")
    assert not failures, "\n" + "\n".join(failures)


def test_page_rows_of_the_matrix(tmp_path):
    """orientation, size, margins and page numbers on the files that declare
    them; a slide background on the deck and its preview."""
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    doc = _document()
    doc.body.style = ST.StyleSpec(page=ST.PageStyle(orientation="landscape", size="Letter", margins="narrow"), header_footer=ST.HeaderFooter(page_numbers=False))
    files, _ = _render(doc, ["docx", "pdf"], tmp_path, "doc")
    section = Document(files["docx"]).sections[0]
    assert round(section.page_width.mm) == 279 and round(section.page_height.mm) == 216 and round(section.left_margin.mm, 1) == 12.7
    import zipfile

    with zipfile.ZipFile(files["docx"]) as z:
        assert b"NUMPAGES" not in b"".join(z.read(n) for n in z.namelist() if n.startswith("word/footer"))
    objects = FR.pdf_objects(files["pdf"])
    assert round(objects[0]["page_width"]) == 792 and round(objects[0]["page_height"]) == 612
    assert not any(o["kind"] == "text" and o["text"].startswith("Page ") for o in objects)

    wb = _workbook()
    wb.body.style = ST.StyleSpec(page=ST.PageStyle(orientation="landscape", size="Legal", margins="wide"), header_footer=ST.HeaderFooter(page_numbers=True))
    files, _ = _render(wb, ["xlsx", "docx", "pdf"], tmp_path, "wb")
    ws = load_workbook(files["xlsx"])["Data"]
    assert ws.page_setup.orientation == "landscape" and int(ws.page_setup.paperSize) == 5 and ws.oddFooter.right.text == "Page &P of &N"
    assert ws.page_margins.left == 1.0
    section = Document(files["docx"]).sections[0]
    assert round(section.page_width.mm) == 356 and round(section.page_height.mm) == 216 and round(section.left_margin.mm, 1) == 25.4
    with zipfile.ZipFile(files["docx"]) as z:
        assert b"NUMPAGES" in b"".join(z.read(n) for n in z.namelist() if n.startswith("word/footer"))
    objects = FR.pdf_objects(files["pdf"])
    assert round(objects[0]["page_width"]) == 1008 and round(objects[0]["page_height"]) == 612
    assert any(o["kind"] == "text" and o["text"].startswith("Page 1 of") for o in objects)

    dk = _deck()
    dk.body.style = ST.StyleSpec(page=ST.PageStyle(background="light yellow"), header_footer=ST.HeaderFooter(page_numbers=False))
    files, _ = _render(dk, ["pptx", "pdf"], tmp_path, "deck")
    prs = Presentation(files["pptx"])
    slide = prs.slides[1]
    assert str(slide.background.fill.fore_color.rgb) == "FFF1C7"
    assert not any(shape.has_text_frame and shape.text_frame.text.strip() == "2 / 5" for shape in slide.shapes)
    objects = FR.pdf_objects(files["pdf"])
    assert any(o["kind"] == "path" and o["page"] == 1 and o["fill"] == "#FFF1C7" and o["bounds"][2] - o["bounds"][0] > 900 for o in objects)
