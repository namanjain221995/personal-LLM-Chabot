"""Styling on the produced files, format by format: the professional XLSX
default and its conditional-format priority, the classy DOCX (clean title,
theme fonts, footer fields, schema-ordered XML), the PDF's fonts, colours and
page box, the PPTX, the styled sheet preview, file-level DOCX↔PDF parity and
back-compat with the pre-styling renderers."""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pytest

from app.artifacts import spec as S
from app.artifacts import style as ST
from app.artifacts.render import render_version, theme
from tests import artifact_file_readers as FR
from tests.test_artifact_render_samples import deck, document, workbook

pytest.importorskip("openpyxl")
pytest.importorskip("docx")


def _tracker(rows: int = 12) -> S.ArtifactSpec:
    statuses = ["Done", "In Progress", "Blocked", "Pass", "Fail", "Not Started"]
    data = [[f"T-{i}", statuses[i % 6], f"2026-0{1 + i % 8}-1{i % 9}", 1000.5 * i, round(0.1 * (i % 10), 2), 40 + i * 5 % 60] for i in range(1, rows + 1)]
    return S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="Tracker", sheets=[S.Sheet(
        name="Tasks", columns=[S.Column(name="ID"), S.Column(name="Status"), S.Column(name="Due date", type="date"),
                               S.Column(name="Budget", type="currency", format=S.ColumnFormat(currency="INR")),
                               S.Column(name="Completion", type="percent"), S.Column(name="Score", type="integer")],
        rows=data, totals=[S.Total(column="Budget"), S.Total(column="Score", fn="average")])]))


# ------------------------------------------------------------------- XLSX --


def test_xlsx_professional_default(tmp_path):
    from openpyxl import load_workbook

    from app.artifacts.render.xlsx import render_xlsx

    path = render_xlsx(_tracker().body, tmp_path / "t.xlsx")
    wb = load_workbook(str(path))
    ws = wb["Tasks"]
    a1 = ws["A1"]
    assert a1.fill.fgColor.rgb.endswith("1F3864") and a1.font.bold and a1.font.color.rgb.endswith("FFFFFF") and a1.font.sz == 11
    assert ws.row_dimensions[1].height == 24 and a1.alignment.vertical == "center" and a1.border.bottom.style == "medium"
    assert ws.freeze_panes == "A2"
    assert ws["D2"].number_format.startswith('[>=10000000]"₹"') and ws["E2"].number_format == "0.0%" and ws["F2"].number_format == "#,##0"
    assert ws["C2"].number_format == "dd-mmm-yyyy"
    assert ws["D14"].value == "=SUBTOTAL(109,D2:D13)" and ws["F14"].value == "=SUBTOTAL(101,F2:F13)"
    assert ws["A14"].value == "Total" and ws["A14"].fill.fgColor.rgb.endswith("DCE6F2") and ws["A14"].font.color.rgb.endswith("1F3864")
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0 and ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.print_title_rows == "$1:$1" or ws.print_title_rows == "1:1"
    assert ws.oddFooter.right.text == "Page &P of &N" and ws.sheet_properties.tabColor.rgb.endswith("1F3864")
    # Never a black border, on any cell or conditional format.
    for row in ws.iter_rows():
        for cell in row:
            for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom):
                assert side is None or side.color is None or not str(side.color.rgb).endswith("000000"), cell.coordinate
    # Status colours are per VALUE: the success rule names "pass"; nothing paints the whole column.
    rules = [(r.priority, str(cf.sqref), r.formula[0] if r.formula else r.type, r.dxf) for cf in ws.conditional_formatting for r in cf.rules]
    success = next(r for r in rules if '"pass"' in str(r[2]))
    assert success[3].fill.bgColor.rgb.endswith("E3F2E6") and "fail" not in success[2]
    danger = next(r for r in rules if '"fail"' in str(r[2]))
    assert danger[3].fill.bgColor.rgb.endswith("FDE4E4")
    assert not any(r[2] == "TRUE" for r in rules), "no rule fills a status column wholesale"
    assert any(r[2] == "colorScale" and r[1] == "F2:F13" for r in rules) and any(r[2] == "colorScale" and r[1] == "E2:E13" for r in rules)
    band = next(r for r in rules if r[2] == "MOD(ROW(),2)=0")
    assert band[3].fill.bgColor.rgb.endswith("F3F6FA") and band[0] == max(r[0] for r in rules), "banding is the lowest priority"


def test_xlsx_conditional_format_priority_order(tmp_path):
    """user fills (stopIfTrue) > user conditions > automatic colours > banding,
    read from the rule priorities in the file."""
    from openpyxl import load_workbook

    from app.artifacts.render.xlsx import render_xlsx

    spec = _tracker()
    spec.body.style = ST.merge(None, ST.parse_style_request("row 5 yellow, cells B2:C4 light blue, Status column red for Blocked", "workbook")[0])
    path = render_xlsx(spec.body, tmp_path / "p.xlsx")
    ws = load_workbook(str(path))["Tasks"]
    rules = sorted(((r.priority, r) for cf in ws.conditional_formatting for r in cf.rules), key=lambda x: x[0])

    def level(r):
        f = r.formula[0] if r.formula else ""
        if f == "TRUE":
            return 0
        if '"blocked"' in f and r.dxf.fill.bgColor.rgb.endswith("C62828"):
            return 1
        if f == "MOD(ROW(),2)=0":
            return 3
        return 2

    levels = [level(r) for _, r in rules]
    assert levels == sorted(levels) and levels[0] == 0 and levels[-1] == 3
    fills = [r for _, r in rules if level(r) == 0]
    assert all(r.stopIfTrue for r in fills)
    assert {str(cf.sqref) for cf in ws.conditional_formatting for r in cf.rules if r.formula == ["TRUE"]} == {"A5:F5", "B2:C4"}
    assert ws["B3"].fill.fgColor.rgb.endswith("DCE6F2") and ws["A5"].fill.fgColor.rgb.endswith("FFD54F")


def test_styled_grid_evaluates_the_workbook_s_rules_like_excel(tmp_path):
    from app.artifacts.render import preview as P
    from app.artifacts.render.xlsx import render_xlsx

    spec = _tracker()
    spec.body.style = ST.merge(None, ST.parse_style_request("row 5 yellow, Status column red for Blocked", "workbook")[0])
    path = render_xlsx(spec.body, tmp_path / "g.xlsx")
    grid = P.grid_for(str(path), "xlsx")
    assert grid["header_styles"][0] == {"fill": "#1F3864", "color": "#FFFFFF", "bold": True}
    styles = grid["cell_styles"]
    by_value = {row[1]: i for i, row in enumerate(grid["rows"]) if isinstance(row[1], str)}
    assert styles[f"{by_value['Pass']}:1"]["fill"] == "#E3F2E6", "Pass is green, never red"
    assert styles[f"{by_value['In Progress']}:1"]["fill"] == "#FFF1C7"
    blocked = [i for i, row in enumerate(grid["rows"]) if row[1] == "Blocked" and i != 3]
    assert styles[f"{blocked[0]}:1"]["fill"] == "#C62828" and styles[f"{blocked[0]}:1"]["color"] == "#FFFFFF"
    assert styles["3:0"]["fill"] == "#FFD54F" and styles["3:1"]["fill"] == "#FFD54F", "the requested row outranks status colours and banding"
    # Banding: =MOD(ROW(),2)=0 paints spreadsheet rows 2, 4, 6 (window rows 0, 2, 4), not 3 or 5.
    assert styles["0:0"]["fill"] == "#F3F6FA" and styles["2:0"]["fill"] == "#F3F6FA"
    assert styles.get("1:0", {}).get("fill") != "#F3F6FA"
    assert grid["display"][0][3] == "₹1,000.50" and grid["display"][0][4] == "10.0%" and grid["display"][-1][0] == "Total"
    assert grid["display"][-1][3] == "₹78,039.00", "the SUBTOTAL is computed for display; the formula stays in rows"
    assert grid["rows"][-1][3] == "=SUBTOTAL(109,D2:D13)"


def test_formula_evaluator_grammar():
    from app.artifacts.render.preview import display_value, evaluate_formula

    cells = {(1, 2): "  Done ", (2, 2): 85.0, (3, 2): None, (1, 3): 'x"),HYPERLINK("y'}
    at = lambda c, r: cells.get((c, r))  # noqa: E731
    assert evaluate_formula('OR(TRIM($A2)="done",TRIM($A2)="pass")', at, anchor=(1, 2), at=(1, 2)) is True
    assert evaluate_formula("AND(ISNUMBER($B2),$B2>80)", at, anchor=(2, 2), at=(2, 2)) is True
    assert evaluate_formula("LEN(TRIM($C2))=0", at, anchor=(3, 2), at=(3, 2)) is True
    assert evaluate_formula("MOD(ROW(),2)=0", at, at=(1, 4)) is True
    assert evaluate_formula('TRIM($A2)="""),HYPERLINK(""y"', at, anchor=(1, 2), at=(1, 3)) is False
    assert evaluate_formula('TRIM($A2)="x""),HYPERLINK(""y"', at, anchor=(1, 2), at=(1, 3)) is True
    assert evaluate_formula('ISNUMBER(SEARCH("on",$A2))', at, anchor=(1, 2), at=(1, 2)) is True
    assert evaluate_formula("SUBTOTAL(109,B2:B2)", at) == 85.0
    with pytest.raises(ValueError):
        evaluate_formula("WEBSERVICE(\"http://x\")", at)
    assert display_value(1234567.891, "#,##0.00") == "1,234,567.89" and display_value(0.125, "0.0%") == "12.5%"


def test_sheet_notes_go_on_the_notes_sheet(tmp_path):
    from openpyxl import load_workbook

    from app.artifacts.render.xlsx import render_xlsx

    spec = _tracker(3)
    spec.body.sheets[0] = spec.body.sheets[0].model_copy(update={"notes": "=HYPERLINK(\"x\") is text"})
    wb = load_workbook(str(render_xlsx(spec.body, tmp_path / "n.xlsx")))
    notes = wb["Notes"]
    assert notes["A2"].value == "Tasks" and notes["B2"].value.startswith("=HYPERLINK") and notes["B2"].data_type == "s" and notes["B2"].quotePrefix


# ------------------------------------------------------------------- DOCX --


def _docx_xml(path, part="word/document.xml") -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf-8")


def test_docx_classy_default_and_requested_style(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn

    spec = document("generic", sections=3)
    spec, _, notes, unparsed = ST.apply_request(spec, "headings dark green in Georgia 17pt, landscape, table header maroon", formats=["docx"])
    assert unparsed == []
    report = render_version(spec, ["docx"], str(tmp_path), title_slug="d", version=1)
    path = tmp_path / report.files[0].filename
    d = Document(str(path))
    h1 = d.styles["Heading 1"]
    assert (str(h1.font.color.rgb), h1.font.name, h1.font.size.pt) == ("1E6B34", "Georgia", 17.0)
    assert h1.element.rPr.find(qn("w:rFonts")).get(qn("w:ascii")) == "Georgia"
    title = d.styles["Title"]
    assert title.element.pPr is None or title.element.pPr.find(qn("w:pBdr")) is None, "the template's title border is removed"
    assert title.element.rPr.find(qn("w:spacing")) is None, "no letter-spaced title"
    section = d.sections[0]
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (297, 210)
    fields = re.findall(r"<w:instrText[^>]*>([^<]*)<", "".join(_docx_xml(path, n) for n in zipfile.ZipFile(path).namelist() if n.startswith("word/footer")))
    assert "PAGE" in fields and "NUMPAGES" in fields
    table = next(t for t in d.tables if t.rows[0].cells[0].text == "Region")
    assert table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "7B1E1E"
    theme_xml = _docx_xml(path, "word/theme/theme1.xml")
    assert re.search(r'<a:minorFont>\s*<a:latin typeface="Calibri"', theme_xml) and re.search(r'<a:majorFont>\s*<a:latin typeface="Georgia"', theme_xml)
    with zipfile.ZipFile(path) as z:
        rels = b"".join(z.read(n) for n in z.namelist() if n.endswith(".rels"))
    for target in re.findall(rb'Target="([^"]+)"[^>]*TargetMode="External"', rels):
        assert target.startswith((b"http://", b"https://", b"mailto:")), target


def _assert_ordered(parent, order, where):
    names = [c.tag.split("}", 1)[-1] for c in parent]
    ranks = [order.index(n) for n in names if n in order]
    assert ranks == sorted(ranks), f"{where}: {names}"


def test_docx_property_children_follow_the_schema_sequence(tmp_path):
    """Word refuses a part whose pPr/tcPr/tblPr/trPr children are out of
    order; every element the writer adds is placed by the sequence."""
    from docx import Document

    from app.artifacts.render import docx as DX

    spec = document("executive_report", sections=3)
    spec.body.style = ST.merge(None, ST.parse_style_request("title background light blue, captions centered, table header dark green", "document")[0])
    report = render_version(spec, ["docx"], str(tmp_path), title_slug="o", version=1)
    d = Document(str(tmp_path / report.files[0].filename))
    body = d.element.body
    for ppr in body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr"):
        _assert_ordered(ppr, DX._PPR_ORDER, "pPr")
    for tcpr in body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr"):
        _assert_ordered(tcpr, DX._TCPR_ORDER, "tcPr")
    for tblpr in body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblPr"):
        _assert_ordered(tblpr, DX._TBLPR_ORDER, "tblPr")
    for trpr in body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}trPr"):
        _assert_ordered(trpr, DX._TRPR_ORDER, "trPr")
    for st in d.styles.element.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr"):
        _assert_ordered(st, DX._PPR_ORDER, "style pPr")
    for borders in body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders"):
        _assert_ordered(borders, DX._BORDER_ORDER, "tcBorders")


def test_docx_header_and_footer_text_are_plain_runs_and_indic_gets_a_complex_script_font(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn

    spec = S.ArtifactSpec(kind="document", document=S.DocumentSpec(title="सारांश report", author="Team", blocks=[S.Paragraph(text="ऑडिट पूरा हुआ। ઓડિટ પૂર્ણ થયું.")]))
    spec.body.style = ST.StyleSpec(header_footer=ST.HeaderFooter(header_text='PAGE \\* MERGEFORMAT {HYPERLINK "x"}', footer_text="NUMPAGES & co"))
    report = render_version(spec, ["docx"], str(tmp_path), title_slug="i", version=1)
    path = tmp_path / report.files[0].filename
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n).decode("utf-8") for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml", n)}
    instr = [t for xml in parts.values() for t in re.findall(r"<w:instrText[^>]*>([^<]*)<", xml)]
    assert sorted(instr) == ["NUMPAGES", "PAGE"], "only the fixed fields; the person's text is never a field"
    d = Document(str(path))
    assert d.sections[0].header.paragraphs[0].runs[0].text.startswith("PAGE \\* MERGEFORMAT")
    rfonts = d.styles["Normal"].element.rPr.find(qn("w:rFonts"))
    assert rfonts.get(qn("w:cs")) == "Nirmala UI" and rfonts.get(qn("w:ascii")) == "Calibri"


# -------------------------------------------------------------------- PDF --


def test_pdf_fonts_colours_and_page_box(tmp_path):
    import shutil
    import subprocess

    spec = document("generic", sections=2, toc=False)
    spec.body.style = ST.merge(None, ST.parse_style_request("headings #7B1E1E, body font DejaVu Serif, landscape", "document")[0])
    report = render_version(spec, ["pdf"], str(tmp_path), title_slug="p", version=1)
    path = str(tmp_path / report.files[0].filename)
    objects = FR.pdf_objects(path)
    h1 = [o for o in objects if o["kind"] == "text" and o["text"].startswith("Section 1") and o["size"] >= 17]
    assert h1 and all(o["fill"] == "#7B1E1E" for o in h1)
    assert objects[0]["page_width"] > objects[0]["page_height"], "landscape media box"
    if theme.font_installed("DejaVu Serif") and shutil.which("pdffonts"):
        listing = subprocess.run(["pdffonts", path], capture_output=True, text=True, check=True).stdout
        assert "DejaVuSerif" in listing.replace("-", "").replace(" ", "") or "DejaVu-Serif" in listing
    # A requested font the server lacks is named in exactly one sentence.
    spec2 = document("generic", sections=1)
    spec2.body.style = ST.StyleSpec(fonts=ST.FontSpec(body="Garamond"))
    second = tmp_path / "garamond"
    second.mkdir()
    report2 = render_version(spec2, ["pdf", "docx"], str(second), title_slug="p", version=1)
    if not theme.installed_family(ST.font_face("Garamond").installed_file_candidates):
        assert sum(1 for w in report2.warnings if w.startswith("Garamond is not installed")) == 1


def _pdf_font_matches(face: ST.FontFace, pdf_font: str) -> bool:
    name = pdf_font.split("+", 1)[-1].replace("-", "").replace(" ", "").casefold()
    chosen = theme.resolve_font(face)
    if chosen.family:
        return chosen.family.replace(" ", "").casefold() in name
    fallback = {"sans": "liberationsans", "serif": "liberationserif", "mono": "liberationmono"}[face.generic]
    return fallback in name or "dejavu" in name


PARITY_REQUESTS = [
    "headings dark blue",
    "headings #7B1E1E in Georgia, landscape",
    "body font DejaVu Serif, table header dark green",
    "modern look, table header teal with white text",
    "boardroom style, portrait",
    "headings navy bold, body font Arial 12pt, landscape",
    "table header orange, headings purple",
    "minimal look",
    "body font Times New Roman, headings maroon, landscape",
    "teal theme, table header #0B5563",
]


@pytest.mark.parametrize("request_text", PARITY_REQUESTS)
def test_docx_and_its_pdf_agree_on_the_produced_files(tmp_path, request_text):
    """Heading colour, body font family, table header fill and orientation
    read back from the DOCX equal those read back from the PDF."""
    from docx import Document

    spec = document("generic", sections=3, toc=False)
    spec, _, _, _ = ST.apply_request(spec, request_text, formats=["docx", "pdf"])
    report = render_version(spec, ["docx", "pdf"], str(tmp_path), title_slug="parity", version=1)
    files = {f.format: str(tmp_path / f.filename) for f in report.files}
    d = Document(files["docx"])
    objects = FR.pdf_objects(files["pdf"])
    docx_heading = FR.docx_observe(d, "Section 2:")
    pdf_heading = FR.pdf_observe(files["pdf"], "Section 2:", objects=objects)
    assert docx_heading["color"] == pdf_heading["color"]
    docx_body = FR.docx_observe(d, "This section explains")
    pdf_body = FR.pdf_observe(files["pdf"], "This section explains", objects=objects)
    assert _pdf_font_matches(ST.font_face(docx_body["font"]), pdf_body["font"]), (docx_body["font"], pdf_body["font"])
    docx_header = FR.docx_observe(d, "Revenue")
    pdf_header = FR.pdf_observe(files["pdf"], "Revenue", objects=objects)
    assert docx_header["background"] == pdf_header["background"]
    section = d.sections[0]
    assert (section.page_width > section.page_height) == (objects[0]["page_width"] > objects[0]["page_height"])


# ------------------------------------------------------------------- PPTX --


def test_pptx_slide_titles_and_table_header_follow_the_request(tmp_path):
    from pptx import Presentation

    spec = deck("generic")
    spec, _, _, unparsed = ST.apply_request(spec, "slide titles dark green in Cambria, table header maroon", formats=["pptx"])
    assert unparsed == []
    report = render_version(spec, ["pptx"], str(tmp_path), title_slug="k", version=1)
    prs = Presentation(str(tmp_path / report.files[0].filename))
    title = FR.pptx_observe(prs, "What happened")
    assert title["font"] == "Cambria" and title["color"] == "#1E6B34"
    header = FR.pptx_observe(prs, "Margin %")
    assert header["background"] == "#7B1E1E" and header["color"] == "#FFFFFF"


# ------------------------------------------------------------- back-compat --

GOLDEN = Path(__file__).parent / "fixtures" / "styling" / "legacy_golden.json"


def test_specs_without_a_style_render_the_same_text_and_cell_values_as_before(tmp_path):
    """Captured from the renderers BEFORE styling (cb0b4e3): every paragraph
    and table text of the document and deck, every cell of the workbook.
    The one intended difference: totals are SUBTOTAL, which evaluate to the
    same numbers as the SUM/AVERAGE they replace."""
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    from app.artifacts.render.preview import evaluate_formula

    golden = json.loads(GOLDEN.read_text(encoding="utf-8"))
    samples = {"document": (document("generic", sections=4), "docx"), "deck": (deck("generic"), "pptx"), "workbook": (workbook("generic", rows=12), "xlsx")}
    for label, (spec, fmt) in samples.items():
        out = tmp_path / label
        out.mkdir()
        report = render_version(spec, [fmt], str(out), title_slug=label, version=1)
        path = str(out / report.files[0].filename)
        if fmt == "docx":
            d = Document(path)
            texts = [p.text for p in d.paragraphs] + [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
            assert texts == golden[label]["docx"]
        elif fmt == "pptx":
            prs = Presentation(path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        texts.append(shape.text_frame.text)
                    if getattr(shape, "has_table", False) and shape.has_table:
                        texts.extend(" | ".join(c.text for c in r.cells) for r in shape.table.rows)
            assert texts == golden[label]["pptx"]
        else:
            wb = load_workbook(path)
            for title, rows in golden[label]["xlsx"].items():
                ws = wb[title]
                got = [[(c.value.isoformat() if hasattr(c.value, "isoformat") else c.value) for c in row] for row in ws.iter_rows()]
                assert len(got) == len(rows), title

                def value_at(sheet_rows):
                    return lambda col, r: sheet_rows[r - 1][col - 1] if 0 < r <= len(sheet_rows) and 0 < col <= len(sheet_rows[r - 1]) else None

                for r, (old, new) in enumerate(zip(rows, got), start=1):
                    for c, (a, b) in enumerate(zip(old, new), start=1):
                        if isinstance(a, str) and a.startswith("=") and isinstance(b, str) and b.startswith("=SUBTOTAL"):
                            assert evaluate_formula(a, value_at(rows), at=(c, r)) == evaluate_formula(b, value_at(got), at=(c, r)), (title, r, c)
                        else:
                            assert a == b, (title, r, c)


# ------------------------------------------------ verifier regressions --


def _owner_tracker(rows: int = 12) -> S.ArtifactSpec:
    statuses = ["Done", "In Progress", "Blocked", "Pass", "Fail", "Not Started"]
    data = [[f"T-{i}", "Asha" if i % 2 else "Ravi", statuses[i % 6], 1000.5 * i] for i in range(1, rows + 1)]
    return S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="Tracker", sheets=[S.Sheet(
        name="Tasks", columns=[S.Column(name="ID"), S.Column(name="Owner"), S.Column(name="Status"), S.Column(name="Budget", type="currency")],
        rows=data, totals=[S.Total(column="Budget")])]))


def test_a_font_only_column_rule_keeps_banding_and_status_colours(tmp_path):
    """'Owner column bold' must not carry stopIfTrue: it would hide the
    banding (and any status colour) under that column in Excel."""
    from openpyxl import load_workbook

    from app.artifacts.render import preview as P
    from app.artifacts.render.xlsx import render_xlsx

    spec = _owner_tracker()
    spec.body.style = ST.merge(None, ST.parse_style_request("Owner column bold, Status column text dark blue", "workbook")[0])
    path = render_xlsx(spec.body, tmp_path / "b.xlsx")
    ws = load_workbook(str(path))["Tasks"]
    true_rules = [r for cf in ws.conditional_formatting for r in cf.rules if r.formula == ["TRUE"]]
    assert true_rules and not any(r.stopIfTrue for r in true_rules)
    styles = P.grid_for(str(path), "xlsx")["cell_styles"]
    assert styles["0:1"]["bold"] and styles["0:1"]["fill"] == styles["0:0"]["fill"] == "#F3F6FA", "banding under the bold column"
    assert styles["0:2"]["color"] == "#1F3864" and styles["0:2"]["fill"] == "#FFF1C7", "user text colour over the status fill"


def test_a_requested_condition_is_not_hidden_by_a_requested_column_fill(tmp_path):
    from app.artifacts.render import preview as P
    from app.artifacts.render.xlsx import render_xlsx

    spec = _owner_tracker()
    spec.body.style = ST.merge(None, ST.parse_style_request("Status column light grey, Status red for Blocked", "workbook")[0])
    path = render_xlsx(spec.body, tmp_path / "c.xlsx")
    grid = P.grid_for(str(path), "xlsx")
    blocked = next(i for i, row in enumerate(grid["rows"]) if row[2] == "Blocked")
    done = next(i for i, row in enumerate(grid["rows"]) if row[2] == "Done")
    assert grid["cell_styles"][f"{blocked}:2"]["fill"] == "#C62828"
    assert grid["cell_styles"][f"{done}:2"]["fill"] == ST.resolve_color("light grey"), "the column fill still outranks automatic colours"


def test_the_styled_grid_evaluates_only_the_requested_page(tmp_path, monkeypatch):
    from app.artifacts.render import preview as P
    from app.artifacts.render.xlsx import render_xlsx

    spec = _tracker(1200)
    spec.body.style = ST.merge(None, ST.parse_style_request("row 1105 yellow", "workbook")[0])
    path = render_xlsx(spec.body, tmp_path / "w.xlsx")
    seen = []
    real = P.styled_window

    def spy(ws, first_row, last_row, width):
        seen.append((first_row, last_row))
        return real(ws, first_row, last_row, width)

    monkeypatch.setattr(P, "styled_window", spy)
    page = P.grid_for(str(path), "xlsx", offset=1100, limit=50)
    assert seen == [(1102, 1151)]
    assert len(page["display"]) == 50 and page["display"][0][0] == page["rows"][0][0] == "T-1101"
    assert page["cell_styles"]["3:0"]["fill"] == "#FFD54F", "sheet row 1105 is the 4th row of the page"
    assert "50:0" not in page["cell_styles"]
