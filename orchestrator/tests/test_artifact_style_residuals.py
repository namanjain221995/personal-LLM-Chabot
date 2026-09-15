"""AS3 fix round B3 (2026-09-15): the styling residuals the integrator saw.

- "highlight Priority Critical in orange" named a column and a value and
  was left for the model (nothing applied);
- a status word with no column ("Critical in orange") was aimed at Status
  even when the value lives in another column: the columns are searched;
- "bold total row" on a sheet with no totals row was refused with "there is
  no the totals row": a workbook with a quantity column gets the renderer's
  own SUBTOTAL formulas (computed by the spreadsheet from the rows), and
  everything else is refused in a plain sentence.

The data here is written for the test; no model is called.
"""
from __future__ import annotations

from openpyxl import load_workbook

from app.artifacts import edits as E
from app.artifacts import spec as S
from app.artifacts import style as ST
from app.artifacts.render import render_version

ORANGE = "#E07B00"


def tickets(*, amount_type: str = "number") -> S.ArtifactSpec:
    return S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="Tickets", sheets=[
        S.Sheet(name="Tickets", columns=[S.Column(name="Ticket ID", type="integer"), S.Column(name="Priority"), S.Column(name="Status"),
                                         S.Column(name="Year", type="integer"), S.Column(name="Amount", type=amount_type)],
                rows=[[101, "Critical", "Open", 2026, 10], [102, "Low", "Done", 2026, 20], [103, "Critical", "Blocked", 2025, 5]]),
    ]))


def labels_only() -> S.ArtifactSpec:
    return S.ArtifactSpec(kind="workbook", workbook=S.WorkbookSpec(title="People", sheets=[
        S.Sheet(name="People", columns=[S.Column(name="Name"), S.Column(name="Team")], rows=[["Asha", "Ops"], ["Ravi", "Sales"]]),
    ]))


def doc_with_table() -> S.ArtifactSpec:
    return S.parse_body("document", {"title": "Risks", "blocks": [
        {"type": "heading", "level": 1, "text": "Open items"},
        {"type": "table", "table": {"columns": ["Item", "Priority", "Status"], "rows": [["Backup", "Critical", "Open"], ["Docs", "Low", "Done"]]}},
    ]})


def _xlsx(spec: S.ArtifactSpec, tmp_path) -> "object":
    report = render_version(spec, ["xlsx"], str(tmp_path), title_slug="t", version=1)
    path = next(p for p in report.paths.values() if str(p).endswith(".xlsx"))
    return report, load_workbook(path)


def _fills_for(ws, column_letter: str):
    """(formula, fill rgb) of the conditional formats over one column."""
    out = []
    for cf in ws.conditional_formatting:
        if str(cf.sqref).startswith(column_letter):
            for rule in cf.rules:
                fill = rule.dxf.fill.fgColor.rgb if rule.dxf is not None and rule.dxf.fill is not None else None
                out.append((" ".join(rule.formula or []), fill))
    return out


# ------------------------------------------------------------- parser --


def test_a_column_then_a_value_then_a_colour_is_a_condition_on_that_column():
    for text in ("highlight Priority Critical in orange", "make Priority Critical orange", "Critical priority in orange"):
        patch, unparsed = ST.parse_style_request(text, "workbook")
        assert unparsed == [], text
        assert [(c.column.casefold(), c.op, str(c.value).casefold(), c.style.background) for c in patch.conditional] == [
            ("priority", "eq", "critical", ORANGE)], text


def test_element_words_are_never_read_as_a_column_name():
    patch, _ = ST.parse_style_request("highlight the header row in orange", "workbook")
    assert patch.conditional == []
    assert [r.target.kind for r in patch.rules] == ["table_header"]


# ------------------------------------------------- value → its column --


def test_a_status_word_with_no_column_is_bound_to_the_column_that_holds_it(tmp_path):
    spec = tickets()
    spec, _formats, notes, unparsed = ST.apply_request(spec, "Critical in orange", formats=["xlsx"])
    assert unparsed == []
    assert [c.column for c in spec.body.style.conditional] == ["Priority"]
    assert any("'Critical' is in the 'Priority' column" in n for n in notes), notes
    _, wb = _xlsx(spec, tmp_path)
    ws = wb["Tickets"]
    # Priority is column B: its first rule is the orange fill on "critical".
    assert _fills_for(ws, "B")[0] == ('TRIM($B2)="critical"', "00E07B00")
    assert all(fill != "00E07B00" for _, fill in _fills_for(ws, "C"))


def test_a_value_that_is_in_the_status_column_stays_there():
    spec = tickets()
    spec, _f, notes, _u = ST.apply_request(spec, "Blocked in red", formats=["xlsx"])
    assert [c.column for c in spec.body.style.conditional] == ["Status"]
    assert not any("column is coloured" in n for n in notes)


def test_a_value_found_nowhere_is_left_where_the_parser_put_it():
    spec = tickets()
    spec, _f, notes, _u = ST.apply_request(spec, "Rejected in red", formats=["xlsx"])
    assert [c.column for c in spec.body.style.conditional] == ["Status"]
    assert not any("column is coloured" in n for n in notes)


def test_the_highlight_named_on_its_column_renders_on_that_column(tmp_path):
    spec = tickets()
    spec, _f, _n, unparsed = ST.apply_request(spec, "highlight Priority Critical in orange", formats=["xlsx"])
    assert unparsed == []
    _, wb = _xlsx(spec, tmp_path)
    assert ('TRIM($B2)="critical"', "00E07B00") in _fills_for(wb["Tickets"], "B")


# ------------------------------------------------------- totals row --


def test_bold_total_row_adds_code_computed_totals_over_quantity_columns(tmp_path):
    spec = tickets()
    spec, _f, notes, _u = ST.apply_request(spec, "bold total row", formats=["xlsx"])
    sheet = spec.body.sheets[0]
    # Amount only: an ID and a year are integers but not quantities.
    assert [(t.column, t.fn) for t in sheet.totals] == [(4, "sum")]
    assert [r.target.kind for r in spec.body.style.rules] == ["table_total"]
    assert any("had no totals row, so one was added" in n and "Amount" in n for n in notes), notes
    assert not any("no the" in n for n in notes)
    _, wb = _xlsx(spec, tmp_path)
    ws = wb["Tickets"]
    total_row = [c for c in ws[5]]
    assert total_row[0].value == "Total" and total_row[0].font.b
    assert total_row[4].value == "=SUBTOTAL(109,E2:E4)" and total_row[4].font.b


def test_normalising_twice_adds_the_totals_row_once():
    spec = tickets()
    ST.apply_request(spec, "bold total row", formats=["xlsx"])
    _, notes = ST.normalize_spec_style(spec)
    assert len(spec.body.sheets[0].totals) == 1
    assert notes == []


def test_bold_total_row_with_no_quantity_column_is_refused_in_a_plain_sentence():
    spec = labels_only()
    spec, _f, notes, _u = ST.apply_request(spec, "bold total row", formats=["xlsx"])
    assert spec.body.sheets[0].totals == []
    assert notes == ["Not applied: the totals row (bold on) — the workbook has no totals row, and no quantity column (a number or an amount) to add one over."]


def test_bold_total_row_in_a_document_says_its_tables_have_no_totals_row():
    spec = doc_with_table()
    spec, _f, notes, _u = ST.apply_request(spec, "bold total row", formats=["docx"])
    assert notes == ["Not applied: the totals row (bold on) — the tables in this document have no totals row."]


# ------------------------------------------------------ prompt edits --


def test_edit_highlight_on_a_named_column_applies_and_says_which_cells():
    parent = tickets()
    text = "highlight Priority Critical in orange"
    out = E.apply(parent, E.preplan(text, parent), instruction=text)
    assert out.applied == ["“Critical” in Priority on an orange fill (#E07B00)"]
    assert [(c.column.casefold(), c.style.background) for c in out.spec.body.style.conditional] == [("priority", ORANGE)]
    assert E.preservation_guard(parent, out.spec, out.touched) == []


def test_edit_status_word_moves_to_the_column_that_holds_it():
    parent = tickets()
    text = "Critical in orange"
    out = E.apply(parent, E.preplan(text, parent), instruction=text)
    assert [c.column for c in out.spec.body.style.conditional] == ["Priority"]
    assert out.applied == ["“Critical” in Priority on an orange fill (#E07B00)"]
    assert any("'Critical' is in the 'Priority' column" in n for n in out.notes)


def test_edit_bold_total_row_adds_the_totals_row_and_marks_it_touched(tmp_path):
    parent = tickets()
    text = "bold total row"
    out = E.apply(parent, E.preplan(text, parent), instruction=text)
    assert out.not_applied == []
    assert out.applied == ["total row bold; a totals row added to 'Tickets' (spreadsheet formulas that sum Amount)"]
    assert "sheet0.totals" in out.touched
    assert [(t.column, t.fn) for t in out.spec.body.sheets[0].totals] == [(4, "sum")]
    assert E.preservation_guard(parent, out.spec, out.touched) == []
    _, wb = _xlsx(out.spec, tmp_path)
    assert wb["Tickets"]["E5"].value == "=SUBTOTAL(109,E2:E4)" and wb["Tickets"]["E5"].font.b


def test_edit_bold_total_row_is_refused_with_the_reason_when_nothing_can_be_totalled():
    parent = labels_only()
    out = E.apply(parent, E.preplan("bold total row", parent), instruction="bold total row")
    assert out.applied == []
    assert out.not_applied == [{"op": "set_style", "reason": "the workbook has no totals row, and no quantity column (a number or an amount) to add one over"}]
    doc = doc_with_table()
    out = E.apply(doc, E.preplan("bold total row", doc), instruction="bold total row")
    assert out.not_applied == [{"op": "set_style", "reason": "the tables in this document have no totals row"}]
