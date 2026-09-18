"""A spreadsheet is profiled like a CSV (B15, 2026-09-18).

Before this, profile_excel reported only each sheet's name, row count, column
NAMES and five sample rows: no types, nulls, ranges, top values, full rows or
computed totals. The same 50 rows were answerable uploaded as .csv and refused
uploaded as .xlsx. Each sheet is now streamed to a scratch CSV and handed to
profile_tabular, AFTER the zip-bomb caps, and the scratch is removed.
"""
from __future__ import annotations

import csv
import json
import os
import zipfile
from datetime import date, datetime

import pytest

from app.core import archive
from app.core import profile as profiler

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures", "dataset_truth")


def _first_rows(n: int):
    with open(os.path.join(FIXTURES, "sales_200.csv"), newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        return header, [row for _, row in zip(range(n), reader)]


def _typed(row):
    """The cells as a spreadsheet holds them: numbers, real dates, text."""
    order_id, order_date, region, channel, quantity, revenue = row
    return [int(order_id), date.fromisoformat(order_date), region, channel, int(quantity),
            float(revenue) if revenue else None]


def _write_xlsx(path, sheets):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    wb.save(path)
    wb.close()


@pytest.fixture()
def same_rows(tmp_path):
    header, rows = _first_rows(50)
    csv_path = tmp_path / "sales.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)
    xlsx_path = tmp_path / "sales.xlsx"
    _write_xlsx(xlsx_path, [("Orders", [header] + [_typed(r) for r in rows])])
    return str(csv_path), str(xlsx_path)


def test_the_same_50_rows_profile_identically_as_csv_and_xlsx(same_rows):
    csv_path, xlsx_path = same_rows
    as_csv = profiler.profile_tabular(csv_path)
    book = profiler.profile_excel(xlsx_path)
    assert book["kind"] == "spreadsheet" and len(book["sheets"]) == 1
    sheet = book["sheets"][0]
    assert sheet["name"] == "Orders" and sheet["rows"] == 50

    assert sheet["columns"] == as_csv["columns"], "types, nulls, ranges, top values, sums"
    assert sheet["aggregates"] == as_csv["aggregates"]
    assert sheet["aggregates"]["measures"] == ["quantity", "revenue"]
    assert sheet["full_rows"] == as_csv["full_rows"] and sheet["full_content"] is True
    assert sheet["sample_rows"] == as_csv["sample_rows"]
    assert {c["name"]: c["dtype"] for c in sheet["columns"]}["order_date"] == "DATE"


def test_every_sheet_is_profiled_and_a_timestamp_column_stays_a_timestamp(tmp_path):
    path = tmp_path / "two.xlsx"
    stamps = [datetime(2024, 1, 1, 0, 0), datetime(2024, 1, 2, 9, 30), datetime(2024, 2, 3, 0, 0)]
    _write_xlsx(
        path,
        [
            ("Visits", [["at", "day", "minutes"]]
             + [[stamps[i % 3], date(2024, 1, 1 + i % 28), i % 7] for i in range(90)]),
            ("Costs", [["item", "cost"]] + [[f"item {i % 4}", i * 1.25] for i in range(12)]),
        ],
    )
    book = profiler.profile_excel(str(path))
    visits, costs = book["sheets"]
    types = {c["name"]: c["dtype"] for c in visits["columns"]}
    # Midnight and non-midnight datetimes in one column: written as they
    # print, the CSV reader would sniff that column as text.
    assert types == {"at": "TIMESTAMP", "day": "DATE", "minutes": "BIGINT"}
    assert [r["month"] for r in visits["aggregates"]["by_month"][0]["rows"]] == ["2024-01", "2024-02"]
    assert costs["rows"] == 12 and costs["aggregates"]["measures"] == ["cost"]
    assert costs["columns"][1]["sum"] == sum(i * 1.25 for i in range(12))


def test_the_scratch_dir_is_gone_afterwards(same_rows, monkeypatch):
    _csv, xlsx_path = same_rows
    made = []
    real_mkdtemp = profiler.tempfile.mkdtemp
    monkeypatch.setattr(profiler.tempfile, "mkdtemp", lambda **kw: made.append(real_mkdtemp(**kw)) or made[-1])
    book = profiler.profile_excel(xlsx_path)
    assert "aggregates" in book["sheets"][0]
    assert made, "each sheet went through a scratch CSV"
    assert not any(os.path.exists(d) for d in made)


def test_a_zip_bomb_xlsx_raises_before_openpyxl_and_leaves_no_scratch(tmp_path, monkeypatch):
    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", b"<Types/>")
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (40 * 1024 * 1024))
    opened = []
    import openpyxl

    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: opened.append(a) or pytest.fail("opened"))
    made = []
    real_mkdtemp = profiler.tempfile.mkdtemp
    monkeypatch.setattr(profiler.tempfile, "mkdtemp", lambda **kw: made.append(real_mkdtemp(**kw)) or made[-1])
    with pytest.raises(archive.ArchiveError):
        profiler.profile_excel(str(bomb))
    assert opened == [], "the caps run before any reader opens the package"
    assert not any(os.path.exists(d) for d in made)


def test_without_a_scratch_dir_a_sheet_keeps_its_old_minimal_profile(same_rows, monkeypatch):
    _csv, xlsx_path = same_rows
    monkeypatch.setattr(profiler, "_scratch_dir", lambda: None)
    sheet = profiler.profile_excel(xlsx_path)["sheets"][0]
    assert sheet["name"] == "Orders" and sheet["rows"] == 50
    assert [c["name"] for c in sheet["columns"]][:2] == ["order_id", "order_date"]
    assert len(sheet["sample_rows"]) == 5
    json.dumps(sheet, default=str)
