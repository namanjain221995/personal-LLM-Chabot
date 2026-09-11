#!/usr/bin/env python3
"""End-to-end smoke for the Artifact Studio, against a RUNNING deployment.

    VIDEO_SMOKE_EMAIL=… VIDEO_SMOKE_PASSWORD=… \\
    scripts/artifact_smoke.py --base http://127.0.0.1:8081 --effort fast

It does what the browser does, in order: sign in, ask in chat for a
document, a deck and a workbook (and, with --follow-ups, an edit and a
conversion of the document), follow the SSE stream — printing every `step`
with a clock, then the sentence — read `meta.artifacts`, then exercise the
API exactly as the cards and the panel do: the job status, the version, an
inline preview with a Range request, every page image the viewer would lazy
load, the sheet grid for the workbook, and an `attachment` download of every
file, REOPENING each downloaded file with its own library (pypdfium2,
python-docx, python-pptx, openpyxl) so a file that is not a file fails here.

Credentials come from the environment (the same two variables
scripts/video_smoke.py uses) so a password is never on a command line. The
report is written next to the script's --out (default: .runtime/) as JSON
with per-stage timings, sizes, page counts and the verdict.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import sys
import time
from typing import Dict, List, Optional

import httpx

_PATHS = {
    "orchestrator": {"chat": "/chat", "auth": "/auth", "artifacts": "/artifacts"},
    "frontend": {"chat": "/api/chat", "auth": "/api/auth", "artifacts": "/api/artifacts"},
}
PATHS = _PATHS["orchestrator"]

PROMPTS = {
    "document": "Create a professional one-page executive brief about moving our team plan to $59 a month, for the CEO. Use the facts in this conversation; where you have to assume, say so.",
    "presentation": "Make a short CEO deck (about six slides) about the same pricing change, with one chart slide comparing the old and new monthly revenue at 120 team accounts.",
    "workbook": "Turn this into an Excel tracker: one sheet listing the three plans (Free $0, Team $59, Enterprise $199), their seats (5, 25, unlimited) and monthly revenue at 120/40/6 accounts, with a total row and a bar chart.",
}
FOLLOW_UPS = [
    ("edit", "Make the brief shorter and add a warning callout about churn risk."),
    ("convert", "Also give me the brief as Word."),
]


def _clock(t0: float) -> str:
    return f"{time.perf_counter() - t0:5.1f}s"


def login(client: httpx.Client, base: str, email: str, password: str) -> dict:
    r = client.post(f"{base}{PATHS['auth']}/login", json={"email": email, "password": password, "remember": True})
    if r.status_code != 200:
        raise SystemExit(f"login failed: {r.status_code} {r.text[:200]}")
    # The session cookie is `Secure` on this deployment; httpx's jar will not
    # replay it over plain http://, so it is pinned as a header — exactly the
    # bytes the browser would send through the HTTPS front door
    # (scripts/video_smoke.py does the same).
    raw = r.headers.get("set-cookie", "")
    pair = raw.split(";", 1)[0].strip()
    if "=" in pair:
        client.headers["Cookie"] = pair
    me = client.get(f"{base}{PATHS['auth']}/me")
    return me.json() if me.status_code == 200 else {}


def chat(client: httpx.Client, base: str, *, conversation_id: str, message: str, history: List[dict], effort: str, t0: float) -> dict:
    body = {
        "message": message,
        "messages": [*history, {"role": "user", "content": message}],
        "session_id": conversation_id,
        "conversation_id": conversation_id,
        "mode": "assistant",
        "model": "smart",
        "effort": effort,
        "web_search": "off",
    }
    steps: Dict[int, dict] = {}
    tokens: List[str] = []
    meta: dict = {}
    started = time.perf_counter()
    with client.stream("POST", f"{base}{PATHS['chat']}", json=body, timeout=httpx.Timeout(10.0, read=3600.0)) as r:
        if r.status_code != 200:
            raise SystemExit(f"/chat -> {r.status_code}: {r.read()[:400]!r}")
        event = None
        for line in r.iter_lines():
            if line.startswith("event:"):
                event = line[6:].strip()
                continue
            if not line.startswith("data:"):
                continue
            try:
                data = json.loads(line[5:].strip() or "{}")
            except json.JSONDecodeError:
                continue
            if event == "step":
                sid = int(data.get("id", 0))
                prev = steps.get(sid) or {}
                steps[sid] = {**prev, **data, "seen_at": round(time.perf_counter() - started, 1)}
                if data.get("status") != "running" or prev.get("status") != "running" or data.get("detail", "") != prev.get("detail", ""):
                    print(f"  [{_clock(t0)}] step {sid} {data.get('title')}: {data.get('status')} {data.get('detail', '')}", flush=True)
            elif event == "token":
                tokens.append(str(data.get("text", "")))
            elif event == "meta":
                meta = data
            elif event == "error":
                print(f"  [{_clock(t0)}] ERROR event: {data}", flush=True)
    return {"answer": "".join(tokens), "meta": meta, "steps": steps, "seconds": round(time.perf_counter() - started, 1)}


def _reopen(fmt: str, data: bytes) -> dict:
    """Open the downloaded bytes with the library that would read them."""
    if fmt == "pdf":
        import pypdfium2 as pdfium

        doc = pdfium.PdfDocument(data)
        try:
            return {"pages": len(doc)}
        finally:
            doc.close()
    if fmt == "docx":
        from docx import Document

        d = Document(io.BytesIO(data))
        return {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}
    if fmt == "pptx":
        from pptx import Presentation

        p = Presentation(io.BytesIO(data))
        return {"slides": len(p.slides)}
    if fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        formulas = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        formulas += 1
        return {"sheets": wb.sheetnames, "formulas": formulas}
    return {}


def exercise(client: httpx.Client, base: str, ref: dict, t0: float) -> dict:
    """The API the way the card and the panel use it."""
    out: dict = {"artifact_id": ref["artifact_id"], "version": ref["version"], "files": []}
    a = PATHS["artifacts"]
    st = client.get(f"{base}{a}/jobs/{ref['job_id']}")
    out["job_status"] = st.json().get("status") if st.status_code == 200 else f"http {st.status_code}"
    v = client.get(f"{base}{a}/{ref['artifact_id']}/v/{ref['version']}")
    out["version_ok"] = v.status_code == 200
    if ref.get("preview_kind") == "pages":
        pv = client.get(f"{base}{a}/{ref['artifact_id']}/v/{ref['version']}/preview", headers={"Range": "bytes=0-1023"})
        out["preview_range"] = pv.status_code
        pages = []
        for n in range(1, int(ref.get("preview_pages") or 0) + 1):
            t = time.perf_counter()
            img = client.get(f"{base}{a}/{ref['artifact_id']}/v/{ref['version']}/preview/{n}.png", params={"w": 1400})
            pages.append({"page": n, "http": img.status_code, "bytes": len(img.content), "ms": int((time.perf_counter() - t) * 1000)})
        out["pages"] = pages
        thumb = client.get(f"{base}{a}/{ref['artifact_id']}/v/{ref['version']}/preview/1.png", params={"w": 240})
        out["thumbnail"] = {"http": thumb.status_code, "bytes": len(thumb.content)}
    if ref.get("preview_kind") == "grid":
        g = client.get(f"{base}{a}/{ref['artifact_id']}/v/{ref['version']}/sheets")
        out["grid"] = {"http": g.status_code, "sheets": [s.get("name") for s in g.json().get("sheets", [])] if g.status_code == 200 else None,
                       "formulas": len((g.json().get("sheet") or {}).get("formulas") or {}) if g.status_code == 200 else None}
    for f in ref.get("files", []):
        t = time.perf_counter()
        d = client.get(f"{base}{a}/{ref['artifact_id']}/v/{ref['version']}/file/{f['format']}", params={"disposition": "attachment"})
        entry = {"format": f["format"], "http": d.status_code, "bytes": len(d.content), "ms": int((time.perf_counter() - t) * 1000),
                 "disposition": d.headers.get("content-disposition", "")[:40], "content_type": d.headers.get("content-type", "")}
        if d.status_code == 200:
            try:
                entry["reopened"] = _reopen(f["format"], d.content)
            except Exception as exc:  # noqa: BLE001 — that IS the finding
                entry["reopened"] = f"FAILED: {type(exc).__name__}: {exc}"
        out["files"].append(entry)
        print(f"  [{_clock(t0)}] {f['format']}: {entry['http']} {entry['bytes']} bytes → {entry.get('reopened')}", flush=True)
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--base", default=os.environ.get("VIDEO_SMOKE_BASE", "http://127.0.0.1:8081"))
    ap.add_argument("--via-frontend", default="", help="a Next.js origin, e.g. http://127.0.0.1:3001 — the browser's path")
    ap.add_argument("--effort", default="fast", choices=["fast", "think", "max"])
    ap.add_argument("--kinds", default="document,presentation,workbook")
    ap.add_argument("--follow-ups", action="store_true", help="after the document: an edit, then a conversion")
    ap.add_argument("--out", default=".runtime/artifact-smoke.json")
    args = ap.parse_args()
    global PATHS
    base = args.base
    if args.via_frontend:
        PATHS = _PATHS["frontend"]
        base = args.via_frontend
    email, password = os.environ.get("VIDEO_SMOKE_EMAIL"), os.environ.get("VIDEO_SMOKE_PASSWORD")
    if not email or not password:
        print("set VIDEO_SMOKE_EMAIL and VIDEO_SMOKE_PASSWORD", file=sys.stderr)
        return 2

    t0 = time.perf_counter()
    report: dict = {"base": base, "effort": args.effort, "turns": [], "verdict": "PASS"}
    with httpx.Client(follow_redirects=True) as client:
        me = login(client, base, email, password)
        print(f"[{_clock(t0)}] signed in as {me.get('username') or email}")
        conversation_id = f"artifact-smoke-{int(time.time())}"
        history: List[dict] = []
        kinds = [k.strip() for k in args.kinds.split(",") if k.strip()]
        plan = [(k, PROMPTS[k]) for k in kinds]
        if args.follow_ups and "document" in kinds:
            plan = [("document", PROMPTS["document"]), *FOLLOW_UPS, *[(k, PROMPTS[k]) for k in kinds if k != "document"]]
        for label, prompt in plan:
            print(f"\n[{_clock(t0)}] {label} ({args.effort}): {prompt[:80]}…")
            turn = chat(client, base, conversation_id=conversation_id, message=prompt, history=history, effort=args.effort, t0=t0)
            print(f"[{_clock(t0)}] answered in {turn['seconds']}s: {turn['answer'][:160]}")
            refs = (turn["meta"] or {}).get("artifacts") or []
            entry = {"label": label, "prompt": prompt, "seconds": turn["seconds"], "answer": turn["answer"][:400], "route": (turn["meta"] or {}).get("route"), "steps": turn["steps"], "artifacts": []}
            if not refs:
                print(f"  NO artifact in meta (route={entry['route']})")
                report["verdict"] = "FAIL"
            for ref in refs:
                print(f"  artifact {ref['artifact_id'][:8]} v{ref['version']} {ref['status']} {ref['kind']} files={[f['format'] for f in ref['files']]} pages={ref.get('preview_pages')} warnings={ref.get('warnings')}")
                ex = exercise(client, base, ref, t0)
                ex["ref"] = ref
                entry["artifacts"].append(ex)
                if ref["status"] not in ("completed", "completed_with_warnings") or any(not isinstance(f.get("reopened"), dict) for f in ex["files"]):
                    report["verdict"] = "FAIL"
            report["turns"].append(entry)
            history += [{"role": "user", "content": prompt}, {"role": "assistant", "content": turn["answer"]}]
    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(report, fh, indent=1, default=str)
    print(f"\n[{_clock(t0)}] verdict: {report['verdict']} · report: {args.out}")
    return 0 if report["verdict"] == "PASS" else 1


if __name__ == "__main__":
    sys.exit(main())
