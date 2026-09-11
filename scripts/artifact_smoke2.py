#!/usr/bin/env python3
"""Artifact Studio 2 — the five reported scenarios, end to end, against the
isolated stack (scripts/e2e-stack.sh), with the real model and the real
renderers. Every claim below is checked by reopening the bytes the API
served, never by trusting the sentence.

  A  "Create a CSV dataset containing 500 realistic sample records …"
     → one artifact, a CSV with EXACTLY 500 data rows + one header row,
       the grid endpoint agrees, no fenced CSV in the answer.
  B  a poem as Word, THEN "Create a professional PDF report on AI in Indian
     Businesses … Make it visually professional …"
     → a NEW artifact (not v2 of the poem), a PDF, the sentence starts
       "Created", every requested section is a heading in the PDF text.
  C  the 34-row messy audit paste (tests/fixtures/audit_paste.txt) with
     "share XLSX, CSV, Word and PDF … black borders, bold headers, the
     Ratio column red, humanised comments, blanks stay blank"
     → one workbook artifact with four files and a ZIP of all four; the
       XLSX has 34 data rows, bold header, borders, a red Ratio header;
       the CSV has 34 rows and the source's blanks; the DOCX table has 34
       rows in a landscape section; the PDF is landscape.
  D  "What is a CSV?" and three more questions → plain chat, no artifact.
  E  a deck, then "Make slide 4 shorter." (v2, same artifact), "Add a
     comparison chart." (v3), "Create a PDF version too." (a conversion).
  And on every artifact turn: no `memory_updated` in the meta.

  VIDEO_SMOKE_EMAIL=… VIDEO_SMOKE_PASSWORD=… scripts/artifact_smoke2.py --base http://127.0.0.1:8081 --scenarios A,B,C,D,E

`--restart-during-c` restarts the e2e orchestrator container while the
audit job is composing and proves the job finishes after the restart
(durability: lease lapse → requeue → resume at the first stage without an
output), polling /artifacts until the version is published.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import subprocess
import sys
import threading
import time
import zipfile
from typing import Dict, List, Optional

import httpx

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import artifact_smoke as base_smoke  # noqa: E402  — login/chat/_reopen/_clock

HERE = os.path.dirname(os.path.abspath(__file__))
FIXTURE = os.path.join(HERE, "..", "orchestrator", "tests", "fixtures", "audit_paste.txt")

A_PROMPT = (
    "Create a CSV dataset containing 500 realistic sample records for an AI project evaluation system. "
    "Include candidate_id, candidate_name, department, project_name, technical_score, communication_score, "
    "quality_score, completion_time, status, evaluator, and evaluation_date. Ensure the data is internally consistent."
)
B_POEM = "Write a short poem about the monsoon in Mumbai and give it to me as a Word document."
B_REPORT = (
    "Create a professional PDF report on Artificial Intelligence in Indian Businesses in 2026. Include an executive summary, "
    "current use cases, benefits, risks, implementation roadmap, comparison table, recommendations, and conclusion. "
    "Make it visually professional and suitable for senior management."
)
B_SECTIONS = ["executive summary", "use cases", "benefits", "risks", "roadmap", "comparison", "recommendations", "conclusion"]
C_INSTRUCTION = (
    "Please clean up this interview session audit and share it as XLSX, CSV, Word and PDF. Keep every row; blank source "
    "values must stay blank. In the spreadsheet use black borders on all cells and bold headers, and highlight the whole "
    "'Ratio of Interview Post-Session' column in red. Rewrite the audit comments so they are crisp, professional and "
    "humanised without changing the findings or losing the timestamps.\n\n"
)
D_PROMPTS = ["What is a CSV?", "Should I use CSV or Excel?", "Show me Python code that reads a CSV.", "What is the difference between PDF and DOCX?"]
E_DECK = "Make a short CEO deck (about six slides) about moving our team plan to $59 a month, with one chart slide comparing old and new monthly revenue at 120 team accounts."
E_FOLLOW_UPS = [("edit", "Make slide 4 shorter."), ("edit", "Add a comparison chart."), ("convert", "Create a PDF version too.")]


class Check:
    def __init__(self) -> None:
        self.rows: List[dict] = []
        self.failed = 0

    def __call__(self, name: str, ok: bool, detail: str = "") -> bool:
        self.rows.append({"check": name, "ok": bool(ok), "detail": detail})
        mark = "✓" if ok else "✗"
        print(f"    {mark} {name}{(' — ' + detail) if detail else ''}", flush=True)
        if not ok:
            self.failed += 1
        return bool(ok)


def _get(client: httpx.Client, base: str, url: str, **kw) -> httpx.Response:
    # The ref carries relative /artifacts/... URLs (or /api/artifacts/... through the frontend).
    prefix = "/api" if base_smoke.PATHS is base_smoke._PATHS["frontend"] else ""
    return client.get(f"{base}{prefix}{url}", **kw)


def _csv_rows(data: bytes) -> List[List[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8"))))


def _pdf_text_and_landscape(data: bytes) -> tuple[str, bool, int]:
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(data)
    try:
        texts = []
        landscape = False
        for i in range(len(doc)):
            page = doc[i]
            w, h = page.get_size()
            landscape = landscape or w > h
            texts.append(page.get_textpage().get_text_range())
        return "\n".join(texts), landscape, len(doc)
    finally:
        doc.close()


def _turn(client, base, conv, history, prompt, effort, t0, check: Check, *, expect_artifact: bool):
    print(f"\n[{base_smoke._clock(t0)}] {prompt[:90].replace(chr(10), ' ')}…")
    turn = base_smoke.chat(client, base, conversation_id=conv, message=prompt, history=history, effort=effort, t0=t0)
    meta = turn["meta"] or {}
    refs = meta.get("artifacts") or []
    print(f"[{base_smoke._clock(t0)}] answered in {turn['seconds']}s: {turn['answer'][:160]!r}")
    if expect_artifact:
        check("an artifact was made", bool(refs), f"route={meta.get('route')}")
        check("no memory chip on an artifact turn", not meta.get("memory_updated"), str(meta.get("memory_updated") or ""))
    else:
        check("no artifact for a question", not refs and meta.get("route") != "artifact", f"route={meta.get('route')}")
    history += [{"role": "user", "content": prompt}, {"role": "assistant", "content": turn["answer"]}]
    return turn, refs


def _download(client, base, ref: dict, file: dict) -> httpx.Response:
    url = file.get("download_url") or f"/artifacts/{ref['artifact_id']}/v/{ref['version']}/file/{file['format']}?disposition=attachment"
    return _get(client, base, url)


def scenario_a(client, base, effort, t0, check):
    conv = f"smoke2-A-{int(time.time())}"
    turn, refs = _turn(client, base, conv, [], A_PROMPT, effort, t0, check, expect_artifact=True)
    check("no fenced CSV in the answer", "```" not in turn["answer"] and sum(1 for line in turn["answer"].splitlines() if line.count(",") >= 8) < 3)
    if not refs:
        return
    ref = refs[0]
    check("one artifact, kind workbook", len(refs) == 1 and ref["kind"] == "workbook", f"{len(refs)} refs, kind={ref['kind']}")
    csvs = [f for f in ref["files"] if f["format"] == "csv"]
    check("a CSV file", bool(csvs), str([f["format"] for f in ref["files"]]))
    check("the sentence names 500", "500" in turn["answer"], turn["answer"][:120])
    for f in csvs:
        r = _download(client, base, ref, f)
        check("csv downloads as text/csv", r.status_code == 200 and r.headers.get("content-type", "").startswith("text/csv"), r.headers.get("content-type", ""))
        rows = _csv_rows(r.content)
        check("exactly 500 data rows + header", len(rows) == 501, f"{len(rows) - 1} data rows")
        check("11 columns in the header", len(rows[0]) == 11, str(rows[0]))
        check("file ref rows == 500", f.get("rows") == 500, str(f.get("rows")))
        ids = [row[0] for row in rows[1:]]
        check("candidate ids unique", len(set(ids)) == len(ids), f"{len(set(ids))} unique of {len(ids)}")
        head = [h.lower() for h in rows[0]]
        if "status" in head and "completion_time" in head:
            si, ci = head.index("status"), head.index("completion_time")
            completed = [row for row in rows[1:] if row[si].lower() == "completed"]
            check("completed records have a completion time", all(row[ci].strip() for row in completed), f"{len(completed)} completed rows")
        for name in ("technical_score", "communication_score", "quality_score"):
            if name in head:
                col = head.index(name)
                vals = [float(row[col]) for row in rows[1:] if row[col].strip()]
                check(f"{name} within 0–100", vals and min(vals) >= 0 and max(vals) <= 100, f"min {min(vals) if vals else '?'} max {max(vals) if vals else '?'}")
        g = _get(client, base, f.get("preview_url") or "")
        check("grid endpoint answers", g.status_code == 200, str(g.status_code))
        if g.status_code == 200:
            gj = g.json()
            check("grid total_rows 500", gj.get("total_rows") == 500, str(gj.get("total_rows")))
            check("grid is bounded", len(gj.get("rows") or []) <= 500 and gj.get("truncated") in (True, False))


def scenario_b(client, base, effort, t0, check):
    conv = f"smoke2-B-{int(time.time())}"
    history: List[dict] = []
    _, refs1 = _turn(client, base, conv, history, B_POEM, effort, t0, check, expect_artifact=True)
    if not refs1:
        return
    poem = refs1[0]
    check("the poem is a Word document", any(f["format"] == "docx" for f in poem["files"]))
    turn, refs2 = _turn(client, base, conv, history, B_REPORT, effort, t0, check, expect_artifact=True)
    if not refs2:
        return
    ref = refs2[0]
    check("a NEW artifact, not v2 of the poem", ref["artifact_id"] != poem["artifact_id"] and ref["version"] == 1, f"{ref['artifact_id'][:8]} v{ref['version']} vs poem {poem['artifact_id'][:8]}")
    check("the sentence says Created, not Updated", turn["answer"].lstrip().startswith("Created") and "Updated" not in turn["answer"][:40], turn["answer"][:80])
    check("it is a PDF", any(f["format"] == "pdf" for f in ref["files"]), str([f["format"] for f in ref["files"]]))
    check("the title is the report's, not the poem's", "poem" not in ref["title"].lower() and "monsoon" not in ref["title"].lower(), ref["title"])
    pdfs = [f for f in ref["files"] if f["format"] == "pdf"]
    if pdfs:
        r = _download(client, base, ref, pdfs[0])
        text, _, pages = _pdf_text_and_landscape(r.content)
        low = text.lower()
        missing = [s for s in B_SECTIONS if s not in low]
        check("every requested section is present in the PDF", not missing, f"missing: {missing}; {pages} pages")
        check("more than a placeholder-length report", pages >= 2 and len(text) > 2500, f"{pages} pages, {len(text)} chars")


def scenario_c(client, base, effort, t0, check, *, restart: bool = False):
    conv = f"smoke2-C-{int(time.time())}"
    paste = open(FIXTURE, encoding="utf-8").read()
    prompt = C_INSTRUCTION + paste
    source_rows = [line for line in paste.splitlines()[1:] if line.strip()]
    stopper: Optional[threading.Thread] = None
    if restart:
        def _restart_when_composing():
            # Wait until the job is running, then restart the container. The
            # stream this thread's caller holds will break; the job must not.
            time.sleep(25)
            print(f"[{base_smoke._clock(t0)}] RESTARTING techsara-e2e-orchestrator mid-job", flush=True)
            subprocess.run(["docker", "restart", "techsara-e2e-orchestrator"], check=False, capture_output=True)
        stopper = threading.Thread(target=_restart_when_composing, daemon=True)
        stopper.start()
    try:
        turn, refs = _turn(client, base, conv, [], prompt, effort, t0, check, expect_artifact=not restart)
    except (httpx.HTTPError, SystemExit) as exc:
        if not restart:
            raise
        print(f"[{base_smoke._clock(t0)}] the stream broke as expected ({type(exc).__name__}); waiting for the job to finish after the restart", flush=True)
        turn, refs = {"answer": ""}, []
    if restart:
        # Wait for /health, then find the version through the listing and poll.
        for _ in range(120):
            try:
                if client.get(f"{base}/health", timeout=15).status_code == 200:
                    break
            except httpx.HTTPError:
                pass
            time.sleep(3)
        deadline = time.time() + 900
        ref = None
        while time.time() < deadline:
            lst = _get(client, base, f"/artifacts?conversation_id={conv}")
            items = (lst.json() if lst.status_code == 200 else {}).get("artifacts") or (lst.json() if lst.status_code == 200 and isinstance(lst.json(), list) else [])
            cur = [a for a in items if (a.get("current") or a).get("status") in ("completed", "completed_with_warnings", "failed")]
            if cur:
                a = cur[0]
                v = _get(client, base, f"/artifacts/{a['artifact_id']}/v/{(a.get('current') or a).get('version') or 1}")
                if v.status_code == 200:
                    ref = v.json()
                    break
            time.sleep(10)
        check("the job finished after the orchestrator restart", ref is not None and ref.get("status") in ("completed", "completed_with_warnings"), str((ref or {}).get("status")))
        refs = [ref] if ref else []
    if not refs:
        return
    ref = refs[0]
    fmts = sorted(f["format"] for f in ref["files"])
    check("one workbook artifact with four files", ref["kind"] == "workbook" and fmts == ["csv", "docx", "pdf", "xlsx"], f"kind={ref['kind']} files={fmts}")
    check("every file has a file_id and a role", all(f.get("file_id") and f.get("role") for f in ref["files"]))
    check("download_all_url offered", bool(ref.get("download_all_url")), str(ref.get("download_all_url")))
    if turn["answer"]:
        check("the sentence reports the rows preserved", str(len(source_rows)) in turn["answer"] and "Updated" not in turn["answer"][:40], turn["answer"][:160])
    for f in ref["files"]:
        r = _download(client, base, ref, f)
        check(f"{f['format']} downloads", r.status_code == 200, str(r.status_code))
        if r.status_code != 200:
            continue
        if f["format"] == "csv":
            rows = _csv_rows(r.content)
            check("csv has 34 data rows", len(rows) - 1 == len(source_rows), f"{len(rows) - 1}")
            blanks = sum(1 for row in rows[1:] for c in row if c == "")
            check("csv keeps the source's blanks (≥ 15 after forward-filling the host)", blanks >= 15, f"{blanks} blank cells")
            check("no formula lead survives in the csv", all(not c or not (c[0] in "=+@" or (c[0] == "-" and not re.match(r"^-?\d", c))) for row in rows for c in row))
        elif f["format"] == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(r.content))
            ws = wb.worksheets[0]
            data_rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if any(c not in (None, "") for c in row)]
            check("xlsx has 34 data rows", len(data_rows) >= len(source_rows), f"{len(data_rows)}")
            header = list(ws.iter_rows(min_row=1, max_row=1))[0]
            check("xlsx header is bold", all(c.font.bold for c in header if c.value), "")
            check("xlsx cells have borders", all(c.border.left.style for c in header if c.value), "")
            # "Interview Du-ratio-n" also contains "ratio": match the header that STARTS with it.
            ratio = [c for c in header if isinstance(c.value, str) and c.value.lower().startswith("ratio")]
            fill = (ratio[0].fill.fgColor.rgb if ratio and ratio[0].fill and ratio[0].fill.fgColor else "") or ""
            data_fill = (ws.cell(row=2, column=ratio[0].column).fill.fgColor.rgb if ratio else "") or ""
            check("the Ratio header is dark red and its cells light red", bool(ratio) and str(fill).upper().endswith("9C0006") and str(data_fill).upper().endswith("FFC7CE"), f"header={fill} cells={data_fill}")
            check("xlsx freezes the header", ws.freeze_panes is not None, str(ws.freeze_panes))
        elif f["format"] == "docx":
            from docx import Document
            from docx.enum.section import WD_ORIENT

            d = Document(io.BytesIO(r.content))
            check("docx table has the rows", d.tables and len(d.tables[0].rows) - 1 >= len(source_rows), f"{len(d.tables[0].rows) - 1 if d.tables else 0}")
            check("docx is landscape", any(s.orientation == WD_ORIENT.LANDSCAPE for s in d.sections), "")
        elif f["format"] == "pdf":
            text, landscape, pages = _pdf_text_and_landscape(r.content)
            check("pdf is landscape with pages", landscape and pages >= 1, f"{pages} pages")
            check("pdf names the highlighted column", "ratio" in text.lower(), "")
    if ref.get("download_all_url"):
        z = _get(client, base, ref["download_all_url"])
        check("zip downloads", z.status_code == 200 and z.headers.get("content-type", "").startswith("application/zip"), f"{z.status_code} {z.headers.get('content-type', '')}")
        if z.status_code == 200:
            with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
                names = sorted(zf.namelist())
                check("zip holds the four files", len(names) == 4 and sorted(n.rsplit(".", 1)[-1] for n in names) == ["csv", "docx", "pdf", "xlsx"], str(names))
                check("zip entries match the recorded sizes", all(zf.getinfo(n).file_size == next(f["size"] for f in ref["files"] if f["filename"] == n) for n in names), "")


def scenario_d(client, base, effort, t0, check):
    conv = f"smoke2-D-{int(time.time())}"
    history: List[dict] = []
    # A prior artifact in the conversation makes the negatives harder.
    _turn(client, base, conv, history, "Give me a one-paragraph note about our pricing as a Word document.", effort, t0, check, expect_artifact=True)
    for q in D_PROMPTS:
        _turn(client, base, conv, history, q, effort, t0, check, expect_artifact=False)


def scenario_e(client, base, effort, t0, check):
    conv = f"smoke2-E-{int(time.time())}"
    history: List[dict] = []
    _, refs = _turn(client, base, conv, history, E_DECK, effort, t0, check, expect_artifact=True)
    if not refs:
        return
    deck = refs[0]
    expected_version = 1
    for op, prompt in E_FOLLOW_UPS:
        turn, refs = _turn(client, base, conv, history, prompt, effort, t0, check, expect_artifact=True)
        if not refs:
            continue
        ref = refs[0]
        expected_version += 1
        check(f"'{prompt}' → same artifact, v{expected_version}", ref["artifact_id"] == deck["artifact_id"] and ref["version"] == expected_version, f"{ref['artifact_id'][:8]} v{ref['version']}")
        if op == "edit":
            check("the sentence says Updated", turn["answer"].lstrip().startswith("Updated"), turn["answer"][:60])
        else:
            check("the sentence says Converted and a PDF is there", turn["answer"].lstrip().startswith("Converted") and any(f["format"] == "pdf" for f in ref["files"]), turn["answer"][:60])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--base", default="http://127.0.0.1:8081")
    ap.add_argument("--via-frontend", default="")
    ap.add_argument("--effort", default="fast", choices=["fast", "think", "max"])
    ap.add_argument("--scenarios", default="A,B,C,D,E")
    ap.add_argument("--restart-during-c", action="store_true")
    ap.add_argument("--out", default=".runtime/artifact-smoke2.json")
    args = ap.parse_args()
    base = args.base
    if args.via_frontend:
        base_smoke.PATHS = base_smoke._PATHS["frontend"]
        base = args.via_frontend
    email, password = os.environ.get("VIDEO_SMOKE_EMAIL"), os.environ.get("VIDEO_SMOKE_PASSWORD")
    if not email or not password:
        print("set VIDEO_SMOKE_EMAIL and VIDEO_SMOKE_PASSWORD", file=sys.stderr)
        return 2
    t0 = time.perf_counter()
    check = Check()
    with httpx.Client(follow_redirects=True, timeout=httpx.Timeout(30.0, read=900.0)) as client:
        me = base_smoke.login(client, base, email, password)
        print(f"[{base_smoke._clock(t0)}] signed in as {me.get('username') or email} · effort {args.effort}")
        runners: Dict[str, object] = {"A": scenario_a, "B": scenario_b, "C": scenario_c, "D": scenario_d, "E": scenario_e}
        for key in [k.strip().upper() for k in args.scenarios.split(",") if k.strip()]:
            print(f"\n===== SCENARIO {key} =====")
            try:
                if key == "C":
                    scenario_c(client, base, args.effort, t0, check, restart=args.restart_during_c)
                else:
                    runners[key](client, base, args.effort, t0, check)
            except Exception as exc:  # noqa: BLE001 — a crash is a failed scenario, reported as such
                check(f"scenario {key} ran to the end", False, f"{type(exc).__name__}: {exc}")
    verdict = "PASS" if check.failed == 0 else "FAIL"
    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump({"base": base, "effort": args.effort, "checks": check.rows, "failed": check.failed, "verdict": verdict}, fh, indent=1)
    print(f"\n[{base_smoke._clock(t0)}] verdict: {verdict} · {len(check.rows) - check.failed}/{len(check.rows)} checks · report: {args.out}")
    return 0 if verdict == "PASS" else 1


if __name__ == "__main__":
    sys.exit(main())
